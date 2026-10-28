"""Shared builders for individual downloads, Export All and release delivery."""
from io import BytesIO
import json
import os
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Service, PoamEntry, Finding, FindingObservation, PolicyFinding
from .service_evidence import current_execution, current_execution_id
from .native_vex import observation_excluded
from .risk_profiles import utcnow
from .architecture import build_architecture_graph
from .architecture_export import build_architecture_svg
from .helm_diagram import build_helm_diagram
from .purpose_exports import template_for_service, workbook_for

def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="17312B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = max(width, 12)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _focused_export_book(headers: list[str], rows: list[list]) -> Workbook:
    from .exchange import literal
    book = Workbook()
    sheet = book.active
    sheet.title = "CATS Export"
    sheet.append(headers)
    for row_number, row in enumerate(rows, 2):
        sheet.append([None] * len(headers))
        for index, value in enumerate(row, 1):
            literal(sheet.cell(row_number, index), value)
    format_sheet(sheet)
    return book


def _service_mitigations_book(service_key: str, db: Session):
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if service is None:
        raise HTTPException(404)
    entries = db.execute(select(PoamEntry.id, PoamEntry.service_version, PoamEntry.item_type,
        PoamEntry.title, PoamEntry.remediation, PoamEntry.status, PoamEntry.due_date, PoamEntry.ticket)
        .where(PoamEntry.service_id == service.id, PoamEntry.remediation.is_not(None), PoamEntry.remediation != "")
        .order_by(PoamEntry.id).execution_options(yield_per=500))
    rows = ([item.id, item.service_version, item.item_type, item.title, item.remediation,
             item.status, item.due_date.isoformat() if item.due_date else "", item.ticket or ""] for item in entries)
    return _focused_export_book(
        ["Entry ID", "Service Version", "Type", "Title", "Mitigation", "Status", "Due Date", "Ticket"], rows)


def _service_findings_book(service_key: str, db: Session):
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if service is None:
        raise HTTPException(404)
    latest_id = current_execution_id(db, service)
    now = utcnow()
    def export_rows():
        observations = db.execute(select(Finding.cve, Finding.severity, Finding.active,
            FindingObservation.image, FindingObservation.package, FindingObservation.installed_version,
            FindingObservation.fixed_version, Finding.first_seen, Finding.last_seen,
            observation_excluded(FindingObservation.id, now).label("vex_excluded"))
            .join(FindingObservation, FindingObservation.finding_id == Finding.id)
            .where(Finding.service_id == service.id, FindingObservation.execution_id == latest_id)
            .order_by(Finding.id, FindingObservation.id).execution_options(yield_per=500))
        for item in observations:
            yield ["Vulnerability", *item[:7], item.first_seen.isoformat(), item.last_seen.isoformat(),
                   "Resolved" if not item.active else "VEX Excluded" if item.vex_excluded else "Existing Policy"]
        policies = db.execute(select(PolicyFinding.finding, PolicyFinding.severity, PolicyFinding.active,
            PolicyFinding.target, PolicyFinding.first_seen, PolicyFinding.last_seen)
            .where(PolicyFinding.service_id == service.id).order_by(PolicyFinding.id)
            .execution_options(yield_per=500))
        for item in policies:
            yield ["Configuration", item.finding, item.severity, item.active, item.target or "",
                   "", "", "", item.first_seen.isoformat(), item.last_seen.isoformat(),
                   "Existing Policy" if item.active else "Resolved"]
    rows = export_rows()
    return _focused_export_book(
        ["Type", "Finding", "Severity", "Active", "Image or Target", "Package", "Installed Version",
         "Fixed Version", "First Seen", "Last Seen", "Effective Disposition"], rows)


def _service_diagrams_bytes(service_key: str, db: Session):
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if service is None:
        raise HTTPException(404)
    latest = current_execution(db, service)
    if latest is None or not latest.raw_payload:
        raise HTTPException(422, detail="No architecture evidence is available to diagram")
    graph = build_architecture_graph(latest.raw_payload or {})
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for view_name in ("all", "configuration", "containers", "flow", "network", "storage"):
            archive.writestr(f"architecture-{view_name}.svg", build_architecture_svg(graph, view_name))
        archive.writestr("helm-diagram.svg", build_helm_diagram(service, latest))
    return output.getvalue()


def _service_sbom_components(service_key: str, db: Session):
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if service is None:
        raise HTTPException(404)
    latest = current_execution(db, service)
    if latest is None or not (latest.raw_payload or {}).get("sbom_components"):
        raise HTTPException(422, detail="No retained SBOM component evidence is available")
    return {"format": "cats-retained-sbom-components-v1", "service": service_key,
        "execution_id": latest.execution_key, "components": latest.raw_payload["sbom_components"]}


def _service_evidence_zip(service_key: str, db: Session, category: str):
    """Download retained bytes from exactly the current version's assessment."""
    from .release_exports import _BoundedFile
    from .retained_documents import (EvidenceDocument, UnavailableDocument, SBOM_FORMATS,
                                     SCANNER_FORMATS, decode_document, validate_documents)

    if category not in {"scanner", "sbom"}:
        raise ValueError("Unknown evidence export category")
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if service is None:
        raise HTTPException(404)
    latest = current_execution(db, service)
    raw = (latest.raw_payload or {}) if latest else {}
    try:
        documents = [EvidenceDocument.model_validate(value) for value in raw.get("evidence_documents", [])]
        unavailable = [UnavailableDocument.model_validate(value) for value in raw.get("evidence_unavailable", [])]
        validate_documents(documents, unavailable)
    except (ValueError, TypeError) as exc:
        raise HTTPException(409, detail="Retained export evidence failed its integrity or size check") from exc
    documents = [document for document in documents if document.category == category]
    unavailable = [document.model_dump() for document in unavailable if document.category == category]
    if not documents:
        label = "original Grype/Trivy outputs" if category == "scanner" else "SBOM documents"
        reason = (" " + unavailable[0]["reason"]) if unavailable else " Run a new scan to retain this evidence."
        raise HTTPException(422, detail=f"No retained {label} are available for the current assessment.{reason}")
    supported = SCANNER_FORMATS if category == "scanner" else SBOM_FORMATS
    manifest = {"format": "cats-retained-evidence-v1", "service": service_key,
                "execution_id": latest.execution_key, "category": category,
                "supported_formats": list(supported), "entries": [], "unavailable": unavailable,
                "missing_formats": [name for name in supported if not any(doc.format == name for doc in documents)]}
    output = BytesIO()
    limit = max(1, int(os.getenv("CATS_EXPORT_MAX_BYTES", str(256 * 1024 * 1024))))
    with ZipFile(_BoundedFile(output, limit), "w", ZIP_DEFLATED) as archive:
        for document in documents:
            name = f"{document.format}/{document.name}"
            archive.writestr(name, decode_document(document))
            manifest["entries"].append({**document.model_dump(exclude={"content", "encoding"}),
                "archive_name": name,
                "provenance": "converted-from-syft" if document.derived_from_sha256 else "original"})
        archive.writestr("manifest.json", json.dumps(manifest, indent=2))
    return output.getvalue()


def _service_raw_scanner_bytes(service_key: str, db: Session):
    return _service_evidence_zip(service_key, db, "scanner")


def _service_sboms_bytes(service_key: str, db: Session):
    return _service_evidence_zip(service_key, db, "sbom")



def write_all(output, service, db, auth, *, bundle_builder=None, max_bytes=None):
    """Write the same permission-scoped ZIP for downloads and retained releases."""
    from .release_exports import write_export_archive
    from .service_transfer import export_service

    if not auth.has("service.export", service.id):
        raise PermissionError("Service export permission required")
    workbooks = {}
    for kind in ("ppsm", "poam", "asset_list"):
        columns, _, _ = template_for_service(db, kind, service)
        workbooks[f"{kind}.xlsx"] = lambda kind=kind, columns=columns: workbook_for(db, service, kind, columns)
    workbooks.update({f"{name}.xlsx": lambda build=build: build(service.service_key, db)
        for name, build in (("mitigations", _service_mitigations_book), ("findings", _service_findings_book))})
    # Generate the exchange bundle outside the missing-evidence handlers: export
    # errors must not silently turn a requested bundle into an unavailable note.
    supplemental = {}
    if auth.has("bundle.export", service.id):
        supplemental["cats-service.zip"] = (bundle_builder() if bundle_builder else
            export_service(db, service, auth.user.username))
    return write_export_archive(output, workbook_builders=workbooks,
        evidence_builders={"diagrams.zip": lambda: _service_diagrams_bytes(service.service_key, db),
                           "sbom-components.json": lambda: _service_sbom_components(service.service_key, db),
                           "raw-scanner-outputs.zip": lambda: _service_raw_scanner_bytes(service.service_key, db),
                           "sboms.zip": lambda: _service_sboms_bytes(service.service_key, db)},
        supplemental_entries=supplemental, max_bytes=max_bytes)
