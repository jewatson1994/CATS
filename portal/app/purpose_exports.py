"""Server-controlled, configurable operational exports; independent of exchange imports."""
from datetime import datetime
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .models import (Execution, FindingObservation, Group, InventoryRecord, PoamEntry,
                     PortalSetting, ServiceArtifact, ServiceImage)
from .overview import normalize_overview


CATALOG = {
    "poam": [
        ("service", "Service"), ("service_id", "Service ID"), ("finding_id", "Finding / CVE ID"),
        ("finding_type", "Finding Type"), ("finding_title", "Finding Title"),
        ("description", "Description"), ("severity", "Severity"), ("finding_status", "Finding Status"),
        ("first_seen", "First Seen"), ("last_seen", "Last Seen"), ("package", "Package"),
        ("installed_version", "Installed Version"), ("fixed_version", "Fixed Version"),
        ("image_artifact", "Image / Artifact"), ("item_type", "POA&M Item Type"),
        ("poam_title", "POA&M Title"), ("poam_status", "POA&M Status"),
        ("created_at", "POA&M Created"), ("due_date", "POA&M Due Date"),
        ("mitigation", "Mitigation"), ("exception_status", "Exception Status"),
        ("scanner", "Scanner / Source"), ("assessment_status", "Assessment Status"),
    ],
    "ppsm": [
        ("service", "Service"), ("service_id", "Service ID"), ("port", "Port"),
        ("protocol", "Protocol"), ("application", "Application / Service"),
        ("declared_by", "Declared By"), ("source_file", "Source File"),
        ("provenance", "Provenance"), ("boundary", "Boundary"), ("fqdn", "FQDN"),
        ("purpose", "Purpose"), ("service_version", "Service Version"),
    ],
    "asset_list": [
        ("service", "Service"), ("service_id", "Service ID"), ("asset_type", "Asset Type"),
        ("asset_name", "Asset Name"), ("image", "Image"), ("registry", "Registry"),
        ("repository", "Repository"), ("tag", "Tag"), ("digest", "Digest"),
        ("chart", "Chart"), ("chart_version", "Chart Version"),
        ("lifecycle_status", "Lifecycle Status"),
        ("assessment_status", "Assessment Status"), ("service_version", "Service Version"),
    ],
}
NAMES = {"poam": "POA&M", "ppsm": "PPSM", "asset_list": "Asset List"}
DEFAULT_ENABLED = {
    "poam": {"service", "service_id", "finding_id", "finding_type", "finding_title", "severity",
             "poam_title", "poam_status", "due_date", "mitigation", "exception_status"},
    "ppsm": {"service", "service_id", "port", "protocol", "application", "declared_by", "source_file", "provenance"},
    "asset_list": {"service", "service_id", "asset_type", "asset_name", "image", "digest",
                   "chart", "chart_version", "lifecycle_status", "assessment_status"},
}


def default_template(kind):
    if kind not in CATALOG:
        raise ValueError("Unknown export template")
    return [{"field": field, "heading": label, "enabled": field in DEFAULT_ENABLED[kind]}
            for field, label in CATALOG[kind]]


def validate_template(kind, columns):
    allowed = {field for field, _ in CATALOG.get(kind, [])}
    if not isinstance(columns, list) or len(columns) != len(allowed):
        raise ValueError("Template must include each available field exactly once")
    seen = set()
    headings = set()
    enabled = 0
    for column in columns:
        if not isinstance(column, dict) or column.get("field") not in allowed or column["field"] in seen:
            raise ValueError("Unknown or duplicate export field")
        seen.add(column["field"])
        heading = column.get("heading")
        if (not isinstance(heading, str) or not heading.strip() or len(heading) > 120
                or any(ord(c) < 32 for c in heading) or heading.lstrip().startswith(("=", "+", "-", "@"))):
            raise ValueError("Headings must be 1–120 printable characters")
        if not isinstance(column.get("enabled"), bool):
            raise ValueError("Enabled state must be a boolean")
        if column["enabled"]:
            enabled += 1
            if heading.strip().casefold() in headings:
                raise ValueError("Enabled headings must be unique")
            headings.add(heading.strip().casefold())
    if seen != allowed or not enabled:
        raise ValueError("Select at least one export field")
    return [{"field": c["field"], "heading": c["heading"].strip(), "enabled": c["enabled"]}
            for c in columns]


def setting_key(kind):
    if kind not in CATALOG:
        raise ValueError("Unknown export template")
    return f"purpose_export_template:{kind}"


def template_for(db, kind):
    setting = db.scalar(select(PortalSetting).where(PortalSetting.key == setting_key(kind)))
    if setting is None:
        return default_template(kind), False
    return validate_template(kind, json.loads(setting.value)), True


def policy_key(kind, group_id):
    return f"group:{group_id}:{setting_key(kind)}"


def template_policy(db, kind, group_id=None):
    """Resolve a group override, its ancestors, then the retained global template.

    A global customized template remains the CATS default for existing installs.
    A child can explicitly stop inheritance without deleting an ancestor override.
    """
    if group_id is None:
        columns, legacy = template_for(db, kind)
        return columns, ("CATS Default (existing template)" if legacy else "CATS Default"), "default"
    current = db.get(Group, group_id)
    visited = set()
    while current is not None:
        if current.id in visited:
            raise ValueError("Group policy inheritance cycle")
        visited.add(current.id)
        setting = db.scalar(select(PortalSetting).where(PortalSetting.key == policy_key(kind, current.id)))
        if setting is not None:
            payload = json.loads(setting.value)
            if payload.get("mode") == "default":
                return default_template(kind), "CATS Default", ("default" if current.id == group_id else "inherit")
            if payload.get("mode") == "custom":
                columns = validate_template(kind, payload.get("columns"))
                return columns, (f"Custom for {current.name}" if current.id == group_id else f"Inherited from {current.name}"), ("custom" if current.id == group_id else "inherit")
        current = db.get(Group, current.parent_id) if current.parent_id is not None else None
    columns, legacy = template_for(db, kind)
    return columns, ("CATS Default (existing template)" if legacy else "CATS Default"), "inherit"


def template_for_service(db, kind, service):
    """Match the existing multi-group policy precedence: highest group ID wins."""
    groups = sorted(service.groups, key=lambda group: group.id)
    return template_policy(db, kind, groups[-1].id if groups else None)


IMPORT_CROSSWALK = {
    "ppsm": {"port": "network.port", "protocol": "network.protocol",
             "application": "network.data_service", "boundary": "network.boundary",
             "fqdn": "network.fqdn", "purpose": "network.purpose"},
    "poam": {"finding_id": "finding.identifier", "description": "finding.description",
             "severity": "finding.severity", "mitigation": "finding.mitigation",
             "poam_status": "finding.status", "due_date": "finding.due_date"},
    "asset_list": {"asset_name": "asset.name", "asset_type": "asset.type", "fqdn": "asset.fqdn"},
}


def import_heading_aliases(db, dataset, service):
    """Only canonical, explicitly crosswalked fields may influence imports."""
    kind = "asset_list" if dataset == "assets" else dataset
    if kind not in IMPORT_CROSSWALK:
        return {}
    columns, _, _ = template_for_service(db, kind, service)
    aliases = {}
    for column in columns:
        field = IMPORT_CROSSWALK[kind].get(column["field"])
        if field and column["enabled"]:
            aliases.setdefault(field, []).append(column["heading"])
    return aliases


def _latest_execution(db, service):
    executions = db.scalars(select(Execution).where(Execution.service_id == service.id)
        .order_by(Execution.scanned_at.desc(), Execution.id.desc())).all()
    return next((e for e in executions if e.scan_scope == "service"), executions[0] if executions else None)


def _version(service, execution):
    payload = execution.raw_payload if execution and isinstance(execution.raw_payload, dict) else {}
    return str((payload.get("service") or {}).get("version") or service.manual_version or "Unknown")


def _poam_rows(db, service):
    rows = []
    entries = db.scalars(select(PoamEntry).where(PoamEntry.service_id == service.id).order_by(PoamEntry.id)).all()
    for entry in entries:
        finding, policy = entry.finding, entry.policy_finding
        observation = (db.scalar(select(FindingObservation).where(FindingObservation.finding_id == finding.id)
            .order_by(FindingObservation.id.desc()).limit(1)) if finding else None)
        exception = (max(finding.exceptions, key=lambda e: e.id, default=None) if finding else
                     max(policy.exceptions, key=lambda e: e.id, default=None) if policy else None)
        rows.append({
            "service": service.name, "service_id": service.service_key,
            "finding_id": finding.cve if finding else policy.finding if policy else "",
            "finding_type": "Vulnerability" if finding else "Configuration" if policy else entry.item_type,
            "finding_title": policy.title if policy else entry.title,
            "description": policy.description if policy else entry.description,
            "severity": finding.severity if finding else policy.severity if policy else "",
            "finding_status": ("Active" if finding.active else "Resolved") if finding else
                              ("Active" if policy.active else "Resolved") if policy else "",
            "first_seen": finding.first_seen if finding else policy.first_seen if policy else None,
            "last_seen": finding.last_seen if finding else policy.last_seen if policy else None,
            "package": observation.package if observation else "",
            "installed_version": observation.installed_version if observation else "",
            "fixed_version": observation.fixed_version if observation else "",
            "image_artifact": observation.image if observation else policy.target if policy else "",
            "item_type": entry.item_type, "poam_title": entry.title, "poam_status": entry.status,
            "created_at": entry.created_at, "due_date": entry.due_date,
            "mitigation": entry.remediation,
            "exception_status": ("Revoked" if exception.revoked_at else
                                 "Expired" if exception.expires_at and exception.expires_at.replace(tzinfo=None) < datetime.utcnow() else
                                 "Active") if exception else "None",
            "scanner": policy.scanner if policy else (observation.evidence or {}).get("data_source", "") if observation else "",
            "assessment_status": service.assessment_status,
        })
    return rows


def _ppsm_rows(db, service, execution):
    payload = execution.raw_payload if execution and isinstance(execution.raw_payload, dict) else {}
    overview = normalize_overview(payload.get("service_overview"))
    version = _version(service, execution)
    rows = [{"service": service.name, "service_id": service.service_key, "port": item.get("port"),
             "protocol": item.get("protocol"), "application": item.get("service"),
             "declared_by": item.get("declared_by"), "source_file": item.get("source_file"),
             "provenance": item.get("provenance"), "service_version": version}
            for item in overview["ports"]]
    for record in db.scalars(select(InventoryRecord).where(InventoryRecord.service_id == service.id,
                              InventoryRecord.version == version, InventoryRecord.dataset == "ppsm")):
        values = record.values or {}
        rows.append({"service": service.name, "service_id": service.service_key,
                     "port": values.get("network.port"), "protocol": values.get("network.protocol"),
                     "application": values.get("network.data_service"), "boundary": values.get("network.boundary"),
                     "fqdn": values.get("network.fqdn"), "purpose": values.get("network.purpose"),
                     "provenance": "Imported inventory", "service_version": version})
    return rows


def _image_parts(reference):
    value = str(reference or "")
    path, tag = (value.rsplit(":", 1) if ":" in value.rsplit("/", 1)[-1] else (value, ""))
    registry, _, repository = path.partition("/")
    return registry if "." in registry or ":" in registry else "", repository if repository else path, tag


def _asset_rows(db, service, execution):
    version = _version(service, execution)
    rows = []
    known_images = set()
    known_charts = set()
    for image in db.scalars(select(ServiceImage).where(ServiceImage.service_id == service.id).order_by(ServiceImage.id)):
        registry, repository, tag = _image_parts(image.image_reference)
        known_images.add((image.image_reference, image.image_digest or ""))
        rows.append({"service": service.name, "service_id": service.service_key, "asset_type": "Container image",
                     "asset_name": image.image_reference, "image": image.image_reference, "registry": registry,
                     "repository": repository, "tag": tag, "digest": image.image_digest,
                     "lifecycle_status": image.lifecycle_status, "assessment_status": image.scan_status,
                     "service_version": version})
    for chart in db.scalars(select(ServiceArtifact).where(ServiceArtifact.service_id == service.id,
                            ServiceArtifact.artifact_type == "helm_chart").order_by(ServiceArtifact.id)):
        known_charts.add((chart.chart_name or chart.artifact_name, chart.chart_version or ""))
        rows.append({"service": service.name, "service_id": service.service_key, "asset_type": "Helm chart",
                     "asset_name": chart.chart_name or chart.artifact_name, "chart": chart.chart_name,
                     "chart_version": chart.chart_version,
                     "lifecycle_status": chart.lifecycle_status, "assessment_status": service.assessment_status,
                     "service_version": version})
    payload = execution.raw_payload if execution and isinstance(execution.raw_payload, dict) else {}
    for artifact in normalize_overview(payload.get("service_overview"))["artifacts"]:
        kind = artifact.get("type")
        name = artifact.get("artifact") or ""
        artifact_version = artifact.get("version") or ""
        if kind == "Image":
            registry = artifact.get("registry") or ""
            repository = "/".join(part for part in (artifact.get("repository"), name) if part and part != "—")
            reference = f"{registry}/{repository}:{artifact_version}" if registry and artifact_version != "—" else repository
            digest = artifact.get("digest") if artifact.get("digest") != "—" else ""
            if (reference, digest) in known_images:
                continue
            rows.append({"service": service.name, "service_id": service.service_key,
                         "asset_type": "Container image", "asset_name": reference, "image": reference,
                         "registry": registry, "repository": repository, "tag": artifact_version,
                         "digest": digest, "assessment_status": service.assessment_status,
                         "service_version": version})
        elif kind == "Chart" and (name, artifact_version) not in known_charts:
            rows.append({"service": service.name, "service_id": service.service_key,
                         "asset_type": "Helm chart", "asset_name": name, "chart": name,
                         "chart_version": artifact_version, "assessment_status": service.assessment_status,
                         "service_version": version})
    for record in db.scalars(select(InventoryRecord).where(InventoryRecord.service_id == service.id,
                              InventoryRecord.version == version, InventoryRecord.dataset == "assets")):
        values = record.values or {}
        rows.append({"service": service.name, "service_id": service.service_key, "asset_type": values.get("asset.type"),
                     "asset_name": values.get("asset.name"), "service_version": version})
    return rows


def rows_for(db, service, kind):
    execution = _latest_execution(db, service)
    return (_poam_rows(db, service) if kind == "poam" else
            _ppsm_rows(db, service, execution) if kind == "ppsm" else
            _asset_rows(db, service, execution))


def workbook_for(db, service, kind, columns):
    columns = [c for c in validate_template(kind, columns) if c["enabled"]]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = NAMES[kind]
    sheet.append([column["heading"] for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17312B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows_for(db, service, kind):
        values = []
        for column in columns:
            value = row.get(column["field"])
            if isinstance(value, datetime):
                value = value.replace(tzinfo=None)
            elif isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                value = "'" + value
            values.append(value)
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for index, column in enumerate(columns, 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(55, max(16, len(column["heading"]) + 5))
    return workbook
