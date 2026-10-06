"""Declarative, offline XLSX exchange using canonical scan evidence.

No template expressions are evaluated. A mapping is a catalog key, direction,
default and presentation properties. Scan evidence always wins over metadata.
"""
from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import re
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .models import Execution, InventoryRecord, ServiceMetadata, PoamEntry
from .overview import normalize_overview
from .exchange_limits import workbook_bytes, workbook_expanded_bytes, workbook_members, workbook_rows, workbook_columns
from .exchange_validation import validate_values


META = {
    "classification": "Classification", "export.generated_at": "Date Exported",
    "export.generated_by": "Exported By", "system.owner": "Information System Owner",
    "system.name": "System Name", "system.poc.name": "POC Name",
    "system.poc.phone": "POC Phone", "system.poc.email": "POC E-Mail",
    "system.reviewed_at": "Date Reviewed / Updated", "system.reviewed_by": "Reviewed / Updated By",
    "system.component": "DoD Component", "system.registration": "DoD IT Registration No.",
    "system.type": "System Type", "system.omb_project": "OMB Project ID",
    "system.security_costs": "Security Costs", "service.name": "Service Name", "service.version": "Service Version",
}
FIELDS = {
    "ppsm": {"row.number": "#", "network.port": "Port", "network.protocol": "Protocol",
             "network.data_service": "Data Service", "network.boundary": "Boundary", "network.fqdn": "FQDN", "network.purpose": "Purpose"},
    "poam": dict(zip(
        ["identifier", "description", "controls", "office", "checks", "resources", "due_date", "milestone_id", "milestones", "milestone_changes", "source", "status", "comments", "raw_severity", "mitigation", "severity", "threat", "likelihood", "impact", "impact_description", "residual_risk", "recommendation"],
        ["POA&M Item ID", "Description", "Controls/APs", "Office/Org", "Security Checks", "Resources Required", "Scheduled Completion Date", "Milestone ID", "Milestone with Completion Dates", "Milestone Changes", "Source Identifying Vulnerability", "Status", "Comments", "Raw Severity", "Mitigations", "Severity", "Relevance of Threat", "Likelihood", "Impact", "Impact Description", "Residual Risk Level", "Recommendations"])),
    "assets": dict(zip(
        ["row.number", "type", "name", "nickname", "ip", "public_facing", "fqdn", "public_ip", "public_url", "virtual", "manufacturer", "model", "serial", "os_version", "memory", "location", "approval_status", "poc", "critical_information"],
        ["#", "Component Type", "Asset Name", "Nickname", "Asset IP Address", "Public Facing", "Public Facing FQDN", "Public Facing IP Address", "Public Facing URL(s)", "Virtual Asset?", "Manufacturer", "Model Number", "Serial Number", "OS/iOS/FW Version", "Memory Size/Type", "Location (P/C/S & state)", "Approval Status", "POC", "Critical Information"])),
}
FIELDS["poam"] = {"finding." + key: value for key, value in FIELDS["poam"].items()}
FIELDS["assets"] = {key if key == "row.number" else "asset." + key: value for key, value in FIELDS["assets"].items()}
IDENTITY = {"ppsm": ("network.port", "network.protocol", "network.data_service"), "assets": ("asset.name", "asset.type"), "poam": ("finding.identifier",)}
REQUIRED = {"ppsm": {"network.port", "network.protocol"}, "assets": {"asset.name"}, "poam": {"finding.identifier", "finding.description"}}


def builtins():
    result = {}
    for dataset, name in [("ppsm", "PPSM"), ("poam", "POA&M"), ("assets", "Asset List")]:
        metadata = list(META)[:10] + ["service.version"]
        if dataset == "poam":
            metadata = ["export.generated_at", "export.generated_by", "system.component", "system.name", "system.registration", "system.type", "system.poc.name", "system.poc.phone", "system.poc.email", "system.omb_project", "system.security_costs", "service.version"]
        result[dataset] = {"name": name, "dataset": dataset, "enabled": True, "block_missing": False,
            "metadata": [{"field": key, "label": META[key], "required": key == "system.name", "direction": "export"} for key in metadata],
            "columns": [{"field": key, "label": label, "required": key in REQUIRED[dataset], "direction": "export" if key == "row.number" else "both", "width": 24} for key, label in FIELDS[dataset].items()],
            "layout": {"sheet": name, "banner": "", "header_color": "17312B", "row_height": 32}}
    return result


def validate_template(value):
    value = deepcopy(value)
    if not isinstance(value, dict) or value.get("dataset") not in FIELDS:
        raise ValueError("Unsupported dataset")
    if not isinstance(value.get("name"), str) or not 1 <= len(value["name"]) <= 120:
        raise ValueError("Template name is required (120 characters maximum)")
    if not isinstance(value.get("description", ""), str) or len(value.get("description", "")) > 2000:
        raise ValueError("Description must be text (2,000 characters maximum)")
    for section, catalog in [("metadata", META), ("columns", {**META, **FIELDS[value["dataset"]]})]:
        fields = value.get(section, [])
        if not isinstance(fields, list) or len(fields) > 80 or (section == "columns" and not fields):
            raise ValueError("Templates require 1–80 columns and at most 80 metadata fields")
        labels = set()
        for field in fields:
            if not isinstance(field, dict) or field.get("field") not in catalog:
                raise ValueError("Unknown field mapping")
            label = field.get("label")
            if not isinstance(label, str) or not label.strip() or len(label) > 120 or label in labels:
                raise ValueError("Labels must be unique, nonempty, and at most 120 characters")
            labels.add(label)
            if field.get("direction", "both") not in {"both", "import", "export", "display"}:
                raise ValueError("Invalid mapping direction")
            if not isinstance(field.get("default", ""), (str, int, float, bool)) or len(str(field.get("default", ""))) > 2000:
                raise ValueError("Defaults must be bounded scalar values")
            if not isinstance(field.get("width", 24), (int, float)) or not 8 <= field.get("width", 24) <= 80:
                raise ValueError("Column width must be 8–80")
    layout = value.setdefault("layout", {})
    if not isinstance(layout, dict) or not re.fullmatch(r"[0-9a-fA-F]{6}", layout.get("header_color", "17312B")):
        raise ValueError("Header color must be a six-digit hex color")
    if not isinstance(layout.get("row_height", 32), (int, float)) or not 15 <= layout.get("row_height", 32) <= 120:
        raise ValueError("Row height must be 15–120")
    if len(str(layout.get("banner", ""))) > 200:
        raise ValueError("Banner is too long")
    return value


def version_of(execution):
    return str((execution.raw_payload.get("service") or {}).get("version") or "Unknown")


def versions(db, service):
    # Version menus need only the version scalar, read from execution summaries.
    from .evidence_reads import version_rows
    values = [value for _, value in version_rows(db, service.id)]
    return list(dict.fromkeys(str(value or "Unknown") for value in values)) or [service.manual_version or "Unknown"]


def selected_evidence(db, service, version):
    # Retain Python's legacy version coercion, but inspect only scalar metadata.
    # Other versions' large scan payloads must never be hydrated for this request.
    from .evidence_reads import version_rows
    metadata = version_rows(db, service.id)
    available = list(dict.fromkeys(str(value or "Unknown") for _, value in metadata)) or [service.manual_version or "Unknown"]
    version = version or available[0]
    if version not in available:
        raise ValueError("Service version not found")
    selected_ids = [identifier for identifier, value in metadata if str(value or "Unknown") == version]
    executions = []
    for offset in range(0, len(selected_ids), 400):
        executions.extend(db.scalars(select(Execution).where(
            Execution.service_id == service.id, Execution.id.in_(selected_ids[offset:offset + 400]),
        ).order_by(Execution.scanned_at.desc(), Execution.id.desc())).all())
    return version, executions


def record_key(dataset, values):
    return sha256(json.dumps([str(values.get(k) if values.get(k) is not None else "").strip() for k in IDENTITY[dataset]], separators=(",", ":")).encode()).hexdigest()


def context_and_rows(db, service, version, dataset, username):
    version, executions = selected_evidence(db, service, version)
    # Prefer the latest full-service scan for inventory; image-only scans must
    # never erase the remainder of a version's architecture.
    execution = next((e for e in executions if e.scan_scope == "service"), executions[0] if executions else None)
    payload = execution.raw_payload if execution else {}
    identity = payload.get("service") or {}
    metadata = db.get(ServiceMetadata, service.id)
    context = dict(metadata.values if metadata else {})
    authoritative = {"export.generated_at": datetime.now(timezone.utc).isoformat(), "export.generated_by": username,
        "system.name": identity.get("name") or service.name, "service.name": identity.get("name") or service.name,
        "system.owner": identity.get("owner"), "system.poc.name": identity.get("poc"), "service.version": version}
    if not executions:
        authoritative.update({"system.owner": service.owner, "system.poc.name": service.poc})
    context.update({k: v for k, v in authoritative.items() if v is not None and v != ""})
    overview = normalize_overview(payload.get("service_overview"), findings_images=payload.get("sbom_images", []))
    rows = []
    if dataset == "ppsm":
        for item in overview["ports"]:
            rows.append({"network.port": item.get("port"), "network.protocol": item.get("protocol"), "network.data_service": item.get("service"), "network.boundary": item.get("boundary"), "network.fqdn": item.get("fqdn"), "network.purpose": item.get("purpose")})
    elif dataset == "assets":
        for item in (payload.get("service_overview") or {}).get("assets", []):
            if isinstance(item, dict):
                rows.append({"asset." + key: value for key, value in item.items() if "asset." + key in FIELDS[dataset]})
        # Container images are evidence-backed assets, never fictional hardware.
        for image in overview["artifacts"]:
            if image.get("type") == "Image":
                reference = "/".join(str(image[k]) for k in ("registry", "repository", "artifact") if image.get(k) not in (None, "", "—"))
                rows.append({"asset.name": reference, "asset.os_version": image.get("version"), "asset.type": "Container image", "asset.virtual": True})
    else:
        for finding in payload.get("findings", []):
            evidence = finding.get("evidence") or {}
            rows.append({"finding.identifier": finding.get("cve"), "finding.description": evidence.get("description"), "finding.severity": finding.get("severity"), "finding.raw_severity": finding.get("severity"), "finding.source": evidence.get("data_source"), "finding.recommendation": evidence.get("recommendation") or evidence.get("remediation")})
        for item in payload.get("policy_findings", []):
            rows.append({"finding.identifier": item.get("finding"), "finding.description": item.get("description") or item.get("title"), "finding.controls": item.get("framework"), "finding.severity": item.get("severity"), "finding.raw_severity": item.get("severity"), "finding.source": item.get("scanner"), "finding.recommendation": item.get("remediation")})
    merged = {record_key(dataset, row): row for row in rows}
    if dataset == "poam":
        for entry in db.scalars(select(PoamEntry).where(PoamEntry.service_id == service.id, PoamEntry.service_version == version)):
            values = {**(entry.supplemental_fields or {}), "finding.identifier": entry.title,
                "finding.description": entry.description, "finding.mitigation": entry.remediation,
                "finding.status": entry.status, "finding.due_date": entry.due_date.isoformat() if entry.due_date else ""}
            key = entry.exchange_key or record_key(dataset, values)
            merged[key] = {**values, **{k: v for k, v in merged.get(key, {}).items() if v is not None and v != ""}}
    for item in db.scalars(select(InventoryRecord).where(InventoryRecord.service_id == service.id, InventoryRecord.version == version, InventoryRecord.dataset == dataset)):
        # Imported metadata fills gaps; immutable scanner data stays authoritative.
        merged[item.record_key] = {**item.values, **{k: v for k, v in merged.get(item.record_key, {}).items() if v is not None and v != ""}}
    return context, list(merged.values())


def resolve(mapping, context, row=None):
    key = mapping["field"]
    for source in (row or {}, context, {key: mapping.get("default")}):
        if key in source and source[key] is not None and source[key] != "":
            return source[key]
    return ""


def missing_fields(template, context, rows):
    missing = []
    for mapping in template.get("metadata", []):
        if mapping.get("direction") != "import" and mapping.get("required") and resolve(mapping, context) == "":
            missing.append(mapping["label"])
    for number, row in enumerate(rows, 1):
        for mapping in template["columns"]:
            if mapping.get("direction") != "import" and mapping.get("required") and resolve(mapping, context, {**row, "row.number": number}) == "":
                missing.append(f"Row {number}: {mapping['label']}")
    return missing


def literal(cell, value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Never execute imported/formula-looking values.


def workbook(template, context, rows):
    template = validate_template(template)
    result = Workbook()
    sheet = result.active
    layout = template["layout"]
    sheet.title = re.sub(r"[\\/*?:\[\]]", "-", str(layout.get("sheet") or template["name"]))[:31] or "Export"
    columns = [item for item in template["columns"] if item.get("direction") != "import"]
    if not columns:
        raise ValueError("Template has no exportable columns")
    width = max(2, len(columns))
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    literal(sheet.cell(1, 1), layout.get("banner") or context.get("classification") or "Classification not supplied")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    literal(sheet.cell(2, 1), template["name"])
    sheet.cell(2, 1).font = Font(size=18, bold=True)
    row_number = 4
    for mapping in template.get("metadata", []):
        if mapping.get("direction") == "import":
            continue
        literal(sheet.cell(row_number, 1), mapping["label"] + (" *" if mapping.get("required") else ""))
        sheet.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=width)
        literal(sheet.cell(row_number, 2), resolve(mapping, context))
        row_number += 1
    row_number += 1
    header = row_number
    for column, mapping in enumerate(columns, 1):
        cell = sheet.cell(header, column)
        literal(cell, mapping["label"])
        cell.fill = PatternFill("solid", fgColor=layout.get("header_color", "17312B"))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = mapping.get("width", 24)
    sheet.row_dimensions[header].height = 45
    for number, row in enumerate(rows, 1):
        for column, mapping in enumerate(columns, 1):
            cell = sheet.cell(header + number, column)
            literal(cell, resolve(mapping, context, {**row, "row.number": number}))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[header + number].height = layout.get("row_height", 32)
    sheet.freeze_panes = f"A{header + 1}"
    sheet.auto_filter.ref = f"A{header}:{get_column_letter(len(columns))}{max(header, sheet.max_row)}"
    sheet.print_title_rows = f"1:{header}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return result


def parse_workbook(data, template, extra_aliases=None):
    """Bounded workbook parser. It never evaluates formulas or trusts hidden IDs."""
    limit = workbook_bytes()
    if len(data) > limit:
        raise ValueError(f"Workbook exceeds configured {limit:,}-byte limit")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > workbook_members() or sum(m.file_size for m in members) > workbook_expanded_bytes():
                raise ValueError("Workbook exceeds expanded/member safety limits")
            if len({m.filename for m in members}) != len(members):
                raise ValueError("Workbook contains duplicate ZIP entries")
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, ParseError) as exc:
        raise ValueError("Malformed XLSX workbook") from exc
    mappings = [m for m in template["columns"] if m.get("direction", "both") in {"both", "import"} and m["field"] in FIELDS[template["dataset"]] and m["field"] != "row.number"]
    domain_required = REQUIRED[template["dataset"]]
    if not domain_required.issubset({m["field"] for m in mappings}):
        book.close()
        raise ValueError("Template lacks required import identity fields")
    required = {m["field"] for m in mappings if m.get("required") or m["field"] in domain_required}
    def normalized(value):
        return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())
    headings = {}
    for mapping in mappings:
        for heading in (mapping["label"], FIELDS[template["dataset"]][mapping["field"]], mapping["field"],
                        *mapping.get("aliases", []), *(extra_aliases or {}).get(mapping["field"], [])):
            headings.setdefault(normalized(heading), set()).add(mapping["field"])
    by_field = {mapping["field"]: mapping for mapping in mappings}
    rows, errors, ignored, seen = [], [], [], set()
    try:
        candidates = []
        for sheet in book.worksheets:
            if sheet.max_row > workbook_rows() + 100 or sheet.max_column > workbook_columns():
                raise ValueError("Workbook exceeds configured row or column limit")
            for number, cells in enumerate(sheet.iter_rows(max_row=min(sheet.max_row, 100)), 1):
                labels = [str(c.value or "").strip() for c in cells]
                fields = [headings.get(normalized(label), set()) if label else set() for label in labels]
                matched = {next(iter(item)) for item in fields if len(item) == 1}
                if required.issubset(matched):
                    candidates.append((len(matched), sheet, number, labels, fields))
        if not candidates:
            raise ValueError("Required columns were not found. Check workbook headings or configure aliases.")
        candidates.sort(key=lambda item: item[0], reverse=True)
        if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
            raise ValueError("Multiple possible tables were found. Use a workbook with one clear table.")
        _, sheet, header_number, headers, matched_fields = candidates[0]
        if any(len(item) > 1 for item in matched_fields):
            raise ValueError("Ambiguous column heading; configure a unique alias before importing")
        selected_fields = [next(iter(item)) if item else None for item in matched_fields]
        used = [field for field in selected_fields if field]
        if len(used) != len(set(used)):
            raise ValueError("Multiple columns map to the same field; remove or rename one column")
        ignored = [label for label, field in zip(headers, selected_fields) if label and not field]
        iterator = sheet.iter_rows(min_row=header_number + 1)
        for number, cells in enumerate(iterator, header_number + 1):
            if all(c.value is None for c in cells):
                continue
            values = {}
            for field, cell in zip(selected_fields, cells):
                if field:
                    if cell.data_type == "f":
                        errors.append(f"Row {number}: formulas are not importable")
                    value = cell.value
                    if isinstance(value, datetime):
                        value = value.isoformat()
                    if value is not None and len(str(value)) > 8000:
                        errors.append(f"Row {number}: cell exceeds 8,000 characters")
                    values[field] = value
            for mapping in mappings:
                if (mapping.get("required") or mapping["field"] in domain_required) and values.get(mapping["field"]) in (None, ""):
                    errors.append(f"Row {number}: missing {mapping['label']}")
            if template["dataset"] == "ppsm":
                try:
                    raw_port = values.get("network.port")
                    port = int(raw_port)
                    if isinstance(raw_port, bool) or str(raw_port).strip() not in {str(port), str(float(port))}:
                        raise ValueError()
                    if not 1 <= port <= 65535:
                        raise ValueError()
                    values["network.port"] = port
                except (ValueError, TypeError):
                    errors.append(f"Row {number}: invalid port")
                if str(values.get("network.protocol", "")).upper() not in {"TCP", "UDP", "SCTP"}:
                    errors.append(f"Row {number}: unsupported protocol")
                else:
                    values["network.protocol"] = str(values["network.protocol"]).upper()
            if template["dataset"] == "poam" and values.get("finding.due_date"):
                try:
                    datetime.fromisoformat(str(values["finding.due_date"]))
                except ValueError:
                    errors.append(f"Row {number}: scheduled completion date must be an Excel date or ISO date")
            if len(str(values.get("finding.identifier") or "")) > 240:
                errors.append(f"Row {number}: POA&M identifier exceeds 240 characters")
            errors.extend(f"Row {number}: {error}" for error in validate_values(values))
            key = record_key(template["dataset"], values)
            if key in seen:
                errors.append(f"Row {number}: duplicate record identity")
            seen.add(key)
            rows.append({"key": key, "values": values})
            if len(rows) > workbook_rows():
                raise ValueError("Workbook exceeds configured row limit")
    finally:
        book.close()
    return {"rows": rows, "errors": errors, "ignored": ignored, "mapped": [by_field[field]["label"] for field in used], "recognized": len(rows), "worksheet": sheet.title, "header_row": header_number}
