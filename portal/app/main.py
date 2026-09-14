import os
import re
import logging
import secrets
import json
import hashlib
import urllib.parse
import urllib.request
import urllib.error
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import tarfile
import uuid
from functools import lru_cache
from concurrent.futures import ThreadPoolExecutor
from zipfile import ZIP_DEFLATED, ZipFile
from io import BytesIO
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo, available_timezones

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile, status
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, PlainTextResponse, RedirectResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.gzip import GZipMiddleware
from sqlalchemy import and_, case, delete, false, func, inspect, or_, select, text, true
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import yaml

from .database import Base, engine, get_db, SessionLocal
from .auth import (
    PERMISSIONS, SESSION_COOKIE, AuthContext, hash_password, record_audit,
    require_permission, require_user, optional_user, seed_auth, token_hash, utcnow as auth_utcnow,
    verify_password, oidc_enabled, oidc_authorization_url, oidc_exchange_code,
    provision_oidc_user, verify_oidc_id_token,
)
from .models import (
    ExceptionRecord, Execution, Finding, FindingObservation, PolicyFinding,
    PolicyExceptionRecord, Service, Group,
    ServiceArchiveEvent, ServiceDeletionAudit, ServiceImage, User, UserSession, Role,
    UserRoleAssignment, WorkflowRequest, AuditEvent, PortalSetting, PoamEntry,
    ServiceGroup,
    PoamHistory, PoamChangeRequest,
    PatchExecution,
    RemediationExecution,
)
from .schemas import ExecutionPayload
from .policy_data import epss_scores, kev_cves, risk_metadata
from .overview import normalize_overview
from .service_export import build_service_workbook, service_export_filename
from .helm_diagram import build_helm_diagram
from .architecture import build_architecture_graph
from .architecture_export import build_architecture_svg
from .report_html import build_public_scan_report
from .patching import PATCH_PHASES, advance_patch_stages, initial_patch_stages, redact, registry_host, safe_job_config
from .secrets import encrypt_secret, decrypt_secret, secret_configured
from . import signing
from .admin_config import OS_DEFINITIONS, PACKAGE_MANAGERS, certificate_bundle_metadata, merge_certificate_metadata, normalize_policy, parse_json, policy_bool, test_repository, validate_os_definition, validate_policy
from .remediation import associate_patch_results, build_plan, candidate_files, classify_policy_finding, plan_yaml, static_validation


Base.metadata.create_all(bind=engine)
# Lightweight additive upgrade for deployments created before service POC
# metadata existed. This keeps the existing create_all-based deployment model
# upgrade-safe without touching or deleting evidence.
with engine.begin() as connection:
    if "poc" not in {column["name"] for column in inspect(connection).get_columns("services") }:
        connection.execute(text("ALTER TABLE services ADD COLUMN poc VARCHAR(240)"))
    if "manual_version" not in {column["name"] for column in inspect(connection).get_columns("services") }:
        connection.execute(text("ALTER TABLE services ADD COLUMN manual_version VARCHAR(120)"))
    if "description" not in {column["name"] for column in inspect(connection).get_columns("services") }:
        connection.execute(text("ALTER TABLE services ADD COLUMN description TEXT"))
    if "lifecycle_status" not in {column["name"] for column in inspect(connection).get_columns("services") }:
        connection.execute(text("ALTER TABLE services ADD COLUMN lifecycle_status VARCHAR(20) DEFAULT 'active'"))
        connection.execute(text("UPDATE services SET lifecycle_status = 'staged' WHERE name LIKE 'Staged %' AND (lifecycle_status IS NULL OR lifecycle_status = 'active')"))
    if "theme" not in {column["name"] for column in inspect(connection).get_columns("users") }:
        connection.execute(text("ALTER TABLE users ADD COLUMN theme VARCHAR(30) DEFAULT 'cats'"))
    if "last_login_at" not in {column["name"] for column in inspect(connection).get_columns("users") }:
        connection.execute(text("ALTER TABLE users ADD COLUMN last_login_at TIMESTAMP"))
    if "group_id" not in {column["name"] for column in inspect(connection).get_columns("user_role_assignments") }:
        connection.execute(text("ALTER TABLE user_role_assignments ADD COLUMN group_id INTEGER"))
    if "group_id" not in {column["name"] for column in inspect(connection).get_columns("portal_settings") }:
        connection.execute(text("ALTER TABLE portal_settings ADD COLUMN group_id INTEGER"))
    if "poam_id" not in {column["name"] for column in inspect(connection).get_columns("workflow_requests") }:
        connection.execute(text("ALTER TABLE workflow_requests ADD COLUMN poam_id INTEGER"))
    if "policy_finding_id" not in {column["name"] for column in inspect(connection).get_columns("workflow_requests") }:
        connection.execute(text("ALTER TABLE workflow_requests ADD COLUMN policy_finding_id INTEGER"))
    if "bulk_group_id" not in {column["name"] for column in inspect(connection).get_columns("workflow_requests") }:
        connection.execute(text("ALTER TABLE workflow_requests ADD COLUMN bulk_group_id INTEGER"))
    if "scan_scope" not in {column["name"] for column in inspect(connection).get_columns("executions") }:
        connection.execute(text("ALTER TABLE executions ADD COLUMN scan_scope VARCHAR(20) DEFAULT 'service'"))
    if "scope_image" not in {column["name"] for column in inspect(connection).get_columns("executions") }:
        connection.execute(text("ALTER TABLE executions ADD COLUMN scope_image TEXT"))
    if "service_image_id" not in {column["name"] for column in inspect(connection).get_columns("workflow_requests") }:
        connection.execute(text("ALTER TABLE workflow_requests ADD COLUMN service_image_id INTEGER"))
    if "replacement_reference" not in {column["name"] for column in inspect(connection).get_columns("workflow_requests") }:
        connection.execute(text("ALTER TABLE workflow_requests ADD COLUMN replacement_reference TEXT"))
    if "policy_finding_id" not in {column["name"] for column in inspect(connection).get_columns("poam_entries") }:
        connection.execute(text("ALTER TABLE poam_entries ADD COLUMN policy_finding_id INTEGER"))
    # Composite indexes match the summary dashboard's bulk filters and latest-
    # execution lookup. They are additive and safe for existing deployments.
    for statement in (
        "CREATE INDEX IF NOT EXISTS ix_findings_service_active ON findings (service_id, active)",
        "CREATE INDEX IF NOT EXISTS ix_policy_findings_service_active ON policy_findings (service_id, active)",
        "CREATE INDEX IF NOT EXISTS ix_executions_service_scanned ON executions (service_id, scanned_at)",
        "CREATE INDEX IF NOT EXISTS ix_poam_service_status_due ON poam_entries (service_id, status, due_date)",
        "CREATE INDEX IF NOT EXISTS ix_finding_observations_finding_id_id ON finding_observations (finding_id, id)",
        "CREATE INDEX IF NOT EXISTS ix_exceptions_finding_active ON exceptions (finding_id, revoked_at, starts_at, expires_at)",
        "CREATE INDEX IF NOT EXISTS ix_policy_exceptions_finding_active ON policy_exceptions (policy_finding_id, revoked_at, starts_at, expires_at)",
    ):
        connection.execute(text(statement))
seed_auth()
app = FastAPI(title="Continuous Assessment & Tracking System", version="2.0.0")
app.add_middleware(GZipMiddleware, minimum_size=1000, compresslevel=5)
snapshot_logger = logging.getLogger("cats.snapshot")


@app.middleware("http")
async def snapshot_timing_middleware(request: Request, call_next):
    started = time.perf_counter()
    response = await call_next(request)
    if request.url.path == "/" and snapshot_logger.isEnabledFor(logging.DEBUG):
        timings = getattr(request.state, "snapshot_timings", {})
        timings["http_total_ms"] = round((time.perf_counter() - started) * 1000, 2)
        snapshot_logger.debug("service_snapshot_timing", extra={"path": request.url.path, "status": response.status_code, "timings_ms": timings})
    return response
root = Path(__file__).parent
app.mount("/static", StaticFiles(directory=root / "static"), name="static")
templates = Jinja2Templates(directory=root / "templates")

# Public scan jobs are deliberately ephemeral. The worker writes only to a
# temporary job directory and the in-memory index is lost on process restart.
PUBLIC_JOB_ROOT = Path(os.getenv("CATS_PUBLIC_JOB_ROOT", Path(tempfile.gettempdir()) / "cats-public-jobs"))
PUBLIC_JOB_ROOT.mkdir(parents=True, exist_ok=True)
PUBLIC_JOBS: dict[str, dict] = {}
PUBLIC_PROCESSES: dict[str, subprocess.Popen] = {}
PUBLIC_JOB_LOCK = threading.Lock()
PUBLIC_WORKERS = ThreadPoolExecutor(max_workers=max(1, int(os.getenv("CATS_PUBLIC_WORKERS", "2"))))
SBOM_OUTPUT_FORMATS = {
    "syft-json": "Syft JSON",
    "cyclonedx-json": "CycloneDX JSON",
    "cyclonedx-xml": "CycloneDX XML",
    "spdx-json": "SPDX JSON",
}
CYCLONEDX_SPEC_VERSIONS = ("1.4", "1.5", "1.6")
PATCH_JOB_ROOT = Path(os.getenv("CATS_PATCH_JOB_ROOT", Path(tempfile.gettempdir()) / "cats-patch-jobs"))
PATCH_JOB_ROOT.mkdir(parents=True, exist_ok=True)
PATCH_JOBS: dict[str, dict] = {}
PATCH_PROCESSES: dict[str, subprocess.Popen] = {}
PATCH_JOB_LOCK = threading.Lock()
PATCH_WORKER_URL = os.getenv("CATS_PATCH_WORKER_URL", "").rstrip("/")
PATCH_WORKER_TOKEN = os.getenv("CATS_PATCH_WORKER_TOKEN", "")
REMEDIATION_JOB_ROOT = Path(os.getenv("CATS_REMEDIATION_JOB_ROOT", Path(tempfile.gettempdir()) / "cats-remediation-jobs"))
REMEDIATION_JOB_ROOT.mkdir(parents=True, exist_ok=True)
REMEDIATION_WORKERS = ThreadPoolExecutor(max_workers=max(1, int(os.getenv("CATS_REMEDIATION_WORKERS", "2"))))


@lru_cache(maxsize=1024)
def _resolve_manifest_digest(reference: str) -> str | None:
    """Best-effort OCI enrichment; failure never changes assessment state."""
    if os.getenv("CATS_RESOLVE_IMAGE_DIGESTS", "true").lower() not in {"1", "true", "yes"}:
        return None
    try:
        completed = subprocess.run(
            ["docker", "manifest", "inspect", "--verbose", reference],
            capture_output=True, text=True, timeout=4, check=False,
        )
        if completed.returncode:
            return None
        candidate = json.loads(completed.stdout)
        candidates = candidate if isinstance(candidate, list) else [candidate]
        for manifest in candidates:
            if not isinstance(manifest, dict):
                continue
            digest = (manifest.get("Descriptor") or {}).get("digest") or manifest.get("digest")
            if str(digest or "").startswith("sha256:"):
                return str(digest)
    except (OSError, ValueError, subprocess.SubprocessError):
        return None
    return None


def utcnow():
    return datetime.now(timezone.utc)


def aware(value):
    return value.replace(tzinfo=timezone.utc) if value and value.tzinfo is None else value


def cybersecurity_group(service: Service):
    """Choose the service's Cybersecurity scope for an optional bulk request."""
    groups = list(getattr(service, "groups", []) or [])
    named = [group for group in groups if "cyber" in group.name.lower()]
    if len(named) == 1:
        return named[0]
    return groups[0] if len(groups) == 1 else None


def policy_finding_values(item) -> dict:
    raw = item.model_dump(mode="json") if hasattr(item, "model_dump") else dict(item)
    optional_fields = ("scanner", "framework", "target", "namespace", "title", "description", "remediation", "fingerprint")
    field_limits = {"scanner": 120, "framework": 240, "namespace": 240, "title": 500, "fingerprint": 240}
    finding_name = str(raw.get("finding") or "Unknown").strip() or "Unknown"
    values = {
        "finding": finding_name[:120],
        "severity": str(raw.get("severity") or "Unknown").strip()[:30],
    }
    for field in optional_fields:
        value = raw.get(field)
        values[field] = str(value).strip() if value is not None and str(value).strip() else None
        if values[field] and field in field_limits:
            values[field] = values[field][:field_limits[field]]
    return values


def policy_finding_identity(values: dict) -> str:
    if values.get("fingerprint"):
        source = {"fingerprint": values["fingerprint"]}
    else:
        source = {
            key: values.get(key)
            for key in ("scanner", "framework", "finding", "target", "namespace")
        }
    encoded = json.dumps(source, sort_keys=True, separators=(",", ":"), ensure_ascii=True).casefold()
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def sync_policy_findings(db: Session, service: Service, items, scanned_at: datetime, complete: bool) -> None:
    existing = {
        item.identity_key: item
        for item in db.scalars(select(PolicyFinding).where(PolicyFinding.service_id == service.id))
    }
    observed = set()
    for payload_item in items:
        values = policy_finding_values(payload_item)
        identity_key = policy_finding_identity(values)
        observed.add(identity_key)
        finding = existing.get(identity_key)
        if not finding:
            finding = PolicyFinding(
                service_id=service.id, identity_key=identity_key,
                first_seen=scanned_at, episode_started=scanned_at,
                last_seen=scanned_at, active=True, **values,
            )
            db.add(finding)
            existing[identity_key] = finding
        else:
            if not finding.active:
                finding.active = True
                finding.episode_started = scanned_at
                finding.resolved_at = None
                finding.recurrence_count += 1
            finding.last_seen = scanned_at
            for field, value in values.items():
                setattr(finding, field, value)
    if complete:
        for finding in existing.values():
            if finding.active and finding.identity_key not in observed:
                finding.active = False
                finding.resolved_at = scanned_at


def backfill_policy_findings() -> None:
    """Materialize the latest legacy JSON policy findings without changing evidence."""
    with SessionLocal() as db:
        services = db.scalars(select(Service).options(selectinload(Service.executions))).all()
        changed = False
        for service in services:
            latest = max(service.executions, key=lambda value: aware(value.scanned_at), default=None)
            if not latest:
                continue
            items = latest.raw_payload.get("policy_findings", []) if isinstance(latest.raw_payload, dict) else []
            if items:
                sync_policy_findings(db, service, items, latest.scanned_at, complete=False)
                changed = True
        if changed:
            db.commit()


backfill_policy_findings()


def promote_staged_services_with_evidence() -> int:
    """Repair staged records that were populated before lifecycle promotion existed.

    Staging is an empty pre-ingest state. A service that already has an
    execution, finding, policy finding, or image is no longer staged, even if
    the original ingest happened before the automatic promotion logic.
    """
    promoted = 0
    with SessionLocal() as db:
        services = db.scalars(
            select(Service)
            .where(Service.lifecycle_status == "staged")
            .options(
                selectinload(Service.executions),
                selectinload(Service.findings),
                selectinload(Service.policy_findings),
                selectinload(Service.images),
            )
        ).all()
        for service in services:
            if not (service.executions or service.findings or service.policy_findings or service.images):
                continue
            service.lifecycle_status = "active"
            db.add(AuditEvent(
                action="service.promoted", target_type="service", target_id=str(service.id),
                detail={"service_id": service.id, "reason": "existing_ingested_evidence"},
            ))
            promoted += 1
        if promoted:
            db.commit()
    return promoted


promote_staged_services_with_evidence()


def optional_form_datetime(value: str, field_name: str) -> datetime | None:
    if not value.strip():
        return None
    try:
        return aware(datetime.fromisoformat(value.strip()))
    except ValueError as exc:
        raise HTTPException(422, detail=f"{field_name} must be a valid date and time") from exc


def excel_datetime(value):
    return aware(value).astimezone(timezone.utc).replace(tzinfo=None) if value else None


def page_context(auth: AuthContext, **values):
    # Keep the global navigation useful without making each individual page
    # route calculate workflow state.  The count is limited to pending
    # requests for services visible to the signed-in account.
    pending_request_count = 0
    pending_poam_count = 0
    actionable_notifications = []
    try:
        allowed_services = auth.accessible_service_ids("service.view")
        with SessionLocal() as nav_db:
            request_query = select(
                WorkflowRequest.request_type, WorkflowRequest.service_id,
                WorkflowRequest.requested_by_id, Service.name,
            ).join(Service, Service.id == WorkflowRequest.service_id).where(WorkflowRequest.status == "pending")
            if allowed_services is not None:
                request_query = request_query.where(
                    WorkflowRequest.service_id.in_(allowed_services)
                )
            permission_by_type = {"exception": "exception.review", "poam": "poam.review", "poam_update": "poam.review", "poam_complete": "poam.review", "poam_reopen": "poam.review", "archive": "archive.review", "image_remove": "archive.review", "image_replace": "archive.review"}
            pending_items = nav_db.execute(request_query).all()
            for item in pending_items:
                if item.requested_by_id == auth.user.id:
                    continue
                permission = permission_by_type.get(item.request_type, "exception.review")
                if not auth.has(permission, item.service_id):
                    continue
                actionable_notifications.append({"label": item.request_type.replace("_", " ").title(), "service": item.name or "Service"})
            pending_request_count = len(actionable_notifications)
            # The pending rows already use the same status and service scope
            # as the navigation count, so counting them in memory avoids a
            # second query against workflow_requests.
            pending_poam_count = sum(str(item.request_type).startswith("poam") for item in pending_items)
    except Exception:
        # Navigation must never prevent a page from rendering if an older
        # deployment is still applying its database migrations.
        pending_request_count = 0
    return {
        "current_user": auth.user,
        "csrf_token": auth.csrf_token,
        "can": auth.has,
        "themes": THEMES,
        "pending_request_count": pending_request_count,
        "pending_poam_count": pending_poam_count,
        "actionable_notifications": actionable_notifications[:8],
        **values,
    }


CONFIG_DEFAULTS = {
    "overdue_days": "90",
    "hardening_overdue_days": "90",
    "hardening_noncompliant": "true",
    "warning_days": "14",
    "exception_max_days": "365",
    "skipped_images_incomplete": "true",
    "incomplete_noncompliant": "true",
    "kev_enabled": "true",
    "kev_noncompliant": "true",
    "epss_enabled": "true",
    "epss_threshold": "0.90",
    "epss_rules": '[{"severity":"Critical","threshold":0.90,"noncompliant":true},{"severity":"High","threshold":0.80,"noncompliant":true},{"severity":"Medium","threshold":0.95,"noncompliant":true}]',
    "minimum_severity": "None",
    "compliance_mode": "raw",
    "raw_due_rules": '[{"severity":"Critical","days":30},{"severity":"High","days":60},{"severity":"Medium","days":90},{"severity":"Low","days":120}]',
    "display_timezone": "UTC",
    "date_format": "%d %b %Y",
    "time_format": "%H:%M UTC",
    "log_level": "INFO",
    "audit_retention_days": "365",
    "identity_mode": "local",
    "trusted_ca_certificates": "[]",
    "repository_policies": "{}",
    "os_definitions": "{}",
    "oidc_configuration": "{}",
    "oci_registries": "[]",
    "image_signing": "{}",
}

# These settings describe the CATS runtime itself.  They are intentionally
# global; service/group governance remains scoped in configuration_for_service
# and the policy pages.
GLOBAL_CONFIGURATION_KEYS = {
    "display_timezone", "date_format", "time_format", "identity_mode",
    "trusted_ca_certificates", "repository_policies", "os_definitions", "log_level",
    "oidc_configuration", "oci_registries", "image_signing",
}

DISPLAY_TIMEZONE_FALLBACKS = (
    "UTC", "America/New_York", "America/Chicago", "America/Denver",
    "America/Los_Angeles", "Europe/London", "Europe/Paris", "Asia/Tokyo",
)


def supported_display_timezones() -> list[str]:
    zones = sorted(zone for zone in available_timezones() if "/" in zone or zone == "UTC")
    return zones or list(DISPLAY_TIMEZONE_FALLBACKS)


def scoped_setting_key(key: str, group_id: int | None = None) -> str:
    return f"group:{group_id}:{key}" if group_id is not None else key


def get_configuration(db: Session, group_id: int | None = None) -> dict[str, str]:
    configuration = dict(CONFIG_DEFAULTS)
    for setting in db.scalars(select(PortalSetting)):
        if group_id is None and setting.group_id is None and not setting.key.startswith("group:"):
            configuration[setting.key] = setting.value
        elif group_id is not None and setting.group_id == group_id:
            configuration[setting.key.split(":", 2)[-1]] = setting.value
    return configuration


def get_global_configuration(db: Session) -> dict[str, str]:
    """Return runtime configuration and safely promote legacy group values.

    Older deployments stored runtime settings under group scopes.  We retain
    those rows for audit/history, but promote the first deterministic value to
    the global key so changing the UI scope cannot silently discard it.
    """
    configuration = get_configuration(db, None)
    global_keys = {
        setting.key for setting in db.scalars(select(PortalSetting).where(PortalSetting.group_id.is_(None)))
    }
    migrated = False
    for key in GLOBAL_CONFIGURATION_KEYS:
        if key in global_keys:
            continue
        legacy = sorted(
            (setting for setting in db.scalars(select(PortalSetting)).all()
             if setting.group_id is not None and setting.key.split(":", 2)[-1] == key and setting.value not in (None, "")),
            key=lambda setting: (setting.group_id or 0, setting.id or 0),
        )
        if legacy:
            source = legacy[0]
            setting = PortalSetting(key=key, group_id=None, value=source.value,
                                    updated_by_id=source.updated_by_id, updated_at=utcnow())
            db.add(setting)
            configuration[key] = source.value
            migrated = True
    if migrated:
        db.commit()
    return configuration


def configuration_for_service(db: Session, service: Service) -> dict[str, str]:
    """Apply explicit policies for every group assigned to the service.

    Services can belong to more than one group. Using only the numerically
    smallest group caused a policy saved for another assigned group to be
    silently ignored. Later group IDs win deterministically when settings
    overlap.
    """
    configuration = get_configuration(db, None)
    for group_id in sorted(group.id for group in service.groups):
        for setting in db.scalars(select(PortalSetting).where(PortalSetting.group_id == group_id)):
            key = setting.key.split(":", 2)[-1]
            # Runtime settings are global. Ignore legacy group-scoped copies so
            # old rows cannot override the instance-wide configuration.
            if key in GLOBAL_CONFIGURATION_KEYS:
                continue
            configuration[key] = setting.value
    return configuration


def configurations_for_services(
    db: Session,
    services: list[Service],
    global_configuration: dict[str, str] | None = None,
) -> dict[int, dict[str, str]]:
    """Resolve service configuration overrides in bulk.

    The Services overview renders several services in one request. Loading
    the same global settings and then querying each service's groups inside
    the render loop created a database round trip per service. Fetch the
    relevant group settings once and apply the same deterministic override
    rules in memory instead.
    """
    base = dict(global_configuration) if global_configuration is not None else get_configuration(db)
    group_ids = sorted({group.id for service in services for group in service.groups})
    settings_by_group: dict[int, list[PortalSetting]] = {group_id: [] for group_id in group_ids}
    if group_ids:
        for setting in db.scalars(
            select(PortalSetting).where(PortalSetting.group_id.in_(group_ids))
        ):
            settings_by_group.setdefault(setting.group_id, []).append(setting)

    resolved: dict[int, dict[str, str]] = {}
    for service in services:
        configuration = dict(base)
        for group_id in sorted(group.id for group in service.groups):
            for setting in settings_by_group.get(group_id, []):
                key = setting.key.split(":", 2)[-1]
                if key in GLOBAL_CONFIGURATION_KEYS:
                    continue
                configuration[key] = setting.value
        resolved[service.id] = configuration
    return resolved


def manageable_groups(db: Session, auth: AuthContext) -> list[Group]:
    groups = db.scalars(select(Group).order_by(Group.name)).all()
    if any(a.service_id is None and a.group_id is None and a.role.name == "Administrator" for a in auth.user.role_assignments):
        return groups
    allowed = {
        a.group_id for a in auth.user.role_assignments
        if a.group_id is not None and a.role.name == "Cybersecurity" and "config.manage" in (a.role.permissions or [])
    }
    return [group for group in groups if group.id in allowed]


def requested_group_scope(group_id: str | int | None, db: Session, auth: AuthContext) -> int | None:
    try:
        selected = int(group_id) if group_id not in (None, "", "global") else None
    except (TypeError, ValueError) as exc:
        raise HTTPException(422, detail="Invalid group scope") from exc
    if selected is None and auth.can_manage_group(None):
        return None
    if selected is None or not auth.can_manage_group(selected):
        raise HTTPException(403, detail="You cannot administer this group scope")
    return selected


def require_config_scope(request: Request, auth: AuthContext = Depends(require_user)) -> AuthContext:
    """Authorize global or group-scoped administration before form parsing."""
    raw_group = request.query_params.get("group_id")
    try:
        group_id = int(raw_group) if raw_group else None
    except ValueError as exc:
        raise HTTPException(422, detail="Invalid group scope") from exc
    if group_id is None and auth.can_manage_group(None):
        return auth
    if group_id is None or not auth.can_manage_group(group_id):
        raise HTTPException(403, detail="You cannot administer this group scope")
    return auth


def require_global_config_scope(auth: AuthContext = Depends(require_user)) -> AuthContext:
    if not auth.can_manage_group(None):
        raise HTTPException(403, detail="Only global administrators may change runtime configuration")
    return auth


def _json_setting(configuration: dict, key: str, fallback):
    value = parse_json(configuration.get(key), fallback)
    return value if isinstance(value, type(fallback)) else fallback


def configured_oidc(configuration: dict) -> dict:
    value = _json_setting(configuration, "oidc_configuration", {})
    if not isinstance(value, dict):
        value = {}
    safe = {key: str(value.get(key) or "") for key in (
        "provider_name", "issuer", "client_id", "scopes", "username_claim",
        "email_claim", "groups_claim", "roles_claim", "redirect_uri",
        "post_logout_redirect_uri", "browser_issuer",
    )}
    safe["client_secret_configured"] = secret_configured(str(value.get("client_secret") or ""))
    return safe


def configured_registries(configuration: dict) -> list[dict]:
    value = _json_setting(configuration, "oci_registries", [])
    if not isinstance(value, list):
        return []
    result = []
    for item in value:
        if not isinstance(item, dict):
            continue
        endpoint = str(item.get("endpoint") or "").strip().rstrip("/")
        namespace = str(item.get("namespace") or "").strip().strip("/")
        parsed = urllib.parse.urlparse(endpoint if "://" in endpoint else f"https://{endpoint}")
        resolved_path = "/".join(part for part in (parsed.netloc, namespace) if part)
        result.append({
            "id": str(item.get("id") or ""),
            "display_name": str(item.get("display_name") or item.get("endpoint") or "Registry"),
            "endpoint": endpoint,
            "namespace": namespace,
            "resolved_path": resolved_path,
            "auth_mode": str(item.get("auth_mode") or "none"),
            "username": str(item.get("username") or ""),
            "secret_configured": secret_configured(str(item.get("password") or "")),
            "status": str(item.get("status") or "Not tested"),
            "status_detail": str(item.get("status_detail") or ""),
            "use_for_remediation": item.get("use_for_remediation") is True,
        })
    return result


def _configured_registry_for_image(reference: str, registries: list[dict]) -> dict | None:
    """Find centrally configured connection settings for an OCI image host."""
    try:
        host = registry_host(reference).lower()
    except ValueError:
        return None
    for registry in registries:
        if not isinstance(registry, dict):
            continue
        endpoint = str(registry.get("endpoint") or "").strip()
        parsed = urllib.parse.urlparse(endpoint if "://" in endpoint else f"https://{endpoint}")
        if (parsed.netloc or "").lower() == host:
            return registry
    return None


def configured_ca_bundle(configuration: dict) -> str | None:
    """Return administrator-added PEM trust material without replacing system trust."""
    certificates = _json_setting(configuration, "trusted_ca_certificates", [])
    pem = [str(item.get("pem") or "") for item in certificates if isinstance(item, dict)]
    value = "\n".join(item for item in pem if "BEGIN CERTIFICATE" in item)
    return value + ("\n" if value else "") or None


def _validate_endpoint(value: str) -> str:
    raw = value.strip().rstrip("/")
    if "://" not in raw:
        raw = f"https://{raw}"
    parsed = urllib.parse.urlparse(raw)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(422, detail="Registry endpoint must be an HTTP or HTTPS URL")
    return f"{parsed.scheme}://{parsed.netloc}{parsed.path.rstrip('/')}"


def configured_time(
    value,
    include_time: bool = False,
    configuration: dict[str, str] | None = None,
) -> str:
    if not value:
        return ""
    try:
        if configuration is None:
            with SessionLocal() as db:
                configuration = get_configuration(db)
        zone = ZoneInfo(configuration.get("display_timezone", "UTC"))
        date_pattern = configuration.get("date_format", "%d %b %Y")
        if include_time:
            pattern = f"{date_pattern}, {configuration.get('time_format', '%H:%M UTC')}"
        else:
            pattern = date_pattern
        return aware(value).astimezone(zone).strftime(pattern)
    except Exception:
        return aware(value).strftime("%d %b %Y, %H:%M UTC" if include_time else "%d %b %Y")


templates.env.globals["cats_date"] = lambda value: configured_time(value)
templates.env.globals["cats_datetime"] = lambda value: configured_time(value, include_time=True)


def check_csrf(auth: AuthContext, supplied: str):
    if not secrets.compare_digest(auth.csrf_token, supplied):
        raise HTTPException(403, detail="Invalid CSRF token")


def require_pipeline(authorization: str | None = Header(default=None)):
    expected = os.getenv("PIPELINE_API_TOKEN", "development-token")
    supplied = authorization.removeprefix("Bearer ").strip() if authorization else ""
    if not secrets.compare_digest(supplied.encode(), expected.encode()):
        raise HTTPException(status_code=401, detail="Invalid pipeline token")


def active_exception(finding: Finding, now: datetime):
    return next((e for e in finding.exceptions if not e.revoked_at and aware(e.starts_at) <= now < aware(e.expires_at)), None)


def archive_state(service: Service):
    latest = max(service.archive_events, key=lambda event: aware(event.created_at), default=None)
    return latest if latest and latest.action == "archive" else None


def service_lifecycle(service: Service) -> str:
    """Return the lifecycle used by the shared service snapshot filters."""
    if archive_state(service):
        return "archived"
    status = str(getattr(service, "lifecycle_status", "active") or "active").lower()
    return status if status in {"active", "staged"} else "active"


def workbook_response(workbook: Workbook, filename: str):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="17312B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = max(width, 12)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def service_view(service: Service, now: datetime, configuration: dict[str, str] | None = None):
    configuration = configuration or CONFIG_DEFAULTS
    overdue_days = max(1, int(configuration.get("overdue_days", "90")))
    warning_days = max(1, int(configuration.get("warning_days", "14")))
    active_all = [f for f in service.findings if f.active]
    resolved = [f for f in service.findings if not f.active]
    try:
        raw_due_rules = json.loads(configuration.get("raw_due_rules", "[]"))
    except (TypeError, ValueError):
        raw_due_rules = []
    raw_due_by_severity = {str(rule.get("severity", "")).lower(): max(1, int(rule.get("days", overdue_days))) for rule in raw_due_rules}
    def due_days_for(finding):
        return raw_due_by_severity.get(finding.severity.lower(), overdue_days) if configuration.get("compliance_mode") == "raw" else overdue_days
    due_dates = {f.id: aware(f.episode_started) + timedelta(days=due_days_for(f)) for f in active_all}
    overdue = [f for f in active_all if now >= due_dates[f.id] and not active_exception(f, now)]
    warning_items = []
    warning_cutoff = now + timedelta(days=warning_days)
    for finding in active_all:
        exception = active_exception(finding, now)
        due_date = due_dates[finding.id]
        if not exception and now < due_date <= warning_cutoff:
            warning_items.append({"type": "CVE", "item": finding.cve, "reason": "Due date approaching", "due": due_date})
        elif exception and now < aware(exception.expires_at) <= warning_cutoff:
            warning_items.append({"type": "Exception", "item": finding.cve, "reason": "Exception expires soon", "due": aware(exception.expires_at)})
    last_execution = max((aware(e.scanned_at) for e in service.executions), default=None)
    latest_execution = max(service.executions, key=lambda e: aware(e.scanned_at), default=None)
    version = service.manual_version or (latest_execution.raw_payload.get("service", {}).get("version") if latest_execution else None)
    skipped_images = latest_execution.raw_payload.get("skipped_images", []) if latest_execution else []
    skipped_charts = latest_execution.raw_payload.get("skipped_charts", []) if latest_execution else []
    policy_active_all = [finding for finding in service.policy_findings if finding.active]
    hardening_overdue_days = max(1, int(configuration.get("hardening_overdue_days", "90")))
    hardening_noncompliant = configuration.get("hardening_noncompliant", "true") == "true"
    policy_due_dates = {
        finding.id: aware(finding.episode_started) + timedelta(days=hardening_overdue_days)
        for finding in policy_active_all
    }


    policy_excepted = [finding for finding in policy_active_all if active_exception(finding, now)]
    policy_noncompliant = [
        finding for finding in policy_active_all
        if hardening_noncompliant and not active_exception(finding, now) and now >= policy_due_dates[finding.id]
    ]
    policy_findings = [
        finding for finding in policy_active_all
        if not active_exception(finding, now) and finding not in policy_noncompliant
    ]
    policy_resolved = [finding for finding in service.policy_findings if not finding.active]
    poam_active = [entry for entry in service.poam_entries if entry.status == "active"]
    incomplete = bool(latest_execution and (not latest_execution.complete or skipped_images or skipped_charts))
    try:
        epss_rules = json.loads(configuration.get("epss_rules", "[]"))
    except (TypeError, ValueError):
        epss_rules = []
    if not epss_rules:
        epss_rules = [{"severity": "Any", "threshold": float(configuration.get("epss_threshold", "0.90")), "noncompliant": True}]
    compliance_mode = configuration.get("compliance_mode", "risk_based")
    # Raw mode evaluates every fixable finding. Risk-based mode first builds the
    # overlay-visible set from KEV, EPSS, and minimum-severity matches; age only
    # determines when a visible finding becomes non-compliant.
    risk_eligible = set(f.id for f in active_all) if compliance_mode == "raw" else set()
    risk_findings = set(f.id for f in overdue) if compliance_mode == "raw" else set()
    risk_metadata_by_finding = {}
    severity_rank = {"unknown": 0, "negligible": 1, "low": 2, "medium": 3, "high": 4, "critical": 5}
    minimum_severity = configuration.get("minimum_severity", "None").lower()
    if compliance_mode != "raw" and minimum_severity in severity_rank:
        severity_matches = {
            finding.id for finding in active_all
            if severity_rank.get(finding.severity.lower(), 0) >= severity_rank[minimum_severity]
        }
        risk_eligible.update(severity_matches)
        risk_findings.update(
            finding.id for finding in active_all
            if finding.id in severity_matches and now >= due_dates[finding.id]
            and not active_exception(finding, now)
        )
    for finding in active_all:
        observation = max(finding.observations, key=lambda item: item.id, default=None)
        evidence = observation.evidence if observation else {}
        catalog_kev, catalog_epss = risk_metadata(finding.cve)
        kev = bool(evidence.get("kev", evidence.get("known_exploited", catalog_kev))) or catalog_kev
        try:
            epss = float(evidence.get("epss", evidence.get("epss_score", catalog_epss or 0)) or 0)
        except (TypeError, ValueError):
            epss = 0
        risk_metadata_by_finding[finding.id] = {"kev": kev, "epss": epss}
        if compliance_mode != "raw" and configuration.get("kev_enabled") == "true" and kev:
            risk_eligible.add(finding.id)
            if configuration.get("kev_noncompliant") == "true" and now >= due_dates[finding.id] and not active_exception(finding, now):
                risk_findings.add(finding.id)
        if compliance_mode != "raw" and configuration.get("epss_enabled") == "true" and epss is not None and now >= due_dates[finding.id] and not active_exception(finding, now):
            for rule in epss_rules:
                severity = str(rule.get("severity", "Any")).lower()
                threshold = float(rule.get("threshold", 1))
                if severity in {"any", finding.severity.lower()} and epss >= threshold:
                    risk_eligible.add(finding.id)
                    if bool(rule.get("noncompliant", True)):
                        risk_findings.add(finding.id)
                    break
        elif compliance_mode != "raw" and configuration.get("epss_enabled") == "true" and epss is not None:
            for rule in epss_rules:
                severity = str(rule.get("severity", "Any")).lower()
                threshold = float(rule.get("threshold", 1))
                if severity in {"any", finding.severity.lower()} and epss >= threshold:
                    risk_eligible.add(finding.id)
                    break
    excepted = [f for f in active_all if f.id in risk_eligible and active_exception(f, now)]
    active = [
        finding for finding in active_all
        if finding.id in risk_eligible and finding.id not in risk_findings
        and not active_exception(finding, now)
    ]
    eligible_cves = {finding.cve for finding in active_all if finding.id in risk_eligible}
    warning_items = [
        item for item in warning_items
        if item.get("type") not in {"CVE", "Exception"} or item.get("item") in eligible_cves
    ]
    evidence_state = "No evidence"
    if latest_execution:
        evidence_state = "Complete" if latest_execution.complete else "Incomplete"
    if skipped_images and evidence_state == "Incomplete":
        evidence_state = f"Incomplete · {len(skipped_images)} skipped"
    if skipped_charts and evidence_state == "Incomplete":
        if skipped_images:
            evidence_state = f"Incomplete · {len(skipped_images)} skipped images, {len(skipped_charts)} skipped charts"
        else:
            evidence_state = f"Incomplete · {len(skipped_charts)} skipped charts"
    if incomplete and configuration.get("incomplete_noncompliant") != "true":
        warning_items.append({"type": "Evidence", "item": "Incomplete evidence", "reason": "Latest assessment is incomplete", "due": None})
    evidence_noncompliant = bool(incomplete and configuration.get("incomplete_noncompliant") == "true")
    noncompliant = [finding for finding in active_all if finding.id in risk_findings and not active_exception(finding, now)]
    noncompliance_items = [
        {
            "type": "CVE",
            "item": finding.cve,
            "finding_id": finding.id,
            "reason": "Overdue fixable vulnerability",
            "due": due_dates.get(finding.id),
            "status": "Non-Compliant",
        }
        for finding in noncompliant
    ]
    noncompliance_items.extend({
        "type": "Configuration",
        "item": finding.finding,
        "policy_finding_id": finding.id,
        "evidence_image": finding.target or "No target reported",
        "reason": finding.title or "Overdue configuration finding",
        "due": policy_due_dates[finding.id],
        "status": "Non-Compliant",
    } for finding in policy_noncompliant)
    # Keep the actionable evidence observations as the source of truth.  A
    # generic assessment roll-up is useful only when the incomplete assessment
    # has no concrete skipped/missing item to explain it.
    latest_overview = (latest_execution.raw_payload.get("service_overview", {})
                       if latest_execution and isinstance(latest_execution.raw_payload, dict)
                       else {})
    specific_missing_evidence = bool(
        isinstance(latest_overview, dict)
        and (latest_overview.get("missing_evidence") or latest_overview.get("evidence"))
    )
    if evidence_noncompliant:
        if skipped_images:
            noncompliance_items.extend(
                {
                    "type": "Evidence",
                    "item": "Image",
                    "evidence_image": image,
                    "reason": "Unavailable for assessment",
                    "due": None,
                    "status": "Non-Compliant",
                }
                for image in skipped_images
            )
        elif not skipped_charts and not specific_missing_evidence:
            noncompliance_items.append(
                {
                    "type": "Evidence",
                    "item": "Assessment",
                    "evidence_image": "",
                    "reason": "Latest assessment reported incomplete evidence",
                    "due": None,
                    "status": "Non-Compliant",
                }
            )
        if skipped_charts:
            noncompliance_items.extend(
                {
                    "type": "Evidence",
                    "item": "Chart",
                    "evidence_image": chart,
                    "reason": "Unavailable for assessment",
                    "due": None,
                    "status": "Non-Compliant",
                }
                for chart in skipped_charts
            )
    noncompliance_items.sort(key=lambda item: (str(item.get("item", "")).casefold(), str(item.get("type", "")).casefold()))
    return {
        "service": service, "compliant": not risk_findings and not policy_noncompliant and not (incomplete and configuration.get("incomplete_noncompliant") == "true"), "overdue": overdue,
        "excepted": excepted, "active": active, "resolved": resolved,
        "last_execution": last_execution,
        "evidence_state": evidence_state, "version": version, "skipped_images": skipped_images, "skipped_charts": skipped_charts,
        "incomplete": incomplete, "risk_findings": risk_findings, "risk_metadata": risk_metadata_by_finding,
        "risk_eligible": risk_eligible,
        "noncompliant": noncompliant, "evidence_noncompliant": evidence_noncompliant,
        "policy_findings": policy_findings,
        "policy_excepted": policy_excepted, "policy_noncompliant": policy_noncompliant,
        "policy_resolved": policy_resolved,
        "policy_due_dates": policy_due_dates,
        "noncompliance_items": noncompliance_items,
        "due_dates": due_dates,
        "compliance_mode": compliance_mode,
        "archive": archive_state(service),
        "oldest_age": max(((now - aware(f.episode_started)).days for f in [*active_all, *policy_active_all]), default=None),
        "overdue_days": overdue_days, "hardening_overdue_days": hardening_overdue_days,
        "hardening_noncompliant": hardening_noncompliant,
        "poam_active": poam_active,
        "warning_days": warning_days, "warning_items": warning_items,
    }

def service_overview_rows_detailed(
    db: Session,
    auth: AuthContext,
    now: datetime,
    configuration: dict[str, str],
) -> tuple[list[dict], dict[int, dict[str, int]]]:
    """Build the Services page from summary columns instead of full ORM graphs.

    The detail and export paths still use ``service_view``.  The dashboard only
    needs counts and a few service metadata fields, so loading complete
    findings, observations, POA&M bodies, and scan payload history here is both
    unnecessary and a major source of latency as evidence grows.
    """
    allowed = auth.accessible_service_ids("service.view")
    if allowed == set():
        raise HTTPException(403, detail="Permission denied")
    service_query = select(Service).order_by(Service.name, Service.service_key)
    if allowed is not None:
        service_query = service_query.where(Service.id.in_(allowed))
    services = list(db.scalars(service_query))
    if not services:
        return [], {}
    service_ids = [service.id for service in services]
    service_id_set = set(service_ids)

    # Resolve group-scoped settings with two bulk queries; never query a group
    # while rendering an individual service row.
    service_group_rows = db.execute(
        select(ServiceGroup.service_id, ServiceGroup.group_id)
        .where(ServiceGroup.service_id.in_(service_ids))
    ).all()
    groups_by_service: dict[int, list[int]] = {}
    group_ids = set()
    for row in service_group_rows:
        groups_by_service.setdefault(row.service_id, []).append(row.group_id)
        group_ids.add(row.group_id)
    settings_by_group: dict[int, list[PortalSetting]] = {group_id: [] for group_id in group_ids}
    if group_ids:
        for setting in db.scalars(select(PortalSetting).where(PortalSetting.group_id.in_(group_ids))):
            settings_by_group.setdefault(setting.group_id, []).append(setting)
    configurations: dict[int, dict[str, str]] = {}
    for service_id in service_ids:
        resolved = dict(configuration)
        for group_id in sorted(groups_by_service.get(service_id, [])):
            for setting in settings_by_group.get(group_id, []):
                key = setting.key.split(":", 2)[-1]
                if key not in GLOBAL_CONFIGURATION_KEYS:
                    resolved[key] = setting.value
        configurations[service_id] = resolved

    # Fetch only the latest execution payload per service.  Raw history is not
    # needed by this page and is intentionally excluded from the query.
    latest_execution_time = (
        select(Execution.service_id, func.max(Execution.scanned_at).label("latest_scanned_at"))
        .where(Execution.service_id.in_(service_ids))
        .group_by(Execution.service_id)
        .subquery()
    )
    latest_executions = db.execute(
        select(Execution.service_id, Execution.scanned_at, Execution.complete, Execution.raw_payload)
        .join(
            latest_execution_time,
            and_(Execution.service_id == latest_execution_time.c.service_id,
                 Execution.scanned_at == latest_execution_time.c.latest_scanned_at),
        )
    ).all()
    latest_by_service = {row.service_id: row for row in latest_executions}

    # The dashboard needs only the active finding fields used by compliance,
    # plus the latest observation's risk metadata.  No package/CVE description
    # or historical observation payload is hydrated.
    finding_rows = db.execute(select(
        Finding.id, Finding.service_id, Finding.cve, Finding.severity, Finding.episode_started,
    ).where(Finding.service_id.in_(service_ids), Finding.active.is_(True))).all()
    finding_by_service: dict[int, list] = {}
    finding_ids = []
    for row in finding_rows:
        finding_by_service.setdefault(row.service_id, []).append(row)
        finding_ids.append(row.id)
    latest_observation_by_finding: dict[int, dict] = {}
    if finding_ids:
        latest_observation = (
            select(FindingObservation.finding_id, func.max(FindingObservation.id).label("latest_id"))
            .join(Finding, Finding.id == FindingObservation.finding_id)
            .where(Finding.service_id.in_(service_ids), Finding.active.is_(True))
            .group_by(FindingObservation.finding_id)
            .subquery()
        )
        for row in db.execute(
            select(FindingObservation.finding_id, FindingObservation.evidence)
            .join(latest_observation, FindingObservation.id == latest_observation.c.latest_id)
        ):
            latest_observation_by_finding[row.finding_id] = row.evidence or {}
    active_exception_ids = set()
    if finding_ids:
        active_exception_ids = set(db.scalars(select(ExceptionRecord.finding_id).join(
            Finding, Finding.id == ExceptionRecord.finding_id,
        ).where(
            Finding.service_id.in_(service_ids), Finding.active.is_(True),
            ExceptionRecord.revoked_at.is_(None),
            ExceptionRecord.starts_at <= now,
            ExceptionRecord.expires_at > now,
        )))

    policy_rows = db.execute(select(
        PolicyFinding.id, PolicyFinding.service_id, PolicyFinding.severity, PolicyFinding.episode_started,
    ).where(PolicyFinding.service_id.in_(service_ids), PolicyFinding.active.is_(True))).all()
    policy_by_service: dict[int, list] = {}
    policy_ids = []
    for row in policy_rows:
        policy_by_service.setdefault(row.service_id, []).append(row)
        policy_ids.append(row.id)
    active_policy_exception_ids = set()
    if policy_ids:
        active_policy_exception_ids = set(db.scalars(select(PolicyExceptionRecord.policy_finding_id).join(
            PolicyFinding, PolicyFinding.id == PolicyExceptionRecord.policy_finding_id,
        ).where(
            PolicyFinding.service_id.in_(service_ids), PolicyFinding.active.is_(True),
            PolicyExceptionRecord.revoked_at.is_(None),
            PolicyExceptionRecord.starts_at <= now,
            PolicyExceptionRecord.expires_at > now,
        )))

    # Archive and POA&M summaries are grouped once for all visible services.
    archive_events = db.execute(select(
        ServiceArchiveEvent.service_id, ServiceArchiveEvent.action, ServiceArchiveEvent.created_at,
    ).where(ServiceArchiveEvent.service_id.in_(service_ids)).order_by(ServiceArchiveEvent.created_at.desc())).all()
    archive_by_service: dict[int, str] = {}
    for row in archive_events:
        archive_by_service.setdefault(row.service_id, row.action)
    poam_counts: dict[int, dict[str, int]] = {}
    for row in db.execute(select(
        PoamEntry.service_id, PoamEntry.status, func.count(PoamEntry.id),
    ).where(PoamEntry.service_id.in_(service_ids)).group_by(PoamEntry.service_id, PoamEntry.status)):
        counts = poam_counts.setdefault(row.service_id, {"active": 0, "pending": 0, "overdue": 0})
        status_name = row.status
        count = int(row[2] or 0)
        if status_name == "active":
            counts["active"] += count
        elif status_name == "pending_approval":
            counts["pending"] += count
    overdue_poams = db.execute(select(PoamEntry.service_id).where(
        PoamEntry.service_id.in_(service_ids), PoamEntry.status == "active",
        PoamEntry.due_date.is_not(None), PoamEntry.due_date < now,
    )).all()
    for row in overdue_poams:
        poam_counts.setdefault(row.service_id, {"active": 0, "pending": 0, "overdue": 0})["overdue"] += 1
    severity_rank = {"unknown": 0, "negligible": 1, "low": 2, "medium": 3, "high": 4, "critical": 5}
    rows: list[dict] = []
    for service in services:
        cfg = configurations[service.id]
        overdue_days = max(1, int(cfg.get("overdue_days", "90")))
        warning_days = max(1, int(cfg.get("warning_days", "14")))
        try:
            raw_due_rules = json.loads(cfg.get("raw_due_rules", "[]"))
        except (TypeError, ValueError):
            raw_due_rules = []
        raw_due_by_severity = {
            str(rule.get("severity", "")).lower(): max(1, int(rule.get("days", overdue_days)))
            for rule in raw_due_rules
        }
        active_findings = finding_by_service.get(service.id, [])
        due_dates = {
            row.id: aware(row.episode_started) + timedelta(days=(
                raw_due_by_severity.get(str(row.severity or "").lower(), overdue_days)
                if cfg.get("compliance_mode") == "raw" else overdue_days
            )) for row in active_findings
        }
        exception_ids = {row.id for row in active_findings if row.id in active_exception_ids}
        overdue_ids = {row.id for row in active_findings if now >= due_dates[row.id] and row.id not in exception_ids}
        compliance_mode = cfg.get("compliance_mode", "risk_based")
        risk_eligible = {row.id for row in active_findings} if compliance_mode == "raw" else set()
        risk_findings = set(overdue_ids) if compliance_mode == "raw" else set()
        try:
            epss_rules = json.loads(cfg.get("epss_rules", "[]"))
        except (TypeError, ValueError):
            epss_rules = []
        if not epss_rules:
            epss_rules = [{"severity": "Any", "threshold": float(cfg.get("epss_threshold", "0.90")), "noncompliant": True}]
        minimum_severity = cfg.get("minimum_severity", "None").lower()
        severity_matches = {
            row.id for row in active_findings
            if minimum_severity in severity_rank and severity_rank.get(str(row.severity or "").lower(), 0) >= severity_rank[minimum_severity]
        }
        if compliance_mode != "raw":
            risk_eligible.update(severity_matches)
            risk_findings.update(row.id for row in active_findings if row.id in severity_matches and now >= due_dates[row.id] and row.id not in exception_ids)
        for row in active_findings:
            evidence = latest_observation_by_finding.get(row.id, {})
            catalog_kev, catalog_epss = risk_metadata(row.cve)
            kev = bool(evidence.get("kev", evidence.get("known_exploited", catalog_kev))) or catalog_kev
            try:
                epss = float(evidence.get("epss", evidence.get("epss_score", catalog_epss or 0)) or 0)
            except (TypeError, ValueError):
                epss = 0
            if compliance_mode != "raw" and cfg.get("kev_enabled") == "true" and kev:
                risk_eligible.add(row.id)
                if cfg.get("kev_noncompliant") == "true" and now >= due_dates[row.id] and row.id not in exception_ids:
                    risk_findings.add(row.id)
            if compliance_mode != "raw" and cfg.get("epss_enabled") == "true" and epss is not None and now >= due_dates[row.id]:
                for rule in epss_rules:
                    severity = str(rule.get("severity", "Any")).lower()
                    if severity in {"any", str(row.severity or "").lower()} and epss >= float(rule.get("threshold", 1)):
                        risk_eligible.add(row.id)
                        if bool(rule.get("noncompliant", True)) and row.id not in exception_ids:
                            risk_findings.add(row.id)
                        break
            elif compliance_mode != "raw" and cfg.get("epss_enabled") == "true" and epss is not None:
                for rule in epss_rules:
                    severity = str(rule.get("severity", "Any")).lower()
                    if severity in {"any", str(row.severity or "").lower()} and epss >= float(rule.get("threshold", 1)):
                        risk_eligible.add(row.id)
                        break
        excepted_count = len(risk_eligible & exception_ids)
        active_count = len(risk_eligible - risk_findings - exception_ids)
        policy_active = policy_by_service.get(service.id, [])
        hardening_overdue_days = max(1, int(cfg.get("hardening_overdue_days", "90")))
        hardening_noncompliant = cfg.get("hardening_noncompliant", "true") == "true"
        policy_due = {row.id: aware(row.episode_started) + timedelta(days=hardening_overdue_days) for row in policy_active}
        policy_excepted_count = sum(row.id in active_policy_exception_ids for row in policy_active)
        policy_noncompliant_count = sum(
            hardening_noncompliant and row.id not in active_policy_exception_ids and now >= policy_due[row.id]
            for row in policy_active
        )
        policy_visible_count = len(policy_active) - policy_excepted_count - policy_noncompliant_count
        latest = latest_by_service.get(service.id)
        raw_payload = latest.raw_payload if latest and isinstance(latest.raw_payload, dict) else {}
        skipped_images = raw_payload.get("skipped_images", []) or []
        skipped_charts = raw_payload.get("skipped_charts", []) or []
        incomplete = bool(latest and (not latest.complete or skipped_images or skipped_charts))
        evidence_state = "No evidence" if not latest else ("Complete" if latest.complete else "Incomplete")
        if skipped_images and evidence_state == "Incomplete":
            evidence_state = f"Incomplete · {len(skipped_images)} skipped"
        if skipped_charts and evidence_state == "Incomplete":
            evidence_state = f"Incomplete · {len(skipped_images)} skipped images, {len(skipped_charts)} skipped charts" if skipped_images else f"Incomplete · {len(skipped_charts)} skipped charts"
        evidence_noncompliant = bool(incomplete and cfg.get("incomplete_noncompliant") == "true")
        compliant = not risk_findings and not policy_noncompliant_count and not evidence_noncompliant
        oldest_dates = [row.episode_started for row in active_findings] + [row.episode_started for row in policy_active]
        oldest_age = max(((now - aware(value)).days for value in oldest_dates), default=None)
        rows.append({
            "service": service,
            "compliant": compliant,
            "version": service.manual_version or (raw_payload.get("service", {}).get("version") if isinstance(raw_payload.get("service"), dict) else None),
            "evidence_state": evidence_state,
            "active": [None] * (active_count + policy_visible_count),
            "policy_findings": [],
            "noncompliant": [None] * len(risk_findings),
            "policy_noncompliant": [None] * policy_noncompliant_count,
            "excepted": [None] * excepted_count,
            "policy_excepted": [None] * policy_excepted_count,
            "oldest_age": oldest_age,
            "archive": archive_by_service.get(service.id) == "archive",
            "poam": poam_counts.get(service.id, {"active": 0, "pending": 0, "overdue": 0}),
            "overdue": [None] * len(overdue_ids),
            "warning_days": warning_days,
        })
    return rows, poam_counts


def _overview_services_and_configurations(db: Session, auth: AuthContext, configuration: dict[str, str]):
    allowed = auth.accessible_service_ids("service.view")
    if allowed == set():
        raise HTTPException(403, detail="Permission denied")
    query = select(Service).order_by(Service.name, Service.service_key)
    if allowed is not None:
        query = query.where(Service.id.in_(allowed))
    services = list(db.scalars(query))
    service_ids = [service.id for service in services]
    groups_by_service: dict[int, list[int]] = {}
    group_ids: set[int] = set()
    if service_ids:
        for row in db.execute(select(ServiceGroup.service_id, ServiceGroup.group_id).where(ServiceGroup.service_id.in_(service_ids))):
            groups_by_service.setdefault(row.service_id, []).append(row.group_id)
            group_ids.add(row.group_id)
    settings_by_group: dict[int, list[PortalSetting]] = {group_id: [] for group_id in group_ids}
    if group_ids:
        for setting in db.scalars(select(PortalSetting).where(PortalSetting.group_id.in_(group_ids))):
            settings_by_group.setdefault(setting.group_id, []).append(setting)
    configurations: dict[int, dict[str, str]] = {}
    for service_id in service_ids:
        resolved = dict(configuration)
        for group_id in sorted(groups_by_service.get(service_id, [])):
            for setting in settings_by_group.get(group_id, []):
                key = setting.key.split(":", 2)[-1]
                if key not in GLOBAL_CONFIGURATION_KEYS:
                    resolved[key] = setting.value
        configurations[service_id] = resolved
    return services, configurations


def _raw_due_groups(configurations: dict[int, dict[str, str]], now: datetime):
    grouped: dict[tuple, list[int]] = {}
    for service_id, cfg in configurations.items():
        overdue_days = max(1, int(cfg.get("overdue_days", "90")))
        try:
            rules = json.loads(cfg.get("raw_due_rules", "[]"))
        except (TypeError, ValueError):
            rules = []
        due_by_severity = tuple(sorted(
            (str(rule.get("severity", "")).lower(), max(1, int(rule.get("days", overdue_days))))
            for rule in rules if isinstance(rule, dict)
        ))
        grouped.setdefault((overdue_days, due_by_severity), []).append(service_id)
    return grouped


def _raw_overdue_expression(model, configurations: dict[int, dict[str, str]], now: datetime):
    parts = []
    for (default_days, due_by_severity), service_ids in _raw_due_groups(configurations, now).items():
        known = []
        for severity, days in due_by_severity:
            known.append(severity)
            parts.append(and_(
                model.service_id.in_(service_ids), model.severity == severity,
                model.episode_started <= now - timedelta(days=days),
            ))
        parts.append(and_(
            model.service_id.in_(service_ids),
            (~model.severity.in_(known) if known else true()),
            model.episode_started <= now - timedelta(days=default_days),
        ))
    return or_(*parts) if parts else false()


def _hardening_overdue_expression(configurations: dict[int, dict[str, str]], now: datetime):
    parts = []
    grouped: dict[tuple[int, bool], list[int]] = {}
    for service_id, cfg in configurations.items():
        grouped.setdefault((max(1, int(cfg.get("hardening_overdue_days", "90"))), cfg.get("hardening_noncompliant", "true") == "true"), []).append(service_id)
    for (days, enabled), service_ids in grouped.items():
        if enabled:
            parts.append(and_(PolicyFinding.service_id.in_(service_ids), PolicyFinding.episode_started <= now - timedelta(days=days)))
    return or_(*parts) if parts else false()


def _risk_finding_expressions(
    configurations: dict[int, dict[str, str]],
    now: datetime,
    finding_exception,
):
    """Build database expressions for risk-based finding summaries.

    The detailed dashboard historically evaluated these rules in Python after
    loading every active finding.  The expressions below keep the same rule
    precedence while allowing the database to count eligible/non-compliant
    rows directly.  Catalog KEV/EPSS data is represented as CVE membership
    predicates; observation evidence remains JSON-native in PostgreSQL and
    SQLite.
    """
    eligible_parts = []
    noncompliant_parts = []
    needs_observations = False
    catalog_kev = kev_cves()
    catalog_epss = epss_scores()
    severity_rank = {"unknown": 0, "negligible": 1, "low": 2, "medium": 3, "high": 4, "critical": 5}
    for service_id, cfg in configurations.items():
        if cfg.get("compliance_mode", "raw") == "raw":
            eligible_parts.append(Finding.service_id == service_id)
            noncompliant_parts.append(and_(Finding.service_id == service_id, _raw_overdue_expression(Finding, {service_id: cfg}, now)))
            continue
        service_scope = Finding.service_id == service_id
        eligible = []
        noncompliant = []
        minimum = str(cfg.get("minimum_severity", "None")).lower()
        if minimum in severity_rank:
            severity_values = [name for name, rank in severity_rank.items() if rank >= severity_rank[minimum]]
            if severity_values:
                match = Finding.severity.in_(severity_values)
                eligible.append(match)
                noncompliant.append(and_(match, _raw_overdue_expression(Finding, {service_id: cfg}, now)))
        if cfg.get("kev_enabled") == "true":
            needs_observations = True
            kev_match = []
            if catalog_kev:
                kev_match.append(Finding.cve.in_(catalog_kev))
            kev_match.extend((
                FindingObservation.evidence["kev"].as_boolean().is_(True),
                FindingObservation.evidence["known_exploited"].as_boolean().is_(True),
            ))
            if kev_match:
                match = or_(*kev_match)
                eligible.append(match)
                if cfg.get("kev_noncompliant") == "true":
                    noncompliant.append(and_(match, _raw_overdue_expression(Finding, {service_id: cfg}, now)))
        try:
            epss_rules = json.loads(cfg.get("epss_rules", "[]"))
        except (TypeError, ValueError):
            epss_rules = []
        if not epss_rules:
            epss_rules = [{"severity": "Any", "threshold": float(cfg.get("epss_threshold", "0.90")), "noncompliant": True}]
        if cfg.get("epss_enabled") == "true":
            needs_observations = True
            prior_match = false()
            for rule in epss_rules:
                try:
                    threshold = float(rule.get("threshold", 1))
                except (TypeError, ValueError):
                    threshold = 1.0
                severity = str(rule.get("severity", "Any")).lower()
                severity_match = severity in {"any", ""} or Finding.severity == severity
                catalog_match = {cve for cve, score in catalog_epss.items() if score >= threshold}
                score_match = [
                    FindingObservation.evidence["epss"].as_float() >= threshold,
                    FindingObservation.evidence["epss_score"].as_float() >= threshold,
                ]
                if catalog_match:
                    score_match.append(Finding.cve.in_(catalog_match))
                match = and_(severity_match, or_(*score_match), ~prior_match)
                eligible.append(match)
                if bool(rule.get("noncompliant", True)):
                    noncompliant.append(and_(match, _raw_overdue_expression(Finding, {service_id: cfg}, now)))
                prior_match = or_(prior_match, and_(severity_match, or_(*score_match)))
        if eligible:
            eligible_parts.append(and_(service_scope, or_(*eligible)))
        if noncompliant:
            noncompliant_parts.append(and_(service_scope, or_(*noncompliant)))
    return (or_(*eligible_parts) if eligible_parts else false(), or_(*noncompliant_parts) if noncompliant_parts else false(), needs_observations)


def service_overview_rows_aggregated(
    db: Session,
    auth: AuthContext,
    now: datetime,
    configuration: dict[str, str],
    services: list[Service] | None = None,
    configurations: dict[int, dict[str, str]] | None = None,
) -> tuple[list[dict], dict[int, dict[str, int]]]:
    """Summary path that never transfers individual findings to Python.

    PostgreSQL/SQLite perform the counts, exception checks, due-date and
    risk-overlay comparisons, and oldest-finding calculation.  This keeps
    page cost tied to the number of services rather than finding history.
    """
    if services is None or configurations is None:
        services, configurations = _overview_services_and_configurations(db, auth, configuration)
    if not services:
        return [], {}
    service_ids = [service.id for service in services]
    latest_execution_time = (
        select(Execution.service_id, func.max(Execution.scanned_at).label("latest_scanned_at"))
        .where(Execution.service_id.in_(service_ids)).group_by(Execution.service_id).subquery()
    )
    latest_by_service = {
        row.service_id: row for row in db.execute(select(
            Execution.service_id, Execution.scanned_at, Execution.complete, Execution.raw_payload,
        ).join(latest_execution_time, and_(
            Execution.service_id == latest_execution_time.c.service_id,
            Execution.scanned_at == latest_execution_time.c.latest_scanned_at,
        )))
    }
    finding_exception = select(ExceptionRecord.id).where(
        ExceptionRecord.finding_id == Finding.id,
        ExceptionRecord.revoked_at.is_(None), ExceptionRecord.starts_at <= now,
        ExceptionRecord.expires_at > now,
    ).exists()
    finding_eligible, finding_noncompliant, needs_observations = _risk_finding_expressions(configurations, now, finding_exception)
    finding_query = select(
        Finding.service_id,
        func.sum(case((finding_eligible, 1), else_=0)).label("total"),
        func.sum(case((and_(finding_eligible, finding_exception), 1), else_=0)).label("excepted"),
        func.sum(case((and_(finding_noncompliant, ~finding_exception), 1), else_=0)).label("noncompliant"),
        func.min(Finding.episode_started).label("oldest"),
    ).where(Finding.service_id.in_(service_ids), Finding.active.is_(True)).group_by(Finding.service_id)
    if needs_observations:
        latest_observation = select(
            FindingObservation.finding_id,
            func.max(FindingObservation.id).label("latest_id"),
        ).group_by(FindingObservation.finding_id).subquery()
        finding_query = finding_query.outerjoin(
            latest_observation, latest_observation.c.finding_id == Finding.id,
        ).outerjoin(
            FindingObservation, FindingObservation.id == latest_observation.c.latest_id,
        )
    finding_aggregate = {row.service_id: row for row in db.execute(finding_query)}
    policy_exception = select(PolicyExceptionRecord.id).where(
        PolicyExceptionRecord.policy_finding_id == PolicyFinding.id,
        PolicyExceptionRecord.revoked_at.is_(None), PolicyExceptionRecord.starts_at <= now,
        PolicyExceptionRecord.expires_at > now,
    ).exists()
    policy_overdue = _hardening_overdue_expression(configurations, now)
    policy_aggregate = {
        row.service_id: row for row in db.execute(select(
            PolicyFinding.service_id,
            func.count(PolicyFinding.id).label("total"),
            func.sum(case((policy_exception, 1), else_=0)).label("excepted"),
            func.sum(case((and_(policy_overdue, ~policy_exception), 1), else_=0)).label("noncompliant"),
            func.min(PolicyFinding.episode_started).label("oldest"),
        ).where(PolicyFinding.service_id.in_(service_ids), PolicyFinding.active.is_(True)).group_by(PolicyFinding.service_id))
    }
    archive_by_service: dict[int, str] = {}
    for row in db.execute(select(
        ServiceArchiveEvent.service_id, ServiceArchiveEvent.action,
    ).where(ServiceArchiveEvent.service_id.in_(service_ids)).order_by(ServiceArchiveEvent.created_at.desc())):
        archive_by_service.setdefault(row.service_id, row.action)
    poam_counts: dict[int, dict[str, int]] = {}
    for row in db.execute(select(
        PoamEntry.service_id, PoamEntry.status, func.count(PoamEntry.id),
    ).where(PoamEntry.service_id.in_(service_ids)).group_by(PoamEntry.service_id, PoamEntry.status)):
        counts = poam_counts.setdefault(row.service_id, {"active": 0, "pending": 0, "overdue": 0})
        if row.status == "active":
            counts["active"] += int(row[2] or 0)
        elif row.status == "pending_approval":
            counts["pending"] += int(row[2] or 0)
    for row in db.execute(select(PoamEntry.service_id).where(
        PoamEntry.service_id.in_(service_ids), PoamEntry.status == "active",
        PoamEntry.due_date.is_not(None), PoamEntry.due_date < now,
    )):
        poam_counts.setdefault(row.service_id, {"active": 0, "pending": 0, "overdue": 0})["overdue"] += 1
    rows = []
    for service in services:
        cfg = configurations[service.id]
        findings = finding_aggregate.get(service.id)
        policies = policy_aggregate.get(service.id)
        finding_total = int(findings.total or 0) if findings else 0
        finding_excepted = int(findings.excepted or 0) if findings else 0
        finding_noncompliant = int(findings.noncompliant or 0) if findings else 0
        policy_total = int(policies.total or 0) if policies else 0
        policy_excepted = int(policies.excepted or 0) if policies else 0
        policy_noncompliant = int(policies.noncompliant or 0) if policies else 0
        latest = latest_by_service.get(service.id)
        raw_payload = latest.raw_payload if latest and isinstance(latest.raw_payload, dict) else {}
        skipped_images = raw_payload.get("skipped_images", []) or []
        skipped_charts = raw_payload.get("skipped_charts", []) or []
        incomplete = bool(latest and (not latest.complete or skipped_images or skipped_charts))
        evidence_state = "No evidence" if not latest else ("Complete" if latest.complete else "Incomplete")
        if skipped_images and evidence_state == "Incomplete":
            evidence_state = f"Incomplete · {len(skipped_images)} skipped"
        if skipped_charts and evidence_state == "Incomplete":
            evidence_state = f"Incomplete · {len(skipped_images)} skipped images, {len(skipped_charts)} skipped charts" if skipped_images else f"Incomplete · {len(skipped_charts)} skipped charts"
        evidence_noncompliant = bool(incomplete and cfg.get("incomplete_noncompliant") == "true")
        oldest_values = [value for value in ((findings.oldest if findings else None), (policies.oldest if policies else None)) if value is not None]
        oldest_age = max(((now - aware(value)).days for value in oldest_values), default=None)
        rows.append({
            "service": service,
            "compliant": not finding_noncompliant and not policy_noncompliant and not evidence_noncompliant,
            "version": service.manual_version or (raw_payload.get("service", {}).get("version") if isinstance(raw_payload.get("service"), dict) else None),
            "evidence_state": evidence_state,
            "active": [None] * (finding_total - finding_excepted - finding_noncompliant + policy_total - policy_excepted - policy_noncompliant),
            "policy_findings": [],
            "noncompliant": [None] * finding_noncompliant,
            "policy_noncompliant": [None] * policy_noncompliant,
            "excepted": [None] * finding_excepted,
            "policy_excepted": [None] * policy_excepted,
            "oldest_age": oldest_age,
            "archive": archive_by_service.get(service.id) == "archive",
            "poam": poam_counts.get(service.id, {"active": 0, "pending": 0, "overdue": 0}),
            "overdue": [None] * finding_noncompliant,
        })
    return rows, poam_counts


def service_overview_rows(db: Session, auth: AuthContext, now: datetime, configuration: dict[str, str]):
    services, configurations = _overview_services_and_configurations(db, auth, configuration)
    if not services:
        return [], {}
    return service_overview_rows_aggregated(db, auth, now, configuration, services, configurations)

@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, db: Session = Depends(get_db)):
    configuration = get_global_configuration(db)
    mode = configuration.get("identity_mode", os.getenv("CATS_IDENTITY_MODE", "local"))
    return templates.TemplateResponse(request, "login.html", {
        "error": request.query_params.get("error"), "next": request.query_params.get("next", "/"),
        "oidc_available": oidc_enabled(mode), "local_available": mode in {"local", "both"},
        "oidc_provider_name": configured_oidc(configuration).get("provider_name") or "OIDC",
    })


@app.get("/auth/oidc/login")
def oidc_login(request: Request, next: str = "/", db: Session = Depends(get_db)):
    configuration = get_global_configuration(db)
    if not oidc_enabled(configuration.get("identity_mode")):
        raise HTTPException(404, detail="OIDC login is not enabled")
    state = secrets.token_urlsafe(32)
    nonce = secrets.token_urlsafe(32)
    try:
        oidc_config = parse_json(configuration.get("oidc_configuration"), {})
        if isinstance(oidc_config, dict) and oidc_config.get("client_secret"):
            try: oidc_config["client_secret"] = decrypt_secret(oidc_config["client_secret"])
            except ValueError: raise HTTPException(503, detail="OIDC client secret is unavailable")
        ca_bundle = configured_ca_bundle(configuration)
        location = oidc_authorization_url(state, nonce, oidc_config, ca_bundle)
    except Exception as exc:
        raise HTTPException(503, detail=f"OIDC is not configured: {redact(exc)}") from exc
    response = RedirectResponse(location, status_code=303)
    response.set_cookie("cats_oidc_state", state, httponly=True, secure=os.getenv("SESSION_COOKIE_SECURE", "false").lower() == "true", samesite="lax", max_age=600)
    response.set_cookie("cats_oidc_nonce", nonce, httponly=True, secure=os.getenv("SESSION_COOKIE_SECURE", "false").lower() == "true", samesite="lax", max_age=600)
    response.set_cookie("cats_oidc_next", next if next.startswith("/") and not next.startswith("//") else "/", httponly=True, secure=os.getenv("SESSION_COOKIE_SECURE", "false").lower() == "true", samesite="lax", max_age=600)
    return response


@app.get("/auth/oidc/callback")
def oidc_callback(request: Request, db: Session = Depends(get_db)):
    if request.query_params.get("state") != request.cookies.get("cats_oidc_state"):
        raise HTTPException(400, detail="Invalid OIDC state")
    error = request.query_params.get("error")
    if error:
        return RedirectResponse(f"/login?error={urllib.parse.quote(redact(error))}", status_code=303)
    code = request.query_params.get("code", "")
    if not code:
        raise HTTPException(400, detail="OIDC callback did not include an authorization code")
    try:
        oidc_config = parse_json(get_global_configuration(db).get("oidc_configuration"), {})
        if isinstance(oidc_config, dict) and oidc_config.get("client_secret"):
            oidc_config["client_secret"] = decrypt_secret(oidc_config["client_secret"])
        configuration = get_global_configuration(db)
        ca_bundle = configured_ca_bundle(configuration)
        tokens, discovery = oidc_exchange_code(code, oidc_config, ca_bundle)
        id_claims = verify_oidc_id_token(tokens, discovery, request.cookies.get("cats_oidc_nonce"), oidc_config)
        access_token = tokens.get("access_token")
        if not access_token:
            raise ValueError("OIDC token response did not include an access token")
        userinfo_request = urllib.request.Request(discovery["userinfo_endpoint"], headers={"Authorization": f"Bearer {access_token}"})
        try:
            if ca_bundle:
                import ssl
                context = ssl.create_default_context()
                context.load_verify_locations(cadata=ca_bundle)
                with urllib.request.urlopen(userinfo_request, timeout=15, context=context) as response:
                    claims = {**id_claims, **json.load(response)}
            else:
                with urllib.request.urlopen(userinfo_request, timeout=15) as response:
                    claims = {**id_claims, **json.load(response)}
        except urllib.error.HTTPError as exc:
            if exc.code == 401:
                # UserInfo is supplementary. The ID token has already been
                # signature-, audience-, issuer-, and nonce-validated, so its
                # claims are a safe fallback for providers that do not permit
                # this client to call UserInfo.
                claims = id_claims
            else:
                raise ValueError(f"OIDC user-info endpoint returned HTTP {exc.code}") from exc
        user = provision_oidc_user(db, claims, oidc_config)
    except Exception as exc:
        db.rollback()
        # Never reflect provider responses, token payloads, or configuration
        # values into the browser.  They may contain credentials or claims.
        return RedirectResponse(f"/login?error={urllib.parse.quote(f'OIDC login failed: {redact(exc)}')}", status_code=303)
    now = utcnow()
    user.last_login_at = now
    raw_token = secrets.token_urlsafe(48)
    session = UserSession(token_hash=token_hash(raw_token), csrf_token=secrets.token_urlsafe(32), user=user,
                          expires_at=now + timedelta(hours=int(os.getenv("SESSION_HOURS", "8"))),
                          user_agent=request.headers.get("user-agent"), source_ip=request.client.host if request.client else None)
    db.add(session)
    record_audit(db, AuthContext(user, session), "auth.oidc_login", "user", user.id)
    db.commit()
    response = RedirectResponse(request.cookies.get("cats_oidc_next", "/"), status_code=303)
    response.set_cookie(SESSION_COOKIE, raw_token, httponly=True, samesite="strict", secure=os.getenv("SESSION_COOKIE_SECURE", "false").lower() == "true", max_age=int(os.getenv("SESSION_HOURS", "8")) * 3600)
    response.delete_cookie("cats_oidc_state")
    response.delete_cookie("cats_oidc_nonce")
    response.delete_cookie("cats_oidc_next")
    return response


@app.post("/login")
def login(
    request: Request, username: str = Form(), password: str = Form(),
    next: str = Form(default="/"), db: Session = Depends(get_db),
):
    user = db.scalar(select(User).where(User.username == username.strip().lower()))
    now = utcnow()
    configuration = get_global_configuration(db)
    if configuration.get("identity_mode", os.getenv("CATS_IDENTITY_MODE", "local")) not in {"local", "both"}:
        return templates.TemplateResponse(request, "login.html", {
            "error": "Local account login is disabled. Use the organizational sign-in button.", "next": next,
            "oidc_available": True, "local_available": False,
            "oidc_provider_name": configured_oidc(configuration).get("provider_name") or "OIDC",
        }, status_code=403)
    valid = bool(user and user.enabled and (not user.locked_until or aware(user.locked_until) <= now))
    valid = valid and verify_password(password, user.password_hash)
    if not valid:
        if user and user.enabled:
            user.failed_login_count += 1
            if user.failed_login_count >= 5:
                user.locked_until = now + timedelta(minutes=15)
                user.failed_login_count = 0
            db.commit()
        return templates.TemplateResponse(request, "login.html", {
            "error": "Invalid username or password.", "next": next,
            "oidc_available": oidc_enabled(configuration.get("identity_mode")),
            "local_available": configuration.get("identity_mode") in {"local", "both"},
            "oidc_provider_name": configured_oidc(configuration).get("provider_name") or "OIDC",
        }, status_code=401)
    user.failed_login_count = 0
    user.locked_until = None
    user.last_login_at = now
    raw_token = secrets.token_urlsafe(48)
    session = UserSession(
        token_hash=token_hash(raw_token), csrf_token=secrets.token_urlsafe(32), user=user,
        expires_at=now + timedelta(hours=int(os.getenv("SESSION_HOURS", "8"))),
        user_agent=request.headers.get("user-agent"),
        source_ip=request.client.host if request.client else None,
    )
    db.add(session)
    record_audit(db, AuthContext(user, session), "auth.login", "user", user.id)
    db.commit()
    destination = "/account/password" if user.must_change_password else (next if next.startswith("/") and not next.startswith("//") else "/")
    response = RedirectResponse(destination, status_code=303)
    response.set_cookie(
        SESSION_COOKIE, raw_token, httponly=True, samesite="strict",
        secure=os.getenv("SESSION_COOKIE_SECURE", "false").lower() == "true",
        max_age=int(os.getenv("SESSION_HOURS", "8")) * 3600,
    )
    return response


@app.post("/logout")
def logout(csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    check_csrf(auth, csrf_token)
    record_audit(db, auth, "auth.logout", "user", auth.user.id)
    db.delete(auth.session)
    db.commit()
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie(SESSION_COOKIE)
    return response


@app.get("/account/password", response_class=HTMLResponse)
def password_page(request: Request, auth: AuthContext = Depends(require_user)):
    return templates.TemplateResponse(request, "password.html", page_context(auth, error=None))


THEMES = {
    "cats": "CATS Green",
    "blue": "Ocean Blue",
    "red": "Warm Red",
    "gray": "Slate Gray",
    "light": "High-contrast Light",
}


@app.get("/account/appearance", response_class=HTMLResponse)
def appearance_page(request: Request, auth: AuthContext = Depends(require_user)):
    return templates.TemplateResponse(request, "appearance.html", page_context(auth, themes=THEMES, saved=request.query_params.get("saved") == "1" or request.query_params.get("theme_saved") == "1"))


@app.post("/account/appearance")
def save_appearance(
    theme: str = Form(), csrf_token: str = Form(), next_path: str = Form("/"), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    if theme not in THEMES:
        raise HTTPException(422, detail="Unsupported theme")
    auth.user.theme = theme
    record_audit(db, auth, "user.theme_changed", "user", auth.user.id, theme=theme)
    db.commit()
    if not next_path.startswith("/") or next_path.startswith("//"):
        next_path = "/"
    separator = "&" if "?" in next_path else "?"
    return RedirectResponse(f"{next_path}{separator}theme_saved=1", status_code=303)


@app.post("/account/password")
def change_password(
    current_password: str = Form(), new_password: str = Form(), confirmation: str = Form(),
    csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    if not verify_password(current_password, auth.user.password_hash):
        raise HTTPException(422, detail="Current password is incorrect")
    if new_password != confirmation:
        raise HTTPException(422, detail="New passwords do not match")
    try:
        auth.user.password_hash = hash_password(new_password)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc
    auth.user.must_change_password = False
    db.execute(delete(UserSession).where(UserSession.user_id == auth.user.id, UserSession.id != auth.session.id))
    record_audit(db, auth, "user.password_changed", "user", auth.user.id)
    db.commit()
    return RedirectResponse("/", status_code=303)


def _image_identity(reference: str | None, digest: str | None = None) -> tuple[str, str | None]:
    value = str(reference or "").strip()
    digest_value = str(digest or "").strip() or None
    if "@" in value:
        value, embedded_digest = value.rsplit("@", 1)
        digest_value = digest_value or embedded_digest.strip() or None
    return value, digest_value


def _ensure_service_image(db: Session, service: Service, reference: str | None, digest: str | None = None) -> ServiceImage | None:
    image_reference, image_digest = _image_identity(reference, digest)
    if not image_reference:
        return None
    # A scan often observes the tag first and learns the manifest digest from
    # a later source.  Reconcile that occurrence in place so the service has
    # one canonical image and remediation cannot patch it twice.
    if image_digest:
        resolved = db.scalar(select(ServiceImage).where(
            ServiceImage.service_id == service.id,
            ServiceImage.image_reference == image_reference,
            ServiceImage.image_digest == image_digest,
        ))
        if resolved:
            stale = db.scalars(select(ServiceImage).where(
                ServiceImage.service_id == service.id,
                ServiceImage.image_reference == image_reference,
                ServiceImage.image_digest.is_(None),
            )).all()
            for occurrence in stale:
                db.delete(occurrence)
            return resolved
        unresolved = db.scalar(select(ServiceImage).where(
            ServiceImage.service_id == service.id,
            ServiceImage.image_reference == image_reference,
            ServiceImage.image_digest.is_(None),
        ))
        if unresolved:
            unresolved.image_digest = image_digest
            unresolved.updated_at = utcnow()
            return unresolved
    else:
        resolved = db.scalar(select(ServiceImage).where(
            ServiceImage.service_id == service.id,
            ServiceImage.image_reference == image_reference,
            ServiceImage.image_digest.is_not(None),
        ))
        if resolved:
            return resolved
    image = db.scalar(select(ServiceImage).where(
        ServiceImage.service_id == service.id,
        ServiceImage.image_reference == image_reference,
        ServiceImage.image_digest == image_digest,
    ))
    if not image:
        image = ServiceImage(service=service, image_reference=image_reference, image_digest=image_digest)
        db.add(image)
        db.flush()
    return image


def _reconcile_image_scope(db: Session, service: Service, target_image: str, observed_cves: set[str], scanned_at: datetime) -> None:
    """Resolve only findings whose evidence belongs to the scanned image."""
    target = target_image.strip()
    if not target:
        return
    active_images = {
        image.image_reference for image in service.images
        if image.lifecycle_status == "active"
    }
    observations = list(db.scalars(
        select(FindingObservation).join(Finding).where(Finding.service_id == service.id)
    ))
    observations_by_finding: dict[int, list[FindingObservation]] = {}
    for observation in observations:
        observations_by_finding.setdefault(observation.finding_id, []).append(observation)
    for finding in db.scalars(select(Finding).where(Finding.service_id == service.id, Finding.active.is_(True))):
        if finding.cve in observed_cves:
            continue
        has_other_active_image = any(
            observation.image != target and observation.image in active_images
            for observation in observations_by_finding.get(finding.id, [])
        )
        if not has_other_active_image:
            finding.active = False
            finding.resolved_at = scanned_at


@app.post("/api/v1/pipeline-results", status_code=201, dependencies=[Depends(require_pipeline)])
def ingest(payload: ExecutionPayload, db: Session = Depends(get_db)):
    if not payload.fixable_only:
        raise HTTPException(status_code=422, detail="Portal accepts fixable-only assessments")
    existing = db.scalar(select(Execution).where(Execution.execution_key == payload.execution_id))
    if existing:
        if existing.raw_payload != payload.model_dump(mode="json"):
            raise HTTPException(status_code=409, detail="Execution ID already exists with different evidence")
        return {"accepted": True, "duplicate": True, "execution_id": existing.id}

    service = db.scalar(select(Service).where(Service.service_key == payload.service.id))
    if not service:
        service = Service(service_key=payload.service.id, name=payload.service.name, lifecycle_status="active")
        db.add(service)
        db.flush()
    elif archive_state(service):
        raise HTTPException(status_code=409, detail="Service is archived; restore it before accepting new evidence")
    service.name = payload.service.name
    if payload.service.description is not None:
        service.description = payload.service.description
    service.owner = payload.service.owner
    service.poc = payload.service.poc
    promoted_from_staged = service_lifecycle(service) == "staged"
    if promoted_from_staged:
        service.lifecycle_status = "active"
    known_service_image_ids = set(db.scalars(select(ServiceImage.id).where(ServiceImage.service_id == service.id)).all())
    requested_groups = {name.strip() for name in payload.service.groups if name.strip()}
    existing_groups = {group.name: group for group in service.groups}
    for group_name in requested_groups:
        group = db.scalar(select(Group).where(Group.name == group_name))
        if not group:
            group = Group(name=group_name)
            db.add(group)
            db.flush()
        if group_name not in existing_groups:
            service.groups.append(group)
    scan_scope = payload.scan_scope or "service"
    scope_image = (payload.scope_image or "").strip() or None
    finding_images = {str(item.image).strip() for item in payload.findings if str(item.image).strip()}
    if scan_scope == "image" and not scope_image:
        if len(finding_images) == 1:
            scope_image = next(iter(finding_images))
        else:
            raise HTTPException(status_code=422, detail="Image-scoped scans must identify exactly one image")
    configuration = configuration_for_service(db, service)
    effective_complete = payload.complete and not (
        (bool(payload.skipped_images)
         and configuration.get("skipped_images_incomplete", "true") == "true")
        or bool(payload.skipped_charts)
    )
    if payload.helm_source_files and isinstance(payload.service_overview, dict):
        _enrich_values_source_mappings({"service_overview": payload.service_overview}, payload.helm_source_files)
    execution = Execution(
        execution_key=payload.execution_id, service=service, scanned_at=payload.scanned_at,
        complete=effective_complete, scan_scope=scan_scope, scope_image=scope_image,
        pipeline_url=payload.pipeline_url,
        commit_sha=payload.commit_sha, scanner_db_built_at=payload.scanner_db_built_at,
        raw_payload=payload.model_dump(mode="json"),
    )
    db.add(execution)
    db.flush()
    observed = set()
    for item in payload.findings:
        observed.add(item.cve)
        finding = db.scalar(select(Finding).where(Finding.service_id == service.id, Finding.cve == item.cve))
        if not finding:
            finding = Finding(service=service, cve=item.cve, severity=item.severity,
                              first_seen=payload.scanned_at, episode_started=payload.scanned_at,
                              last_seen=payload.scanned_at, active=True)
            db.add(finding)
            db.flush()
        elif not finding.active:
            finding.active = True
            finding.episode_started = payload.scanned_at
            finding.resolved_at = None
            finding.recurrence_count += 1
        finding.last_seen = payload.scanned_at
        finding.severity = item.severity
        evidence = dict(item.evidence or {})
        evidence.setdefault("kev", item.kev)
        if item.epss is not None:
            evidence.setdefault("epss", item.epss)
        db.add(FindingObservation(
            finding=finding, execution_id=execution.id, image=item.image,
            image_digest=item.image_digest, package=item.package,
            installed_version=item.installed_version, fixed_version=item.fixed_version,
            evidence=evidence,
        ))
        _ensure_service_image(db, service, item.image, item.image_digest)
    overview_payload = payload.service_overview if isinstance(payload.service_overview, dict) else {}
    for raw_image in [*(overview_payload.get("images") or []), *(overview_payload.get("container_images") or [])]:
        if isinstance(raw_image, dict):
            _ensure_service_image(db, service, raw_image.get("image") or raw_image.get("reference") or raw_image.get("name"), raw_image.get("digest") or raw_image.get("image_digest"))
        else:
            _ensure_service_image(db, service, str(raw_image), None)
    for artifact in overview_payload.get("artifacts") or []:
        if isinstance(artifact, dict) and str(artifact.get("type") or "").lower() == "image":
            registry = str(artifact.get("registry") or "").strip()
            repository = str(artifact.get("repository") or "").strip("/")
            name = str(artifact.get("artifact") or "").strip()
            version = str(artifact.get("version") or "").strip()
            reference = "/".join(part for part in (registry, repository, name) if part and part != "—")
            if version and version != "—":
                reference = f"{reference}:{version}"
            _ensure_service_image(db, service, reference, artifact.get("digest"))
    if effective_complete:
        if scan_scope == "image":
            _ensure_service_image(db, service, scope_image)
            _reconcile_image_scope(db, service, scope_image or "", observed, payload.scanned_at)
        else:
            for finding in db.scalars(select(Finding).where(Finding.service_id == service.id, Finding.active.is_(True))):
                if finding.cve not in observed:
                    finding.active = False
                    finding.resolved_at = payload.scanned_at
    sync_policy_findings(db, service, payload.policy_findings, payload.scanned_at, effective_complete)
    for image in db.scalars(select(ServiceImage).where(ServiceImage.service_id == service.id)).all():
        if image.id not in known_service_image_ids:
            db.add(AuditEvent(
                action="image_added", target_type="service_image", target_id=str(image.id),
                detail={"service_id": service.id, "image": image.image_reference, "digest": image.image_digest},
            ))
    db.add(AuditEvent(
        action="scan.ingested", target_type="execution", target_id=str(execution.id),
        detail={"service_id": service.id, "scan_scope": scan_scope, "scope_image": scope_image,
                "complete": effective_complete},
    ))
    if promoted_from_staged:
        db.add(AuditEvent(
            action="service.promoted", target_type="service", target_id=str(service.id),
            detail={"service_id": service.id, "execution_id": execution.id, "reason": "ingested_evidence"},
        ))
    db.commit()
    return {"accepted": True, "duplicate": False, "execution_id": execution.id}


@app.get("/home", response_class=HTMLResponse)
def public_home(request: Request, auth: AuthContext | None = Depends(optional_user)):
    if auth:
        return RedirectResponse("/", status_code=303)
    return templates.TemplateResponse(request, "home.html", {
        "current_user": auth.user if auth else None,
        "csrf_token": auth.csrf_token if auth else "",
        "can": auth.has if auth else (lambda _permission: False),
        "themes": THEMES,
        "pending_request_count": 0,
        "pending_poam_count": 0,
    })


def self_service_context(request: Request, mode: str, image_list: str = "", chart_url: str = "", status_message: str | None = None, job_id: str | None = None, auth: AuthContext | None = None, services: list[Service] | None = None, ingest_service_id: str = "", archive_names: list[str] | None = None, sbom_formats: list[str] | None = None, cyclonedx_spec_version: str = "1.5"):
    descriptions = {
        "scan": "Review public container images without creating a persistent service record.",
        "sbom": "Generate one or more standard SBOM documents from a single image inventory.",
        "patch": "Request patched image outputs for public container images without changing a persistent service record.",
    }
    progress_phases = (
        [("prepare_inputs", "Prepare"), ("generate_sboms", "Generate")]
        if mode == "sbom"
        else [("prepare_inputs", "Prepare"), ("generate_sboms", "SBOM"), ("scan_sboms", "Scan"), ("configuration_scan", "Configuration"), ("report_results", "Report")]
    )
    return templates.TemplateResponse(request, "self_service.html", {
        "current_user": auth.user if auth else None, "csrf_token": auth.csrf_token if auth else "", "can": auth.has if auth else (lambda _permission: False),
        "themes": THEMES, "pending_request_count": 0, "pending_poam_count": 0,
        "mode": mode, "description": descriptions[mode], "image_list": image_list, "chart_url": chart_url,
        "archive_names": archive_names or [],
        "authenticated_services": services or [], "authenticated_ingest": bool(auth and auth.accessible_service_ids("scan.ingest") != set()),
        "ingest_service_id": ingest_service_id,
        "status_message": status_message, "job_id": job_id,
        "sbom_output_formats": SBOM_OUTPUT_FORMATS,
        "selected_sbom_formats": sbom_formats or ["cyclonedx-json"],
        "cyclonedx_spec_versions": CYCLONEDX_SPEC_VERSIONS,
        "cyclonedx_spec_version": cyclonedx_spec_version,
        "progress_phases": progress_phases,
        "progress_order": ["queued", *(phase for phase, _label in progress_phases)],
    })


def _public_job_update(job_id: str, **values):
    with PUBLIC_JOB_LOCK:
        if job_id in PUBLIC_JOBS:
            PUBLIC_JOBS[job_id].update(values)


def _safe_chart_member(name: str) -> bool:
    normalized = name.replace("\\", "/")
    return bool(normalized) and not normalized.startswith("/") and ".." not in normalized.split("/")


def _stage_public_chart(input_dir: Path, archive: bytes, filename: str) -> None:
    """Extract one uploaded Helm archive into the ephemeral charts directory."""
    max_bytes = int(os.getenv("CATS_PUBLIC_MAX_CHART_BYTES", str(50 * 1024 * 1024)))
    if len(archive) > max_bytes:
        raise HTTPException(status_code=413, detail="Helm chart archive is too large")
    charts_dir = input_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)
    name = (filename or "chart.tgz").lower()
    try:
        if name.endswith(".zip"):
            with ZipFile(BytesIO(archive)) as bundle:
                members = bundle.infolist()
                if any(not _safe_chart_member(member.filename) for member in members):
                    raise HTTPException(status_code=400, detail="Helm archive contains an unsafe path")
                bundle.extractall(charts_dir)
        elif name.endswith((".tgz", ".tar.gz", ".tar")):
            with tarfile.open(fileobj=BytesIO(archive), mode="r:*") as bundle:
                members = bundle.getmembers()
                if any(not _safe_chart_member(member.name) for member in members):
                    raise HTTPException(status_code=400, detail="Helm archive contains an unsafe path")
                bundle.extractall(charts_dir, filter="data")
        else:
            raise HTTPException(status_code=400, detail="Helm chart must be a .tgz, .tar.gz, .tar, or .zip archive")
    except (tarfile.TarError, OSError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Helm chart archive could not be read") from exc
    if not any(charts_dir.rglob("Chart.yaml")):
        raise HTTPException(status_code=400, detail="Helm archive does not contain a Chart.yaml")


def _fetch_public_url(url: str) -> tuple[bytes, str]:
    """Fetch a public URL and return its bytes plus the final URL."""
    parsed = urllib.parse.urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(status_code=400, detail="Chart URL must use http or https")
    max_bytes = int(os.getenv("CATS_PUBLIC_MAX_CHART_BYTES", str(50 * 1024 * 1024)))
    try:
        request = urllib.request.Request(url.strip(), headers={"User-Agent": "CATS/standalone-scanner"})
        with urllib.request.urlopen(request, timeout=30) as response:
            content_length = int(response.headers.get("Content-Length") or 0)
            if content_length > max_bytes:
                raise HTTPException(status_code=413, detail="Helm chart archive is too large")
            data = response.read(max_bytes + 1)
            final_url = response.geturl()
        if len(data) > max_bytes:
            raise HTTPException(status_code=413, detail="Helm chart archive is too large")
    except HTTPException:
        raise
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Chart URL could not be downloaded") from exc
    return data, final_url


def _download_oci_chart(reference: str) -> list[tuple[bytes, str]]:
    """Pull one public OCI Helm chart using the bundled Helm executable."""
    helm = shutil.which("helm") or "/usr/local/bin/helm"
    if not Path(helm).exists() and shutil.which(helm) is None:
        raise HTTPException(status_code=400, detail="OCI Helm charts require Helm in the scanner image")
    max_bytes = int(os.getenv("CATS_PUBLIC_MAX_CHART_BYTES", str(50 * 1024 * 1024)))
    try:
        with tempfile.TemporaryDirectory(prefix="cats-oci-chart-") as destination:
            result = subprocess.run(
                [helm, "pull", reference, "--destination", destination],
                capture_output=True, text=True,
                timeout=int(os.getenv("CATS_PUBLIC_HELM_PULL_TIMEOUT", "180")),
                check=False,
            )
            if result.returncode != 0:
                detail = (result.stderr or result.stdout or "Helm could not pull the OCI chart").strip()
                raise HTTPException(status_code=400, detail=f"OCI Helm chart could not be pulled: {detail[-500:]}")
            archives = sorted(Path(destination).glob("*.tgz")) + sorted(Path(destination).glob("*.tar.gz"))
            if not archives:
                raise HTTPException(status_code=400, detail="Helm did not produce an OCI chart archive")
            output: list[tuple[bytes, str]] = []
            for archive_path in archives:
                data = archive_path.read_bytes()
                if len(data) > max_bytes:
                    raise HTTPException(status_code=413, detail="Helm chart archive is too large")
                output.append((data, archive_path.name))
            return output
    except HTTPException:
        raise
    except (OSError, subprocess.SubprocessError) as exc:
        raise HTTPException(status_code=400, detail="OCI Helm chart could not be pulled") from exc


def _download_public_chart(url: str) -> list[tuple[bytes, str]]:
    """Download a chart archive or expand a Helm repository/index URL."""
    raw_url = url.strip()
    if raw_url.lower().startswith("oci://"):
        return _download_oci_chart(raw_url)
    parsed = urllib.parse.urlparse(raw_url)
    selector = parsed.fragment.strip()
    fetch_url = urllib.parse.urlunparse(parsed._replace(fragment=""))
    data, final_url = _fetch_public_url(fetch_url)
    final_path = urllib.parse.urlparse(final_url).path.lower()
    if final_path.endswith((".tgz", ".tar.gz", ".tar", ".zip")):
        return [(data, Path(final_path).name or "chart.tgz")]

    index_data, index_url = data, final_url
    try:
        index = yaml.safe_load(index_data.decode("utf-8-sig"))
    except (UnicodeDecodeError, yaml.YAMLError):
        index = None
    if not isinstance(index, dict) or not isinstance(index.get("entries"), dict):
        candidate_index_url = urllib.parse.urljoin(final_url, "index.yaml")
        if candidate_index_url != final_url:
            try:
                index_data, index_url = _fetch_public_url(candidate_index_url)
                index = yaml.safe_load(index_data.decode("utf-8-sig"))
            except (UnicodeDecodeError, yaml.YAMLError):
                index = None
    if not isinstance(index, dict) or not isinstance(index.get("entries"), dict):
        raise HTTPException(status_code=400, detail="Helm URL must point to a chart archive or Helm repository index.yaml")

    entries = index["entries"]
    names = [selector] if selector else list(entries)
    if selector and selector not in entries:
        raise HTTPException(status_code=400, detail=f"Helm repository does not contain chart: {selector}")
    max_charts = max(1, int(os.getenv("CATS_PUBLIC_MAX_REPOSITORY_CHARTS", "25")))
    archives: list[tuple[bytes, str]] = []
    for chart_name in names[:max_charts]:
        versions = entries.get(chart_name)
        if not isinstance(versions, list) or not versions:
            continue
        version = next((item for item in versions if isinstance(item, dict) and item.get("urls")), None)
        if not version:
            continue
        base_url = index_url if index_url.endswith("/") else index_url.rsplit("/", 1)[0] + "/"
        archive_url = urllib.parse.urljoin(base_url, str((version.get("urls") or [])[0]))
        archive_data, archive_final_url = _fetch_public_url(archive_url)
        archive_name = Path(urllib.parse.urlparse(archive_final_url).path).name
        if not archive_name.lower().endswith((".tgz", ".tar.gz", ".tar", ".zip")):
            archive_name = f"{chart_name}-{version.get('version', 'latest')}.tgz"
        archives.append((archive_data, archive_name))
    if not archives:
        raise HTTPException(status_code=400, detail="Helm repository index contains no downloadable chart archives")
    return archives


def _run_public_scan(job_id: str, image_list: str):
    job_dir = PUBLIC_JOB_ROOT / job_id
    input_dir = job_dir / "input"
    output_dir = job_dir / "output"
    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    if image_list.strip():
        input_dir.joinpath("images.txt").write_text(image_list, encoding="utf-8")
    runner = os.getenv("CATS_SCANNER_RUNNER", "cats-scan")
    with PUBLIC_JOB_LOCK:
        job_configuration = dict(PUBLIC_JOBS.get(job_id, {}))
        if job_configuration.get("status") == "cancelled":
            return
    job_kind = str(job_configuration.get("job_kind") or "scan")
    process_environment = os.environ.copy()
    process_environment["CATS_JOB_MODE"] = job_kind
    if job_kind == "sbom":
        process_environment["SBOM_FORMATS"] = ",".join(job_configuration.get("sbom_formats") or ["cyclonedx-json"])
        process_environment["SBOM_CYCLONEDX_SPEC_VERSION"] = str(job_configuration.get("cyclonedx_spec_version") or "1.5")
    _public_job_update(job_id, status="running", phase="prepare_inputs")
    try:
        log_path = output_dir / "worker.log"
        with log_path.open("w", encoding="utf-8") as log_file:
            process = subprocess.Popen(
                [runner, str(input_dir), str(output_dir)],
                stdout=log_file, stderr=subprocess.STDOUT, text=True, env=process_environment,
            )
            with PUBLIC_JOB_LOCK:
                if PUBLIC_JOBS.get(job_id, {}).get("status") == "cancelled":
                    process.terminate()
                else:
                    PUBLIC_PROCESSES[job_id] = process
            phases = (("prepare_inputs", "generate_sboms") if job_kind == "sbom" else
                      ("prepare_inputs", "generate_sboms", "scan_sboms", "configuration_scan", "report_results"))
            deadline = time.monotonic() + int(os.getenv("CATS_PUBLIC_JOB_TIMEOUT", "3600"))
            while process.poll() is None:
                if time.monotonic() > deadline:
                    process.kill()
                    process.wait()
                    raise TimeoutError("public scan exceeded its timeout")
                for phase in phases:
                    phase_path = output_dir / f"phase-{phase}.json"
                    if phase_path.exists():
                        try:
                            phase_state = json.loads(phase_path.read_text(encoding="utf-8"))
                        except (OSError, ValueError):
                            continue
                        if phase_state.get("status") == "running":
                            _public_job_update(job_id, status="running", phase=phase)
                            break
                time.sleep(0.25)
            returncode = process.returncode
        with PUBLIC_JOB_LOCK:
            PUBLIC_PROCESSES.pop(job_id, None)
            cancelled = PUBLIC_JOBS.get(job_id, {}).get("status") == "cancelled"
        if cancelled:
            return
        completed_returncode = returncode
        summary = {}
        summary_path = output_dir / "scan-summary.json"
        if summary_path.exists():
            try:
                summary = json.loads(summary_path.read_text(encoding="utf-8"))
            except (ValueError, OSError):
                summary = {}
        _public_job_update(
            job_id,
            status="complete" if completed_returncode == 0 else "incomplete",
            phase="generate_sboms" if job_kind == "sbom" else "report_results", returncode=completed_returncode,
            summary=summary,
        )
    except Exception as exc:
        with PUBLIC_JOB_LOCK:
            PUBLIC_PROCESSES.pop(job_id, None)
            cancelled = PUBLIC_JOBS.get(job_id, {}).get("status") == "cancelled"
        if cancelled:
            return
        _public_job_update(job_id, status="error", phase="worker", error=str(exc))


def _public_chart_skip_entry(source: str, error: object) -> str:
    """Return a compact, human-readable missing-evidence entry.

    The value is deliberately stored as one line because the same file is
    consumed by the GitLab report-to-portal script and by the standalone
    worker.  Keep the source first so it remains useful even when the error
    text is abbreviated.
    """
    detail = str(getattr(error, "detail", error)).strip().replace("\r", " ").replace("\n", " ")
    return f"{source} :: {detail[-500:] or 'chart could not be retrieved'}"


def _collect_helm_source_files(charts_dir: Path) -> dict[str, str]:
    """Collect editable text sources with bounded size and stable paths."""
    if not charts_dir.is_dir():
        return {}
    limit = int(os.getenv("CATS_REMEDIATION_SOURCE_MAX_BYTES", str(10 * 1024 * 1024)))
    total = 0
    files: dict[str, str] = {}
    allowed_names = {"Chart.yaml", "Chart.lock", "values.yaml", "values.yml"}
    allowed_suffixes = {".yaml", ".yml", ".tpl", ".txt"}
    for path in sorted(charts_dir.rglob("*")):
        if not path.is_file() or (path.name not in allowed_names and path.suffix.lower() not in allowed_suffixes):
            continue
        raw = path.read_bytes()
        total += len(raw)
        if total > limit:
            return {}
        try:
            files[path.relative_to(charts_dir).as_posix()] = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            continue
    return files


def _enrich_values_source_mappings(data: dict, source_files: dict[str, str]) -> None:
    """Prove simple template-to-values mappings without guessing composites."""
    overview = data.get("service_overview") if isinstance(data.get("service_overview"), dict) else {}
    resources = overview.get("rendered_resources") if isinstance(overview.get("rendered_resources"), list) else []

    def value_exists(document: object, dotted: str) -> bool:
        cursor = document
        for key in dotted.split("."):
            if not isinstance(cursor, dict) or key not in cursor:
                return False
            cursor = cursor[key]
        return True

    for resource in resources:
        if not isinstance(resource, dict):
            continue
        template_name = str(resource.get("_cats_source_file") or "").replace("\\", "/").lstrip("/")
        template_path = next((path for path in source_files if path == template_name or path.endswith("/" + template_name)
                              or template_name.endswith("/" + path)), None)
        if not template_path or "/templates/" not in "/" + template_path:
            continue
        prefix = template_path.rsplit("/templates/", 1)[0] if "/templates/" in template_path else ""
        values_path = f"{prefix}/values.yaml" if prefix else "values.yaml"
        if values_path not in source_files:
            continue
        try:
            values_document = yaml.safe_load(source_files[values_path]) or {}
        except yaml.YAMLError:
            continue
        template = source_files[template_path]
        mappings = [item for item in (resource.get("_cats_source_mappings") or []) if isinstance(item, dict)]
        for yaml_key, field_path in (
            ("allowPrivilegeEscalation", "securityContext.allowPrivilegeEscalation"),
            ("privileged", "securityContext.privileged"),
            ("readOnlyRootFilesystem", "securityContext.readOnlyRootFilesystem"),
            ("runAsNonRoot", "securityContext.runAsNonRoot"),
        ):
            pattern = rf"(?m)^\s*{re.escape(yaml_key)}\s*:\s*['\"]?\s*{{{{-?\s*\.Values\.([A-Za-z0-9_.]+)(?:\s*\|[^}}]*)?\s*-?}}}}\s*['\"]?\s*$"
            matches = sorted(set(re.findall(pattern, template)))
            if len(matches) == 1 and value_exists(values_document, matches[0]):
                mappings.append({"field_path": field_path, "template": template_path, "values_file": values_path,
                                 "values_key": f".Values.{matches[0]}", "ambiguous": False})
        image_pattern = r"(?m)^\s*image\s*:\s*['\"]?\s*{{-?\s*\.Values\.([A-Za-z0-9_.]+)(?:\s*\|[^}]*)?\s*-?}}\s*['\"]?\s*$"
        image_matches = sorted(set(re.findall(image_pattern, template)))
        if len(image_matches) == 1 and value_exists(values_document, image_matches[0]):
            mappings.append({"field_path": "spec.template.spec.containers[].image", "template": template_path,
                             "values_file": values_path, "values_key": f".Values.{image_matches[0]}", "ambiguous": False})
        resource["_cats_source_mappings"] = mappings


def _start_public_scan(image_list: str, chart_archives: list[tuple[bytes, str]] | None = None, chart_urls: list[str] | None = None, image_archive: tuple[bytes, str] | None = None, ingest_service_id: str | None = None, job_kind: str = "scan", sbom_formats: list[str] | None = None, cyclonedx_spec_version: str = "1.5") -> str:
    lines = [line.strip() for line in image_list.splitlines() if line.strip()]
    chart_archives = chart_archives or []
    chart_urls = [url.strip() for url in (chart_urls or []) if url.strip()]
    if job_kind not in {"scan", "sbom"}:
        raise HTTPException(status_code=400, detail="Unsupported analysis job type")
    format_values = sbom_formats if sbom_formats is not None else (["cyclonedx-json"] if job_kind == "sbom" else ["syft-json"])
    requested_sbom_formats = list(dict.fromkeys(format_values))
    if any(format_name not in SBOM_OUTPUT_FORMATS for format_name in requested_sbom_formats):
        raise HTTPException(status_code=400, detail="Select only supported SBOM formats")
    if job_kind == "sbom" and not requested_sbom_formats:
        raise HTTPException(status_code=400, detail="Select at least one SBOM format")
    if cyclonedx_spec_version not in CYCLONEDX_SPEC_VERSIONS:
        raise HTTPException(status_code=400, detail="Unsupported CycloneDX specification version")
    if job_kind == "sbom" and (chart_archives or chart_urls):
        raise HTTPException(status_code=400, detail="SBOM generation accepts container images or image archives")
    if not lines and not chart_archives and not chart_urls and not image_archive:
        detail = "Provide an image reference or image archive" if job_kind == "sbom" else "Provide an image reference, image archive, or Helm chart"
        raise HTTPException(status_code=400, detail=detail)
    if len(lines) > int(os.getenv("CATS_PUBLIC_MAX_IMAGES", "50")):
        raise HTTPException(status_code=413, detail="Too many image references")
    job_id = uuid.uuid4().hex
    job_input = PUBLIC_JOB_ROOT / job_id / "input"
    job_input.mkdir(parents=True, exist_ok=True)
    skipped_charts: list[str] = []
    for chart_url in chart_urls:
        # A repository or OCI endpoint is external evidence.  Its failure
        # must not discard otherwise usable images/charts from this scan.
        # Preserve the failed source for report-to-portal, which turns it
        # into a Missing Evidence item on authenticated ingest.
        try:
            chart_archives.extend(_download_public_chart(chart_url))
        except Exception as exc:
            skipped_charts.append(_public_chart_skip_entry(chart_url, exc))
    for archive, filename in chart_archives:
        try:
            _stage_public_chart(job_input, archive, filename)
        except Exception as exc:
            skipped_charts.append(_public_chart_skip_entry(filename or "uploaded chart", exc))
    if skipped_charts:
        (job_input / "skipped_charts.txt").write_text("\n".join(skipped_charts) + "\n", encoding="utf-8")
    if image_archive:
        image_dir = job_input / "image-archives"
        image_dir.mkdir(parents=True, exist_ok=True)
        image_bytes, image_filename = image_archive
        max_bytes = int(os.getenv("CATS_PUBLIC_MAX_IMAGE_ARCHIVE_BYTES", str(2 * 1024 * 1024 * 1024)))
        if len(image_bytes) > max_bytes:
            raise HTTPException(status_code=413, detail="Docker image archive is too large")
        safe_name = Path(image_filename or "images.tar").name
        if not safe_name.lower().endswith((".tar", ".tar.gz", ".tgz")):
            raise HTTPException(status_code=400, detail="Docker image upload must be a .tar, .tar.gz, or .tgz archive")
        (image_dir / safe_name).write_bytes(image_bytes)
    with PUBLIC_JOB_LOCK:
        PUBLIC_JOBS[job_id] = {"job_id": job_id, "job_kind": job_kind, "status": "queued", "phase": "queued", "summary": {}, "skipped_charts": skipped_charts, "ingest_service_id": ingest_service_id or "", "image_list": "\n".join(lines), "chart_url": "\n".join(chart_urls), "chart_names": [name for _, name in chart_archives], "archive_names": [name for _, name in chart_archives] + ([image_archive[1]] if image_archive else []), "sbom_formats": requested_sbom_formats, "cyclonedx_spec_version": cyclonedx_spec_version}
    PUBLIC_WORKERS.submit(_run_public_scan, job_id, "\n".join(lines) + "\n")
    return job_id


@app.get("/scan", response_class=HTMLResponse)
def public_scan(request: Request, job_id: str | None = None, db: Session = Depends(get_db), auth: AuthContext | None = Depends(optional_user)):
    image_list = ""
    chart_url = ""
    archive_names: list[str] = []
    ingest_service_id = ""
    if job_id:
        job_input = PUBLIC_JOB_ROOT / job_id / "input" / "images.txt"
        if job_input.exists():
            try:
                image_list = job_input.read_text(encoding="utf-8")
            except OSError:
                image_list = ""
        with PUBLIC_JOB_LOCK:
            job = PUBLIC_JOBS.get(job_id, {})
            image_list = str(job.get("image_list") or image_list)
            chart_url = str(job.get("chart_url") or "")
            archive_names = [str(name) for name in (job.get("archive_names") or [])]
            ingest_service_id = str(job.get("ingest_service_id") or "")
    services = []
    if auth and auth.accessible_service_ids("scan.ingest") != set():
        scoped = auth.accessible_service_ids("scan.ingest")
        query = select(Service).order_by(Service.name)
        services = list(db.scalars(query))
        if scoped is not None:
            services = [service for service in services if service.id in scoped]
    return self_service_context(request, "scan", image_list=image_list, chart_url=chart_url, archive_names=archive_names, job_id=job_id, auth=auth, services=services, ingest_service_id=ingest_service_id)


@app.post("/scan", response_class=HTMLResponse)
async def public_scan_submit(request: Request, image_list: str = Form(""), chart_url: str = Form(""), chart_archive: UploadFile | None = File(None), ingest_service_id: str = Form(""), db: Session = Depends(get_db), auth: AuthContext | None = Depends(optional_user)):
    try:
        chart_uploads = [chart_archive] if chart_archive and chart_archive.filename else []
        # FastAPI accepts repeated chart_archive fields as a list; inspect the
        # request form as well so the HTML multi-file input remains compatible
        # with clients that submit multiple parts.
        form = await request.form()
        chart_uploads = [value for value in form.getlist("chart_archives") if hasattr(value, "read") and getattr(value, "filename", None)] or chart_uploads
        chart_archives = [(await upload.read(), upload.filename or "chart.tgz") for upload in chart_uploads]
        image_upload = form.get("image_archive")
        image_archive = None
        if hasattr(image_upload, "read") and getattr(image_upload, "filename", None):
            image_archive = (await image_upload.read(), image_upload.filename)
        if ingest_service_id:
            if not auth or auth.accessible_service_ids("scan.ingest") == set():
                raise HTTPException(status_code=403, detail="Sign in with scan-ingest permission to select a service")
            service = db.scalar(select(Service).where(Service.service_key == ingest_service_id))
            if not service:
                raise HTTPException(status_code=404, detail="Selected service was not found")
            if not auth.has("scan.ingest", service.id):
                raise HTTPException(status_code=403, detail="Selected service is outside your scope")
        job_id = _start_public_scan(image_list, chart_archives, chart_url.splitlines(), image_archive, ingest_service_id or None)
    except HTTPException as exc:
        return self_service_context(request, "scan", image_list, chart_url, str(exc.detail), auth=auth)
    # Redirect after a successful submission so refreshing the browser only
    # reloads the existing job rather than replaying the POST and starting a
    # second scan.
    return RedirectResponse(url=f"/scan?job_id={urllib.parse.quote(job_id)}", status_code=303)


@app.get("/sbom", response_class=HTMLResponse)
def public_sbom(request: Request, job_id: str | None = None, auth: AuthContext | None = Depends(optional_user)):
    image_list = ""
    archive_names: list[str] = []
    formats = ["cyclonedx-json"]
    spec_version = "1.5"
    if job_id:
        with PUBLIC_JOB_LOCK:
            job = dict(PUBLIC_JOBS.get(job_id, {}))
        image_list = str(job.get("image_list") or "")
        archive_names = [str(name) for name in (job.get("archive_names") or [])]
        formats = [str(value) for value in (job.get("sbom_formats") or formats)]
        spec_version = str(job.get("cyclonedx_spec_version") or spec_version)
    return self_service_context(
        request, "sbom", image_list=image_list, archive_names=archive_names,
        job_id=job_id, auth=auth, sbom_formats=formats,
        cyclonedx_spec_version=spec_version,
    )


@app.post("/sbom", response_class=HTMLResponse)
async def public_sbom_submit(
    request: Request,
    image_list: str = Form(""),
    auth: AuthContext | None = Depends(optional_user),
):
    form = await request.form()
    formats = [str(value).strip() for value in form.getlist("sbom_formats") if str(value).strip()]
    spec_version = str(form.get("cyclonedx_spec_version") or "1.5").strip()
    image_upload = form.get("image_archive")
    image_archive = None
    if hasattr(image_upload, "read") and getattr(image_upload, "filename", None):
        image_archive = (await image_upload.read(), image_upload.filename)
    try:
        job_id = _start_public_scan(
            image_list,
            image_archive=image_archive,
            job_kind="sbom",
            sbom_formats=formats,
            cyclonedx_spec_version=spec_version,
        )
    except HTTPException as exc:
        return self_service_context(
            request, "sbom", image_list=image_list, status_message=str(exc.detail),
            auth=auth, sbom_formats=formats, cyclonedx_spec_version=spec_version,
        )
    return RedirectResponse(url=f"/sbom?job_id={urllib.parse.quote(job_id)}", status_code=303)


@app.post("/api/public/jobs")
def create_public_scan_job(payload: dict):
    chart_urls = payload.get("chart_urls") or ([payload.get("chart_url")] if payload.get("chart_url") else [])
    job_id = _start_public_scan(str(payload.get("images", "")), chart_urls=[str(value) for value in chart_urls])
    return {"job_id": job_id, "status_url": f"/api/public/jobs/{job_id}", "results_url": f"/api/public/jobs/{job_id}/results"}


@app.get("/api/public/jobs/{job_id}")
def public_scan_job_status(job_id: str):
    with PUBLIC_JOB_LOCK:
        job = dict(PUBLIC_JOBS.get(job_id, {}))
    if not job:
        raise HTTPException(status_code=404, detail="Job not found or expired")
    return job


@app.post("/api/public/jobs/{job_id}/cancel")
def cancel_public_scan_job(job_id: str):
    with PUBLIC_JOB_LOCK:
        job = PUBLIC_JOBS.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found or expired")
        if job.get("status") in {"complete", "incomplete", "error", "cancelled"}:
            return {"job_id": job_id, "status": job.get("status")}
        job["status"] = "cancelled"
        job["phase"] = "cancelled"
        process = PUBLIC_PROCESSES.get(job_id)
    if process and process.poll() is None:
        process.terminate()
    return {"job_id": job_id, "status": "cancelled"}


@app.get("/api/public/jobs/{job_id}/results")
def public_scan_job_results(job_id: str):
    with PUBLIC_JOB_LOCK:
        job = dict(PUBLIC_JOBS.get(job_id, {}))
    if not job:
        raise HTTPException(status_code=404, detail="Job not found or expired")
    summary_path = PUBLIC_JOB_ROOT / job_id / "output" / "scan-summary.json"
    if summary_path.exists():
        try:
            return {**job, "summary": json.loads(summary_path.read_text(encoding="utf-8"))}
        except (ValueError, OSError):
            pass
    return job


@app.get("/api/public/jobs/{job_id}/results/view", response_class=HTMLResponse)
def public_scan_job_results_view(request: Request, job_id: str):
    """Render ephemeral results in a read-only, portal-style view.

    Public scans intentionally bypass the authenticated service view: there is
    no service record, policy gating, exception state, or action endpoint to
    mutate.  The page therefore displays the raw vulnerability payload and
    configuration findings without Request Exception/POA&M/Mitigation buttons.
    """
    with PUBLIC_JOB_LOCK:
        job = dict(PUBLIC_JOBS.get(job_id, {}))
    if not job:
        raise HTTPException(status_code=404, detail="Job not found or expired")
    output_dir = PUBLIC_JOB_ROOT / job_id / "output"
    result_path = output_dir / "portal-result.json"
    if not result_path.exists():
        raise HTTPException(status_code=409, detail="Scan results are not available yet")
    try:
        payload = json.loads(result_path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=500, detail="Scan results are invalid") from exc
    # Public jobs can be interrupted between result fragments.  Treat a
    # malformed fragment as an empty section so the human-readable view stays
    # usable and never turns a partial scan into an unhandled 500.
    if not isinstance(payload, dict):
        payload = {}
    service = payload.get("service") if isinstance(payload.get("service"), dict) else {}
    overview_data = {}
    overview_path = output_dir / "service-overview.json"
    if overview_path.exists():
        try:
            candidate = json.loads(overview_path.read_text(encoding="utf-8"))
            if isinstance(candidate, dict):
                overview_data = candidate
        except (OSError, ValueError):
            overview_data = {}
    vulnerabilities = [item for item in (payload.get("findings") or []) if isinstance(item, dict)]
    configurations = [item for item in (payload.get("policy_findings") or []) if isinstance(item, dict)]
    skipped_images = payload.get("skipped_images") or []
    skipped_charts = payload.get("skipped_charts") or []
    if not isinstance(skipped_images, list):
        skipped_images = [skipped_images]
    if not isinstance(skipped_charts, list):
        skipped_charts = [skipped_charts]
    overview_data.setdefault("source", "Helm rendered manifests" if overview_data else "Submitted scan evidence")
    overview_data.setdefault("description", service.get("description"))
    overview_data = normalize_overview(
        overview_data,
        skipped_images=skipped_images,
        skipped_charts=skipped_charts,
        findings_images=[{"image": item.get("image"), "digest": item.get("image_digest"), "discovered_from": item.get("discovered_from") or "Submitted"} for item in vulnerabilities if item.get("image")],
        submitted_charts=job.get("chart_names") or [],
        incomplete=job.get("status") == "incomplete",
        # Rendering must remain read-only and fast.  Digest enrichment is
        # performed by scan processing and persisted in service-overview.json;
        # never invoke Docker/registry lookups from this request path.
        digest_resolver=None,
    )
    rows = []
    for item in vulnerabilities:
        rows.append({
            "type": "Vulnerability", "finding": item.get("cve") or item.get("finding") or "Unknown",
            "title": "", "severity": item.get("severity") or "Unknown", "state": "Raw",
            "scanner": item.get("scanner") or "Grype",
            "image": item.get("image") or "—", "target": item.get("package") or "—",
            "details": (item.get("evidence") or {}).get("description") or "",
            "remediation": item.get("fixed_version") or "—",
            "kev": "Yes" if item.get("kev") else "No", "epss": item.get("epss"),
        })
    for item in configurations:
        rows.append({
            "type": "Configuration", "finding": item.get("finding") or "Unknown",
            "title": item.get("title") or "", "severity": item.get("severity") or "Unknown", "state": "Raw",
            "scanner": item.get("scanner") or ("Dockle" if str(item.get("framework") or "").lower().startswith("docker") else "Trivy"),
            "image": item.get("target") or "—", "target": item.get("framework") or "—",
            "details": item.get("description") or "", "remediation": item.get("remediation") or "—",
            "kev": "—", "epss": "—",
        })
    grouped = {}
    simplified_rows = []
    for index, item in enumerate(rows):
        key = (item["type"], item["image"], item["severity"], item["remediation"], item["scanner"])
        if key not in grouped:
            grouped[key] = {**item, "finding_ids": [item["finding"]], "finding_indices": [index], "count": 1}
            simplified_rows.append(grouped[key])
        else:
            # A CVE can occur in several packages/images. Show it once in the
            # simplified row while retaining the first detail record for the
            # clickable finding link.
            if item["finding"] not in grouped[key]["finding_ids"]:
                grouped[key]["finding_ids"].append(item["finding"])
                grouped[key]["finding_indices"].append(index)
                grouped[key]["count"] += 1
    return templates.TemplateResponse(request, "public_results.html", {
        "current_user": None, "csrf_token": "", "can": lambda *_permission: False,
        "themes": THEMES, "pending_request_count": 0, "pending_poam_count": 0,
        "job": job, "service": service, "view": {"service": {"id": None, "service_key": ""}},
        "service_images": [], "rows": rows, "simplified_rows": simplified_rows,
        "vulnerability_count": len(vulnerabilities), "configuration_count": len(configurations),
        "overview_data": overview_data, "skipped_images": skipped_images, "skipped_charts": skipped_charts,
    })


@app.get("/api/public/jobs/{job_id}/logs", response_class=PlainTextResponse)
def public_scan_job_logs(job_id: str):
    with PUBLIC_JOB_LOCK:
        if job_id not in PUBLIC_JOBS:
            raise HTTPException(status_code=404, detail="Job not found or expired")
    log_path = PUBLIC_JOB_ROOT / job_id / "output" / "worker.log"
    if not log_path.exists():
        return "The worker has not produced a log yet."
    return log_path.read_text(encoding="utf-8", errors="replace")


def _write_public_scan_workbook(output_dir: Path, workbook_path: Path) -> None:
    """Create an Excel export for an ephemeral scan using portal-style sheets."""
    result_path = output_dir / "portal-result.json"
    payload = {}
    if result_path.exists():
        try:
            payload = json.loads(result_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            payload = {}
    findings = [item for item in (payload.get("findings") or []) if isinstance(item, dict)]
    policy_findings = [item for item in (payload.get("policy_findings") or []) if isinstance(item, dict)]
    service = payload.get("service") if isinstance(payload.get("service"), dict) else {}
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary_rows = [
        ("Service ID", service.get("id") or "Standalone scan"),
        ("Service", service.get("name") or "Standalone scan"),
        ("Version", service.get("version") or "Not provided"),
        ("Owner", service.get("owner") or "Not provided"),
        ("Active Findings", len(findings) + len(policy_findings)),
        ("Vulnerability Findings", len(findings)),
        ("Configuration Findings", len(policy_findings)),
        ("Evidence", "Complete" if payload.get("complete") else "Incomplete"),
        ("Exported At", excel_datetime(utcnow())),
    ]
    for row in summary_rows:
        summary.append(list(row))
    format_sheet(summary)

    vuln = workbook.create_sheet("Findings")
    vuln.append(["Finding", "Severity", "Image", "Package", "Installed Version", "Fixed Version", "Description", "References"])
    for item in findings:
        evidence = item.get("evidence") if isinstance(item.get("evidence"), dict) else {}
        vuln.append([
            item.get("cve") or item.get("finding"), item.get("severity"), item.get("image"),
            item.get("package"), item.get("installed_version"), item.get("fixed_version"),
            evidence.get("description"), "\n".join(evidence.get("urls") or []),
        ])
    format_sheet(vuln)

    config = workbook.create_sheet("Configuration Findings")
    config.append(["Finding", "Title", "Severity", "Framework", "Target", "Namespace", "Description", "Remediation"])
    for item in policy_findings:
        config.append([
            item.get("finding"), item.get("title"), item.get("severity"), item.get("framework"),
            item.get("target"), item.get("namespace"), item.get("description"), item.get("remediation"),
        ])
    format_sheet(config)
    workbook.save(workbook_path)


def _write_public_scan_html(output_dir: Path, report_path: Path, job: dict) -> None:
    payload = {}
    result_path = output_dir / "portal-result.json"
    if result_path.exists():
        try:
            candidate = json.loads(result_path.read_text(encoding="utf-8"))
            if isinstance(candidate, dict):
                payload = candidate
        except (OSError, ValueError):
            payload = {}
    output_root = output_dir.resolve()
    report_target = report_path.resolve()
    artifacts = []
    for path in sorted(output_dir.rglob("*")):
        try:
            resolved = path.resolve()
            if not path.is_file() or resolved == report_target or not resolved.is_relative_to(output_root):
                continue
            artifacts.append({
                "path": resolved.relative_to(output_root).as_posix(),
                "size": resolved.stat().st_size,
            })
        except OSError:
            continue
    report_path.write_text(build_public_scan_report(payload, job, artifacts), encoding="utf-8")


@app.get("/api/public/jobs/{job_id}/overview.html", response_class=HTMLResponse)
def public_scan_job_html_overview(job_id: str):
    with PUBLIC_JOB_LOCK:
        job = dict(PUBLIC_JOBS.get(job_id, {}))
    if not job:
        raise HTTPException(status_code=404, detail="Job not found or expired")
    output_dir = PUBLIC_JOB_ROOT / job_id / "output"
    if not (output_dir / "portal-result.json").exists():
        raise HTTPException(status_code=409, detail="Scan results are not available yet")
    _write_public_scan_workbook(output_dir, output_dir / "scan-results.xlsx")
    report_path = output_dir / "scan-overview.html"
    _write_public_scan_html(output_dir, report_path, job)
    return HTMLResponse(report_path.read_text(encoding="utf-8"))


@app.get("/api/public/jobs/{job_id}/artifacts")
def public_scan_job_artifacts(job_id: str):
    with PUBLIC_JOB_LOCK:
        job = dict(PUBLIC_JOBS.get(job_id, {}))
        if not job:
            raise HTTPException(status_code=404, detail="Job not found or expired")
    output_dir = PUBLIC_JOB_ROOT / job_id / "output"
    if not output_dir.exists():
        raise HTTPException(status_code=409, detail="Job has not produced results yet")
    # Keep the self-service download useful even without a persistent service:
    # materialize the same workbook-style export used by the authenticated UI.
    _write_public_scan_workbook(output_dir, output_dir / "scan-results.xlsx")
    _write_public_scan_html(output_dir, output_dir / "scan-overview.html", job)
    archive = PUBLIC_JOB_ROOT / f"{job_id}.zip"
    with ZipFile(archive, "w", ZIP_DEFLATED) as bundle:
        for path in output_dir.rglob("*"):
            if path.is_file():
                bundle.write(path, path.relative_to(output_dir))
    return StreamingResponse(open(archive, "rb"), media_type="application/zip", headers={"Content-Disposition": f"attachment; filename={job_id}-results.zip"})


@app.get("/api/public/jobs/{job_id}/sboms")
def public_sbom_job_artifacts(job_id: str):
    """Download only the SBOM formats recorded by the serializer manifest."""
    with PUBLIC_JOB_LOCK:
        if job_id not in PUBLIC_JOBS:
            raise HTTPException(status_code=404, detail="Job not found or expired")
    output_dir = PUBLIC_JOB_ROOT / job_id / "output"
    manifest_path = output_dir / "sboms" / "formats" / "manifest.json"
    if not manifest_path.exists():
        raise HTTPException(status_code=409, detail="SBOM documents are not available yet")
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=409, detail="SBOM manifest is invalid") from exc
    reports = manifest.get("reports") if isinstance(manifest, dict) else None
    if not isinstance(reports, list) or not reports:
        raise HTTPException(status_code=409, detail="No SBOM documents were generated")

    output_root = output_dir.resolve()
    archive = PUBLIC_JOB_ROOT / f"{job_id}-sboms.zip"
    written: set[str] = set()
    with ZipFile(archive, "w", ZIP_DEFLATED) as bundle:
        bundle.write(manifest_path, "manifest.json")
        for report in reports:
            if not isinstance(report, dict) or not report.get("path"):
                continue
            relative = Path(str(report["path"]).replace("\\", "/"))
            if relative.is_absolute() or ".." in relative.parts:
                continue
            candidate = (output_dir / relative).resolve()
            if not candidate.is_relative_to(output_root) or not candidate.is_file():
                continue
            archive_name = candidate.relative_to(output_root).as_posix()
            if archive_name not in written:
                bundle.write(candidate, archive_name)
                written.add(archive_name)
    if not written:
        archive.unlink(missing_ok=True)
        raise HTTPException(status_code=409, detail="Generated SBOM documents are missing")
    return FileResponse(archive, media_type="application/zip", filename=f"{job_id}-sboms.zip")


@app.get("/api/public/jobs/{job_id}/export.xlsx")
def public_scan_job_export(job_id: str):
    with PUBLIC_JOB_LOCK:
        if job_id not in PUBLIC_JOBS:
            raise HTTPException(status_code=404, detail="Job not found or expired")
    output_dir = PUBLIC_JOB_ROOT / job_id / "output"
    if not output_dir.exists():
        raise HTTPException(status_code=409, detail="Job has not produced results yet")
    workbook_path = output_dir / "scan-results.xlsx"
    _write_public_scan_workbook(output_dir, workbook_path)
    return FileResponse(
        workbook_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"{job_id}-scan-results.xlsx",
    )


@app.get("/api/public/jobs/{job_id}/results-export")
def public_scan_job_results_export(job_id: str):
    """Download the native scanner JSON bundle for manual ingestion."""
    with PUBLIC_JOB_LOCK:
        if job_id not in PUBLIC_JOBS:
            raise HTTPException(status_code=404, detail="Job not found or expired")
    archive = PUBLIC_JOB_ROOT / job_id / "output" / "results-export.tar.gz"
    if not archive.exists():
        raise HTTPException(status_code=409, detail="Scanner results export is not available yet")
    return FileResponse(archive, media_type="application/gzip", filename=f"{job_id}-results.tar.gz")


@app.post("/api/public/jobs/{job_id}/ingest")
def ingest_public_scan(
    job_id: str,
    service_id: str,
    db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_user),
):
    """Attach an ephemeral scan to an existing service after login.

    A service selection on the authenticated scan form invokes this endpoint
    automatically after the temporary scan completes. It still requires an
    authenticated service administrator and an existing in-scope service, so
    a public request cannot create or overwrite portal scope by itself.
    """
    if auth.accessible_service_ids("scan.ingest") == set():
        raise HTTPException(status_code=403, detail="Scan ingest permission is required")
    with PUBLIC_JOB_LOCK:
        if job_id not in PUBLIC_JOBS:
            raise HTTPException(status_code=404, detail="Job not found or expired")
    result_path = PUBLIC_JOB_ROOT / job_id / "output" / "portal-result.json"
    if not result_path.exists():
        raise HTTPException(status_code=409, detail="The scan has not produced a portal result")
    service = db.scalar(select(Service).where(Service.service_key == service_id))
    if not service:
        raise HTTPException(status_code=404, detail="Service not found")
    if not auth.has("scan.ingest", service.id):
        raise HTTPException(status_code=403, detail="The service is outside your scope")
    try:
        data = json.loads(result_path.read_text(encoding="utf-8"))
        # Public scans may include raw (non-fixable) vulnerabilities for
        # display, but the authenticated portal ingest contract stores only
        # actionable findings. Keep skipped evidence and policy findings
        # intact while normalizing the vulnerability list for ingest.
        data["findings"] = [
            item
            for item in (data.get("findings") or [])
            if isinstance(item, dict)
            and str(item.get("fixed_version") or "").strip() not in {"", "—", "-"}
        ]
        data["fixable_only"] = True
        data["execution_id"] = f"public:{job_id}"
        source_files = _collect_helm_source_files(PUBLIC_JOB_ROOT / job_id / "input" / "charts")
        if source_files:
            data["artifact_type"] = "helm"
            data["helm_source_files"] = source_files
            _enrich_values_source_mappings(data, source_files)
        with PUBLIC_JOB_LOCK:
            job_input_images = [line.strip() for line in str(PUBLIC_JOBS.get(job_id, {}).get("image_list") or "").splitlines() if line.strip()]
        candidate_images = {
            str(item.get("image") or "").strip()
            for item in (data.get("findings") or [])
            if isinstance(item, dict) and str(item.get("image") or "").strip()
        }
        # A one-image public scan is authoritative only for that image. A
        # multi-image result remains a service-scope execution unless its
        # producer explicitly supplied a scope.
        if data.get("scan_scope") not in {"service", "image", "evidence"}:
            data["scan_scope"] = "service"
        if not candidate_images and len(job_input_images) == 1:
            candidate_images = set(job_input_images)
        if data.get("scan_scope") == "service" and len(candidate_images) == 1 and not data.get("skipped_images") and not data.get("skipped_charts"):
            data["scan_scope"] = "image"
            data["scope_image"] = next(iter(candidate_images))
        data["service"] = {
            **data.get("service", {}),
            "id": service.service_key,
            "name": service.name,
            "version": service.manual_version or data.get("service", {}).get("version") or "Unknown",
            "owner": service.owner,
            "poc": service.poc,
            "groups": [group.name for group in service.groups],
        }
        payload = ExecutionPayload.model_validate(data)
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=f"Invalid scan result: {exc}") from exc
    return ingest(payload, db)


def _enforce_required_signature(current: dict, values: dict) -> None:
    if current.get("signing_enabled") and values.get("status") == "complete":
        result = values.get("result") if "result" in values else current.get("result")
        if result is None and "result" not in values:
            # Completion and the result can arrive in separate polls.
            values.update(status="running", phase="signing_image")
        else:
            result = result if isinstance(result, dict) else {}
            signature = result.get("signature")
            signature = signature if isinstance(signature, dict) else {}
            verified = result.get("signature_status") == "verified"
            verified = verified and signature.get("image") == result.get("immutable_destination") and bool(signature.get("image"))
            verified = verified and signature.get("key_fingerprint") == current.get("signing_fingerprint")
            if not verified:
                message = "Required signature verification was not reported by the patch worker. Upgrade both portal and worker and check signing configuration."
                values.update(status="failed", phase="signing_image", error=message,
                              failed_stage="signing_image",
                              stages=advance_patch_stages(values.get("stages") or current.get("stages"), "signing_image", "failed", "push"),
                              result={**result, "status": "failed", "delivery_status": "failed", "delivery_error": message, "signature_status": "failed"})


def _patch_job_update(job_id: str, **values):
    snapshot = None
    with PATCH_JOB_LOCK:
        if job_id in PATCH_JOBS:
            current = PATCH_JOBS[job_id]
            _enforce_required_signature(current, values)
            phase = values.get("phase") or current.get("phase")
            status = values.get("status") or current.get("status")
            if phase in PATCH_PHASES and status in {"running", "complete", "failed"}:
                values["stages"] = advance_patch_stages(
                    values.get("stages") or current.get("stages"), phase, status,
                    values.get("output_mode") or current.get("output_mode") or "download",
                )
                if status == "failed":
                    values.setdefault("failed_stage", phase)
            PATCH_JOBS[job_id].update(values)
            snapshot = dict(PATCH_JOBS[job_id])
    if snapshot:
        job_root = PATCH_JOB_ROOT / job_id
        job_root.mkdir(parents=True, exist_ok=True)
        temp = job_root / "public-job.json.tmp"
        temp.write_text(json.dumps(snapshot, indent=2), encoding="utf-8")
        temp.replace(job_root / "public-job.json")
    if snapshot and snapshot.get("patch_execution_id"):
        with SessionLocal() as db:
            record = db.get(PatchExecution, int(snapshot["patch_execution_id"]))
            if record:
                record.status = str(snapshot.get("status") or record.status)
                record.phase = str(snapshot.get("phase") or record.phase)
                record.error = str(snapshot.get("error"))[:4000] if snapshot.get("error") else None
                result = snapshot.get("result") or {}
                summary = result if isinstance(result, dict) else {}
                record.summary = {
                    key: value for key, value in summary.items()
                    if key in {
                        "vulnerabilities_before", "vulnerabilities_after",
                        "vulnerabilities_removed", "vulnerabilities_remaining",
                        "could_not_be_patched", "patched_image", "destination_image",
                        "patch_status", "reason", "image_changed", "delivery_status", "delivery_error",
                        "immutable_destination", "signature_status", "signature", "artifact_sha256",
                        "stages", "failed_stage",
                    }
                }
                if snapshot.get("stages"):
                    record.summary["stages"] = snapshot["stages"]
                if snapshot.get("failed_stage"):
                    record.summary["failed_stage"] = snapshot["failed_stage"]
                if result and record.status == "complete" and str(result.get("patch_status") or "") in {"PATCHED", "PARTIALLY_PATCHED", "NO_APPLICABLE_FIXES"}:
                    associated_ref = str(result.get("destination_image") or result.get("patched_image") or "").strip()
                    service = db.get(Service, record.service_id)
                    if service and associated_ref:
                        _ensure_service_image(db, service, associated_ref)
                        existing_audit = db.scalar(select(func.count(AuditEvent.id)).where(
                            AuditEvent.action == "patch.completed", AuditEvent.target_id == str(record.id)
                        ))
                        if not existing_audit:
                            db.add(AuditEvent(actor_user_id=record.requested_by_id, action="patch.completed", target_type="patch_execution", target_id=str(record.id), detail={"service_id": record.service_id, "image": associated_ref, "patch_status": result.get("patch_status")}))
                audit_action = None
                if record.status == "failed":
                    audit_action = "patch.failed"
                elif result and record.status == "complete":
                    patch_status = str(result.get("patch_status") or "")
                    audit_action = {
                        "FAILED": "patch.failed",
                        "UNSUPPORTED": "patch.unsupported",
                    }.get(patch_status)
                if audit_action:
                    existing_audit = db.scalar(select(func.count(AuditEvent.id)).where(
                        AuditEvent.action == audit_action, AuditEvent.target_id == str(record.id)
                    ))
                    if not existing_audit:
                        db.add(AuditEvent(
                            actor_user_id=record.requested_by_id, action=audit_action,
                            target_type="patch_execution", target_id=str(record.id),
                            detail={"service_id": record.service_id, "reason": (result or {}).get("reason") or record.error},
                        ))
                if result and record.status == "complete" and str(result.get("output_mode") or "") == "push":
                    delivery_status = str(result.get("delivery_status") or "")
                    delivery_action = {
                        "delivered": "oci.publication.succeeded",
                        "failed": "oci.publication.failed",
                    }.get(delivery_status)
                    if delivery_action:
                        existing_audit = db.scalar(select(func.count(AuditEvent.id)).where(
                            AuditEvent.action == delivery_action, AuditEvent.target_id == str(record.id)
                        ))
                        if not existing_audit:
                            db.add(AuditEvent(
                                actor_user_id=record.requested_by_id, action=delivery_action,
                                target_type="patch_execution", target_id=str(record.id),
                                detail={"service_id": record.service_id, "destination_image": result.get("destination_image"),
                                        "error": result.get("delivery_error")},
                            ))
                record.updated_at = utcnow()
                if record.status == "complete":
                    record.completed_at = utcnow()
                db.commit()


    if snapshot and (snapshot.get("status") in {"failed", "cancelled"} or (snapshot.get("status") == "complete" and snapshot.get("result"))):
        with SessionLocal() as db:
            requested = db.scalar(select(AuditEvent).where(AuditEvent.action == "signing.requested", AuditEvent.target_type == "patch_job", AuditEvent.target_id == job_id))
            if requested:
                result = snapshot.get("result") or {}
                action = "signing.verified" if result.get("signature_status") == "verified" else "signing.failed"
                exists = db.scalar(select(AuditEvent.id).where(AuditEvent.action.in_(["signing.verified", "signing.failed"]), AuditEvent.target_type == "patch_job", AuditEvent.target_id == job_id))
                if not exists:
                    db.add(AuditEvent(actor_user_id=requested.actor_user_id, action=action, target_type="patch_job", target_id=job_id,
                                      detail={"image": result.get("immutable_destination"), "signature": result.get("signature"),
                                              "key_fingerprint": requested.detail.get("key_fingerprint"), "job_status": snapshot.get("status")}))
                    db.commit()


def _portal_signing_material(db: Session, auth: AuthContext | None, output_mode: str, service: Service | None = None):
    settings = get_global_configuration(db)
    try:
        if output_mode == "push" and signing.configuration(settings).get("enabled") is True:
            if not auth or not auth.user.enabled or not auth.has("artifact.sign", service.id if service else None):
                raise HTTPException(status_code=403, detail="Sign in with image-signing permission to publish using the configured key")
        return signing.job_material(settings, output_mode)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


def _audit_signing_request(db: Session, job_id: str, actor_id: int, config: dict):
    if config.get("signing_enabled"):
        db.add(AuditEvent(actor_user_id=actor_id, action="signing.requested", target_type="patch_job", target_id=job_id,
                          detail={"destination": config.get("destination_image"), "key_fingerprint": config["signing_fingerprint"]}))
        db.commit()


def _load_patch_job(job_id: str) -> dict:
    """Load a nonsensitive patch snapshot from memory and durable job files."""
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,64}", str(job_id or "")):
        return {}
    job_root = PATCH_JOB_ROOT / job_id
    snapshot = {}
    with PATCH_JOB_LOCK:
        snapshot.update(PATCH_JOBS.get(job_id, {}))
    for path in (job_root / "public-job.json", job_root / "output" / "patch-state.json"):
        if path.exists():
            try:
                value = json.loads(path.read_text(encoding="utf-8"))
                if isinstance(value, dict):
                    snapshot.update(value)
            except (OSError, ValueError):
                pass
    result_path = job_root / "output" / "patch-result.json"
    if result_path.exists():
        try:
            result = json.loads(result_path.read_text(encoding="utf-8"))
            if isinstance(result, dict):
                snapshot["result"] = result
        except (OSError, ValueError):
            pass
    if snapshot:
        _enforce_required_signature(snapshot, snapshot)
        snapshot.setdefault("job_id", job_id)
        with PATCH_JOB_LOCK:
            PATCH_JOBS[job_id] = dict(snapshot)
    return snapshot


def _patch_worker_values(state: dict) -> tuple[dict, object | None]:
    """Normalize worker responses before forwarding them to the job store.

    The worker status endpoint includes its path identifier as ``job_id``.
    That identifier is already supplied positionally to ``_patch_job_update``;
    forwarding it again via ``**state`` raises Python's duplicate-argument
    error and leaves the UI stuck in a failed job state.
    """
    values = dict(state or {})
    values.pop("job_id", None)
    result = values.pop("result", None)
    return values, result


def _patch_worker_request(path: str, *, method: str = "GET", payload: dict | None = None) -> dict:
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    headers = {"Accept": "application/json"}
    if body is not None:
        headers["Content-Type"] = "application/json"
    if PATCH_WORKER_TOKEN:
        headers["X-CATS-Worker-Token"] = PATCH_WORKER_TOKEN
    request = urllib.request.Request(f"{PATCH_WORKER_URL}{path}", data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=15) as response:
            return json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        raise RuntimeError("Internal patch worker is unavailable") from exc


def _patch_page_context(
    request: Request, auth: AuthContext | None, db: Session,
    job_id: str = "", error: str = "", selected_service_id: str = "",
):
    job = _load_patch_job(job_id) if job_id else {}
    services = []
    configuration = get_global_configuration(db)
    if auth and auth.accessible_service_ids("scan.ingest") != set():
        scoped = auth.accessible_service_ids("scan.ingest")
        services = list(db.scalars(select(Service).order_by(Service.name)))
        if scoped is not None:
            services = [service for service in services if service.id in scoped]
    return templates.TemplateResponse(request, "patch.html", {
        "current_user": auth.user if auth else None,
        "csrf_token": auth.csrf_token if auth else "",
        "can": auth.has if auth else (lambda _permission: False),
        "themes": THEMES,
        "pending_request_count": 0,
        "pending_poam_count": 0,
        "job_id": job_id,
        "job": job,
        "error": error,
        "patch_phases": PATCH_PHASES,
        "authenticated_services": services,
        "selected_service_id": selected_service_id,
        "configured_registries": configured_registries(configuration),
        "signing": signing.public_metadata(configuration),
    })


def _run_patch_job(job_id: str, credential_env: dict[str, str]) -> None:
    job_root = PATCH_JOB_ROOT / job_id
    output_dir = job_root / "output"
    state_path = output_dir / "patch-state.json"
    result_path = output_dir / "patch-result.json"
    process: subprocess.Popen | None = None
    timeout = int(os.getenv("CATS_PATCH_JOB_TIMEOUT", "3600"))
    started = time.monotonic()
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
        if PATCH_WORKER_URL:
            config = json.loads((job_root / "job-config.json").read_text(encoding="utf-8"))
            _patch_worker_request(
                f"/internal/patch-jobs/{job_id}", method="POST",
                payload={"config": config, "credentials": credential_env},
            )
            credential_env.clear()
            while True:
                with PATCH_JOB_LOCK:
                    cancelled = PATCH_JOBS.get(job_id, {}).get("status") == "cancelled"
                if cancelled:
                    _patch_worker_request(f"/internal/patch-jobs/{job_id}/cancel", method="POST")
                    return
                if time.monotonic() - started > timeout:
                    _patch_worker_request(f"/internal/patch-jobs/{job_id}/cancel", method="POST")
                    raise TimeoutError("Patch job exceeded its configured time limit")
                state = _patch_worker_request(f"/internal/patch-jobs/{job_id}")
                state, result = _patch_worker_values(state)
                _patch_job_update(job_id, **state, **({"result": result} if isinstance(result, dict) else {}))
                if state.get("status") == "complete":
                    _patch_job_update(job_id, status="complete", phase="completed", result=result or {}, error=None)
                    return
                if state.get("status") in {"failed", "cancelled"}:
                    return
                time.sleep(0.5)

        env = dict(os.environ)
        env.update(credential_env)
        _patch_job_update(job_id, status="running", phase="queued")
        process = subprocess.Popen(
            [sys.executable, "-m", "app.patch_worker", str(job_root / "job-config.json"), str(output_dir)],
            env=env, cwd=str(root.parent), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        with PATCH_JOB_LOCK:
            PATCH_PROCESSES[job_id] = process
        last_state = None
        while process.poll() is None:
            with PATCH_JOB_LOCK:
                cancelled = PATCH_JOBS.get(job_id, {}).get("status") == "cancelled"
            if cancelled:
                process.terminate()
                try:
                    process.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    process.kill()
                return
            if time.monotonic() - started > timeout:
                process.kill()
                raise TimeoutError("Patch job exceeded its configured time limit")
            if state_path.exists():
                try:
                    state = json.loads(state_path.read_text(encoding="utf-8"))
                    if state != last_state:
                        last_state = state
                        _patch_job_update(job_id, **state)
                except (OSError, ValueError):
                    pass
            time.sleep(0.5)
        if state_path.exists():
            state = json.loads(state_path.read_text(encoding="utf-8"))
            _patch_job_update(job_id, **state)
        if process.returncode == 0 and result_path.exists():
            result = json.loads(result_path.read_text(encoding="utf-8"))
            _patch_job_update(job_id, status="complete", phase="completed", result=result, error=None)
        else:
            with PATCH_JOB_LOCK:
                cancelled = PATCH_JOBS.get(job_id, {}).get("status") == "cancelled"
            if cancelled:
                return
            state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
            if state.get("status") in {"failed", "cancelled"}:
                if result_path.exists():
                    _patch_job_update(job_id, result=json.loads(result_path.read_text(encoding="utf-8")))
                _patch_job_update(job_id, **state)
                return
            raise RuntimeError(state.get("error") or "Patch worker failed")
    except Exception as exc:
        with PATCH_JOB_LOCK:
            existing = dict(PATCH_JOBS.get(job_id, {}))
        _patch_job_update(
            job_id, status="failed", phase=existing.get("phase") or "queued",
            stages=existing.get("stages"), failed_stage=existing.get("phase") or "queued",
            error=redact(exc),
        )
    finally:
        credential_env.clear()
        shutil.rmtree(job_root / "input", ignore_errors=True)
        (job_root / "job-config.json").unlink(missing_ok=True)
        with PATCH_JOB_LOCK:
            PATCH_PROCESSES.pop(job_id, None)


@app.get("/patch", response_class=HTMLResponse)
def public_patch(
    request: Request, job_id: str = "", db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(optional_user),
):
    return _patch_page_context(request, auth, db, job_id)


@app.post("/patch")
def public_patch_submit(
    request: Request,
    csrf_token: str = Form(""),
    source_mode: str = Form("oci"),
    source_image: str = Form(""),
    source_registry_id: str = Form(""),
    source_username: str = Form(""),
    source_password: str = Form(""),
    source_archive: UploadFile | None = File(None),
    output_mode: str = Form("download"),
    destination_image: str = Form(""),
    destination_registry_id: str = Form(""),
    destination_username: str = Form(""),
    destination_password: str = Form(""),
    reuse_source_credentials: bool = Form(False),
    service_id: str = Form(""),
    db: Session = Depends(get_db),
    auth: AuthContext | None = Depends(optional_user),
):
    source_mode = source_mode.strip().lower()
    output_mode = output_mode.strip().lower()
    # Registry credentials are administrator-managed. Ignore legacy form
    # fields so a Patch workspace submission can never override them.
    source_username = source_password = ""
    destination_username = destination_password = ""
    source_registry_id = source_registry_id.strip()
    destination_registry_id = destination_registry_id.strip()
    registry_rows = parse_json(get_global_configuration(db).get("oci_registries"), [])
    registry_map = {str(row.get("id")): row for row in registry_rows if isinstance(row, dict) and row.get("id")}
    selected_source_registry = registry_map.get(source_registry_id) if source_registry_id else None
    selected_destination_registry = registry_map.get(destination_registry_id) if destination_registry_id else None
    def resolve_registry_image(image: str, registry: dict | None) -> str:
        image = image.strip()
        if not registry or not image:
            return image
        endpoint = str(registry.get("endpoint") or "").strip().rstrip("/")
        host = urllib.parse.urlparse(endpoint if "://" in endpoint else f"https://{endpoint}").netloc
        prefix = str(registry.get("namespace") or "").strip().strip("/")
        first = image.split("/", 1)[0]
        already_qualified = "." in first or ":" in first or first == "localhost"
        return image if already_qualified else "/".join(part for part in (host, prefix, image.lstrip("/")) if part)
    if source_registry_id and not selected_source_registry:
        return _patch_page_context(request, auth, db, error="Selected source registry is not configured.", selected_service_id=service_id)
    if destination_registry_id and not selected_destination_registry:
        return _patch_page_context(request, auth, db, error="Selected destination registry is not configured.", selected_service_id=service_id)
    if source_mode == "oci" and not selected_source_registry:
        selected_source_registry = _configured_registry_for_image(source_image, list(registry_map.values()))
    def apply_registry_credentials(registry: dict | None, username: str, password: str) -> tuple[str, str]:
        if not registry:
            return username, password
        username = username or str(registry.get("username") or "")
        if not password and registry.get("password"):
            try:
                password = decrypt_secret(str(registry.get("password")))
            except Exception as exc:
                raise ValueError("Configured registry secret could not be decrypted.") from exc
        return username, password
    try:
        if source_mode == "oci":
            source_username, source_password = apply_registry_credentials(selected_source_registry, source_username, source_password)
        if output_mode == "push":
            destination_username, destination_password = apply_registry_credentials(selected_destination_registry, destination_username, destination_password)
    except ValueError as exc:
        return _patch_page_context(request, auth, db, error=str(exc), selected_service_id=service_id)
    if source_mode == "oci" and selected_source_registry:
        source_image = resolve_registry_image(source_image, selected_source_registry)
    if output_mode == "push" and selected_destination_registry:
        destination_image = resolve_registry_image(destination_image, selected_destination_registry)
    if source_mode not in {"oci", "upload"} or output_mode not in {"download", "push"}:
        return _patch_page_context(request, auth, db, error="Choose a valid image source and output mode.", selected_service_id=service_id)
    if source_mode == "oci" and not source_image.strip():
        return _patch_page_context(request, auth, db, error="Image URI is required for an OCI registry source.", selected_service_id=service_id)
    if source_mode == "upload" and (not source_archive or not source_archive.filename):
        return _patch_page_context(request, auth, db, error="Choose a Docker or OCI image .tar archive.", selected_service_id=service_id)
    if output_mode == "push" and not destination_image.strip():
        return _patch_page_context(request, auth, db, error="Destination image / repository / tag is required for OCI push.", selected_service_id=service_id)
    if output_mode == "push" and not selected_destination_registry:
        return _patch_page_context(request, auth, db, error="Failed to upload, OCI registry not configured. Please contact Cybersecurity.", selected_service_id=service_id)
    if bool(source_username) != bool(source_password):
        return _patch_page_context(request, auth, db, error="Provide both source registry username and password/token.", selected_service_id=service_id)
    if bool(destination_username) != bool(destination_password):
        return _patch_page_context(request, auth, db, error="Provide both destination registry username and password/token.", selected_service_id=service_id)

    selected_service = None
    if service_id:
        if not auth or auth.accessible_service_ids("scan.ingest") == set():
            raise HTTPException(status_code=403, detail="Sign in with scan-ingest permission to retain patch history")
        selected_service = db.scalar(select(Service).where(Service.service_key == service_id))
        if not selected_service:
            raise HTTPException(status_code=404, detail="Selected service was not found")
        if not auth.has("scan.ingest", selected_service.id):
            raise HTTPException(status_code=403, detail="Selected service is outside your scope")

    signing_config, signing_credentials = _portal_signing_material(db, auth, output_mode, selected_service)
    if signing_config:
        check_csrf(auth, csrf_token)
    job_id = uuid.uuid4().hex
    job_root = PATCH_JOB_ROOT / job_id
    input_dir = job_root / "input"
    input_dir.mkdir(parents=True, exist_ok=False)
    policy_configuration = configuration_for_service(db, selected_service) if selected_service else get_configuration(db)
    config = {
        "job_id": job_id,
        **signing_config,
        "source_mode": source_mode,
        "source_image": source_image.strip() if source_mode == "oci" else None,
        "output_mode": output_mode,
        "destination_image": destination_image.strip() if output_mode == "push" else None,
        "reuse_source_credentials": bool(reuse_source_credentials),
        # Non-secret administrative policy is copied into the isolated worker
        # job. Credentials continue to arrive only through its environment.
        "trusted_ca_certificates": parse_json(policy_configuration.get("trusted_ca_certificates"), []),
        "repository_policies": parse_json(policy_configuration.get("repository_policies"), {}),
        "os_definitions": parse_json(policy_configuration.get("os_definitions"), {}),
    }
    if source_mode == "upload":
        max_bytes = int(os.getenv("CATS_PATCH_MAX_UPLOAD_BYTES", str(4 * 1024 * 1024 * 1024)))
        archive_path = input_dir / "source-image.tar"
        total = 0
        with archive_path.open("wb") as target:
            while chunk := source_archive.file.read(1024 * 1024):
                total += len(chunk)
                if total > max_bytes:
                    shutil.rmtree(job_root, ignore_errors=True)
                    return _patch_page_context(request, auth, db, error="Uploaded image archive is too large.", selected_service_id=service_id)
                target.write(chunk)
        config["archive_path"] = str(archive_path)
    (job_root / "job-config.json").write_text(json.dumps(safe_job_config(config), indent=2), encoding="utf-8")
    public = {
        "job_id": job_id,
        "signing_enabled": bool(signing_config),
        "signing_fingerprint": signing_config.get("signing_fingerprint"),
        "status": "queued",
        "phase": "queued",
        "stages": initial_patch_stages(output_mode),
        "source_mode": source_mode,
        "source_image": source_image.strip() if source_mode == "oci" else "Uploaded image archive",
        "output_mode": output_mode,
        "destination_image": destination_image.strip() if output_mode == "push" else None,
        "created_at": utcnow().isoformat(),
    }
    if selected_service:
        record = PatchExecution(
            job_key=job_id, service_id=selected_service.id, requested_by_id=auth.user.id,
            source_mode=source_mode,
            source_image=source_image.strip() if source_mode == "oci" else None,
            output_mode=output_mode,
            destination_image=destination_image.strip() if output_mode == "push" else None,
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        db.add(AuditEvent(actor_user_id=auth.user.id, action="patch.started", target_type="patch_execution", target_id=str(record.id), detail={"service_id": selected_service.id, "source_image": source_image.strip() if source_mode == "oci" else None, "output_mode": output_mode}))
        db.commit()
        public["patch_execution_id"] = record.id
        public["service_id"] = selected_service.service_key
    with PATCH_JOB_LOCK:
        PATCH_JOBS[job_id] = public
    _patch_job_update(job_id)
    credentials = {
        **signing_credentials,
        "CATS_PATCH_SOURCE_USERNAME": source_username,
        "CATS_PATCH_SOURCE_PASSWORD": source_password,
        "CATS_PATCH_DEST_USERNAME": destination_username,
        "CATS_PATCH_DEST_PASSWORD": destination_password,
    }
    _audit_signing_request(db, job_id, auth.user.id if auth else None, config)
    PUBLIC_WORKERS.submit(_run_patch_job, job_id, credentials)
    return RedirectResponse(f"/patch?job_id={job_id}", status_code=303)


@app.get("/api/public/patch-jobs/{job_id}")
def public_patch_job(job_id: str):
    job = _load_patch_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Patch job not found")
    return job


@app.post("/api/public/patch-jobs/{job_id}/cancel")
def cancel_public_patch_job(job_id: str):
    _load_patch_job(job_id)
    with PATCH_JOB_LOCK:
        job = PATCH_JOBS.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Patch job not found")
        if job.get("status") not in {"complete", "failed", "cancelled"}:
            job.update(status="cancelled", phase="cancelled", error=None)
    _patch_job_update(job_id)
    if PATCH_WORKER_URL:
        try:
            _patch_worker_request(f"/internal/patch-jobs/{job_id}/cancel", method="POST")
        except RuntimeError:
            pass
    with PATCH_JOB_LOCK:
        return {"job_id": job_id, "status": PATCH_JOBS[job_id]["status"]}


@app.get("/api/public/patch-jobs/{job_id}/logs", response_class=PlainTextResponse)
def public_patch_logs(job_id: str):
    if not _load_patch_job(job_id):
        raise HTTPException(status_code=404, detail="Patch job not found")
    path = PATCH_JOB_ROOT / job_id / "output" / "patch.log"
    return path.read_text(encoding="utf-8") if path.exists() else "Patch worker has not produced logs yet.\n"


@app.get("/api/public/patch-jobs/{job_id}/results", response_class=HTMLResponse)
def public_patch_results(request: Request, job_id: str, auth: AuthContext | None = Depends(optional_user)):
    job = _load_patch_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Patch job not found")
    if job.get("status") not in {"complete", "failed"} or not isinstance(job.get("result"), dict):
        raise HTTPException(status_code=409, detail="Patch job is not complete")
    return templates.TemplateResponse(request, "patch_results.html", {
        "current_user": auth.user if auth else None, "csrf_token": auth.csrf_token if auth else "",
        "can": auth.has if auth else (lambda _permission: False), "themes": THEMES,
        "pending_request_count": 0, "pending_poam_count": 0, "job": job, "result": job["result"],
    })


@app.get("/api/public/patch-jobs/{job_id}/patched-image.tar")
def download_patched_image(job_id: str):
    job = _load_patch_job(job_id)
    path = PATCH_JOB_ROOT / job_id / "output" / "patched-image.tar"
    result = job.get("result") if isinstance(job, dict) else {}
    if not job or job.get("status") not in {"complete", "failed"} or not path.exists() or not (result or {}).get("artifact_available"):
        raise HTTPException(status_code=404, detail="Patched image archive is not available")
    return FileResponse(path, media_type="application/x-tar", filename=f"cats-patched-{job_id[:12]}.tar")


def _validate_materialized_candidate(candidate_dir: Path, payload: dict, plan: dict, validation: dict) -> dict:
    """Run available Level 1 tools and keep missing tools explicit."""
    checks = validation["checks"]
    chart_roots = sorted({path.parent for path in candidate_dir.rglob("Chart.yaml")})
    rendered_parts: list[str] = []
    helm = shutil.which("helm")
    if chart_roots and helm:
        lint_results, template_results = [], []
        for chart_root in chart_roots:
            lint = subprocess.run([helm, "lint", str(chart_root)], capture_output=True, text=True, timeout=180, check=False)
            lint_results.append(lint)
            rendered = subprocess.run([helm, "template", "cats-remediation", str(chart_root), "--include-crds"],
                                      capture_output=True, text=True, timeout=180, check=False)
            template_results.append(rendered)
            if rendered.returncode == 0:
                rendered_parts.append(rendered.stdout)
        checks["helm_lint"] = {"status": "PASS" if all(item.returncode == 0 for item in lint_results) else "FAIL",
                               "detail": "Validated every discovered chart."}
        checks["helm_template"] = {"status": "PASS" if all(item.returncode == 0 for item in template_results) else "FAIL",
                                   "detail": "Rendered every discovered chart with CRDs."}
    elif chart_roots:
        checks["helm_lint"] = {"status": "FAIL", "detail": "Helm is unavailable in the remediation worker."}
        checks["helm_template"] = {"status": "FAIL", "detail": "Helm is unavailable in the remediation worker."}
    else:
        raw = candidate_dir / "remediated-manifests.yaml"
        if raw.exists():
            rendered_parts.append(raw.read_text(encoding="utf-8"))

    rendered_text = "\n---\n".join(rendered_parts)
    if rendered_text:
        try:
            objects = [item for item in yaml.safe_load_all(rendered_text) if isinstance(item, dict) and item.get("kind")]
            checks["yaml_parsing"] = {"status": "PASS", "detail": f"Parsed {len(objects)} candidate Kubernetes resources."}
            expected_kinds = {str(value).split("/", 1)[0] for value in plan["before"].get("resource_identities", [])}
            actual_kinds = {str(item.get("kind")) for item in objects}
            checks["expected_resources"] = {"status": "PASS" if expected_kinds <= actual_kinds else "FAIL",
                                             "detail": "Expected resource kinds remain present in the candidate render."}
            rendered_path = candidate_dir / ".cats-rendered.yaml"
            rendered_path.write_text(rendered_text, encoding="utf-8")
            kubectl = shutil.which("kubectl")
            if kubectl:
                result = subprocess.run([kubectl, "apply", "--dry-run=client", "--validate=true", "-f", str(rendered_path)],
                                        capture_output=True, text=True, timeout=180, check=False)
                checks["kubernetes_schema"] = {"status": "PASS" if result.returncode == 0 else "FAIL",
                                                "detail": (result.stderr or result.stdout)[-1000:]}
        except (OSError, yaml.YAMLError) as exc:
            checks["yaml_parsing"] = {"status": "FAIL", "detail": str(exc)}
    trivy = shutil.which("trivy")
    if trivy:
        result = subprocess.run([trivy, "config", "--exit-code", "1", str(candidate_dir)], capture_output=True, text=True,
                                timeout=int(os.getenv("CATS_REMEDIATION_SCAN_TIMEOUT", "600")), check=False)
        checks["trivy_config_rescan"] = {"status": "PASS" if result.returncode == 0 else "FAIL",
                                         "detail": (result.stderr or result.stdout)[-1000:]}
    if not plan.get("images"):
        checks["vulnerability_rescan"] = {"status": "PASS", "detail": "No image references require vulnerability scanning."}
    elif all(item.get("candidate") and item.get("patch_status") in {"PATCHED", "PARTIALLY_PATCHED", "NO_APPLICABLE_FIXES"}
             for item in plan.get("images", [])):
        checks["vulnerability_rescan"] = {"status": "PASS", "detail": "Every staged image completed the existing patch worker's before/after vulnerability scan."}
    validation["status"] = "PASS" if all(checks[name]["status"] == "PASS" for name in validation["required_checks"]) else "FAIL"
    kubeconfig = os.getenv("CATS_REMEDIATION_KUBECONFIG", "").strip()
    kubectl = shutil.which("kubectl")
    if validation["status"] == "PASS" and chart_roots and kubeconfig and helm and kubectl:
        namespace = "cats-remediation-" + re.sub(r"[^a-z0-9]", "", str(plan.get("job_id", "")).lower())[-20:]
        env = {**os.environ, "KUBECONFIG": kubeconfig}
        deployment = {"status": "FAIL", "namespace": namespace, "detail": "Deployment validation did not complete."}
        try:
            created = subprocess.run([kubectl, "create", "namespace", namespace], env=env, capture_output=True, text=True, timeout=60, check=False)
            if created.returncode != 0:
                raise RuntimeError((created.stderr or created.stdout)[-1000:])
            for index, chart_root in enumerate(chart_roots, start=1):
                installed = subprocess.run([helm, "upgrade", "--install", f"candidate-{index}", str(chart_root), "--namespace", namespace,
                                            "--wait", "--timeout", os.getenv("CATS_REMEDIATION_DEPLOY_TIMEOUT", "5m")],
                                           env=env, capture_output=True, text=True, timeout=600, check=False)
                if installed.returncode != 0:
                    raise RuntimeError((installed.stderr or installed.stdout)[-1000:])
            ready = subprocess.run([kubectl, "wait", "--for=condition=Ready", "pods", "--all", "--namespace", namespace,
                                    "--timeout", os.getenv("CATS_REMEDIATION_DEPLOY_TIMEOUT", "5m")],
                                   env=env, capture_output=True, text=True, timeout=600, check=False)
            if ready.returncode != 0:
                raise RuntimeError((ready.stderr or ready.stdout)[-1000:])
            deployment = {"status": "PASS", "namespace": namespace, "detail": "Charts installed and all created pods reached Ready."}
        except Exception as exc:
            deployment["detail"] = redact(exc)
        finally:
            subprocess.run([kubectl, "delete", "namespace", namespace, "--wait=false"], env=env,
                           capture_output=True, text=True, timeout=60, check=False)
        validation["deployment"] = deployment
    return validation


def _run_remediation_image_patches(db: Session, record: RemediationExecution, service: Service, plan: dict) -> None:
    """Run the existing patch worker for images when one staging registry is configured."""
    configuration = get_global_configuration(db)
    registries = [item for item in parse_json(configuration.get("oci_registries"), []) if isinstance(item, dict)]
    staging = [item for item in registries if item.get("use_for_remediation") is True]
    if len(staging) != 1:
        reason = "Configure exactly one OCI registry for remediation staging before image publication."
        for image in plan.get("images", []):
            if not image.get("candidate"):
                image["reason"] = reason
        return
    requester = db.get(User, record.requested_by_id) if record.requested_by_id else None
    signing_config, signing_credentials = _portal_signing_material(db, AuthContext(requester, None) if requester else None, "push", service)
    destination_registry = staging[0]
    endpoint = str(destination_registry.get("endpoint") or "").strip().rstrip("/")
    destination_host = urllib.parse.urlparse(endpoint if "://" in endpoint else f"https://{endpoint}").netloc
    destination_prefix = str(destination_registry.get("namespace") or "").strip("/")

    def credentials(registry: dict | None) -> tuple[str, str]:
        if not registry:
            return "", ""
        username = str(registry.get("username") or "")
        password = decrypt_secret(str(registry.get("password"))) if registry.get("password") else ""
        return username, password

    for image in plan.get("images", []):
        if image.get("candidate"):
            continue
        source = str(image.get("original") or "")
        source_registry = _configured_registry_for_image(source, registries)
        source_path = source.rsplit("@", 1)[0]
        last_slash, last_colon = source_path.rfind("/"), source_path.rfind(":")
        tag = source_path[last_colon + 1:] if last_colon > last_slash else "latest"
        repository = source_path[:last_colon] if last_colon > last_slash else source_path
        first, separator, remainder = repository.partition("/")
        if separator and ("." in first or ":" in first or first == "localhost"):
            repository = remainder
        destination = "/".join(part for part in (destination_host, destination_prefix, repository) if part) + f":{tag}-cats-r1"
        patch_key = uuid.uuid4().hex
        patch_root = PATCH_JOB_ROOT / patch_key
        (patch_root / "input").mkdir(parents=True, exist_ok=False)
        policy = configuration_for_service(db, service)
        config = {"job_id": patch_key, "source_mode": "oci", "source_image": source, "output_mode": "push",
                  **signing_config,
                  "destination_image": destination, "reuse_source_credentials": False,
                  "trusted_ca_certificates": parse_json(policy.get("trusted_ca_certificates"), []),
                  "repository_policies": parse_json(policy.get("repository_policies"), {}),
                  "os_definitions": parse_json(policy.get("os_definitions"), {})}
        (patch_root / "job-config.json").write_text(json.dumps(safe_job_config(config), indent=2), encoding="utf-8")
        patch_record = PatchExecution(job_key=patch_key, service_id=service.id, requested_by_id=record.requested_by_id,
                                      source_mode="oci", source_image=source, output_mode="push",
                                      destination_image=destination, status="queued", phase="queued", summary={})
        db.add(patch_record); db.flush()
        public = {"job_id": patch_key, "patch_execution_id": patch_record.id, "status": "queued", "phase": "queued",
                  "signing_enabled": bool(signing_config), "signing_fingerprint": signing_config.get("signing_fingerprint"),
                  "stages": initial_patch_stages("push"), "source_mode": "oci", "source_image": source,
                  "output_mode": "push", "destination_image": destination, "created_at": utcnow().isoformat()}
        with PATCH_JOB_LOCK:
            PATCH_JOBS[patch_key] = public
        db.commit()
        source_user, source_password = credentials(source_registry)
        destination_user, destination_password = credentials(destination_registry)
        _audit_signing_request(db, patch_key, record.requested_by_id, config)
        _run_patch_job(patch_key, {**signing_credentials, "CATS_PATCH_SOURCE_USERNAME": source_user, "CATS_PATCH_SOURCE_PASSWORD": source_password,
                                   "CATS_PATCH_DEST_USERNAME": destination_user, "CATS_PATCH_DEST_PASSWORD": destination_password})
        result = (_load_patch_job(patch_key).get("result") or {})
        immutable = result.get("immutable_destination")
        if result.get("delivery_status") == "delivered" and immutable:
            image.update(candidate=immutable, digest=str(immutable).split("@", 1)[-1], patch_status=result.get("patch_status"),
                         signature_status=result.get("signature_status", "not_configured"), patch_job_id=patch_key)
            mapping = image.get("source_mapping") or {}
            source_files = plan.get("_source_files") or {}
            editable = not mapping.get("ambiguous", True) and (mapping.get("values_file") or "") in source_files
            image["classification"] = "AUTO-REMEDIABLE" if editable else "REVIEW REQUIRED"
            image["reason"] = "Published and digest-qualified; source mapping is exact." if editable else "Published and digest-qualified, but the Helm source mapping needs review."
        else:
            image.update(classification="REVIEW REQUIRED", patch_job_id=patch_key,
                         reason=result.get("delivery_error") or result.get("reason") or "Image patch/publication did not produce an immutable staged reference.")


def _run_remediation_job(record_id: int) -> None:
    """Create an isolated candidate and persist every terminal outcome."""
    with SessionLocal() as db:
        record = db.get(RemediationExecution, record_id)
        if not record:
            return
        record.status, record.phase, record.started_at, record.updated_at = "running", "snapshot", utcnow(), utcnow()
        record.logs = ["Captured immutable references to the original execution and service revision."]
        db.commit()
        try:
            service = db.scalar(select(Service).where(Service.id == record.service_id).options(
                selectinload(Service.executions), selectinload(Service.policy_findings), selectinload(Service.findings).selectinload(Finding.observations)
            ))
            if not service or not service.executions:
                raise ValueError("The service has no assessment execution to remediate")
            execution = max(service.executions, key=lambda item: aware(item.scanned_at))
            payload = dict(execution.raw_payload) if isinstance(execution.raw_payload, dict) else {}
            if record.finding_type == "configuration":
                findings = [item for item in service.policy_findings if item.id == record.finding_id]
            elif record.finding_type == "vulnerability":
                findings = []
            else:
                findings = [item for item in service.policy_findings if item.active]
            plan = build_plan(payload, findings, record.job_key)
            plan["_source_files"] = payload.get("helm_source_files") or payload.get("source_files") or {}
            patch_records = list(db.scalars(select(PatchExecution).where(PatchExecution.service_id == service.id)))
            associate_patch_results(payload, plan, patch_records)
            _run_remediation_image_patches(db, record, service, plan)
            plan.pop("_source_files", None)
            for image in plan["images"]:
                if image.get("classification") == "AUTO-REMEDIABLE":
                    source = (image.get("source_mapping") or {}).get("values_file") or (image.get("source_mapping") or {}).get("template")
                    artifact = next((item for item in plan["changed_artifacts"] if item["path"] == source), None)
                    if not artifact:
                        artifact = {"path": source, "changes": []}; plan["changed_artifacts"].append(artifact)
                    artifact["changes"].append(f"Image {image['original']} → {image['candidate']}")
            record.phase = "candidate"
            record.before_snapshot = plan["before"]
            record.configuration_changes = plan["configuration_changes"]
            record.changed_artifacts = plan["changed_artifacts"]
            record.patched_images = plan["images"]
            record.rollback_reference = str(plan.get("rollback_reference") or execution.execution_key)
            record.logs = [*record.logs, "Classified findings and images without changing the original artifacts."]
            db.commit()

            files = candidate_files(payload, plan)
            job_root = REMEDIATION_JOB_ROOT / record.job_key
            job_root.mkdir(parents=True, exist_ok=True)
            artifact = job_root / "remediation-candidate.zip"
            candidate_dir = job_root / "candidate"
            candidate_dir.mkdir(parents=True, exist_ok=True)
            for relative, content in sorted(files.items()):
                safe = Path(relative.replace("\\", "/"))
                if safe.is_absolute() or ".." in safe.parts:
                    raise ValueError(f"Unsafe source path in uploaded artifact: {relative}")
                target = candidate_dir / safe
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(content, encoding="utf-8")
            validation = _validate_materialized_candidate(candidate_dir, payload, plan, static_validation(payload, plan))
            with ZipFile(artifact, "w", ZIP_DEFLATED) as bundle:
                bundle.writestr("remediation-plan.yaml", plan_yaml(plan))
                bundle.writestr("before-after.json", json.dumps({"before": plan["before"], "after": plan["after"]}, indent=2))
                for relative, content in sorted(files.items()):
                    safe = Path(relative.replace("\\", "/"))
                    if safe.is_absolute() or ".." in safe.parts:
                        raise ValueError(f"Unsafe source path in uploaded artifact: {relative}")
                    bundle.writestr(f"candidate/{safe.as_posix()}", content)
            record.phase = "validation"
            record.validation_results = validation
            record.after_snapshot = plan["after"]
            record.scan_results = {
                "configuration": validation["checks"]["trivy_config_rescan"],
                "vulnerabilities": validation["checks"]["vulnerability_rescan"],
            }
            record.artifact_path = str(artifact)
            automatic = [item for item in plan["configuration_changes"] if item["classification"] == "AUTO-REMEDIABLE"]
            automatic.extend(item for item in plan["images"] if item["classification"] == "AUTO-REMEDIABLE")
            review = [item for item in plan["configuration_changes"] if item["classification"] == "REVIEW REQUIRED"]
            review.extend(item for item in plan["images"] if item["classification"] == "REVIEW REQUIRED")
            if automatic and not files:
                raise ValueError("An automatic change had no editable source artifact")
            if review:
                record.status = "review_required"
                record.logs = [*record.logs, "Candidate created; ambiguous source or image changes require review before validation and promotion."]
            elif automatic and validation["status"] == "PASS":
                record.status = "validated"
                record.resulting_revision = f"{record.original_revision or service.manual_version or execution.execution_key}-cats-r1"
                record.logs = [*record.logs, "Static validation passed for the isolated candidate."]
            elif not automatic:
                record.status = "not_remediable"
                record.logs = [*record.logs, "No registered deterministic remediation could be applied."]
            else:
                record.status = "failed"
                record.failure_reason = "Required static validation failed"
            record.phase = "complete" if record.status != "failed" else "failed"
            record.completed_at = record.updated_at = utcnow()
            db.add(AuditEvent(actor_user_id=record.requested_by_id, action=f"remediation.{record.status}",
                              target_type="remediation_execution", target_id=str(record.id),
                              detail={"service_id": record.service_id, "job_key": record.job_key, "status": record.status}))
            db.commit()
        except Exception as exc:
            db.rollback()
            record = db.get(RemediationExecution, record_id)
            if record:
                record.status, record.phase = "failed", "failed"
                record.failure_reason = str(exc)[:4000]
                record.logs = [*(record.logs or []), f"Failed safely: {str(exc)[:1000]}"]
                record.completed_at = record.updated_at = utcnow()
                db.commit()


def _queue_remediation(db: Session, auth: AuthContext, service: Service, finding_type: str | None = None,
                       finding_id: int | None = None) -> RemediationExecution:
    latest = db.scalar(select(Execution).where(Execution.service_id == service.id).order_by(Execution.scanned_at.desc()))
    if not latest:
        raise HTTPException(422, detail="The service has no assessment execution to remediate")
    job_key = f"R-{uuid.uuid4().hex[:12].upper()}"
    record = RemediationExecution(job_key=job_key, service_id=service.id, requested_by_id=auth.user.id,
                                  finding_type=finding_type, finding_id=finding_id, status="queued", phase="queued",
                                  original_revision=service.manual_version or latest.commit_sha or latest.execution_key,
                                  rollback_reference=latest.execution_key, logs=["Remediation request queued."])
    db.add(record); db.flush()
    record_audit(db, auth, "remediation.queued", "remediation_execution", record.id,
                 service_id=service.id, job_key=job_key, finding_type=finding_type, finding_id=finding_id)
    db.commit()
    REMEDIATION_WORKERS.submit(_run_remediation_job, record.id)
    return record


@app.post("/services/{service_key}/remediate")
def remediate_service(service_key: str, csrf_token: str = Form(), db: Session = Depends(get_db),
                      auth: AuthContext = Depends(require_permission("remediation.execute", scoped=True))):
    check_csrf(auth, csrf_token)
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if not service:
        raise HTTPException(404)
    record = _queue_remediation(db, auth, service)
    return RedirectResponse(f"/services/{service_key}/remediations/{record.job_key}", status_code=303)


@app.post("/services/{service_key}/policy-findings/{finding_id}/remediate")
def remediate_policy_finding(service_key: str, finding_id: int, csrf_token: str = Form(), db: Session = Depends(get_db),
                              auth: AuthContext = Depends(require_permission("remediation.execute", scoped=True))):
    check_csrf(auth, csrf_token)
    finding = db.scalar(select(PolicyFinding).where(PolicyFinding.id == finding_id).options(selectinload(PolicyFinding.service)))
    if not finding or finding.service.service_key != service_key:
        raise HTTPException(404)
    record = _queue_remediation(db, auth, finding.service, "configuration", finding.id)
    return RedirectResponse(f"/services/{service_key}/remediations/{record.job_key}", status_code=303)


@app.post("/services/{service_key}/findings/{finding_id}/remediate")
def remediate_vulnerability_finding(service_key: str, finding_id: int, csrf_token: str = Form(), db: Session = Depends(get_db),
                                     auth: AuthContext = Depends(require_permission("remediation.execute", scoped=True))):
    check_csrf(auth, csrf_token)
    finding = db.scalar(select(Finding).where(Finding.id == finding_id).options(selectinload(Finding.service)))
    if not finding or finding.service.service_key != service_key:
        raise HTTPException(404)
    record = _queue_remediation(db, auth, finding.service, "vulnerability", finding.id)
    return RedirectResponse(f"/services/{service_key}/remediations/{record.job_key}", status_code=303)


@app.get("/services/{service_key}/remediations/{job_key}", response_class=HTMLResponse)
def remediation_report(service_key: str, job_key: str, request: Request, db: Session = Depends(get_db),
                       auth: AuthContext = Depends(require_permission("service.view", scoped=True))):
    record = db.scalar(select(RemediationExecution).where(RemediationExecution.job_key == job_key).options(selectinload(RemediationExecution.service)))
    if not record or record.service.service_key != service_key:
        raise HTTPException(404)
    return templates.TemplateResponse(request, "remediation_report.html", page_context(auth, job=record, service=record.service))


@app.get("/services/{service_key}/remediations/{job_key}/candidate.zip")
def remediation_candidate(service_key: str, job_key: str, db: Session = Depends(get_db),
                          auth: AuthContext = Depends(require_permission("service.export", scoped=True))):
    record = db.scalar(select(RemediationExecution).join(Service).where(
        RemediationExecution.job_key == job_key, Service.service_key == service_key))
    path = Path(record.artifact_path) if record and record.artifact_path else None
    if not path or not path.is_file() or path.parent != REMEDIATION_JOB_ROOT / job_key:
        raise HTTPException(404, detail="Remediation candidate is not available")
    return FileResponse(path, media_type="application/zip", filename=f"cats-{service_key}-{job_key}.zip")


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, archived: bool = False, lifecycle: str = "active", q: str = "", sort: str = "name",
              page: int = 1, page_size: int = 50, db: Session = Depends(get_db), auth: AuthContext | None = Depends(optional_user)):
    if not auth:
        return templates.TemplateResponse(request, "home.html", {
            "current_user": None, "csrf_token": "", "can": lambda _permission: False,
            "themes": THEMES, "pending_request_count": 0, "pending_poam_count": 0,
        })
    request_started = time.perf_counter()
    stage_timings = {}
    now = utcnow()
    stage_started = time.perf_counter()
    configuration = get_configuration(db)
    stage_timings["configuration_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
    now_display = configured_time(now, include_time=True, configuration=configuration)
    stage_started = time.perf_counter()
    views, poam_counts = service_overview_rows(db, auth, now, configuration)
    stage_timings["overview_aggregation_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
    stage_started = time.perf_counter()
    if archived:
        lifecycle = "archived"
    if lifecycle not in {"active", "staged", "archived"}:
        raise HTTPException(422, detail="Unknown service lifecycle")
    def view_lifecycle(view: dict) -> str:
        if view.get("archive"):
            return "archived"
        status = str(getattr(view["service"], "lifecycle_status", "active") or "active").lower()
        return status if status in {"active", "staged"} else "active"
    lifecycle_counts = {state: sum(view_lifecycle(view) == state for view in views) for state in ("active", "staged", "archived")}
    selected = [view for view in views if view_lifecycle(view) == lifecycle]
    search = q.strip().casefold()
    if search:
        selected = [view for view in selected if search in " ".join(str(value or "") for value in (
            view["service"].name, view["service"].service_key, view["service"].owner, view["service"].poc)).casefold()]
    sort_keys = {
        "name": lambda item: (item["service"].name.casefold(), item["service"].service_key.casefold()),
        "version": lambda item: str(item.get("version") or "").casefold(),
        "owner": lambda item: str(item["service"].owner or "").casefold(),
        "findings": lambda item: len(item["active"]) + len(item["policy_findings"]),
        "overdue": lambda item: len(item["noncompliant"]) + len(item["policy_noncompliant"]),
        "oldest": lambda item: item.get("oldest_age") if item.get("oldest_age") is not None else -1,
    }
    sort = sort if sort in sort_keys else "name"
    selected.sort(key=sort_keys[sort], reverse=request.query_params.get("direction") == "desc")
    total_count = len(selected)
    page_size = max(10, min(page_size, 200))
    total_pages = max(1, (total_count + page_size - 1) // page_size)
    page = max(1, min(page, total_pages))
    visible_all = selected
    views = selected[(page - 1) * page_size:page * page_size]
    visible_ids = [view["service"].id for view in visible_all]
    visible_poam = [poam_counts.get(service_id, {"active": 0, "pending": 0, "overdue": 0}) for service_id in visible_ids]
    stage_timings["filter_sort_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
    stage_started = time.perf_counter()
    stage_groups = manageable_groups(db, auth) if auth.has("user.manage") else []
    stage_timings["authorization_groups_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
    stage_started = time.perf_counter()
    context = page_context(auth,
        views=views, now=now, now_display=now_display,
        compliant_count=sum(v["compliant"] for v in visible_all),
        noncompliant_count=sum(not v["compliant"] for v in visible_all),
        showing_archived=lifecycle == "archived", lifecycle=lifecycle, lifecycle_counts=lifecycle_counts,
        query=q, sort=sort, page=page, page_size=page_size, total_count=total_count, total_pages=total_pages,
        overdue_days=int(configuration["overdue_days"]),
        poam_active_count=sum(item["active"] for item in visible_poam),
        poam_pending_count=sum(item["pending"] for item in visible_poam),
        poam_overdue_count=sum(item["overdue"] for item in visible_poam),
        stage_groups=stage_groups if lifecycle == "active" else [],
    )
    stage_timings["page_context_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
    stage_timings["route_ms"] = round((time.perf_counter() - request_started) * 1000, 2)
    request.state.snapshot_timings = stage_timings
    return templates.TemplateResponse(request, "dashboard.html", context)


@app.get("/exports/services.xlsx")
def export_services(db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    now = utcnow()
    configuration = get_configuration(db)
    services = db.scalars(select(Service).options(
        selectinload(Service.findings).selectinload(Finding.exceptions),
        selectinload(Service.policy_findings).selectinload(PolicyFinding.exceptions),
        selectinload(Service.executions), selectinload(Service.archive_events), selectinload(Service.groups),
    )).all()
    allowed = auth.accessible_service_ids("service.export")
    if allowed == set():
        raise HTTPException(403, detail="Permission denied")
    if allowed is not None:
        services = [service for service in services if service.id in allowed]
    views = [service_view(service, now, configuration_for_service(db, service)) for service in services]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Service Snapshot"
    sheet.append([
        "Service ID", "Service", "Version", "Owner", "POC", "Groups", "Lifecycle", "Compliance",
        "Evidence", "Active Findings", "Overdue Findings", "Active Exceptions",
        "Resolved", "Oldest Active (Days)", "Last Execution",
    ])
    for view in sorted(views, key=lambda item: item["service"].name.lower()):
        sheet.append([
            view["service"].service_key, view["service"].name, view["version"],
            view["service"].owner, view["service"].poc, ", ".join(group.name for group in view["service"].groups), "Archived" if view["archive"] else "Active",
            "Compliant" if view["compliant"] else "Non-Compliant",
            view["evidence_state"], len(view["active"]) + len(view["policy_findings"]),
            len(view["noncompliant"]) + len(view["policy_noncompliant"]),
            len(view["excepted"]) + len(view["policy_excepted"]),
            len(view["resolved"]) + len(view["policy_resolved"]), view["oldest_age"],
            excel_datetime(view["last_execution"]),
        ])
    for cell in sheet["M"][1:]:
        cell.number_format = "yyyy-mm-dd hh:mm"
    format_sheet(sheet)
    return workbook_response(workbook, f"cats-service-posture-{now:%Y%m%d}.xlsx")


@app.get("/services/{service_key}", response_class=HTMLResponse)
def service_detail(
    service_key: str,
    request: Request,
    finding_state: str = "active",
    finding_type: str = "all",
    page: int = 1,
    page_size: int = 50,
    overview: bool = False,
    simplified: bool = False,
    poam: bool = False,
    remediations: bool = False,
    tab: str = "poams",
    activity: bool = False,
    architecture: bool = False,
    layout_width: int | None = None,
    status_filter: str = "all",
    sort_by: str = "newest",
    db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("service.view", scoped=True)),
):
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(
        selectinload(Service.findings).selectinload(Finding.exceptions),
        selectinload(Service.policy_findings).selectinload(PolicyFinding.exceptions),
        selectinload(Service.findings).selectinload(Finding.observations),
        selectinload(Service.executions),
        selectinload(Service.archive_events),
        selectinload(Service.groups),
        selectinload(Service.images),
    ))
    if not service:
        raise HTTPException(404)
    now = utcnow()
    configuration = configuration_for_service(db, service)
    view = service_view(service, now, configuration)
    archive_pending = bool(db.scalar(select(WorkflowRequest.id).where(
        WorkflowRequest.request_type == "archive",
        WorkflowRequest.service_id == service.id,
        WorkflowRequest.status == "pending",
    )))
    if architecture:
        latest_execution = max(service.executions, key=lambda execution: aware(execution.scanned_at), default=None)
        payload = dict(latest_execution.raw_payload) if latest_execution and isinstance(latest_execution.raw_payload, dict) else {}
        if latest_execution:
            payload["complete"] = latest_execution.complete
        if layout_width is not None:
            if not 240 <= layout_width <= 10000:
                raise HTTPException(422, detail="Invalid architecture canvas width")
            return JSONResponse(build_architecture_graph(payload, layout_width)["layouts"])
        return templates.TemplateResponse(request, "service_architecture.html", page_context(auth,
            view=view, architecture_graph=build_architecture_graph(payload), latest_execution=latest_execution,
            now=now, archive_pending=archive_pending,
        ))
    if remediations:
        if tab not in {"pipeline", "poams", "exceptions", "mitigations"}:
            raise HTTPException(422, detail="Unknown remediation tab")
        entries = db.scalars(select(PoamEntry).where(PoamEntry.service_id == service.id).options(
            selectinload(PoamEntry.finding), selectinload(PoamEntry.policy_finding),
            selectinload(PoamEntry.created_by), selectinload(PoamEntry.approved_by),
        ).order_by(PoamEntry.created_at.desc())).all()
        poams = [entry for entry in entries if entry.item_type != "mitigation"]
        mitigations = [entry for entry in entries if entry.item_type == "mitigation"]
        exceptions = []
        for record in db.scalars(select(ExceptionRecord).join(Finding).where(Finding.service_id == service.id).options(
            selectinload(ExceptionRecord.finding),
        ).order_by(ExceptionRecord.created_at.desc())).all():
            finding = record.finding
            state = "Revoked" if record.revoked_at else ("Expired" if aware(record.expires_at) < now else "Active")
            exceptions.append({"kind": "Vulnerability", "item": finding.cve, "severity": finding.severity,
                               "status": state, "expires_at": record.expires_at, "approved_by": record.approved_by,
                               "created_at": record.created_at, "justification": record.justification,
                               "record_id": record.id, "href": f"/services/{service.service_key}/findings/{finding.id}",
                               "revoke_href": f"/exceptions/{record.id}/revoke"})
        for record in db.scalars(select(PolicyExceptionRecord).join(PolicyFinding).where(PolicyFinding.service_id == service.id).options(
            selectinload(PolicyExceptionRecord.policy_finding),
        ).order_by(PolicyExceptionRecord.created_at.desc())).all():
            finding = record.policy_finding
            state = "Revoked" if record.revoked_at else ("Expired" if aware(record.expires_at) < now else "Active")
            exceptions.append({"kind": "Configuration", "item": finding.finding, "severity": finding.severity,
                               "status": state, "expires_at": record.expires_at, "approved_by": record.approved_by,
                               "created_at": record.created_at, "justification": record.justification,
                               "record_id": record.id, "href": f"/services/{service.service_key}?finding_state=exceptions&finding_type=configuration",
                               "revoke_href": f"/policy-exceptions/{record.id}/revoke"})
        for workflow in db.scalars(select(WorkflowRequest).where(
            WorkflowRequest.service_id == service.id, WorkflowRequest.request_type == "exception", WorkflowRequest.status == "pending",
        ).options(selectinload(WorkflowRequest.finding), selectinload(WorkflowRequest.policy_finding)).order_by(WorkflowRequest.created_at.desc())).all():
            target = workflow.finding or workflow.policy_finding
            if not target:
                continue
            is_vulnerability = bool(workflow.finding)
            exceptions.append({"kind": "Vulnerability" if is_vulnerability else "Configuration",
                               "item": target.cve if is_vulnerability else target.finding,
                               "severity": target.severity, "status": "Pending", "expires_at": workflow.requested_expires_at,
                               "approved_by": "Pending review", "created_at": workflow.created_at,
                               "justification": workflow.justification, "record_id": workflow.id,
                               "href": f"/services/{service.service_key}/findings/{target.id}" if is_vulnerability else f"/services/{service.service_key}?finding_state=exceptions&finding_type=configuration"})
        remediation_jobs = db.scalars(select(RemediationExecution).where(
            RemediationExecution.service_id == service.id).order_by(RemediationExecution.created_at.desc()).limit(100)).all()
        return templates.TemplateResponse(request, "service_remediations.html", page_context(auth,
            service=service, view=view, tab=tab, poams=poams, exceptions=exceptions, mitigations=mitigations,
            remediation_jobs=remediation_jobs, can_remediate=auth.has("remediation.execute", service.id),
            can_create_poam=auth.has("poam.request", service.id), now=now,
        ))
    if poam:
        if status_filter not in {"all", "active", "pending_approval", "overdue", "completed", "rejected"}:
            raise HTTPException(422, detail="Unknown POA&M filter")
        entries = db.scalars(select(PoamEntry).where(PoamEntry.service_id == service.id).options(
            selectinload(PoamEntry.finding), selectinload(PoamEntry.policy_finding), selectinload(PoamEntry.created_by), selectinload(PoamEntry.approved_by),
        ).order_by(PoamEntry.created_at.desc())).all()
        if status_filter == "overdue":
            entries = [entry for entry in entries if entry.status == "active" and entry.due_date and aware(entry.due_date) < now]
        elif status_filter != "all":
            entries = [entry for entry in entries if entry.status == status_filter]
        if sort_by == "due":
            entries.sort(key=lambda entry: (aware(entry.due_date) if entry.due_date else datetime.max.replace(tzinfo=timezone.utc), entry.title.lower()))
        elif sort_by == "severity":
            severity_rank = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Negligible": 4}
            entries.sort(key=lambda entry: (severity_rank.get(entry.finding.severity if entry.finding else (entry.policy_finding.severity if entry.policy_finding else ""), 5), entry.title.lower()))
        elif sort_by != "newest":
            raise HTTPException(422, detail="Unknown POA&M sort")
        return templates.TemplateResponse(request, "poam_service.html", page_context(auth,
            service=service, view=view, embedded=True, entries=entries, can_create_poam=auth.has("poam.request", service.id),
            status_filter=status_filter, sort_by=sort_by, now=now,
            overdue_entry_ids={entry.id for entry in entries if entry.status == "active" and entry.due_date and aware(entry.due_date) < now},
        ))
    if activity:
        events = db.scalars(select(AuditEvent).options(selectinload(AuditEvent.actor)).order_by(AuditEvent.created_at.desc()).limit(500)).all()
        service_events = [event for event in events if isinstance(event.detail, dict) and str(event.detail.get("service_id")) == str(service.id)]
        return templates.TemplateResponse(request, "service_activity.html", page_context(auth,
            service=service, view=view, events=service_events, embedded=True, now=now,
        ))
    finding_groups = {
        "active": view["active"],
        "noncompliant": view["noncompliant"],
        # Keep the old query value working for bookmarked links.
        "overdue": view["noncompliant"],
        "exceptions": view["excepted"],
        "resolved": view["resolved"],
        "warnings": view["warning_items"],
    }
    if finding_state not in finding_groups:
        raise HTTPException(400, detail="Unknown finding state")
    if finding_type not in {"all", "vulnerability", "configuration", "evidence"}:
        raise HTTPException(422, detail="Finding type must be all, vulnerability, configuration, or evidence")
    if page_size not in {50, 100, 250}:
        raise HTTPException(422, detail="Page size must be 50, 100, or 250")
    if finding_state == "warnings":
        all_items = list(finding_groups["warnings"])
        if finding_type != "all":
            warning_types = {"vulnerability": "CVE", "configuration": "Configuration", "evidence": "Evidence"}
            all_items = [item for item in all_items if item.get("type") == warning_types[finding_type]]
        total_items = len(all_items)
        total_pages = max(1, (total_items + page_size - 1) // page_size)
        page = max(1, min(page, total_pages))
        page_start = (page - 1) * page_size
        displayed_warning_items = all_items[page_start:page_start + page_size]
        return templates.TemplateResponse(request, "service.html", page_context(auth,
            view=view, findings=[], affected_images={}, policy_findings=[], noncompliance_items=[],
            warning_items=displayed_warning_items, now=now, active_exception=active_exception,
            finding_state=finding_state, groups=db.scalars(select(Group).order_by(Group.name)).all(),
            overdue_days=view["overdue_days"], saved=request.query_params.get("saved") == "1",
            page=page, page_size=page_size, total_items=total_items, total_pages=total_pages,
            finding_type=finding_type,
        ))
    all_findings = sorted(finding_groups[finding_state], key=lambda finding: (aware(finding.episode_started), finding.cve))
    policy_items = {
        "active": view["policy_findings"],
        "exceptions": view["policy_excepted"],
        "resolved": view["policy_resolved"],
    }.get(finding_state, [])
    mixed_state = finding_state in {"active", "exceptions", "resolved"}
    if mixed_state:
        active_entries = [("vulnerability", finding) for finding in all_findings]
        active_entries.extend(("configuration", item) for item in sorted(
            policy_items, key=lambda finding: (aware(finding.episode_started), finding.finding)
        ))
        if finding_type != "all":
            active_entries = [entry for entry in active_entries if entry[0] == finding_type]
        total_items = len(active_entries)
    else:
        all_items = view["noncompliance_items"] if finding_state in {"noncompliant", "overdue"} else all_findings
        if finding_type != "all":
            item_type = {"evidence": "Evidence", "configuration": "Configuration", "vulnerability": "CVE"}[finding_type]
            all_items = [item for item in all_items if (item.get("type") if isinstance(item, dict) else "CVE") == item_type]
        total_items = len(all_items)
    total_pages = max(1, (total_items + page_size - 1) // page_size)
    page = max(1, min(page, total_pages))
    page_start = (page - 1) * page_size
    if mixed_state:
        displayed_entries = active_entries[page_start:page_start + page_size]
        findings = [item for kind, item in displayed_entries if kind == "vulnerability"]
        displayed_policy_findings = [item for kind, item in displayed_entries if kind == "configuration"]
        displayed_items = findings
    else:
        displayed_items = all_items[page_start:page_start + page_size]
        findings = [] if finding_state in {"noncompliant", "overdue"} else displayed_items
        displayed_policy_findings = []
    latest_execution = max(service.executions, key=lambda execution: aware(execution.scanned_at), default=None)
    if overview:
        raw_overview = latest_execution.raw_payload.get("service_overview", {}) if latest_execution and isinstance(latest_execution.raw_payload, dict) else {}
        raw_overview = raw_overview if isinstance(raw_overview, dict) else {}
        latest_payload = latest_execution.raw_payload if latest_execution and isinstance(latest_execution.raw_payload, dict) else {}
        raw_overview.setdefault("source", "Helm rendered manifests" if latest_payload.get("policy_findings") else "Image metadata")
        overview_data = normalize_overview(
            raw_overview,
            skipped_images=latest_payload.get("skipped_images", []) or [],
            skipped_charts=latest_payload.get("skipped_charts", []) or [],
            findings_images=[{"image": item.get("image"), "digest": item.get("image_digest"), "discovered_from": item.get("discovered_from") or "Submitted"} for item in latest_payload.get("findings", []) if isinstance(item, dict) and item.get("image")],
            incomplete=bool(latest_execution and not latest_execution.complete),
            # The persisted scan overview is the source of truth for this
            # page.  Do not run docker manifest inspect while navigating.
            digest_resolver=None,
        )
        return templates.TemplateResponse(request, "service_overview.html", page_context(auth,
            view=view, overview_data=overview_data, latest_execution=latest_execution,
            service_images=service.images,
            now=now, finding_type=finding_type, archive_pending=archive_pending,
            groups=db.scalars(select(Group).order_by(Group.name)).all(),
        ))
    if simplified:
        simplified_groups = {}
        for finding in view["active"]:
            observations = [o for o in finding.observations if latest_execution and o.execution_id == latest_execution.id]
            observation = max(observations or finding.observations, key=lambda item: item.id, default=None)
            evidence = observation.evidence if observation else {}
            package = (observation.package if observation else None) or "Package update"
            fixed = (observation.fixed_version if observation else None) or "Latest fixed version"
            remediation = evidence.get("remediation") or evidence.get("recommendation") or "Update the affected package to the fixed version."
            key = (package, str(remediation))
            group = simplified_groups.setdefault(key, {
                "package": package, "fixed_versions": set(), "remediation": str(remediation),
                "cves": [], "images": set(), "severities": [], "finding_ids": [],
                "due": view["due_dates"].get(finding.id),
            })
            cve = str(finding.cve or "").strip().upper()
            if cve and cve not in group["cves"]:
                group["cves"].append(cve)
                group["finding_ids"].append(finding.id)
            group["fixed_versions"].add(str(fixed))
            group["severities"].append(finding.severity)
            group["due"] = min(group["due"], view["due_dates"].get(finding.id)) if group["due"] and view["due_dates"].get(finding.id) else group["due"]
            group["images"].update(o.image for o in observations if o.image)
        severity_rank = {"unknown": 0, "negligible": 1, "low": 2, "medium": 3, "high": 4, "critical": 5}
        simplified_findings = []
        for group in simplified_groups.values():
            group["cves"] = sorted(group["cves"])
            group["images"] = sorted(group["images"])
            group["fixed_versions"] = sorted(group["fixed_versions"])
            group["fixed_version"] = ", ".join(group["fixed_versions"])
            group["severity"] = max(group["severities"], key=lambda value: severity_rank.get(str(value).lower(), 0), default="Unknown")
            simplified_findings.append(group)
        simplified_findings.sort(key=lambda item: (str(item["package"]).casefold(), str(item["fixed_version"]).casefold()))
        groups = db.scalars(select(Group).order_by(Group.name)).all()
        return templates.TemplateResponse(request, "service_simplified.html", page_context(auth,
            view=view, simplified_findings=simplified_findings, now=now, finding_type=finding_type, archive_pending=archive_pending, groups=groups,
        ))
    groups = db.scalars(select(Group).order_by(Group.name)).all()
    affected_images = {}
    image_finding_ids = {finding.id for finding in findings}
    image_finding_ids.update(item["finding_id"] for item in displayed_items if isinstance(item, dict) and item.get("finding_id"))
    for finding in service.findings:
        if finding.id not in image_finding_ids:
            continue
        if finding.active and latest_execution:
            observations = [
                observation for observation in finding.observations
                if observation.execution_id == latest_execution.id
            ]
        else:
            latest_observation = max(finding.observations, key=lambda observation: observation.id, default=None)
            observations = [
                observation for observation in finding.observations
                if latest_observation and observation.execution_id == latest_observation.execution_id
            ]
        affected_images[finding.id] = sorted({observation.image for observation in observations})
    for item in displayed_items if finding_state in {"noncompliant", "overdue"} else []:
        if item["type"] == "CVE":
            item["images"] = affected_images.get(item["finding_id"], [])
        else:
            item["images"] = [item["evidence_image"]] if item.get("evidence_image") else []
    latest_payload = latest_execution.raw_payload if latest_execution and isinstance(latest_execution.raw_payload, dict) else {}
    remediation_classes = {item.id: classify_policy_finding(item, latest_payload) for item in displayed_policy_findings}
    return templates.TemplateResponse(request, "service.html", page_context(auth,
        view=view, findings=findings, affected_images=affected_images,
        policy_findings=displayed_policy_findings,
        noncompliance_items=displayed_items if finding_state in {"noncompliant", "overdue"} else [],
        now=now, active_exception=active_exception, finding_state=finding_state, groups=groups,
        overdue_days=view["overdue_days"], saved=request.query_params.get("saved") == "1",
        archive_pending=archive_pending,
        page=page, page_size=page_size, total_items=total_items, total_pages=total_pages,
        finding_type=finding_type, remediation_classes=remediation_classes,
    ))


@app.get("/services/{service_key}/export.xlsx")
def export_service(service_key: str, include_diagrams: bool = False, db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("service.export", scoped=True))):
    now = utcnow()
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(
        selectinload(Service.findings).selectinload(Finding.exceptions),
        selectinload(Service.policy_findings).selectinload(PolicyFinding.exceptions),
        selectinload(Service.findings).selectinload(Finding.observations),
        selectinload(Service.executions), selectinload(Service.archive_events), selectinload(Service.groups),
        selectinload(Service.images),
        selectinload(Service.poam_entries).selectinload(PoamEntry.finding),
        selectinload(Service.poam_entries).selectinload(PoamEntry.policy_finding),
        selectinload(Service.poam_entries).selectinload(PoamEntry.created_by),
        selectinload(Service.poam_entries).selectinload(PoamEntry.approved_by),
    ))
    if not service:
        raise HTTPException(404)
    configuration = configuration_for_service(db, service)
    view = service_view(service, now, configuration)
    latest_execution = max(service.executions, key=lambda execution: aware(execution.scanned_at), default=None)
    audit_events = db.scalars(select(AuditEvent).options(selectinload(AuditEvent.actor)).order_by(AuditEvent.created_at)).all()
    service_audit = []
    for event in audit_events:
        detail = event.detail if isinstance(event.detail, dict) else {}
        if event.target_type == "service" and str(event.target_id) == str(service.id):
            service_audit.append(event)
        elif str(detail.get("service_id") or detail.get("service_key")) in {str(service.id), service.service_key}:
            service_audit.append(event)
    histories = db.scalars(
        select(PoamHistory).join(PoamEntry, PoamEntry.id == PoamHistory.poam_id)
        .where(PoamEntry.service_id == service.id)
        .options(selectinload(PoamHistory.actor))
        .order_by(PoamHistory.created_at)
    ).all()
    workbook = build_service_workbook(
        service, view, configuration, now, latest_execution,
        activity_events=service_audit, poam_histories=histories,
    )
    if not include_diagrams:
        return workbook_response(workbook, service_export_filename(service, now))
    archive = BytesIO()
    with ZipFile(archive, "w", ZIP_DEFLATED) as bundle:
        workbook_bytes = BytesIO()
        workbook.save(workbook_bytes)
        bundle.writestr("service-export.xlsx", workbook_bytes.getvalue())
        graph = build_architecture_graph(latest_execution.raw_payload if latest_execution else {})
        for view_name in ("all", "configuration", "containers", "flow", "network", "storage"):
            bundle.writestr(f"architecture-{view_name}.svg", build_architecture_svg(graph, view_name))
        bundle.writestr("legacy-helm-diagram.svg", build_helm_diagram(service, latest_execution))
    archive.seek(0)
    filename = service_export_filename(service, now).rsplit(".", 1)[0] + ".zip"
    return StreamingResponse(archive, media_type="application/zip", headers={"Content-Disposition": f"attachment; filename={filename}"})


@app.get("/services/{service_key}/helm-diagram.svg")
def export_service_helm_diagram(service_key: str, db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("service.view", scoped=True))):
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(selectinload(Service.executions)))
    if not service:
        raise HTTPException(404)
    latest_execution = max(service.executions, key=lambda execution: aware(execution.scanned_at), default=None)
    return Response(build_helm_diagram(service, latest_execution), media_type="image/svg+xml")


@app.get("/services/{service_key}/findings/{finding_id}", response_class=HTMLResponse)
def finding_detail(
    service_key: str, finding_id: int, request: Request,
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("service.view", scoped=True)),
):
    finding = db.scalar(select(Finding).where(Finding.id == finding_id).options(
        selectinload(Finding.service).selectinload(Service.executions),
        selectinload(Finding.service).selectinload(Service.groups),
        selectinload(Finding.observations),
        selectinload(Finding.exceptions),
    ))
    if not finding or finding.service.service_key != service_key:
        raise HTTPException(404)
    now = utcnow()
    configuration = configuration_for_service(db, finding.service)
    risk = service_view(finding.service, now, configuration)["risk_metadata"].get(finding.id, {})
    latest_execution = max(finding.service.executions, key=lambda execution: aware(execution.scanned_at), default=None)
    current_observations = [
        observation for observation in finding.observations
        if finding.active and latest_execution and observation.execution_id == latest_execution.id
    ]
    evidence_observation = current_observations[0] if current_observations else (
        max(finding.observations, key=lambda observation: observation.id, default=None)
    )
    evidence = evidence_observation.evidence if evidence_observation else {}
    history = sorted(finding.observations, key=lambda observation: observation.id, reverse=True)
    due_days = max(1, int(configuration.get("overdue_days", "90")))
    if configuration.get("compliance_mode") == "raw":
        try:
            raw_rules = json.loads(configuration.get("raw_due_rules", "[]"))
            raw_days = {
                str(rule.get("severity", "")).lower(): max(1, int(rule.get("days", due_days)))
                for rule in raw_rules
            }
            due_days = raw_days.get(finding.severity.lower(), due_days)
        except (TypeError, ValueError, json.JSONDecodeError):
            pass
    return templates.TemplateResponse(request, "finding.html", page_context(auth,
        finding=finding, service=finding.service, now=now,
        age=(now - aware(finding.episode_started)).days if finding.active else None,
        exception=active_exception(finding, now), evidence=evidence,
        current_observations=current_observations, history=history,
        risk=risk,
        remediation_classification={"classification": "REVIEW REQUIRED", "reason": "Image patching and its Helm source update require a candidate job."},
        due_date=(aware(finding.episode_started) + timedelta(days=due_days) if finding.active else None),
    ))


@app.post("/findings/{finding_id}/exceptions")
def request_exception(
    finding_id: int, justification: str = Form(min_length=3), expires_at: datetime = Form(),
    ticket: str = Form(default=""), bulk_scope: str = Form(default=""), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("exception.request", scoped=True)),
):
    check_csrf(auth, csrf_token)
    finding = db.scalar(select(Finding).where(Finding.id == finding_id).options(selectinload(Finding.service).selectinload(Service.groups)))
    if not finding:
        raise HTTPException(404)
    bulk_group = cybersecurity_group(finding.service) if bulk_scope == "group" else None
    if bulk_scope == "group" and not bulk_group:
        raise HTTPException(422, detail="This service does not have one unambiguous Cybersecurity group scope")
    expires_at = aware(expires_at)
    now = utcnow()
    if expires_at <= now:
        raise HTTPException(422, detail="Expiration must be in the future")
    maximum_exception = int(configuration_for_service(db, finding.service).get("exception_max_days", "365"))
    if maximum_exception and expires_at > now + timedelta(days=maximum_exception):
        raise HTTPException(422, detail=f"Exception expiration cannot exceed {maximum_exception} days")
    pending = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.request_type == "exception", WorkflowRequest.finding_id == finding.id,
        WorkflowRequest.status == "pending",
    ))
    if pending:
        raise HTTPException(409, detail="An exception request is already pending")
    workflow = WorkflowRequest(
        request_type="exception", service_id=finding.service_id, finding_id=finding.id,
        requested_by_id=auth.user.id, justification=justification,
        requested_expires_at=expires_at, ticket=ticket or None,
        bulk_group_id=bulk_group.id if bulk_group else None,
    )
    db.add(workflow)
    db.flush()
    record_audit(db, auth, "exception.requested", "workflow_request", workflow.id, service_id=finding.service_id)
    db.commit()
    return RedirectResponse(f"/services/{finding.service.service_key}", status_code=status.HTTP_303_SEE_OTHER)


@app.post("/policy-findings/{policy_finding_id}/exceptions")
def request_policy_exception(
    policy_finding_id: int, justification: str = Form(min_length=3), expires_at: datetime = Form(),
    ticket: str = Form(default=""), bulk_scope: str = Form(default=""), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    finding = db.scalar(select(PolicyFinding).where(PolicyFinding.id == policy_finding_id).options(
        selectinload(PolicyFinding.service).selectinload(Service.groups)
    ))
    if not finding:
        raise HTTPException(404)
    if not auth.has("exception.request", finding.service_id):
        raise HTTPException(403, detail="Permission denied")
    bulk_group = cybersecurity_group(finding.service) if bulk_scope == "group" else None
    if bulk_scope == "group" and not bulk_group:
        raise HTTPException(422, detail="This service does not have one unambiguous Cybersecurity group scope")
    expires_at = aware(expires_at)
    now = utcnow()
    if expires_at <= now:
        raise HTTPException(422, detail="Expiration must be in the future")
    maximum_exception = int(configuration_for_service(db, finding.service).get("exception_max_days", "365"))
    if maximum_exception and expires_at > now + timedelta(days=maximum_exception):
        raise HTTPException(422, detail=f"Exception expiration cannot exceed {maximum_exception} days")
    pending = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.request_type == "exception",
        WorkflowRequest.policy_finding_id == finding.id,
        WorkflowRequest.status == "pending",
    ))
    if pending:
        raise HTTPException(409, detail="An exception request is already pending")
    workflow = WorkflowRequest(
        request_type="exception", service_id=finding.service_id,
        policy_finding_id=finding.id, requested_by_id=auth.user.id,
        justification=justification, requested_expires_at=expires_at,
        ticket=ticket.strip() or None, bulk_group_id=bulk_group.id if bulk_group else None,
    )
    db.add(workflow)
    db.flush()
    record_audit(db, auth, "exception.requested", "workflow_request", workflow.id,
                 service_id=finding.service_id, finding_type="configuration")
    db.commit()
    return RedirectResponse(f"/services/{finding.service.service_key}", status_code=status.HTTP_303_SEE_OTHER)


@app.post("/findings/{finding_id}/poams")
def request_poam(
    finding_id: int, title: str = Form(min_length=3), remediation: str = Form(min_length=3),
    due_date: str = Form(default=""), ticket: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("poam.request", scoped=True)),
):
    """Create a pending vulnerability POA&M entry for Cybersecurity review."""
    check_csrf(auth, csrf_token)
    finding = db.scalar(select(Finding).where(Finding.id == finding_id).options(selectinload(Finding.service).selectinload(Service.groups)))
    if not finding:
        raise HTTPException(404)
    due_date = optional_form_datetime(due_date, "POA&M due date")
    if due_date is not None:
        if due_date <= utcnow():
            raise HTTPException(422, detail="POA&M due date must be in the future")
    pending = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.request_type == "poam", WorkflowRequest.finding_id == finding.id,
        WorkflowRequest.status == "pending",
    ))
    if pending:
        raise HTTPException(409, detail="A POA&M request is already pending")
    entry = PoamEntry(
        service_id=finding.service_id, finding_id=finding.id, item_type="vulnerability",
        title=title.strip(), description=f"Vulnerability finding {finding.cve}",
        remediation=remediation.strip(), due_date=due_date, ticket=ticket or None,
        created_by_id=auth.user.id,
    )
    db.add(entry)
    db.flush()
    workflow = WorkflowRequest(
        request_type="poam", service_id=finding.service_id, finding_id=finding.id,
        poam_id=entry.id, requested_by_id=auth.user.id,
        justification=f"{entry.title}\n{entry.description}\nRemediation: {entry.remediation}",
        requested_expires_at=due_date, ticket=ticket or None,
    )
    db.add(workflow)
    db.flush()
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action="requested", note="Submitted for Cybersecurity approval"))
    record_audit(db, auth, "poam.requested", "workflow_request", workflow.id, service_id=finding.service_id)
    db.commit()
    return RedirectResponse(f"/services/{finding.service.service_key}", status_code=status.HTTP_303_SEE_OTHER)


@app.post("/policy-findings/{policy_finding_id}/poams")
def request_policy_poam(
    policy_finding_id: int, title: str = Form(min_length=3), remediation: str = Form(min_length=3),
    due_date: str = Form(default=""), ticket: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    finding = db.scalar(select(PolicyFinding).where(PolicyFinding.id == policy_finding_id).options(
        selectinload(PolicyFinding.service).selectinload(Service.groups)
    ))
    if not finding:
        raise HTTPException(404)
    if not auth.has("poam.request", finding.service_id):
        raise HTTPException(403, detail="Permission denied")
    due_date = optional_form_datetime(due_date, "POA&M due date")
    if due_date is not None and due_date <= utcnow():
        raise HTTPException(422, detail="POA&M due date must be in the future")
    pending = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.request_type == "poam",
        WorkflowRequest.policy_finding_id == finding.id,
        WorkflowRequest.status == "pending",
    ))
    if pending:
        raise HTTPException(409, detail="A POA&M request is already pending")
    source_detail = finding.description or finding.title
    detail = f"Configuration finding {finding.finding}"
    if source_detail:
        detail = f"{detail}: {source_detail}"
    entry = PoamEntry(
        service_id=finding.service_id, policy_finding_id=finding.id,
        item_type="configuration", title=title.strip(), description=detail,
        remediation=remediation.strip(), due_date=due_date,
        ticket=ticket.strip() or None, created_by_id=auth.user.id,
    )
    db.add(entry)
    db.flush()
    workflow = WorkflowRequest(
        request_type="poam", service_id=finding.service_id,
        policy_finding_id=finding.id, poam_id=entry.id,
        requested_by_id=auth.user.id,
        justification=f"{entry.title}\n{entry.description}\nRemediation: {entry.remediation}",
        requested_expires_at=due_date, ticket=entry.ticket,
    )
    db.add(workflow)
    db.flush()
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id,
                       action="requested", note="Submitted for Cybersecurity approval"))
    record_audit(db, auth, "poam.requested", "workflow_request", workflow.id,
                 service_id=finding.service_id, finding_type="configuration")
    db.commit()
    return RedirectResponse(f"/services/{finding.service.service_key}", status_code=status.HTTP_303_SEE_OTHER)


@app.get("/poam", response_class=HTMLResponse)
def poam_page(request: Request, db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    allowed = auth.accessible_service_ids("service.view")
    service_query = select(Service).options(selectinload(Service.groups)).order_by(Service.name)
    if allowed is not None:
        service_query = service_query.where(Service.id.in_(allowed))
    visible_services = db.scalars(service_query).all()
    visible_ids = [service.id for service in visible_services]
    counts = {
        service_id: {"total": 0, "active": 0, "pending": 0, "overdue": 0}
        for service_id in visible_ids
    }
    if visible_ids:
        for entry in db.scalars(select(PoamEntry).where(PoamEntry.service_id.in_(visible_ids))):
            summary = counts[entry.service_id]
            summary["total"] += 1
            if entry.status == "active":
                summary["active"] += 1
                if entry.due_date and aware(entry.due_date) < utcnow():
                    summary["overdue"] += 1
            elif entry.status == "pending_approval":
                summary["pending"] += 1
    service_summaries = [{"service": service, **counts[service.id]} for service in visible_services]
    services = [service for service in visible_services if auth.has("poam.request", service.id)]
    return templates.TemplateResponse(request, "poam.html", page_context(
        auth, service_summaries=service_summaries, poam_services=services,
    ))


@app.get("/remediations", response_class=HTMLResponse)
def remediations_page(
    request: Request, tab: str = "poams", page: int = 1, page_size: int = 50,
    service: str = "", identifier: str = "", title: str = "", status_filter: str = "all",
    owner: str = "", severity: str = "", due_from: str = "", due_to: str = "",
    expiration_from: str = "", expiration_to: str = "", sort: str = "actionable", direction: str = "asc",
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    """Enterprise remediation oversight, scoped and paginated in the database."""
    if tab not in {"poams", "exceptions", "mitigations"}:
        raise HTTPException(422, detail="Unknown remediation tab")
    if page_size not in {25, 50, 100}:
        raise HTTPException(422, detail="Page size must be 25, 50, or 100")
    page = max(1, page)
    allowed = auth.accessible_service_ids("service.view")
    service_query = select(Service).order_by(Service.name)
    if allowed is not None:
        service_query = service_query.where(Service.id.in_(allowed))
    services = db.scalars(service_query).all()
    visible_ids = [item.id for item in services]
    service_filter_ids = set(visible_ids)
    service_term = service.strip()
    if service_term:
        lowered = service_term.casefold()
        service_filter_ids = {item.id for item in services if lowered in item.name.casefold() or lowered in item.service_key.casefold() or service_term == str(item.id)}
    now = utcnow()
    def parse_date(value: str):
        if not value:
            return None
        try:
            return datetime.fromisoformat(value).replace(tzinfo=timezone.utc)
        except ValueError:
            raise HTTPException(422, detail="Dates must use ISO format")
    parsed_due_from, parsed_due_to = parse_date(due_from), parse_date(due_to)
    parsed_exp_from, parsed_exp_to = parse_date(expiration_from), parse_date(expiration_to)
    identifier_term, title_term, owner_term = identifier.strip(), title.strip(), owner.strip()
    severity_term = severity.strip()
    status_values = {"all", "active", "pending_approval", "completed", "rejected", "overdue", "expired", "revoked"}
    if status_filter not in status_values:
        raise HTTPException(422, detail="Unknown remediation status")
    def apply_poam_filters(query, item_type=None):
        query = query.where(PoamEntry.service_id.in_(service_filter_ids or {-1}))
        if item_type:
            query = query.where(PoamEntry.item_type == item_type)
        else:
            query = query.where(PoamEntry.item_type != "mitigation")
        if status_filter not in {"all", "overdue"}:
            query = query.where(PoamEntry.status == status_filter)
        if status_filter == "overdue":
            query = query.where(PoamEntry.status == "active", PoamEntry.due_date.is_not(None), PoamEntry.due_date < now)
        if title_term:
            query = query.where(PoamEntry.title.ilike(f"%{title_term}%"))
        if owner_term:
            query = query.join(User, User.id == PoamEntry.created_by_id).where(User.display_name.ilike(f"%{owner_term}%"))
        if parsed_due_from:
            query = query.where(PoamEntry.due_date >= parsed_due_from)
        if parsed_due_to:
            query = query.where(PoamEntry.due_date <= parsed_due_to)
        if identifier_term or severity_term:
            query = query.outerjoin(Finding, Finding.id == PoamEntry.finding_id).outerjoin(PolicyFinding, PolicyFinding.id == PoamEntry.policy_finding_id)
            if identifier_term:
                query = query.where(or_(Finding.cve.ilike(f"%{identifier_term}%"), PolicyFinding.finding.ilike(f"%{identifier_term}%")))
            if severity_term:
                query = query.where(or_(Finding.severity == severity_term, PolicyFinding.severity == severity_term))
        return query
    def poam_order(query):
        descending = direction.lower() == "desc"
        if sort in {"due", "actionable"}:
            return query.order_by(PoamEntry.due_date.desc() if descending else PoamEntry.due_date.asc(), PoamEntry.created_at.desc())
        if sort == "status":
            return query.order_by(PoamEntry.status.desc() if descending else PoamEntry.status.asc(), PoamEntry.created_at.desc())
        return query.order_by(PoamEntry.created_at.asc() if not descending else PoamEntry.created_at.desc())
    poams, mitigations, exceptions = [], [], []
    total_items = 0
    if tab in {"poams", "mitigations"}:
        base = apply_poam_filters(select(PoamEntry).options(
            selectinload(PoamEntry.service), selectinload(PoamEntry.finding), selectinload(PoamEntry.policy_finding),
            selectinload(PoamEntry.created_by), selectinload(PoamEntry.approved_by),
        ), "mitigation" if tab == "mitigations" else None)
        total_items = db.scalar(select(func.count()).select_from(base.order_by(None).subquery())) or 0
        rows = db.scalars(poam_order(base).offset((page - 1) * page_size).limit(page_size)).all()
        if tab == "mitigations":
            mitigations = rows
        else:
            poams = rows
    else:
        def exception_state_clause(model):
            if status_filter == "active": return and_(model.revoked_at.is_(None), model.starts_at <= now, model.expires_at > now)
            if status_filter == "expired": return and_(model.revoked_at.is_(None), model.expires_at <= now)
            if status_filter == "revoked": return model.revoked_at.is_not(None)
            if status_filter == "pending_approval": return false()
            return true()
        vul_query = select(ExceptionRecord).join(Finding).where(Finding.service_id.in_(service_filter_ids or {-1}), exception_state_clause(ExceptionRecord)).options(selectinload(ExceptionRecord.finding).selectinload(Finding.service)).order_by(ExceptionRecord.created_at.desc())
        cfg_query = select(PolicyExceptionRecord).join(PolicyFinding).where(PolicyFinding.service_id.in_(service_filter_ids or {-1}), exception_state_clause(PolicyExceptionRecord)).options(selectinload(PolicyExceptionRecord.policy_finding).selectinload(PolicyFinding.service)).order_by(PolicyExceptionRecord.created_at.desc())
        if identifier_term:
            vul_query = vul_query.where(Finding.cve.ilike(f"%{identifier_term}%")); cfg_query = cfg_query.where(PolicyFinding.finding.ilike(f"%{identifier_term}%"))
        if severity_term:
            vul_query = vul_query.where(Finding.severity == severity_term); cfg_query = cfg_query.where(PolicyFinding.severity == severity_term)
        if owner_term:
            vul_query = vul_query.where(ExceptionRecord.approved_by.ilike(f"%{owner_term}%")); cfg_query = cfg_query.where(PolicyExceptionRecord.approved_by.ilike(f"%{owner_term}%"))
        if parsed_exp_from:
            vul_query = vul_query.where(ExceptionRecord.expires_at >= parsed_exp_from); cfg_query = cfg_query.where(PolicyExceptionRecord.expires_at >= parsed_exp_from)
        if parsed_exp_to:
            vul_query = vul_query.where(ExceptionRecord.expires_at <= parsed_exp_to); cfg_query = cfg_query.where(PolicyExceptionRecord.expires_at <= parsed_exp_to)
        vul_total = db.scalar(select(func.count()).select_from(vul_query.order_by(None).subquery())) or 0
        cfg_total = db.scalar(select(func.count()).select_from(cfg_query.order_by(None).subquery())) or 0
        total_items = vul_total + cfg_total
        # Fetch only the current page from each underlying record type; both are
        # authorized before filtering and are merged into one operational view.
        for record in db.scalars(vul_query.offset((page - 1) * page_size).limit(page_size)).all():
            finding = record.finding; state = "Revoked" if record.revoked_at else ("Expired" if aware(record.expires_at) < now else "Active")
            exceptions.append({"kind":"Vulnerability", "item":finding.cve, "service":finding.service, "severity":finding.severity, "status":state, "expires_at":record.expires_at, "days_remaining":max(0,(aware(record.expires_at)-now).days) if state == "Active" else 0, "approved_by":record.approved_by, "created_at":record.created_at, "justification":record.justification, "record_id":record.id, "revoke_href":f"/exceptions/{record.id}/revoke", "href":f"/services/{finding.service.service_key}/findings/{finding.id}"})
        for record in db.scalars(cfg_query.offset((page - 1) * page_size).limit(page_size)).all():
            finding = record.policy_finding; state = "Revoked" if record.revoked_at else ("Expired" if aware(record.expires_at) < now else "Active")
            exceptions.append({"kind":"Configuration", "item":finding.finding, "service":finding.service, "severity":finding.severity, "status":state, "expires_at":record.expires_at, "days_remaining":max(0,(aware(record.expires_at)-now).days) if state == "Active" else 0, "approved_by":record.approved_by, "created_at":record.created_at, "justification":record.justification, "record_id":record.id, "revoke_href":f"/policy-exceptions/{record.id}/revoke", "href":f"/services/{finding.service.service_key}?finding_state=exceptions&finding_type=configuration"})
        if status_filter in {"all", "pending_approval"}:
            pending_query = select(WorkflowRequest).where(WorkflowRequest.service_id.in_(service_filter_ids or {-1}), WorkflowRequest.request_type == "exception", WorkflowRequest.status == "pending").options(selectinload(WorkflowRequest.service), selectinload(WorkflowRequest.finding), selectinload(WorkflowRequest.policy_finding)).order_by(WorkflowRequest.created_at.desc())
            if identifier_term:
                pending_query = pending_query.where(or_(WorkflowRequest.finding_id.in_(select(Finding.id).where(Finding.cve.ilike(f"%{identifier_term}%"))), WorkflowRequest.policy_finding_id.in_(select(PolicyFinding.id).where(PolicyFinding.finding.ilike(f"%{identifier_term}%")))))
            pending_total = db.scalar(select(func.count()).select_from(pending_query.order_by(None).subquery())) or 0
            total_items += pending_total
            for workflow in db.scalars(pending_query.offset((page - 1) * page_size).limit(page_size)).all():
                target = workflow.finding or workflow.policy_finding
                if not target:
                    continue
                is_vulnerability = bool(workflow.finding)
                exceptions.append({"kind": "Vulnerability" if is_vulnerability else "Configuration", "item": target.cve if is_vulnerability else target.finding, "service": workflow.service, "severity": target.severity, "status": "Pending", "expires_at": workflow.requested_expires_at, "days_remaining": 0, "approved_by": "Pending review", "created_at": workflow.created_at, "justification": workflow.justification, "record_id": workflow.id, "href": f"/services/{workflow.service.service_key}/findings/{target.id}" if is_vulnerability else f"/services/{workflow.service.service_key}?finding_state=exceptions&finding_type=configuration"})
    active_poams = db.scalar(select(func.count(PoamEntry.id)).where(PoamEntry.service_id.in_(service_filter_ids or {-1}), PoamEntry.item_type != "mitigation", PoamEntry.status == "active")) or 0
    overdue_poams = db.scalar(select(func.count(PoamEntry.id)).where(PoamEntry.service_id.in_(service_filter_ids or {-1}), PoamEntry.item_type != "mitigation", PoamEntry.status == "active", PoamEntry.due_date < now)) or 0
    active_exceptions = db.scalar(select(func.count(ExceptionRecord.id)).join(Finding).where(Finding.service_id.in_(service_filter_ids or {-1}), ExceptionRecord.revoked_at.is_(None), ExceptionRecord.starts_at <= now, ExceptionRecord.expires_at > now)) or 0
    active_exceptions += db.scalar(select(func.count(PolicyExceptionRecord.id)).join(PolicyFinding).where(PolicyFinding.service_id.in_(service_filter_ids or {-1}), PolicyExceptionRecord.revoked_at.is_(None), PolicyExceptionRecord.starts_at <= now, PolicyExceptionRecord.expires_at > now)) or 0
    expiring_exceptions = db.scalar(select(func.count(ExceptionRecord.id)).join(Finding).where(Finding.service_id.in_(service_filter_ids or {-1}), ExceptionRecord.revoked_at.is_(None), ExceptionRecord.starts_at <= now, ExceptionRecord.expires_at > now, ExceptionRecord.expires_at <= now + timedelta(days=30))) or 0
    expiring_exceptions += db.scalar(select(func.count(PolicyExceptionRecord.id)).join(PolicyFinding).where(PolicyFinding.service_id.in_(service_filter_ids or {-1}), PolicyExceptionRecord.revoked_at.is_(None), PolicyExceptionRecord.starts_at <= now, PolicyExceptionRecord.expires_at > now, PolicyExceptionRecord.expires_at <= now + timedelta(days=30))) or 0
    active_mitigations = db.scalar(select(func.count(PoamEntry.id)).where(PoamEntry.service_id.in_(service_filter_ids or {-1}), PoamEntry.item_type == "mitigation", PoamEntry.status == "active")) or 0
    page_count = max(1, (total_items + page_size - 1) // page_size)
    page = min(page, page_count)
    return templates.TemplateResponse(request, "remediations.html", page_context(auth,
        tab=tab, services=services, poams=poams, exceptions=exceptions, mitigations=mitigations, now=now,
        poam_services=[item for item in services if auth.has("poam.request", item.id)],
        summary={"poams": active_poams, "poams_overdue": overdue_poams, "exceptions": active_exceptions, "exceptions_soon": expiring_exceptions, "mitigations": active_mitigations},
        filters={"service": service, "identifier": identifier, "title": title, "status_filter": status_filter, "owner": owner, "severity": severity, "due_from": due_from, "due_to": due_to, "expiration_from": expiration_from, "expiration_to": expiration_to, "sort": sort, "direction": direction, "page_size": page_size},
        page=page, page_count=page_count, total_items=total_items,
    ))


@app.get("/poam/services/{service_key}", response_class=HTMLResponse)
def service_poam_page(
    service_key: str, request: Request, status_filter: str = "all", sort_by: str = "newest", db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_user),
):
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(selectinload(Service.groups)))
    if not service:
        raise HTTPException(404)
    if not auth.has("service.view", service.id):
        raise HTTPException(403, detail="Permission denied for this service")
    # Keep the legacy URL as a compatibility alias; the canonical experience
    # is the POA&M tab embedded in the service workspace.
    return RedirectResponse(f"/services/{service.service_key}?poam=true&status_filter={status_filter}&sort_by={sort_by}", status_code=303)
    entries = db.scalars(select(PoamEntry).where(PoamEntry.service_id == service.id).options(
        selectinload(PoamEntry.finding), selectinload(PoamEntry.policy_finding), selectinload(PoamEntry.created_by),
        selectinload(PoamEntry.approved_by),
    ).order_by(PoamEntry.created_at.desc())).all()
    if status_filter not in {"all", "active", "pending_approval", "overdue", "completed", "rejected"}:
        raise HTTPException(422, detail="Unknown POA&M filter")
    now = utcnow()
    if status_filter == "overdue":
        entries = [entry for entry in entries if entry.status == "active" and entry.due_date and aware(entry.due_date) < now]
    elif status_filter != "all":
        entries = [entry for entry in entries if entry.status == status_filter]
    if sort_by == "due":
        entries.sort(key=lambda entry: (aware(entry.due_date) if entry.due_date else datetime.max.replace(tzinfo=timezone.utc), entry.title.lower()))
    elif sort_by == "severity":
        severity_rank = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Negligible": 4}
        entries.sort(key=lambda entry: (severity_rank.get(
            entry.finding.severity if entry.finding else (entry.policy_finding.severity if entry.policy_finding else ""), 5
        ), entry.title.lower()))
    elif sort_by != "newest":
        raise HTTPException(422, detail="Unknown POA&M sort")
    return templates.TemplateResponse(request, "poam_service.html", page_context(
        auth, service=service, entries=entries, can_create_poam=auth.has("poam.request", service.id),
        status_filter=status_filter, sort_by=sort_by, now=now,
        overdue_entry_ids={entry.id for entry in entries if entry.status == "active" and entry.due_date and aware(entry.due_date) < now},
    ))


def poam_workbook(entries: list[PoamEntry]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "POA&M"
    sheet.append(["Service ID", "Service", "Type", "Title", "Description", "Finding", "Remediation",
                  "Due Date", "Status", "Ticket", "Created By", "Approved By", "Approved At"])
    for entry in entries:
        sheet.append([entry.service.service_key, entry.service.name, entry.item_type.replace("_", " ").title(),
                      entry.title, entry.description,
                      entry.finding.cve if entry.finding else (entry.policy_finding.finding if entry.policy_finding else None),
                      entry.remediation, excel_datetime(entry.due_date), entry.status.replace("_", " ").title(),
                      entry.ticket, entry.created_by.display_name,
                      entry.approved_by.display_name if entry.approved_by else None, excel_datetime(entry.approved_at)])
    for column in ("H", "M"):
        for cell in sheet[column][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
    format_sheet(sheet)
    return workbook


def scoped_poam_export_entries(db: Session, auth: AuthContext, service_id: int | None = None, status_filter: str = "all", service_term: str = "", due_overdue: bool = False) -> list[PoamEntry]:
    allowed = auth.accessible_service_ids("service.export")
    if allowed == set():
        raise HTTPException(403, detail="Permission denied")
    query = select(PoamEntry).options(selectinload(PoamEntry.service), selectinload(PoamEntry.finding), selectinload(PoamEntry.policy_finding),
        selectinload(PoamEntry.created_by), selectinload(PoamEntry.approved_by)).order_by(PoamEntry.service_id, PoamEntry.created_at.desc())
    if allowed is not None:
        query = query.where(PoamEntry.service_id.in_(allowed))
    if service_id is not None:
        query = query.where(PoamEntry.service_id == service_id)
    if service_term:
        matching = db.scalars(select(Service.id).where(or_(Service.name.ilike(f"%{service_term}%"), Service.service_key.ilike(f"%{service_term}%")))).all()
        query = query.where(PoamEntry.service_id.in_(matching or [-1]))
    if status_filter != "all":
        query = query.where(PoamEntry.status == status_filter)
    if due_overdue:
        query = query.where(PoamEntry.status == "active", PoamEntry.due_date.is_not(None), PoamEntry.due_date < utcnow())
    return db.scalars(query).all()


@app.get("/poam/export.xlsx")
def export_all_poams(service: str = "", status_filter: str = "all", due_overdue: bool = False, db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    return workbook_response(poam_workbook(scoped_poam_export_entries(db, auth, status_filter=status_filter, service_term=service, due_overdue=due_overdue)), f"cats-poam-{utcnow():%Y%m%d}.xlsx")


@app.get("/poam/services/{service_key}/export.xlsx")
def export_service_poams(service_key: str, db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if not service:
        raise HTTPException(404)
    if not auth.has("service.export", service.id):
        raise HTTPException(403, detail="Permission denied")
    return workbook_response(poam_workbook(scoped_poam_export_entries(db, auth, service.id)),
                             f"cats-{service.service_key}-poam-{utcnow():%Y%m%d}.xlsx")


@app.post("/poam")
def create_poam(
    service_id: int = Form(), item_type: str = Form(), title: str = Form(min_length=3),
    description: str = Form(min_length=3), remediation: str = Form(min_length=3),
    due_date: str = Form(default=""), ticket: str = Form(default=""),
    finding_id: str = Form(default=""), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    service = db.scalar(select(Service).where(Service.id == service_id).options(selectinload(Service.groups)))
    if not service or not auth.has("poam.request", service.id):
        raise HTTPException(403, detail="Permission denied for this service")
    if item_type not in {"vulnerability", "missing_evidence", "missing_requirement", "mitigation"}:
        raise HTTPException(422, detail="Unsupported POA&M item type")
    try:
        parsed_finding_id = int(finding_id) if finding_id.strip() else None
    except ValueError as exc:
        raise HTTPException(422, detail="Finding ID must be a whole number") from exc
    finding = db.get(Finding, parsed_finding_id) if parsed_finding_id else None
    if finding and finding.service_id != service.id:
        raise HTTPException(422, detail="Finding does not belong to the selected service")
    if item_type == "vulnerability" and not finding:
        raise HTTPException(422, detail="A vulnerability POA&M requires a finding ID")
    due_date = optional_form_datetime(due_date, "POA&M due date")
    if due_date is not None:
        if due_date <= utcnow():
            raise HTTPException(422, detail="POA&M due date must be in the future")
    entry = PoamEntry(
        service_id=service.id, finding_id=finding.id if finding else None, item_type=item_type,
        title=title.strip(), description=description.strip(), remediation=remediation.strip(),
        due_date=due_date, ticket=ticket.strip() or None, created_by_id=auth.user.id,
    )
    db.add(entry)
    db.flush()
    workflow = WorkflowRequest(
        request_type="poam", service_id=service.id, finding_id=entry.finding_id,
        poam_id=entry.id, requested_by_id=auth.user.id,
        justification=f"{entry.title}\n{entry.description}\nRemediation: {entry.remediation}",
        requested_expires_at=due_date, ticket=entry.ticket,
    )
    db.add(workflow)
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action="requested", note="Submitted for Cybersecurity approval"))
    db.flush()
    record_audit(db, auth, "poam.requested", "poam_entry", entry.id, service_id=service.id, item_type=item_type)
    db.commit()
    return RedirectResponse(f"/poam/services/{service.service_key}", status_code=303)


@app.post("/services/{service_key}/mitigations")
def create_mitigation(
    service_key: str, title: str = Form(min_length=3), description: str = Form(min_length=3),
    remediation: str = Form(min_length=3), due_date: str = Form(default=""), ticket: str = Form(default=""),
    finding_id: str = Form(default=""), policy_finding_id: str = Form(default=""), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("poam.request", scoped=True)),
):
    check_csrf(auth, csrf_token)
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(selectinload(Service.groups)))
    if not service:
        raise HTTPException(404)
    def optional_id(raw: str, label: str):
        if not raw.strip():
            return None
        try:
            return int(raw)
        except ValueError as exc:
            raise HTTPException(422, detail=f"{label} must be a whole number") from exc
    parsed_finding_id = optional_id(finding_id, "Finding ID")
    parsed_policy_id = optional_id(policy_finding_id, "Configuration finding ID")
    if parsed_finding_id and (not (finding := db.get(Finding, parsed_finding_id)) or finding.service_id != service.id):
        raise HTTPException(422, detail="Finding does not belong to this service")
    if parsed_policy_id and (not (policy := db.get(PolicyFinding, parsed_policy_id)) or policy.service_id != service.id):
        raise HTTPException(422, detail="Configuration finding does not belong to this service")
    due = optional_form_datetime(due_date, "Mitigation due date")
    if due and due <= utcnow():
        raise HTTPException(422, detail="Mitigation due date must be in the future")
    entry = PoamEntry(service_id=service.id, finding_id=parsed_finding_id, policy_finding_id=parsed_policy_id,
                      item_type="mitigation", title=title.strip(), description=description.strip(),
                      remediation=remediation.strip(), due_date=due, ticket=ticket.strip() or None,
                      created_by_id=auth.user.id)
    db.add(entry); db.flush()
    workflow = WorkflowRequest(request_type="poam", service_id=service.id, finding_id=parsed_finding_id,
                               policy_finding_id=parsed_policy_id, poam_id=entry.id, requested_by_id=auth.user.id,
                               justification=f"{entry.title}\n{entry.description}\nRemediation: {entry.remediation}",
                               requested_expires_at=due, ticket=entry.ticket)
    db.add(workflow); db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action="requested",
                                         note="Mitigation submitted for Cybersecurity approval"))
    record_audit(db, auth, "mitigation.requested", "poam_entry", entry.id, service_id=service.id)
    db.commit()
    return RedirectResponse(f"/services/{service.service_key}", status_code=303)


def scoped_poam(db: Session, auth: AuthContext, poam_id: int, permission: str = "service.view") -> PoamEntry:
    entry = db.scalar(select(PoamEntry).where(PoamEntry.id == poam_id).options(
        selectinload(PoamEntry.service).selectinload(Service.groups), selectinload(PoamEntry.finding), selectinload(PoamEntry.policy_finding),
        selectinload(PoamEntry.created_by), selectinload(PoamEntry.approved_by),
    ))
    if not entry:
        raise HTTPException(404)
    if not auth.has(permission, entry.service_id):
        raise HTTPException(403, detail="Permission denied for this POA&M")
    return entry


@app.get("/poam/entries/{poam_id}", response_class=HTMLResponse)
def poam_entry_page(poam_id: int, request: Request, db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    entry = scoped_poam(db, auth, poam_id)
    history = db.scalars(select(PoamHistory).where(PoamHistory.poam_id == entry.id).options(
        selectinload(PoamHistory.actor)
    ).order_by(PoamHistory.created_at.desc())).all()
    pending_change = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.poam_id == entry.id, WorkflowRequest.status == "pending",
        WorkflowRequest.request_type.in_(["poam_update", "poam_complete", "poam_reopen"]),
    ))
    return templates.TemplateResponse(request, "poam_entry.html", page_context(
        auth, entry=entry, history=history, pending_change=pending_change,
        can_change=auth.has("poam.request", entry.service_id), now=utcnow(),
        is_overdue=bool(entry.status == "active" and entry.due_date and aware(entry.due_date) < utcnow()),
    ))


def create_poam_change(db: Session, auth: AuthContext, entry: PoamEntry, change_type: str, proposed: dict, justification: str):
    pending = db.scalar(select(WorkflowRequest.id).where(
        WorkflowRequest.poam_id == entry.id, WorkflowRequest.status == "pending",
        WorkflowRequest.request_type.in_(["poam_update", "poam_complete", "poam_reopen"]),
    ))
    if pending:
        raise HTTPException(409, detail="A POA&M change is already pending approval")
    workflow = WorkflowRequest(
        request_type=f"poam_{change_type}", service_id=entry.service_id, finding_id=entry.finding_id,
        policy_finding_id=entry.policy_finding_id,
        poam_id=entry.id, requested_by_id=auth.user.id, justification=justification,
    )
    db.add(workflow)
    db.flush()
    db.add(PoamChangeRequest(workflow_id=workflow.id, poam_id=entry.id, change_type=change_type, proposed=proposed))
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action=f"{change_type}_requested", note=justification, detail=proposed))
    record_audit(db, auth, f"poam.{change_type}_requested", "poam_entry", entry.id, service_id=entry.service_id)
    db.commit()


@app.post("/poam/entries/{poam_id}/update")
def request_poam_update(
    poam_id: int, title: str = Form(min_length=3), description: str = Form(min_length=3),
    remediation: str = Form(min_length=3), due_date: str = Form(default=""),
    ticket: str = Form(default=""), reason: str = Form(min_length=3), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    entry = scoped_poam(db, auth, poam_id, "poam.request")
    due_date = optional_form_datetime(due_date, "POA&M due date")
    proposed = {"title": title.strip(), "description": description.strip(), "remediation": remediation.strip(),
                "due_date": due_date.isoformat() if due_date else None, "ticket": ticket.strip() or None}
    create_poam_change(db, auth, entry, "update", proposed, reason.strip())
    return RedirectResponse(f"/poam/entries/{entry.id}", status_code=303)


@app.post("/poam/entries/{poam_id}/complete")
def request_poam_completion(
    poam_id: int, closure_note: str = Form(min_length=3), evidence_reference: str = Form(min_length=3),
    csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    entry = scoped_poam(db, auth, poam_id, "poam.request")
    if entry.status != "active":
        raise HTTPException(409, detail="Only active POA&M entries can be completed")
    proposed = {"closure_note": closure_note.strip(), "evidence_reference": evidence_reference.strip()}
    create_poam_change(db, auth, entry, "complete", proposed, closure_note.strip())
    return RedirectResponse(f"/poam/entries/{entry.id}", status_code=303)


@app.post("/poam/entries/{poam_id}/reopen")
def request_poam_reopen(
    poam_id: int, reason: str = Form(min_length=3), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    entry = scoped_poam(db, auth, poam_id, "poam.request")
    if entry.status != "completed":
        raise HTTPException(409, detail="Only completed POA&M entries can be reopened")
    create_poam_change(db, auth, entry, "reopen", {}, reason.strip())
    return RedirectResponse(f"/poam/entries/{entry.id}", status_code=303)


@app.post("/poam/entries/{poam_id}/comments")
def add_poam_comment(
    poam_id: int, comment: str = Form(min_length=3), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    entry = scoped_poam(db, auth, poam_id)
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action="comment", note=comment.strip()))
    record_audit(db, auth, "poam.comment_added", "poam_entry", entry.id, service_id=entry.service_id)
    db.commit()
    return RedirectResponse(f"/poam/entries/{entry.id}", status_code=303)


def _missing_evidence_key(row: dict) -> tuple[str, str, str]:
    return (str(row.get("type") or "Other").strip().lower(), str(row.get("item") or row.get("reference") or row.get("name") or row.get("image") or row.get("chart") or row.get("resource") or "").strip(), str(row.get("source_file") or row.get("source_path") or row.get("filepath") or row.get("file") or row.get("chart_path") or row.get("values_file") or row.get("template") or row.get("provenance") or row.get("source") or row.get("resource_path") or "").strip())


@app.post("/services/{service_key}/missing-evidence/remove")
def remove_missing_evidence(
    service_key: str, evidence_type: str = Form(), item: str = Form(), source_file: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(selectinload(Service.executions)))
    if not service:
        raise HTTPException(404)
    if not auth.has("evidence.remove", service.id):
        raise HTTPException(403, detail="Permission denied")
    latest = max(service.executions, key=lambda execution: aware(execution.scanned_at), default=None)
    if not latest or not isinstance(latest.raw_payload, dict):
        raise HTTPException(409, detail="No current evidence observation is available")
    payload = dict(latest.raw_payload)
    overview = dict(payload.get("service_overview") or {})
    target = (str(evidence_type).strip().lower(), str(item).strip(), str(source_file).strip())
    removed = False
    # The assessment-level row is a derived observation used only when an
    # incomplete execution has no concrete evidence item.  It has no list
    # entry to delete, so removing it resolves the current execution's
    # incomplete state.  A future execution is authoritative and can recreate
    # the row if it is incomplete again.
    if target[1].casefold() == "assessment" and not (
        overview.get("missing_evidence") or overview.get("evidence")
        or payload.get("skipped_images") or payload.get("skipped_charts")
    ):
        latest.complete = True
        payload["incomplete"] = False
        removed = True
    for container, field in ((overview, "missing_evidence"), (overview, "evidence"), (payload, "skipped_images"), (payload, "skipped_charts")):
        values = container.get(field)
        if not isinstance(values, list):
            continue
        kept = []
        for value in values:
            if isinstance(value, dict):
                candidate = _missing_evidence_key({"type": value.get("type") or ("Image" if field == "skipped_images" else "Chart" if field == "skipped_charts" else "Other"), **value})
            else:
                candidate = ("image" if field == "skipped_images" else "chart" if field == "skipped_charts" else "other", str(value).split(" :: ", 1)[0].strip(), "")
            if candidate == target or (candidate[0] == target[0] and candidate[1] == target[1] and not target[2]):
                removed = True
            else:
                kept.append(value)
        container[field] = kept
    if not removed:
        raise HTTPException(404, detail="That evidence observation is no longer current")
    payload["service_overview"] = overview
    latest.raw_payload = payload
    record_audit(db, auth, "missing_evidence.removed", "service", service.id,
                 service_id=service.id, evidence_type=evidence_type, item=item, source_file=source_file)
    db.commit()
    return RedirectResponse(f"/services/{service.service_key}?overview=true", status_code=303)


def _safe_remediation_return(value: str, fallback: str) -> str:
    """Accept only local remediation/service paths supplied by our forms."""
    raw = (value or "").strip()
    if not raw or raw.startswith("//"):
        return fallback
    parsed = urllib.parse.urlsplit(raw)
    if parsed.scheme or parsed.netloc or not parsed.path.startswith(("/remediations", "/services/")):
        return fallback
    if parsed.path.startswith("/services/") and "remediations=true" not in parsed.query:
        return fallback
    return raw


@app.post("/poam/entries/{poam_id}/revoke")
def revoke_poam_entry(poam_id: int, return_to: str = Form(default=""), csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    check_csrf(auth, csrf_token)
    entry = scoped_poam(db, auth, poam_id, "poam.review")
    if entry.status not in {"active", "pending"}:
        raise HTTPException(409, detail="Only active or pending remediation entries can be revoked")
    entry.status = "revoked"
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action="revoked", note="Revoked by authorized reviewer"))
    record_audit(db, auth, "poam.revoked", "poam_entry", entry.id, service_id=entry.service_id)
    db.commit()
    fallback = f"/services/{entry.service.service_key}?remediations=true&tab={'mitigations' if entry.item_type == 'mitigation' else 'poams'}"
    return RedirectResponse(_safe_remediation_return(return_to, fallback), status_code=303)


@app.post("/poam/entries/{poam_id}/close")
def close_poam_entry(poam_id: int, return_to: str = Form(default=""), closure_note: str = Form(default="Closed by authorized reviewer"), csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    check_csrf(auth, csrf_token)
    entry = scoped_poam(db, auth, poam_id, "poam.review")
    if entry.status != "active":
        raise HTTPException(409, detail="Only active remediation entries can be closed")
    entry.status = "completed"
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action="completed", note=closure_note.strip()[:2000] or "Closed by authorized reviewer"))
    record_audit(db, auth, "poam.completed", "poam_entry", entry.id, service_id=entry.service_id)
    db.commit()
    fallback = f"/services/{entry.service.service_key}?remediations=true&tab={'mitigations' if entry.item_type == 'mitigation' else 'poams'}"
    return RedirectResponse(_safe_remediation_return(return_to, fallback), status_code=303)


@app.post("/exceptions/{exception_id}/revoke")
def revoke_exception(
    exception_id: int, return_to: str = Form(default=""), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("exception.revoke", scoped=True)),
):
    check_csrf(auth, csrf_token)
    record = db.get(ExceptionRecord, exception_id)
    if not record:
        raise HTTPException(404)
    if record.revoked_at is not None:
        raise HTTPException(409, detail="Exception is already revoked")
    record.revoked_at = utcnow()
    record_audit(db, auth, "exception.revoked", "exception", record.id, service_id=record.finding.service_id)
    db.commit()
    fallback = f"/services/{record.finding.service.service_key}?remediations=true&tab=exceptions"
    return RedirectResponse(_safe_remediation_return(return_to, fallback), status_code=303)


@app.post("/policy-exceptions/{policy_exception_id}/revoke")
def revoke_policy_exception(
    policy_exception_id: int, return_to: str = Form(default=""), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    record = db.scalar(select(PolicyExceptionRecord).where(
        PolicyExceptionRecord.id == policy_exception_id
    ).options(selectinload(PolicyExceptionRecord.policy_finding).selectinload(PolicyFinding.service)))
    if not record:
        raise HTTPException(404)
    if not auth.has("exception.revoke", record.policy_finding.service_id):
        raise HTTPException(403, detail="Permission denied")
    if record.revoked_at is not None:
        raise HTTPException(409, detail="Exception is already revoked")
    record.revoked_at = utcnow()
    record_audit(db, auth, "exception.revoked", "policy_exception", record.id,
                 service_id=record.policy_finding.service_id, finding_type="configuration")
    db.commit()
    fallback = f"/services/{record.policy_finding.service.service_key}?remediations=true&tab=exceptions"
    return RedirectResponse(_safe_remediation_return(return_to, fallback), status_code=303)


@app.post("/services/{service_key}/archive")
def request_archive(
    service_key: str, reason: str = Form(min_length=3), ticket: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("archive.request", scoped=True)),
):
    check_csrf(auth, csrf_token)
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(selectinload(Service.archive_events)))
    if not service:
        raise HTTPException(404)
    if archive_state(service):
        raise HTTPException(409, detail="Service is already archived")
    pending = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.request_type == "archive", WorkflowRequest.service_id == service.id,
        WorkflowRequest.status == "pending",
    ))
    if pending:
        raise HTTPException(409, detail="An archive request is already pending")
    workflow = WorkflowRequest(request_type="archive", service_id=service.id,
                               requested_by_id=auth.user.id, justification=reason, ticket=ticket or None)
    db.add(workflow)
    db.flush()
    record_audit(db, auth, "archive.requested", "workflow_request", workflow.id, service_id=service.id)
    db.commit()
    return RedirectResponse(f"/services/{service.service_key}", status_code=303)


def _service_image_for_action(db: Session, service_key: str, image_id: int) -> tuple[Service, ServiceImage]:
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(selectinload(Service.images)))
    image = db.scalar(select(ServiceImage).where(ServiceImage.id == image_id, ServiceImage.service_id == service.id if service else False)) if service else None
    if not service or not image:
        raise HTTPException(404, detail="Container image was not found")
    return service, image


@app.post("/services/{service_key}/images/{image_id}/remove")
def request_image_removal(
    service_key: str, image_id: int, reason: str = Form(min_length=3), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("archive.request", scoped=True)),
):
    check_csrf(auth, csrf_token)
    service, image = _service_image_for_action(db, service_key, image_id)
    if image.lifecycle_status != "active":
        raise HTTPException(409, detail="Only an active image can be removed")
    pending = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.service_image_id == image.id,
        WorkflowRequest.request_type == "image_remove",
        WorkflowRequest.status == "pending",
    ))
    if pending:
        raise HTTPException(409, detail="An image-removal request is already pending")
    workflow = WorkflowRequest(
        request_type="image_remove", service_id=service.id, service_image_id=image.id,
        requested_by_id=auth.user.id, justification=reason.strip(),
    )
    db.add(workflow)
    db.flush()
    record_audit(db, auth, "image_removal.requested", "workflow_request", workflow.id,
                 service_id=service.id, image_id=image.id, image=image.image_reference, reason=reason.strip())
    db.commit()
    return RedirectResponse("/requests", status_code=303)


@app.post("/services/{service_key}/images/{image_id}/replace")
def request_image_replacement(
    service_key: str, image_id: int, replacement_reference: str = Form(min_length=1), reason: str = Form(min_length=3),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("archive.request", scoped=True)),
):
    check_csrf(auth, csrf_token)
    service, image = _service_image_for_action(db, service_key, image_id)
    if image.lifecycle_status != "active":
        raise HTTPException(409, detail="Only an active image can be replaced")
    replacement_reference = replacement_reference.strip()
    if replacement_reference == image.image_reference:
        raise HTTPException(422, detail="Replacement image must be different from the current image")
    pending = db.scalar(select(WorkflowRequest).where(
        WorkflowRequest.service_image_id == image.id,
        WorkflowRequest.request_type == "image_replace",
        WorkflowRequest.status == "pending",
    ))
    if pending:
        raise HTTPException(409, detail="An image-replacement request is already pending")
    workflow = WorkflowRequest(
        request_type="image_replace", service_id=service.id, service_image_id=image.id,
        replacement_reference=replacement_reference, requested_by_id=auth.user.id,
        justification=reason.strip(),
    )
    db.add(workflow)
    db.flush()
    record_audit(db, auth, "image_replacement.requested", "workflow_request", workflow.id,
                 service_id=service.id, image_id=image.id, image=image.image_reference,
                 replacement=replacement_reference, reason=reason.strip())
    db.commit()
    return RedirectResponse("/requests", status_code=303)


@app.post("/services/{service_key}/restore")
def restore_service(
    service_key: str, reason: str = Form(min_length=3), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("service.restore", scoped=True)),
):
    check_csrf(auth, csrf_token)
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(selectinload(Service.archive_events)))
    if not service:
        raise HTTPException(404)
    if not archive_state(service):
        raise HTTPException(409, detail="Service is not archived")
    db.add(ServiceArchiveEvent(service=service, action="restore", reason=reason, performed_by=auth.user.username))
    record_audit(db, auth, "service.restored", "service", service.id, service_id=service.id, reason=reason)
    db.commit()
    return RedirectResponse(f"/services/{service_key}", status_code=303)


@app.post("/services/{service_key}/delete")
def delete_service(
    service_key: str, confirmation: str = Form(), reason: str = Form(min_length=3),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("service.delete", scoped=True)),
):
    check_csrf(auth, csrf_token)
    if os.getenv("ALLOW_SERVICE_DELETE", "false").lower() != "true":
        raise HTTPException(403, detail="Permanent service deletion is disabled")
    service = db.scalar(select(Service).where(Service.service_key == service_key).options(
        selectinload(Service.archive_events)
    ))
    if not service:
        raise HTTPException(404)
    if not archive_state(service):
        raise HTTPException(409, detail="Archive the service before permanently deleting it")
    if confirmation != service.service_key:
        raise HTTPException(422, detail="Service ID confirmation did not match")
    finding_ids = list(db.scalars(select(Finding.id).where(Finding.service_id == service.id)))
    policy_finding_ids = list(db.scalars(select(PolicyFinding.id).where(PolicyFinding.service_id == service.id)))
    poam_ids = list(db.scalars(select(PoamEntry.id).where(PoamEntry.service_id == service.id)))
    if poam_ids:
        db.execute(delete(PoamChangeRequest).where(PoamChangeRequest.poam_id.in_(poam_ids)))
        db.execute(delete(PoamHistory).where(PoamHistory.poam_id.in_(poam_ids)))
    db.execute(delete(WorkflowRequest).where(WorkflowRequest.service_id == service.id))
    db.execute(delete(PoamEntry).where(PoamEntry.service_id == service.id))
    if finding_ids:
        db.execute(delete(FindingObservation).where(FindingObservation.finding_id.in_(finding_ids)))
        db.execute(delete(ExceptionRecord).where(ExceptionRecord.finding_id.in_(finding_ids)))
        db.execute(delete(Finding).where(Finding.id.in_(finding_ids)))
    if policy_finding_ids:
        db.execute(delete(PolicyExceptionRecord).where(PolicyExceptionRecord.policy_finding_id.in_(policy_finding_ids)))
        db.execute(delete(PolicyFinding).where(PolicyFinding.id.in_(policy_finding_ids)))
    db.execute(delete(UserRoleAssignment).where(UserRoleAssignment.service_id == service.id))
    db.execute(delete(RemediationExecution).where(RemediationExecution.service_id == service.id))
    db.execute(delete(PatchExecution).where(PatchExecution.service_id == service.id))
    db.execute(delete(Execution).where(Execution.service_id == service.id))
    db.execute(delete(ServiceImage).where(ServiceImage.service_id == service.id))
    db.execute(delete(ServiceArchiveEvent).where(ServiceArchiveEvent.service_id == service.id))
    db.add(ServiceDeletionAudit(service_key=service.service_key, service_name=service.name,
                                reason=reason, deleted_by=auth.user.username))
    record_audit(db, auth, "service.deleted", "service", service.id, service_key=service.service_key, reason=reason)
    db.delete(service)
    db.commit()
    return RedirectResponse("/?archived=true", status_code=303)


@app.get("/requests", response_class=HTMLResponse)
def requests_page(request: Request, request_type: str = "all", db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    allowed = auth.accessible_service_ids("service.view")
    query = select(WorkflowRequest).options(
        selectinload(WorkflowRequest.service), selectinload(WorkflowRequest.finding),
        selectinload(WorkflowRequest.finding).selectinload(Finding.exceptions),
        selectinload(WorkflowRequest.policy_finding).selectinload(PolicyFinding.exceptions),
        selectinload(WorkflowRequest.requested_by), selectinload(WorkflowRequest.reviewed_by),
        selectinload(WorkflowRequest.service_image),
    ).order_by(WorkflowRequest.created_at.desc())
    if allowed is not None:
        query = query.where(WorkflowRequest.service_id.in_(allowed))
    if request_type == "poam":
        query = query.where(WorkflowRequest.request_type.like("poam%"))
    elif request_type in {"exception", "archive"}:
        query = query.where(WorkflowRequest.request_type == request_type)
    elif request_type != "all":
        raise HTTPException(422, detail="Unknown request filter")
    workflows = db.scalars(query).all()
    now = utcnow()
    workflow_statuses = {}
    for item in workflows:
        display_status = item.status
        exception_subject = item.finding or item.policy_finding
        if item.request_type == "exception" and item.status == "approved" and exception_subject:
            matching = [exception for exception in exception_subject.exceptions
                        if item.reviewed_at and aware(exception.created_at) >= aware(item.reviewed_at)]
            exception = max(matching, key=lambda value: aware(value.created_at), default=None)
            if exception and exception.revoked_at:
                display_status = "revoked"
            elif exception and aware(exception.expires_at) <= now:
                display_status = "expired"
        workflow_statuses[item.id] = display_status
    return templates.TemplateResponse(request, "requests.html", page_context(
        auth, workflows=workflows, workflow_statuses=workflow_statuses, request_type_filter=request_type,
    ))


@app.post("/requests/{workflow_id}/review")
def review_request(
    workflow_id: int, decision: str = Form(), review_reason: str = Form(default=""),
    apply_bulk: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_user),
):
    check_csrf(auth, csrf_token)
    workflow = db.get(WorkflowRequest, workflow_id)
    if not workflow:
        raise HTTPException(404)
    permission = "poam.review" if workflow.request_type.startswith("poam") else {
        "exception": "exception.review", "archive": "archive.review",
        "image_remove": "archive.review", "image_replace": "archive.review",
    }.get(workflow.request_type, "archive.review")
    if not auth.has(permission, workflow.service_id):
        raise HTTPException(403, detail="Permission denied")
    if workflow.status != "pending":
        raise HTTPException(409, detail="Request has already been reviewed")
    if workflow.requested_by_id == auth.user.id:
        raise HTTPException(409, detail="Requesters cannot approve or reject their own requests")
    if decision not in {"approved", "rejected"}:
        raise HTTPException(422, detail="Decision must be approved or rejected")
    if decision == "rejected" and len(review_reason.strip()) < 3:
        raise HTTPException(422, detail="A rejection reason is required")
    if decision == "approved" and workflow.bulk_group_id and apply_bulk != "yes":
        raise HTTPException(422, detail="Confirm the group-wide exception scope before approving")
    workflow.status = decision
    workflow.reviewed_by_id = auth.user.id
    workflow.review_reason = review_reason.strip() or None
    workflow.reviewed_at = utcnow()
    if workflow.request_type == "poam" and workflow.poam_id:
        poam = db.get(PoamEntry, workflow.poam_id)
        if not poam:
            raise HTTPException(409, detail="POA&M entry is missing")
        poam.status = "active" if decision == "approved" else "rejected"
        if decision == "approved":
            poam.approved_by_id = auth.user.id
            poam.approved_at = workflow.reviewed_at
        db.add(PoamHistory(poam_id=poam.id, actor_user_id=auth.user.id, action=decision,
                           note=workflow.review_reason, detail={"request_type": "creation"}))
    elif workflow.request_type in {"poam_update", "poam_complete", "poam_reopen"} and workflow.poam_id:
        poam = db.get(PoamEntry, workflow.poam_id)
        change = db.scalar(select(PoamChangeRequest).where(PoamChangeRequest.workflow_id == workflow.id))
        if not poam or not change:
            raise HTTPException(409, detail="POA&M change request is missing")
        if decision == "approved":
            if change.change_type == "update":
                for field in ("title", "description", "remediation", "ticket"):
                    setattr(poam, field, change.proposed.get(field))
                raw_due = change.proposed.get("due_date")
                poam.due_date = datetime.fromisoformat(raw_due) if raw_due else None
            elif change.change_type == "complete":
                poam.status = "completed"
            elif change.change_type == "reopen":
                poam.status = "active"
            poam.approved_by_id = auth.user.id
            poam.approved_at = workflow.reviewed_at
        db.add(PoamHistory(poam_id=poam.id, actor_user_id=auth.user.id,
                           action=f"{change.change_type}_{decision}", note=workflow.review_reason,
                           detail=change.proposed))
    if decision == "approved" and workflow.request_type == "exception":
        bulk_service_ids = None
        if workflow.bulk_group_id and apply_bulk == "yes":
            bulk_service_ids = db.scalars(select(ServiceGroup.service_id).where(
                ServiceGroup.group_id == workflow.bulk_group_id
            )).all()
        if workflow.policy_finding_id:
            source = db.get(PolicyFinding, workflow.policy_finding_id)
            query = select(PolicyFinding).options(selectinload(PolicyFinding.exceptions)).where(
                PolicyFinding.identity_key == source.identity_key,
                PolicyFinding.active.is_(True),
            )
            if bulk_service_ids is not None:
                query = query.where(PolicyFinding.service_id.in_(bulk_service_ids))
            for target in db.scalars(query).all():
                if not active_exception(target, utcnow()):
                    db.add(PolicyExceptionRecord(
                        policy_finding_id=target.id,
                        justification=workflow.justification, approved_by=auth.user.username,
                        expires_at=workflow.requested_expires_at, ticket=workflow.ticket,
                    ))
        else:
            source = db.get(Finding, workflow.finding_id)
            query = select(Finding).options(selectinload(Finding.exceptions)).where(
                Finding.cve == source.cve, Finding.active.is_(True),
            )
            if bulk_service_ids is not None:
                query = query.where(Finding.service_id.in_(bulk_service_ids))
            for target in db.scalars(query).all():
                if not active_exception(target, utcnow()):
                    db.add(ExceptionRecord(
                        finding_id=target.id, justification=workflow.justification,
                        approved_by=auth.user.username, expires_at=workflow.requested_expires_at,
                        ticket=workflow.ticket,
                    ))
    elif decision == "approved" and workflow.request_type == "archive":
        service = db.scalar(select(Service).where(Service.id == workflow.service_id).options(selectinload(Service.archive_events)))
        if archive_state(service):
            raise HTTPException(409, detail="Service is already archived")
        db.add(ServiceArchiveEvent(
            service_id=workflow.service_id, action="archive", reason=workflow.justification,
            performed_by=auth.user.username, ticket=workflow.ticket,
        ))
    elif decision == "approved" and workflow.request_type in {"image_remove", "image_replace"}:
        image = db.get(ServiceImage, workflow.service_image_id) if workflow.service_image_id else None
        if not image:
            raise HTTPException(409, detail="The requested service image no longer exists")
        image.lifecycle_status = "removed" if workflow.request_type == "image_remove" else "replaced"
        image.lifecycle_reason = workflow.justification
        image.approved_by_id = auth.user.id
        image.updated_at = utcnow()
        if workflow.request_type == "image_replace":
            replacement_reference = (workflow.replacement_reference or "").strip()
            if not replacement_reference:
                raise HTTPException(422, detail="Replacement image reference is missing")
            replacement = ServiceImage(
                service_id=workflow.service_id, image_reference=replacement_reference,
                lifecycle_status="active", replacement_of_id=image.id,
                lifecycle_reason=workflow.justification,
                requested_by_id=workflow.requested_by_id, approved_by_id=auth.user.id,
            )
            db.add(replacement)
            db.flush()
        _reconcile_image_scope(db, image.service, image.image_reference, set(), workflow.reviewed_at or utcnow())
        record_audit(db, auth,
                     "image_removed" if workflow.request_type == "image_remove" else "image_replaced",
                     "service_image", image.id, service_id=workflow.service_id,
                     image=image.image_reference, replacement=workflow.replacement_reference,
                     reason=workflow.justification)
    record_audit(db, auth, f"{workflow.request_type}.{decision}", "workflow_request", workflow.id,
                 service_id=workflow.service_id, review_reason=workflow.review_reason)
    db.commit()
    return RedirectResponse("/requests", status_code=303)


@app.get("/audit")
def audit_legacy_page(auth: AuthContext = Depends(require_permission("audit.view"))):
    return RedirectResponse("/admin/audit", status_code=303)


@app.get("/admin/audit", response_class=HTMLResponse)
def audit_page(request: Request, group_id: str = "", show: int = 10, db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("audit.view"))):
    groups = manageable_groups(db, auth) if auth.has("config.manage") else []
    selected_group_id = requested_group_scope(group_id or (str(groups[0].id) if groups else None), db, auth) if groups else None
    configuration = get_configuration(db, selected_group_id)
    # Logging level is a runtime setting owned by the global Audit Policy.
    # Keep retention/integration group-aware, but always show the global level.
    configuration["log_level"] = get_global_configuration(db).get("log_level", configuration.get("log_level", "INFO"))
    try:
        retention_days = int(configuration.get("audit_retention_days", "365") or "365")
    except ValueError:
        retention_days = 365
    query = select(AuditEvent).options(selectinload(AuditEvent.actor)).order_by(AuditEvent.created_at.desc())
    if retention_days > 0:
        query = query.where(AuditEvent.created_at >= utcnow() - timedelta(days=retention_days))
    retained_count = int(db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0)
    show = max(10, min(show, retained_count or 10))
    events = db.scalars(query.limit(show)).all()
    return templates.TemplateResponse(request, "audit.html", page_context(
        auth, events=events, groups=groups, selected_group_id=selected_group_id,
        configuration=configuration, saved=request.query_params.get("saved") == "1",
        retained_count=retained_count, shown_count=len(events), next_show=min(show + 50, retained_count),
    ))


@app.get("/admin/general-policy", response_class=HTMLResponse)
def general_policy_page(request: Request, group_id: str = "", db: Session = Depends(get_db), auth: AuthContext = Depends(require_user)):
    if not auth.has("audit.view") and not auth.has("config.manage"):
        raise HTTPException(status_code=403, detail="Permission denied")
    groups = manageable_groups(db, auth) if auth.has("config.manage") else []
    selected_group_id = requested_group_scope(group_id or (str(groups[0].id) if groups else None), db, auth) if groups else None
    configuration = get_configuration(db, selected_group_id)
    configuration["log_level"] = get_global_configuration(db).get("log_level", configuration.get("log_level", "INFO"))
    try:
        retention_days = int(configuration.get("audit_retention_days", "365") or "365")
    except ValueError:
        retention_days = 365
    query = select(AuditEvent).options(selectinload(AuditEvent.actor)).order_by(AuditEvent.created_at.desc())
    if retention_days > 0:
        query = query.where(AuditEvent.created_at >= utcnow() - timedelta(days=retention_days))
    retained_count = int(db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0)
    events = db.scalars(query.limit(10)).all()
    return templates.TemplateResponse(request, "general_policy.html", page_context(
        auth, events=events, groups=groups, selected_group_id=selected_group_id,
        configuration=configuration, saved=request.query_params.get("saved") == "1",
        retained_count=retained_count, shown_count=len(events),
    ))


@app.post("/admin/audit-policy")
def save_audit_policy(
    audit_retention_days: str = Form(default="365"), audit_tool_integration: str = Form(default="planned"),
    log_level: str = Form(default="INFO"),
    group_id: str = Form(default=""), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_config_scope),
):
    check_csrf(auth, csrf_token)
    selected_group_id = requested_group_scope(group_id, db, auth)
    try:
        retention = int(audit_retention_days)
    except ValueError as exc:
        raise HTTPException(422, detail="Audit retention must be a number") from exc
    if retention < 0 or retention > 3650:
        raise HTTPException(422, detail="Audit retention must be between 0 and 3650 days")
    if audit_tool_integration not in {"planned"}:
        raise HTTPException(422, detail="Unsupported audit integration")
    if log_level not in {"DEBUG", "INFO", "WARNING", "ERROR"}:
        raise HTTPException(422, detail="Unsupported logging level")
    values = {"audit_retention_days": str(retention), "audit_tool_integration": audit_tool_integration, "log_level": log_level}
    for key, value in values.items():
        # Logging controls are runtime-wide even though retention/tool integration
        # may retain their existing group policy scope for backwards compatibility.
        target_group_id = None if key == "log_level" else selected_group_id
        storage_key = scoped_setting_key(key, target_group_id)
        setting = db.scalar(select(PortalSetting).where(PortalSetting.key == storage_key))
        if not setting:
            setting = PortalSetting(key=storage_key, group_id=target_group_id)
            db.add(setting)
        setting.value = value
        setting.updated_by_id = auth.user.id
        setting.updated_at = utcnow()
    record_audit(db, auth, "audit_policy.updated", "group" if selected_group_id else "portal", str(selected_group_id or "global"), changed=list(values))
    db.commit()
    suffix = f"&group_id={selected_group_id}" if selected_group_id else ""
    return RedirectResponse(f"/admin/general-policy?saved=1{suffix}", status_code=303)


@app.get("/admin", response_class=HTMLResponse)
def admin_page(request: Request, db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("user.manage"))):
    users = db.scalars(select(User).options(
        selectinload(User.role_assignments).selectinload(UserRoleAssignment.role),
        selectinload(User.role_assignments).selectinload(UserRoleAssignment.service),
        selectinload(User.role_assignments).selectinload(UserRoleAssignment.group),
    ).order_by(User.username)).all()
    roles = db.scalars(select(Role).order_by(Role.system.desc(), Role.name)).all()
    services = db.scalars(select(Service).options(
        selectinload(Service.groups), selectinload(Service.archive_events),
    ).order_by(Service.name)).all()
    groups = db.scalars(select(Group).order_by(Group.name)).all()
    stage_groups = manageable_groups(db, auth)
    group_summaries = [
        {
            "group": group,
            "active_services": sum(1 for service in group.services if not archive_state(service)),
            "service_count": len(group.services),
        }
        for group in groups
    ]
    user_rows = []
    for user in users:
        assignments = []
        permissions = set()
        for assignment in user.role_assignments:
            permissions.update(assignment.role.permissions or [])
            assignments.append({
                "role": assignment.role.name,
                "group": assignment.group.name if assignment.group else "Global",
                "service": assignment.service.name if assignment.service else (
                    "All services in group" if assignment.group else "All services"
                ),
                "id": assignment.id,
            })
        user_rows.append({"user": user, "assignments": assignments, "permissions": sorted(permissions)})
    return templates.TemplateResponse(request, "admin.html", page_context(
        auth, users=users, user_rows=user_rows, roles=roles, services=services, groups=groups,
        group_summaries=group_summaries, permission_catalog=PERMISSIONS, stage_groups=stage_groups,
        saved=request.query_params.get("saved") == "1",
    ))


@app.get("/admin/staging", response_class=HTMLResponse)
def staging_page(request: Request, db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("user.manage"))):
    return templates.TemplateResponse(request, "staging.html", page_context(
        auth, stage_groups=manageable_groups(db, auth), saved=request.query_params.get("saved") == "1",
    ))


@app.post("/admin/services/stage")
def stage_service(
    service_id: str = Form(), group_id: str = Form(default=""), csrf_token: str = Form(),
    next_path: str = Form(default="/admin/staging"),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("user.manage")),
):
    check_csrf(auth, csrf_token)
    service_key = service_id.strip()
    if not service_key or len(service_key) > 120:
        raise HTTPException(422, detail="A valid service ID is required")
    if db.scalar(select(Service).where(Service.service_key == service_key)):
        raise HTTPException(409, detail="That service is already staged or registered")
    groups = manageable_groups(db, auth)
    if len(groups) == 1:
        selected_group = groups[0]
    else:
        try:
            selected_id = int(group_id)
        except (TypeError, ValueError) as exc:
            raise HTTPException(422, detail="Select a group for this staged service") from exc
        selected_group = next((group for group in groups if group.id == selected_id), None)
        if not selected_group:
            raise HTTPException(403, detail="You cannot stage a service in that group")
    service = Service(service_key=service_key, name=f"Staged — {service_key}", lifecycle_status="staged")
    service.groups.append(selected_group)
    db.add(service)
    db.flush()
    record_audit(db, auth, "service.staged", "service", service.id,
                 service_key=service_key, group_id=selected_group.id)
    db.commit()
    # Return to the page that opened the staging modal. Only local paths are
    # accepted so this cannot become an open redirect.
    target = next_path.strip() or "/admin/staging"
    if not target.startswith("/") or target.startswith("//"):
        target = "/admin/staging"
    separator = "&" if "?" in target else "?"
    return RedirectResponse(f"{target}{separator}saved=1", status_code=303)


@app.post("/admin/services/{service_key}")
def edit_service(
    service_key: str, name: str = Form(), description: str = Form(default=""), owner: str = Form(default=""), poc: str = Form(default=""),
    manual_version: str = Form(default=""), group_ids: list[int] = Form(default=[]), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("service.edit", scoped=True)),
):
    check_csrf(auth, csrf_token)
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if not service:
        raise HTTPException(404)
    if not name.strip():
        raise HTTPException(422, detail="Service name is required")
    previous = {
        "name": service.name, "description": service.description, "owner": service.owner,
        "poc": service.poc, "manual_version": service.manual_version,
        "group_ids": sorted(group.id for group in service.groups),
    }
    service.name = name.strip()[:240]
    service.description = description.strip()[:2000] or None
    service.owner = owner.strip()[:240] or None
    service.poc = poc.strip()[:240] or None
    service.manual_version = manual_version.strip()[:120] or None
    service.groups = list(db.scalars(select(Group).where(Group.id.in_(group_ids))).all()) if group_ids else []
    current = {
        "name": service.name, "description": service.description, "owner": service.owner,
        "poc": service.poc, "manual_version": service.manual_version,
        "group_ids": sorted(group.id for group in service.groups),
    }
    record_audit(db, auth, "service.metadata_updated", "service", service.id,
                 service_key=service.service_key, changed=[key for key in previous if previous[key] != current[key]])
    db.commit()
    return RedirectResponse(f"/services/{service.service_key}?saved=1", status_code=303)


@app.post("/admin/groups")
def create_group(
    name: str = Form(), description: str = Form(default=""), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("role.manage")),
):
    check_csrf(auth, csrf_token)
    name = name.strip()[:120]
    if not name or db.scalar(select(Group).where(Group.name == name)):
        raise HTTPException(409, detail="Group name is invalid or already exists")
    group = Group(name=name, description=description.strip())
    db.add(group)
    db.flush()
    record_audit(db, auth, "group.created", "group", group.id, name=name)
    db.commit()
    return RedirectResponse("/admin?saved=1", status_code=303)


@app.get("/admin/configuration", response_class=HTMLResponse)
def configuration_page(request: Request, edit_os_id: str = "", db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope)):
    configuration = get_global_configuration(db)
    certificates = parse_json(configuration.get("trusted_ca_certificates"), [])
    raw_repository_policies = parse_json(configuration.get("repository_policies"), {})
    repository_policies = (
        {str(key).lower(): normalize_policy(value) for key, value in raw_repository_policies.items()}
        if isinstance(raw_repository_policies, dict) else {}
    )
    custom_os = parse_json(configuration.get("os_definitions"), {})
    timezones = supported_display_timezones()
    oidc = configured_oidc(configuration)
    public_url = os.getenv("CATS_PUBLIC_URL", str(request.base_url).rstrip("/"))
    oidc["redirect_uri"] = oidc.get("redirect_uri") or f"{public_url}/auth/oidc/callback"
    oidc["post_logout_redirect_uri"] = oidc.get("post_logout_redirect_uri") or f"{public_url}/login"
    return templates.TemplateResponse(request, "configuration.html", page_context(
        auth,
        configuration=configuration,
        saved=request.query_params.get("saved") == "1",
        timezones=timezones,
        date_formats=[("%d %b %Y", "31 Jul 2026"), ("%Y-%m-%d", "2026-07-31"), ("%m/%d/%Y", "07/31/2026")],
        time_formats=[("%H:%M UTC", "24-hour time (UTC)"), ("%I:%M %p UTC", "12-hour time (UTC)")],
        certificates=certificates, repository_policies=repository_policies,
        custom_os=custom_os,
        oidc=oidc, public_url=public_url, registries=configured_registries(configuration),
        signing=signing.public_metadata(configuration),
        os_definitions={**OS_DEFINITIONS, **custom_os}, package_managers=sorted(PACKAGE_MANAGERS),
        edit_os_id=edit_os_id.strip().lower(),
        repository_result=request.query_params.get("repository_result", ""),
        oidc_result=request.query_params.get("oidc_result", ""),
    ))


@app.get("/admin/compliance", response_class=HTMLResponse)
def compliance_page(request: Request, group_id: str = "", db: Session = Depends(get_db), auth: AuthContext = Depends(require_config_scope)):
    groups = manageable_groups(db, auth)
    selected_group_id = requested_group_scope(group_id or (str(groups[0].id) if groups else None), db, auth)
    configuration = get_configuration(db, selected_group_id)
    try:
        epss_rules = json.loads(configuration.get("epss_rules", "[]"))
    except (TypeError, ValueError):
        epss_rules = []
    if not epss_rules:
        epss_rules = [{"severity": "Any", "threshold": configuration.get("epss_threshold", "0.90"), "noncompliant": True}]
    try:
        raw_due_rules = json.loads(configuration.get("raw_due_rules", "[]"))
    except (TypeError, ValueError):
        raw_due_rules = []
    return templates.TemplateResponse(request, "compliance.html", page_context(
        auth, configuration=configuration, epss_rules=epss_rules, raw_due_rules=raw_due_rules,
        saved=request.query_params.get("saved") == "1", groups=groups, selected_group_id=selected_group_id,
    ))


@app.get("/admin/compliance-frameworks", response_class=HTMLResponse)
def compliance_frameworks_page(request: Request, group_id: str = "", db: Session = Depends(get_db), auth: AuthContext = Depends(require_config_scope)):
    groups = manageable_groups(db, auth)
    selected_group_id = requested_group_scope(group_id or (str(groups[0].id) if groups else None), db, auth)
    return templates.TemplateResponse(request, "compliance_frameworks.html", page_context(
        auth, configuration=get_configuration(db, selected_group_id),
        saved=request.query_params.get("saved") == "1", groups=groups,
        selected_group_id=selected_group_id,
    ))


@app.get("/admin/evidence-policy", response_class=HTMLResponse)
def evidence_policy_page(request: Request, group_id: str = "", db: Session = Depends(get_db), auth: AuthContext = Depends(require_config_scope)):
    groups = manageable_groups(db, auth)
    selected_group_id = requested_group_scope(group_id or (str(groups[0].id) if groups else None), db, auth)
    return templates.TemplateResponse(request, "evidence_policy.html", page_context(
        auth, configuration=get_configuration(db, selected_group_id),
        saved=request.query_params.get("saved") == "1", groups=groups,
        selected_group_id=selected_group_id,
    ))


@app.get("/admin/workflow-policy", response_class=HTMLResponse)
def workflow_policy_page(request: Request, group_id: str = "", db: Session = Depends(get_db), auth: AuthContext = Depends(require_config_scope)):
    groups = manageable_groups(db, auth)
    selected_group_id = requested_group_scope(group_id or (str(groups[0].id) if groups else None), db, auth)
    return templates.TemplateResponse(request, "workflow_policy.html", page_context(
        auth, configuration=get_configuration(db, selected_group_id),
        saved=request.query_params.get("saved") == "1", groups=groups,
        selected_group_id=selected_group_id,
    ))


@app.post("/admin/configuration")
def save_configuration(
    display_timezone: str = Form(), date_format: str = Form(), time_format: str = Form(),
    identity_mode: str = Form(), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    if display_timezone not in set(supported_display_timezones()):
        raise HTTPException(422, detail="Unknown display timezone")
    if date_format not in {"%d %b %Y", "%Y-%m-%d", "%m/%d/%Y"}:
        raise HTTPException(422, detail="Unsupported date format")
    if time_format not in {"%H:%M UTC", "%I:%M %p UTC"}:
        raise HTTPException(422, detail="Unsupported time format")
    if identity_mode not in {"local", "oidc", "both"}:
        raise HTTPException(422, detail="Unsupported identity mode")
    values = {
        "display_timezone": display_timezone, "date_format": date_format,
        "time_format": time_format,
        "identity_mode": identity_mode,
    }
    for key, value in values.items():
        storage_key = scoped_setting_key(key, None)
        setting = db.scalar(select(PortalSetting).where(PortalSetting.key == storage_key))
        if not setting:
            setting = PortalSetting(key=storage_key, group_id=None)
            db.add(setting)
        setting.value = value
        setting.updated_by_id = auth.user.id
        setting.updated_at = utcnow()
    record_audit(db, auth, "configuration.updated", "portal", "global", changed=list(values))
    db.commit()
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


@app.post("/admin/configuration/oidc")
def save_oidc_configuration(
    request: Request, csrf_token: str = Form(), provider_name: str = Form(default=""), issuer: str = Form(default=""),
    client_id: str = Form(default=""), client_secret: str = Form(default=""), scopes: str = Form(default="openid profile email"),
    username_claim: str = Form(default="preferred_username"), email_claim: str = Form(default="email"),
    groups_claim: str = Form(default=""), roles_claim: str = Form(default=""),
    redirect_uri: str = Form(default=""), post_logout_redirect_uri: str = Form(default=""), browser_issuer: str = Form(default=""),
    action: str = Form(default="save"),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    issuer = issuer.strip().rstrip("/")
    browser_issuer = browser_issuer.strip().rstrip("/")
    parsed = urllib.parse.urlparse(issuer)
    if not issuer or parsed.scheme not in {"https", "http"} or not parsed.netloc:
        raise HTTPException(422, detail="OIDC issuer must be a valid HTTP(S) URL")
    # Remote HTTP issuers are rejected by default because OIDC discovery and
    # token exchange would otherwise expose credentials on the network.  A
    # local development Keycloak on a LAN address can be explicitly enabled
    # with CATS_OIDC_ALLOW_INSECURE_HTTP=true; production remains HTTPS-only.
    allow_insecure_http = os.getenv("CATS_OIDC_ALLOW_INSECURE_HTTP", "").strip().lower() in {"1", "true", "yes", "on"}
    if parsed.scheme == "http" and parsed.hostname not in {"localhost", "127.0.0.1", "::1"} and not allow_insecure_http:
        raise HTTPException(422, detail="OIDC issuer must use HTTPS outside local development")
    if browser_issuer:
        browser_parsed = urllib.parse.urlparse(browser_issuer)
        if browser_parsed.scheme not in {"https", "http"} or not browser_parsed.netloc:
            raise HTTPException(422, detail="Browser issuer must be a valid HTTP(S) URL")
    if not client_id.strip():
        raise HTTPException(422, detail="OIDC client ID is required")
    scopes = " ".join(scopes.split())
    if "openid" not in scopes.split():
        scopes = "openid " + scopes
    current = parse_json(get_global_configuration(db).get("oidc_configuration"), {})
    if not isinstance(current, dict): current = {}
    public_url = os.getenv("CATS_PUBLIC_URL", str(request.base_url).rstrip("/"))
    redirect_uri = redirect_uri.strip() or f"{public_url}/auth/oidc/callback"
    post_logout_redirect_uri = post_logout_redirect_uri.strip() or f"{public_url}/login"
    current.update({"provider_name": provider_name.strip()[:120], "issuer": issuer, "client_id": client_id.strip()[:240],
                    "scopes": scopes[:500], "username_claim": username_claim.strip()[:120] or "preferred_username",
                    "email_claim": email_claim.strip()[:120] or "email", "groups_claim": groups_claim.strip()[:120] or "groups",
                    "roles_claim": roles_claim.strip()[:180] or "realm_access.roles", "redirect_uri": redirect_uri.strip()[:500],
                    "post_logout_redirect_uri": post_logout_redirect_uri[:500], "browser_issuer": browser_issuer[:500]})
    if client_secret.strip():
        current["client_secret"] = encrypt_secret(client_secret.strip())
    _set_config_value(db, auth, "oidc_configuration", json.dumps(current), None)
    record_audit(db, auth, "configuration.oidc_updated", "portal", "global", provider=current.get("provider_name"), issuer=issuer)
    db.commit()
    if action == "test":
        # Test the values submitted in this form, after saving the non-secret
        # settings.  This keeps the form populated when discovery fails while
        # ensuring the client secret remains encrypted and is never echoed.
        try:
            from .auth import oidc_discovery
            oidc_discovery(current, configured_ca_bundle(get_global_configuration(db)))
            result = "OIDC discovery succeeded"
        except Exception as exc:
            result = f"OIDC discovery failed: {redact(exc)}"
        return RedirectResponse(f"/admin/configuration?saved=1&oidc_result={urllib.parse.quote(result)}", status_code=303)
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


@app.post("/admin/configuration/oidc-test")
def test_oidc_configuration(csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope)):
    check_csrf(auth, csrf_token)
    config = parse_json(get_global_configuration(db).get("oidc_configuration"), {})
    try:
        from .auth import oidc_discovery
        oidc_discovery(config, configured_ca_bundle(get_global_configuration(db)))
        result = "OIDC discovery succeeded"
    except Exception as exc:
        result = f"OIDC discovery failed: {redact(exc)}"
    return RedirectResponse(f"/admin/configuration?saved=1&oidc_result={urllib.parse.quote(result)}", status_code=303)


@app.post("/admin/configuration/registries")
def save_registry_configuration(
    csrf_token: str = Form(), registry_id: str = Form(default=""), display_name: str = Form(default=""),
    endpoint: str = Form(default=""), namespace: str = Form(default=""), auth_mode: str = Form(default="none"),
    use_for_remediation: bool = Form(default=False),
    username: str = Form(default=""), password: str = Form(default=""), action: str = Form(default="save"),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    configuration = get_global_configuration(db)
    registries = parse_json(configuration.get("oci_registries"), [])
    if not isinstance(registries, list): registries = []
    registry_id = registry_id.strip() or uuid.uuid4().hex[:12]
    if action == "delete":
        registries = [item for item in registries if str(item.get("id")) != registry_id]
    else:
        endpoint = _validate_endpoint(endpoint)
        if auth_mode not in {"none", "credentials"}:
            raise HTTPException(422, detail="Unsupported registry authentication mode")
        existing = next((item for item in registries if str(item.get("id")) == registry_id), None)
        if existing is None:
            existing = {"id": registry_id}
            registries.append(existing)
        existing.update({"display_name": display_name.strip()[:120] or endpoint, "endpoint": endpoint,
                         "namespace": namespace.strip()[:240], "auth_mode": auth_mode, "username": username.strip()[:240],
                         "use_for_remediation": bool(use_for_remediation)})
        if use_for_remediation:
            for item in registries:
                if item is not existing:
                    item["use_for_remediation"] = False
        if password.strip(): existing["password"] = encrypt_secret(password.strip())
        existing.setdefault("password", "")
        existing.setdefault("status", "Not tested")
    _set_config_value(db, auth, "oci_registries", json.dumps(registries), None)
    record_audit(db, auth, "configuration.oci_registry_updated", "portal", "global", registry_id=registry_id, registry_action=action)
    db.commit()
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


@app.post("/admin/configuration/registry-test")
def test_registry_configuration(registry_id: str = Form(), csrf_token: str = Form(), db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope)):
    check_csrf(auth, csrf_token)
    configuration = get_global_configuration(db)
    registries = parse_json(configuration.get("oci_registries"), [])
    item = next((entry for entry in registries if str(entry.get("id")) == registry_id), None)
    if not item:
        raise HTTPException(404, detail="Registry was not found")
    status_text, detail = "Reachable", "Registry endpoint responded"
    try:
        request = urllib.request.Request(str(item.get("endpoint")) + "/v2/", method="GET")
        if item.get("auth_mode") == "credentials" and item.get("username") and item.get("password"):
            import base64
            token = base64.b64encode(f"{item['username']}:{decrypt_secret(item['password'])}".encode()).decode()
            request.add_header("Authorization", f"Basic {token}")
        ca_bundle = configured_ca_bundle(configuration)
        if ca_bundle:
            import ssl
            context = ssl.create_default_context()
            context.load_verify_locations(cadata=ca_bundle)
            with urllib.request.urlopen(request, timeout=10, context=context) as response:
                if response.status >= 400: raise RuntimeError(f"HTTP {response.status}")
        else:
            with urllib.request.urlopen(request, timeout=10) as response:
                if response.status >= 400: raise RuntimeError(f"HTTP {response.status}")
    except Exception as exc:
        status_text, detail = "Unavailable", redact(exc)
    for entry in registries:
        if str(entry.get("id")) == registry_id:
            entry["status"], entry["status_detail"] = status_text, detail[:240]
    _set_config_value(db, auth, "oci_registries", json.dumps(registries), None)
    record_audit(db, auth, "configuration.oci_registry_tested", "portal", "global", registry_id=registry_id, status=status_text)
    db.commit()
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


def _set_config_value(db: Session, auth: AuthContext, key: str, value: str, group_id: int | None) -> None:
    storage_key = scoped_setting_key(key, group_id)
    setting = db.scalar(select(PortalSetting).where(PortalSetting.key == storage_key))
    if not setting:
        setting = PortalSetting(key=storage_key, group_id=group_id); db.add(setting)
    setting.value = value; setting.updated_by_id = auth.user.id; setting.updated_at = utcnow()


@app.post("/admin/configuration/signing")
def save_signing_configuration(
    csrf_token: str = Form(), private_key: UploadFile | None = File(None),
    public_key: UploadFile | None = File(None), key_password: str = Form(""),
    enabled: bool = Form(False), action: str = Form("save"),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    try:
        current = {} if action == "remove" or (private_key and private_key.filename) else signing.configuration(get_global_configuration(db))
        if action == "remove":
            current = {}
        elif action == "save":
            if private_key and private_key.filename:
                private = private_key.file.read(signing.MAX_KEY_BYTES + 1)
                public = public_key.file.read(signing.MAX_KEY_BYTES + 1) if public_key and public_key.filename else b""
                current = signing.store_keys(private, public, key_password, enabled)
            elif (public_key and public_key.filename) or key_password:
                raise ValueError("Upload the private key again to change its public key or password.")
            else:
                current["enabled"] = enabled
                if enabled:
                    signing.job_material({signing.SETTING: json.dumps(current)}, "push")
                current["updated_at"] = utcnow().isoformat()
        else:
            raise ValueError("Unknown signing configuration action.")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    _set_config_value(db, auth, signing.SETTING, json.dumps(current), None)
    record_audit(db, auth, "configuration.signing_updated", "portal", "global",
                 enabled=current.get("enabled", False), key_fingerprint=current.get("fingerprint"), removed=action == "remove")
    db.commit()
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


@app.get("/admin/configuration/signing/public-key", response_class=PlainTextResponse)
def download_signing_public_key(db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope)):
    pem = signing.configuration(get_global_configuration(db)).get("public_key")
    if not pem:
        raise HTTPException(status_code=404, detail="No signing public key is configured")
    return PlainTextResponse(pem, headers={"Content-Disposition": 'attachment; filename="cosign.pub"'})


@app.post("/admin/configuration/trusted-certificates")
async def update_trusted_certificates(
    csrf_token: str = Form(), certificate: UploadFile | None = File(default=None),
    remove_fingerprint: str = Form(default=""), db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    configuration = get_global_configuration(db)
    certificates = parse_json(configuration.get("trusted_ca_certificates"), [])
    if certificate and certificate.filename:
        payload = await certificate.read()
        try:
            uploaded = certificate_bundle_metadata(payload)
        except (ValueError, UnicodeDecodeError) as exc:
            raise HTTPException(422, detail=str(exc))
        merged = merge_certificate_metadata(certificates, uploaded)
        if len(merged) == len(certificates):
            raise HTTPException(409, detail="All certificates in that upload are already configured")
        certificates = merged
    if remove_fingerprint:
        certificates = [item for item in certificates if item.get("fingerprint") != remove_fingerprint]
    _set_config_value(db, auth, "trusted_ca_certificates", json.dumps(certificates), None)
    record_audit(db, auth, "configuration.trusted_ca_updated", "portal", "global", changed=["trusted_ca_certificates"])
    db.commit()
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


@app.post("/admin/configuration/repositories")
def update_repository_policies(
    csrf_token: str = Form(), policy_json: str = Form(default="{}"), definitions_json: str = Form(default="{}"),
    os_id: str = Form(default=""), edit_os_id: str = Form(default=""), action: str = Form(default="save"),
    mode: str = Form(default="default"), url: str = Form(default=""), display_name: str = Form(default=""),
    package_manager: str = Form(default=""), remove_os_id: str = Form(default=""),
    verify_tls: list[str] = Form(default=[]), verify_packages: list[str] = Form(default=[]),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    configuration = get_global_configuration(db)
    policies = parse_json(configuration.get("repository_policies"), {})
    definitions = parse_json(configuration.get("os_definitions"), {})
    if not isinstance(policies, dict): policies = {}
    if not isinstance(definitions, dict): definitions = {}
    key = (edit_os_id or os_id).strip().lower()
    if action == "delete" or remove_os_id:
        key = (remove_os_id or key).strip().lower()
        if key in OS_DEFINITIONS:
            raise HTTPException(422, detail="Built-in operating systems cannot be removed")
        definitions.pop(key, None); policies.pop(key, None)
    else:
        if not key: raise HTTPException(422, detail="Operating system is required")
        def posted_bool(values: list[str]) -> bool:
            if not values:
                return True
            decoded = [policy_bool(value, None) for value in values]
            if any(value is None for value in decoded):
                raise HTTPException(422, detail="Repository verification settings are invalid")
            return any(decoded)
        tls_verified = posted_bool(verify_tls)
        packages_verified = posted_bool(verify_packages)
        valid, message = validate_policy(key, mode, url, tls_verified, packages_verified)
        if not valid: raise HTTPException(422, detail=message)
        if key in OS_DEFINITIONS:
            # Built-in IDs and capabilities are immutable, but their mirror
            # policy can be changed independently.
            pass
        else:
            name = display_name.strip()[:120]
            manager = package_manager.strip().lower()
            valid, message = validate_os_definition(key, name, manager)
            if not valid: raise HTTPException(422, detail=message)
            definitions[key] = {"name": name, "package_manager": manager}
        policies[key] = {
            "mode": mode, "url": url.strip() if mode == "custom" else "",
            "verify_tls": tls_verified, "verify_packages": packages_verified,
        }
    _set_config_value(db, auth, "repository_policies", json.dumps(policies), None)
    _set_config_value(db, auth, "os_definitions", json.dumps(definitions), None)
    record_audit(
        db, auth, "configuration.repositories_updated", "portal", "global",
        changed=["repository_policies", "os_definitions"],
        os_id=key, verify_tls=policies.get(key, {}).get("verify_tls", True),
        verify_packages=policies.get(key, {}).get("verify_packages", True),
    )
    db.commit()
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


@app.post("/admin/configuration/os-definitions")
def update_os_definitions(
    csrf_token: str = Form(), group_id: str = Form(default=""), definitions_json: str = Form(default="{}"),
    os_id: str = Form(default=""), display_name: str = Form(default=""), package_manager: str = Form(default=""),
    remove_os_id: str = Form(default=""), db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    selected_group_id = None
    definitions = parse_json(definitions_json, {})
    if not isinstance(definitions, dict):
        definitions = {}
    if remove_os_id:
        key = remove_os_id.strip().lower()
        if key in OS_DEFINITIONS:
            raise HTTPException(422, detail="Built-in operating systems cannot be removed")
        definitions.pop(key, None)
    else:
        key = os_id.strip().lower()
        if key in OS_DEFINITIONS:
            raise HTTPException(422, detail="Built-in operating systems cannot be replaced")
        name = display_name.strip()[:120]
        valid, message = validate_os_definition(key, name, package_manager)
        if not valid:
            raise HTTPException(422, detail=message)
        definitions[key] = {"name": name, "package_manager": package_manager}
    _set_config_value(db, auth, "os_definitions", json.dumps(definitions), selected_group_id)
    record_audit(db, auth, "configuration.os_definition_updated", "portal", "global", changed=["os_definitions"])
    db.commit()
    return RedirectResponse("/admin/configuration?saved=1", status_code=303)


@app.post("/admin/configuration/repository-test")
def repository_test(
    csrf_token: str = Form(), url: str = Form(default=""), os_id: str = Form(default=""),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    if not url.strip():
        ok, message = True, "Image-defined repositories are supplied by the scanned image"
    else:
        ok, message = test_repository(url.strip())
    result = ("ok:" if ok else "error:") + message
    return RedirectResponse(f"/admin/configuration?edit_os_id={urllib.parse.quote(os_id.strip().lower())}&repository_result={urllib.parse.quote(result)}", status_code=303)


@app.post("/admin/evidence-policy")
def save_evidence_policy(
    skipped_images_incomplete: str = Form(default="true"), incomplete_noncompliant: str = Form(default="true"),
    group_id: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_config_scope),
):
    check_csrf(auth, csrf_token)
    selected_group_id = requested_group_scope(group_id, db, auth)
    if skipped_images_incomplete not in {"true", "false"}:
        raise HTTPException(422, detail="Unsupported skipped-image policy")
    if incomplete_noncompliant not in {"true", "false"}:
        raise HTTPException(422, detail="Unsupported incomplete-evidence policy")
    values = {
        "skipped_images_incomplete": skipped_images_incomplete,
        "incomplete_noncompliant": incomplete_noncompliant,
    }
    for key, value in values.items():
        storage_key = scoped_setting_key(key, selected_group_id)
        setting = db.scalar(select(PortalSetting).where(PortalSetting.key == storage_key))
        if not setting:
            setting = PortalSetting(key=storage_key, group_id=selected_group_id)
            db.add(setting)
        setting.value = value
        setting.updated_by_id = auth.user.id
        setting.updated_at = utcnow()
    record_audit(db, auth, "evidence_policy.updated", "group" if selected_group_id else "portal", str(selected_group_id or "global"), changed=list(values))
    db.commit()
    return RedirectResponse("/admin/general-policy?saved=1" + (f"&group_id={selected_group_id}" if selected_group_id else ""), status_code=303)


@app.post("/admin/workflow-policy")
def save_workflow_policy(
    warning_days: str = Form(default="14"), exception_max_days: str = Form(default="365"),
    group_id: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_config_scope),
):
    check_csrf(auth, csrf_token)
    selected_group_id = requested_group_scope(group_id, db, auth)
    try:
        warning_window = int(warning_days)
        maximum_exception = int(exception_max_days)
    except ValueError as exc:
        raise HTTPException(422, detail="Workflow policy values must be numbers") from exc
    if warning_window < 1 or warning_window > 3650:
        raise HTTPException(422, detail="Warning window must be between 1 and 3650 days")
    if maximum_exception < 0 or maximum_exception > 3650:
        raise HTTPException(422, detail="Maximum exception duration must be between 0 and 3650 days")
    values = {"warning_days": str(warning_window), "exception_max_days": str(maximum_exception)}
    for key, value in values.items():
        storage_key = scoped_setting_key(key, selected_group_id)
        setting = db.scalar(select(PortalSetting).where(PortalSetting.key == storage_key))
        if not setting:
            setting = PortalSetting(key=storage_key, group_id=selected_group_id)
            db.add(setting)
        setting.value = value
        setting.updated_by_id = auth.user.id
        setting.updated_at = utcnow()
    record_audit(db, auth, "workflow_policy.updated", "group" if selected_group_id else "portal", str(selected_group_id or "global"), changed=list(values))
    db.commit()
    return RedirectResponse("/admin/general-policy?saved=1" + (f"&group_id={selected_group_id}" if selected_group_id else ""), status_code=303)


@app.post("/admin/compliance")
def save_compliance(
    compliance_mode: str = Form(default="risk_based"), overdue_days: str = Form(),
    kev_enabled: str = Form(default="true"),
    kev_noncompliant: str = Form(default="true"), epss_enabled: str = Form(default="true"),
    epss_threshold: str = Form(default="0.90"), epss_severity: list[str] = Form(default=[]),
    epss_rule_threshold: list[str] = Form(default=[]), epss_rule_noncompliant: list[str] = Form(default=[]),
    minimum_severity: str = Form(default="None"),
    raw_due_severity: list[str] = Form(default=[]), raw_due_days: list[str] = Form(default=[]),
    group_id: str = Form(default=""), csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_config_scope),
):
    check_csrf(auth, csrf_token)
    selected_group_id = requested_group_scope(group_id, db, auth)
    try:
        overdue = int(overdue_days)
    except ValueError as exc:
        raise HTTPException(422, detail="Compliance thresholds must be numbers") from exc
    if overdue < 1 or overdue > 3650:
        raise HTTPException(422, detail="CVE aging threshold must be between 1 and 3650 days")
    if compliance_mode not in {"raw", "risk_based"}:
        raise HTTPException(422, detail="Unsupported compliance mode")
    if any(value not in {"true", "false"} for value in (kev_enabled, kev_noncompliant, epss_enabled)):
        raise HTTPException(422, detail="Risk policy switches must be true or false")
    if minimum_severity not in {"None", "Low", "Medium", "High", "Critical"}:
        raise HTTPException(422, detail="Unsupported minimum severity")
    try:
        epss_value = float(epss_threshold)
    except ValueError as exc:
        raise HTTPException(422, detail="EPSS threshold must be a number") from exc
    if epss_value < 0 or epss_value > 1:
        raise HTTPException(422, detail="EPSS threshold must be between 0 and 1")
    rules = []
    for index, severity in enumerate(epss_severity):
        severity = severity.strip()
        if not severity:
            continue
        threshold_text = epss_rule_threshold[index] if index < len(epss_rule_threshold) else ""
        try:
            threshold = float(threshold_text)
        except ValueError as exc:
            raise HTTPException(422, detail="Each EPSS rule threshold must be a number") from exc
        if threshold < 0 or threshold > 1:
            raise HTTPException(422, detail="Each EPSS rule threshold must be between 0 and 1")
        noncompliant = index < len(epss_rule_noncompliant) and epss_rule_noncompliant[index] == "true"
        rules.append({"severity": severity, "threshold": threshold, "noncompliant": noncompliant})
    if not rules:
        rules = [{"severity": "Any", "threshold": epss_value, "noncompliant": True}]
    raw_rules = []
    for index, severity in enumerate(raw_due_severity):
        severity = severity.strip()
        if not severity:
            continue
        days_text = raw_due_days[index] if index < len(raw_due_days) else ""
        try:
            days = int(days_text)
        except ValueError as exc:
            raise HTTPException(422, detail="Each raw due-date rule must be a whole number") from exc
        if days < 1 or days > 3650:
            raise HTTPException(422, detail="Raw due-date rules must be between 1 and 3650 days")
        raw_rules.append({"severity": severity, "days": days})
    if not raw_rules:
        raw_rules = [{"severity": "Critical", "days": 30}, {"severity": "High", "days": 60}, {"severity": "Medium", "days": 90}, {"severity": "Low", "days": 120}]
    values = {
        "compliance_mode": compliance_mode, "overdue_days": str(overdue),
        "kev_enabled": kev_enabled,
        "kev_noncompliant": kev_noncompliant, "epss_enabled": epss_enabled,
        "epss_threshold": str(epss_value), "epss_rules": json.dumps(rules, separators=(",", ":")),
        "minimum_severity": minimum_severity,
        "raw_due_rules": json.dumps(raw_rules, separators=(",", ":")),
    }
    for key, value in values.items():
        storage_key = scoped_setting_key(key, selected_group_id)
        setting = db.scalar(select(PortalSetting).where(PortalSetting.key == storage_key))
        if not setting:
            setting = PortalSetting(key=storage_key, group_id=selected_group_id)
            db.add(setting)
        setting.value = value
        setting.updated_by_id = auth.user.id
        setting.updated_at = utcnow()
    record_audit(db, auth, "vulnerability_policy.updated", "group" if selected_group_id else "portal", str(selected_group_id or "global"), changed=list(values))
    db.commit()
    return RedirectResponse("/admin/compliance?saved=1" + (f"&group_id={selected_group_id}" if selected_group_id else ""), status_code=303)


@app.post("/admin/compliance-frameworks")
def save_hardening_policy(
    hardening_overdue_days: str = Form(default="90"),
    hardening_noncompliant: str = Form(default="true"),
    group_id: str = Form(default=""), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_global_config_scope),
):
    check_csrf(auth, csrf_token)
    selected_group_id = requested_group_scope(group_id, db, auth)
    try:
        threshold = int(hardening_overdue_days)
    except ValueError as exc:
        raise HTTPException(422, detail="Hardening threshold must be a whole number") from exc
    if threshold < 1 or threshold > 3650:
        raise HTTPException(422, detail="Hardening threshold must be between 1 and 3650 days")
    if hardening_noncompliant not in {"true", "false"}:
        raise HTTPException(422, detail="Hardening treatment must be true or false")
    values = {"hardening_overdue_days": str(threshold), "hardening_noncompliant": hardening_noncompliant}
    for key, value in values.items():
        storage_key = scoped_setting_key(key, selected_group_id)
        setting = db.scalar(select(PortalSetting).where(PortalSetting.key == storage_key))
        if not setting:
            setting = PortalSetting(key=storage_key, group_id=selected_group_id)
            db.add(setting)
        setting.value = value
        setting.updated_by_id = auth.user.id
        setting.updated_at = utcnow()
    record_audit(db, auth, "hardening_policy.updated", "group" if selected_group_id else "portal", str(selected_group_id or "global"), changed=list(values))
    db.commit()
    return RedirectResponse("/admin/compliance-frameworks?saved=1" + (f"&group_id={selected_group_id}" if selected_group_id else ""), status_code=303)


@app.post("/admin/users")
def create_user(
    username: str = Form(), display_name: str = Form(), password: str = Form(),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("user.manage")),
):
    check_csrf(auth, csrf_token)
    username = username.strip().lower()
    if not username or db.scalar(select(User).where(User.username == username)):
        raise HTTPException(409, detail="Username is invalid or already exists")
    try:
        password_hash = hash_password(password)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc
    user = User(username=username, display_name=display_name.strip() or username,
                password_hash=password_hash, must_change_password=True)
    db.add(user)
    db.flush()
    record_audit(db, auth, "user.created", "user", user.id, username=username)
    db.commit()
    return RedirectResponse("/admin", status_code=303)


@app.post("/admin/roles")
def create_role(
    name: str = Form(), description: str = Form(default=""), permissions: list[str] = Form(default=[]),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("role.manage")),
):
    check_csrf(auth, csrf_token)
    selected = sorted(set(permissions))
    if not name.strip() or db.scalar(select(Role).where(Role.name == name.strip())):
        raise HTTPException(409, detail="Role name is invalid or already exists")
    if any(item not in PERMISSIONS for item in selected):
        raise HTTPException(422, detail="Unknown permission")
    role = Role(name=name.strip(), description=description.strip(), permissions=selected, system=False)
    db.add(role)
    db.flush()
    record_audit(db, auth, "role.created", "role", role.id, name=role.name, permissions=selected)
    db.commit()
    return RedirectResponse("/admin", status_code=303)


@app.post("/admin/assignments")
def assign_role(
    user_id: int = Form(), role_id: int = Form(), service_ids: list[int] = Form(default=[]),
    service_id: str = Form(default=""), group_id: str = Form(default=""),
    csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("role.manage")),
):
    check_csrf(auth, csrf_token)
    user, role = db.get(User, user_id), db.get(Role, role_id)
    if not user or not role:
        raise HTTPException(404)
    selected_service_ids = list(dict.fromkeys(service_ids))
    if not selected_service_ids and service_id:
        selected_service_ids = [int(service_id)]
    group_scope = int(group_id) if group_id else None
    if selected_service_ids and group_scope is not None:
        if role.name != "Service Manager":
            raise HTTPException(422, detail="Only Service Manager assignments may combine group and service scope")
        group = db.get(Group, group_scope)
        if not group or any(not db.get(Service, selected_id) or not any(item.id == selected_id for item in group.services) for selected_id in selected_service_ids):
            raise HTTPException(422, detail="Every selected service must belong to the selected group")
    if (selected_service_ids or group_scope is not None) and any(p in role.permissions for p in {"user.manage", "role.manage", "service.delete"}):
        raise HTTPException(422, detail="Administrative permissions must be assigned globally")
    scopes = selected_service_ids or [None]
    try:
        for scope in scopes:
            exists = db.scalar(select(UserRoleAssignment).where(
                UserRoleAssignment.user_id == user_id, UserRoleAssignment.role_id == role_id,
                UserRoleAssignment.service_id == scope,
                UserRoleAssignment.group_id == group_scope,
            ))
            if not exists:
                assignment = UserRoleAssignment(user_id=user_id, role_id=role_id, service_id=scope, group_id=group_scope)
                db.add(assignment)
                db.flush()
                record_audit(db, auth, "role.assigned", "user_role_assignment", assignment.id,
                             user_id=user_id, role_id=role_id, service_id=scope, group_id=group_scope)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, detail="One or more selected assignments already exist or conflict with an existing assignment") from exc
    return RedirectResponse("/admin", status_code=303)


@app.post("/admin/assignments/{assignment_id}/delete")
def remove_assignment(
    assignment_id: int, csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("role.manage")),
):
    check_csrf(auth, csrf_token)
    assignment = db.get(UserRoleAssignment, assignment_id)
    if not assignment:
        raise HTTPException(404)
    if assignment.user_id == auth.user.id and assignment.role.name == "Administrator":
        raise HTTPException(409, detail="You cannot remove your own Administrator role")
    record_audit(db, auth, "role.unassigned", "user_role_assignment", assignment.id,
                 user_id=assignment.user_id, role_id=assignment.role_id, service_id=assignment.service_id)
    db.delete(assignment)
    db.commit()
    return RedirectResponse("/admin", status_code=303)


@app.post("/admin/users/{user_id}/toggle")
def toggle_user(
    user_id: int, csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("user.manage")),
):
    check_csrf(auth, csrf_token)
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(404)
    if user.id == auth.user.id:
        raise HTTPException(409, detail="You cannot disable your own account")
    user.enabled = not user.enabled
    if not user.enabled:
        db.execute(delete(UserSession).where(UserSession.user_id == user.id))
    record_audit(db, auth, "user.enabled" if user.enabled else "user.disabled", "user", user.id)
    db.commit()
    return RedirectResponse("/admin", status_code=303)


@app.post("/admin/users/{user_id}/reset-password")
def reset_user_password(
    user_id: int, temporary_password: str = Form(), csrf_token: str = Form(),
    db: Session = Depends(get_db), auth: AuthContext = Depends(require_permission("user.manage")),
):
    check_csrf(auth, csrf_token)
    user = db.get(User, user_id)
    if not user or user.auth_source != "local":
        raise HTTPException(404)
    try:
        user.password_hash = hash_password(temporary_password)
    except ValueError as exc:
        raise HTTPException(422, detail=str(exc)) from exc
    user.must_change_password = True
    user.failed_login_count = 0
    user.locked_until = None
    db.execute(delete(UserSession).where(UserSession.user_id == user.id))
    record_audit(db, auth, "user.password_reset", "user", user.id)
    db.commit()
    return RedirectResponse("/admin", status_code=303)


@app.post("/admin/users/{user_id}/delete")
def delete_user_account(
    user_id: int, csrf_token: str = Form(), db: Session = Depends(get_db),
    auth: AuthContext = Depends(require_permission("user.manage")),
):
    check_csrf(auth, csrf_token)
    if user_id == auth.user.id:
        raise HTTPException(409, detail="You cannot delete your own account")
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(404, detail="User account not found")
    # Never remove the last global Administrator/break-glass account.
    administrator = db.scalar(select(Role).where(Role.name == "Administrator"))
    if administrator:
        remaining = db.scalar(select(func.count(UserRoleAssignment.id)).where(
            UserRoleAssignment.user_id == user.id,
            UserRoleAssignment.role_id == administrator.id,
            UserRoleAssignment.service_id.is_(None), UserRoleAssignment.group_id.is_(None),
        )) or 0
        if remaining and (db.scalar(select(func.count(UserRoleAssignment.id)).where(
            UserRoleAssignment.role_id == administrator.id,
            UserRoleAssignment.service_id.is_(None), UserRoleAssignment.group_id.is_(None),
            UserRoleAssignment.user_id != user.id,
        )) or 0) == 0:
            raise HTTPException(409, detail="The last global Administrator account cannot be deleted")
    username, display_name = user.username, user.display_name
    for event in db.scalars(select(AuditEvent).where(AuditEvent.actor_user_id == user.id)):
        detail = dict(event.detail or {})
        detail.setdefault("actor_username", username)
        detail.setdefault("actor_display_name", display_name)
        event.detail = detail
        event.actor_user_id = None
    for history in db.scalars(select(PoamHistory).where(PoamHistory.actor_user_id == user.id)):
        history.actor_user_id = None
    # Re-point non-null historical foreign keys to the administrator performing
    # the deletion; audit snapshots above preserve the original actor.
    for workflow in db.scalars(select(WorkflowRequest).where(WorkflowRequest.requested_by_id == user.id)):
        workflow.requested_by_id = auth.user.id
    for workflow in db.scalars(select(WorkflowRequest).where(WorkflowRequest.reviewed_by_id == user.id)):
        workflow.reviewed_by_id = auth.user.id
    for entry in db.scalars(select(PoamEntry).where(PoamEntry.created_by_id == user.id)):
        entry.created_by_id = auth.user.id
    for entry in db.scalars(select(PoamEntry).where(PoamEntry.approved_by_id == user.id)):
        entry.approved_by_id = auth.user.id
    for patch_job in db.scalars(select(PatchExecution).where(PatchExecution.requested_by_id == user.id)):
        patch_job.requested_by_id = auth.user.id
    for setting in db.scalars(select(PortalSetting).where(PortalSetting.updated_by_id == user.id)):
        setting.updated_by_id = auth.user.id
    for image in db.scalars(select(ServiceImage).where(or_(ServiceImage.requested_by_id == user.id, ServiceImage.approved_by_id == user.id))):
        if image.requested_by_id == user.id:
            image.requested_by_id = None
        if image.approved_by_id == user.id:
            image.approved_by_id = None
    db.execute(delete(UserSession).where(UserSession.user_id == user.id))
    db.execute(delete(UserRoleAssignment).where(UserRoleAssignment.user_id == user.id))
    record_audit(db, auth, "user.deleted", "user", user.id, username=username, display_name=display_name)
    db.delete(user)
    db.commit()
    return RedirectResponse("/admin", status_code=303)
