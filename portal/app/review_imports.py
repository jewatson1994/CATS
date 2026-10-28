"""Previewed, atomic imports whose individual rows always require Cyber review."""
import csv
import json
import secrets
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from sqlalchemy import inspect, select, text, update

from .auth import record_audit, require_user
from .database import get_db
from .exchange_limits import workbook_bytes, workbook_columns, workbook_expanded_bytes, workbook_members, workbook_rows
from .models import ExchangePreview, Finding, PolicyFinding, PoamEntry, PoamHistory, Service, WorkflowImportBatch, WorkflowRequest

router = APIRouter()
KINDS = {
    "poams": ("title", "description", "remediation", "due_date", "ticket", "finding_id", "policy_finding_id"),
    "mitigations": ("title", "description", "remediation", "due_date", "ticket", "finding_id", "policy_finding_id"),
    "exceptions": ("finding_id", "policy_finding_id", "justification", "expires_at", "ticket"),
}


def upgrade_workflow_imports(connection):
    columns = {column["name"] for column in inspect(connection).get_columns("workflow_requests")}
    for name in ("import_batch_id", "import_row_number"):
        if name not in columns:
            connection.execute(text(f"ALTER TABLE workflow_requests ADD COLUMN {name} INTEGER"))
    connection.execute(text("CREATE INDEX IF NOT EXISTS ix_workflow_requests_import_batch_id ON workflow_requests (import_batch_id)"))


def permitted(auth, kind, service_id):
    return (auth.has("exception.request", service_id) if kind == "exceptions" else
            auth.has("poam.request", service_id) or auth.has("poam.import", service_id))


def scoped_service(db, auth, service_key, kind):
    if kind not in KINDS:
        raise HTTPException(404, detail="Unknown import type")
    service = db.scalar(select(Service).where(Service.service_key == service_key))
    if service is None:
        raise HTTPException(404)
    if not permitted(auth, kind, service.id):
        raise HTTPException(403, detail="Permission denied for this service")
    return service


def _error(row, field, message):
    return {"row": row, "field": field, "message": message}


def parse_upload(data, filename, kind):
    """One explicit table; formulas, duplicate headings and oversized files fail closed."""
    if len(data) > workbook_bytes():
        raise ValueError("Upload exceeds the configured file size limit")
    suffix = PurePath(filename.lower()).suffix
    book = None
    if suffix == ".csv":
        try:
            table = csv.reader(StringIO(data.decode("utf-8-sig")), strict=True)
        except UnicodeDecodeError as exc:
            raise ValueError("CSV must be saved with UTF-8 encoding") from exc
    elif suffix == ".xlsx":
        try:
            with ZipFile(BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > workbook_members() or sum(m.file_size for m in members) > workbook_expanded_bytes():
                    raise ValueError("Workbook exceeds expanded/member safety limits")
                if len({m.filename for m in members}) != len(members):
                    raise ValueError("Workbook contains duplicate ZIP entries")
            book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            if len(book.worksheets) != 1:
                raise ValueError("Use a workbook with one worksheet and headings in the first row")
            sheet = book.active
            if sheet.max_row > workbook_rows() + 1 or sheet.max_column > workbook_columns():
                raise ValueError("Workbook exceeds the configured row or column limit")
            table = ([cell.value for cell in cells] for cells in sheet.iter_rows())
        except (BadZipFile, KeyError, OSError, ParseError) as exc:
            if book:
                book.close()
            raise ValueError("Malformed XLSX workbook") from exc
        except ValueError:
            if book:
                book.close()
            raise
    else:
        raise ValueError("Choose a .csv or .xlsx file")
    rows, errors, seen = [], [], set()
    try:
        headings = [str(value or "").strip().lower().replace(" ", "_") for value in next(table, [])]
        if len(headings) > workbook_columns():
            raise ValueError("Upload exceeds the configured column limit")
        if not headings or len(headings) != len(set(headings)):
            raise ValueError("Supply one row of unique column headings")
        unknown = set(headings) - set(KINDS[kind])
        if unknown:
            raise ValueError("Unknown columns: " + ", ".join(sorted(unknown)) + ". Download the import template.")
        required = {"justification", "expires_at"} if kind == "exceptions" else {"title", "description", "remediation"}
        if not required.issubset(headings):
            raise ValueError("Missing required columns: " + ", ".join(sorted(required - set(headings))))
        for number, cells in enumerate(table, 2):
            if not any(value not in (None, "") for value in cells):
                continue
            if len(rows) >= workbook_rows():
                raise ValueError("Upload exceeds the configured row limit")
            if len(cells) > len(headings):
                errors.append(_error(number, "", "More values than column headings"))
            values = {}
            for field, value in zip(headings, cells):
                if isinstance(value, datetime):
                    value = value.isoformat()
                value = str(value if value is not None else "").strip()
                if value.startswith("="):
                    errors.append(_error(number, field, "Formulas are not importable"))
                if len(value) > 8000:
                    errors.append(_error(number, field, "Value exceeds 8,000 characters"))
                values[field] = value
            identity = json.dumps(values, sort_keys=True)
            if identity in seen:
                errors.append(_error(number, "", "Duplicate row in this upload"))
            seen.add(identity)
            rows.append({"row": number, "values": values})
    except csv.Error as exc:
        raise ValueError("Malformed CSV: " + str(exc)) from exc
    finally:
        if book:
            book.close()
    if not rows:
        errors.append(_error(1, "", "The upload contains no entries"))
    return rows, errors


def validate_rows(db, service, kind, rows):
    """Repeat domain validation on confirmation while the service write lock is held."""
    from .main import active_exception, configuration_for_service
    from .native_vex import finding_excluded
    errors, targets = [], set()
    now = datetime.now(timezone.utc)
    maximum = int(configuration_for_service(db, service).get("exception_max_days", "365")) if kind == "exceptions" else 0
    for row in rows:
        number, values = row["row"], row["values"]
        required = ("justification",) if kind == "exceptions" else ("title", "description", "remediation")
        for field in required:
            if len(str(values.get(field) or "").strip()) < 3:
                errors.append(_error(number, field, "Supply at least 3 characters"))
        for field in ("title", "ticket"):
            if len(str(values.get(field) or "")) > 240:
                errors.append(_error(number, field, "Value exceeds 240 characters"))
        linked = []
        for field, model in (("finding_id", Finding), ("policy_finding_id", PolicyFinding)):
            raw = values.get(field)
            if not raw:
                values[field] = None
                continue
            try:
                identifier = int(raw)
                if identifier <= 0 or str(raw) != str(identifier):
                    raise ValueError()
            except (ValueError, TypeError):
                errors.append(_error(number, field, "Supply a positive whole number"))
                continue
            values[field] = identifier
            target = db.get(model, identifier)
            if not target or target.service_id != service.id:
                errors.append(_error(number, field, "Finding does not belong to the selected service"))
            elif not target.active:
                errors.append(_error(number, field, "Finding is no longer active"))
            elif kind == "exceptions" and active_exception(target, now):
                errors.append(_error(number, field, "An approved exception is already active for this finding; revoke or let it expire before requesting another"))
            elif field == "finding_id" and kind == "poams" and db.scalar(select(finding_excluded(identifier, now))):
                errors.append(_error(number, field, "Verified VEX Not Affected evidence excludes this finding from actionable POA&M entries"))
            linked.append((field, identifier))
        if len(linked) > 1 or (kind == "exceptions" and not linked):
            errors.append(_error(number, "finding_id", "Supply exactly one finding_id or policy_finding_id" if kind == "exceptions" else "Link at most one vulnerability or configuration finding"))
        if kind == "exceptions" and len(linked) == 1:
            field, identifier = linked[0]
            if (field, identifier) in targets:
                errors.append(_error(number, field, "Duplicate exception target in this upload"))
            targets.add((field, identifier))
            if db.scalar(select(WorkflowRequest.id).where(WorkflowRequest.status == "pending", WorkflowRequest.request_type == "exception", getattr(WorkflowRequest, field) == identifier)):
                errors.append(_error(number, field, "An exception request is already pending for this finding"))
        field = "expires_at" if kind == "exceptions" else "due_date"
        raw = values.get(field)
        if not raw:
            if kind == "exceptions":
                errors.append(_error(number, field, "Supply a future expiration date"))
            values[field] = None
        else:
            try:
                date = datetime.fromisoformat(str(raw))
                date = date.replace(tzinfo=timezone.utc) if date.tzinfo is None else date.astimezone(timezone.utc)
                values[field] = date.isoformat()
                if kind == "exceptions" and date <= now:
                    errors.append(_error(number, field, "Date must be in the future"))
                elif maximum and date > now + timedelta(days=maximum):
                    errors.append(_error(number, field, f"Expiration cannot exceed {maximum} days"))
            except ValueError:
                errors.append(_error(number, field, "Use an ISO date, for example 2027-01-31"))
    return errors


def create_batch(db, service_id, user_id, kind, filename):
    batch = WorkflowImportBatch(service_id=service_id, requested_by_id=user_id, kind=kind, filename=filename[:240])
    db.add(batch)
    db.flush()
    return batch


def add_poam_request(db, auth, entry, batch, row_number, *, request_type="poam", justification=None):
    workflow = WorkflowRequest(request_type=request_type, service_id=entry.service_id,
        finding_id=entry.finding_id, policy_finding_id=entry.policy_finding_id,
        poam_id=entry.id, requested_by_id=auth.user.id,
        import_batch_id=batch.id, import_row_number=row_number,
        justification=justification or f"{entry.title}\n{entry.description}\nRemediation: {entry.remediation}",
        requested_expires_at=entry.due_date, ticket=entry.ticket)
    db.add(workflow)
    db.add(PoamHistory(poam_id=entry.id, actor_user_id=auth.user.id, action="import_requested",
        note="Imported for Cybersecurity approval", detail={"batch_id": batch.id, "row": row_number}))
    return workflow


@router.get("/services/{service_key}/review-imports/{kind}/template.csv")
def import_template(service_key: str, kind: str, db=Depends(get_db), auth=Depends(require_user)):
    scoped_service(db, auth, service_key, kind)
    output = StringIO()
    csv.writer(output).writerow(KINDS[kind])
    return Response(output.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="cats-{kind}-import.csv"'})


@router.post("/services/{service_key}/review-imports/{kind}/preview")
def preview_import(service_key: str, kind: str, csrf_token: str = Form(), upload: UploadFile = File(), db=Depends(get_db), auth=Depends(require_user)):
    from .exchange_routes import csrf
    csrf(auth, csrf_token)
    service = scoped_service(db, auth, service_key, kind)
    filename = (upload.filename or "upload.csv").replace("\\", "/").rsplit("/", 1)[-1][:240]
    try:
        rows, errors = parse_upload(upload.file.read(workbook_bytes() + 1), filename, kind)
    except ValueError as exc:
        return {"token": None, "kind": kind, "rows": [], "recognized": 0, "errors": [_error(1, "", str(exc))]}
    errors.extend(validate_rows(db, service, kind, rows))
    payload = {"kind": kind, "filename": filename, "rows": rows, "errors": errors, "recognized": len(rows)}
    token = None
    if not errors:
        token = secrets.token_hex(24)
        db.add(ExchangePreview(token=token, service_id=service.id, user_id=auth.user.id,
            version="", dataset="review_" + kind, payload=payload,
            expires_at=datetime.now(timezone.utc) + timedelta(minutes=30)))
        db.commit()
    return {"token": token, **payload}


@router.post("/services/{service_key}/review-imports/confirm/{token}")
def confirm_import(service_key: str, token: str, csrf_token: str = Form(), db=Depends(get_db), auth=Depends(require_user)):
    from .exchange_routes import csrf, lock_service
    csrf(auth, csrf_token)
    preview = db.get(ExchangePreview, token)
    if preview is None or preview.user_id != auth.user.id or not preview.dataset.startswith("review_"):
        raise HTTPException(404)
    kind = preview.payload["kind"]
    service = scoped_service(db, auth, service_key, kind)
    now = datetime.now(timezone.utc)
    expiry = preview.expires_at.replace(tzinfo=timezone.utc) if preview.expires_at.tzinfo is None else preview.expires_at
    if preview.service_id != service.id or preview.consumed or expiry <= now:
        raise HTTPException(409, detail="Preview expired or already consumed")
    lock_service(db, service.id)
    rows = preview.payload["rows"]
    errors = validate_rows(db, service, kind, rows)
    if errors:
        detail = "; ".join(f"Row {error['row']} ({error['field']}): {error['message']}" for error in errors[:10])
        raise HTTPException(409, detail="Rows changed since preview; correct them and upload again. " + detail)
    try:
        claimed = db.execute(update(ExchangePreview).where(ExchangePreview.token == token,
            ExchangePreview.consumed.is_(False), ExchangePreview.expires_at > now).values(consumed=True),
            execution_options={"synchronize_session": False})
        if claimed.rowcount != 1:
            raise HTTPException(409, detail="Preview already consumed")
        batch = create_batch(db, service.id, auth.user.id, kind, preview.payload["filename"])
        for row in rows:
            values = row["values"]
            date_field = "expires_at" if kind == "exceptions" else "due_date"
            date = datetime.fromisoformat(values[date_field]) if values.get(date_field) else None
            if kind == "exceptions":
                db.add(WorkflowRequest(request_type="exception", service_id=service.id,
                    finding_id=values.get("finding_id"), policy_finding_id=values.get("policy_finding_id"),
                    requested_by_id=auth.user.id, justification=values["justification"],
                    requested_expires_at=date, ticket=values.get("ticket") or None,
                    import_batch_id=batch.id, import_row_number=row["row"]))
            else:
                entry = PoamEntry(service_id=service.id, finding_id=values.get("finding_id"),
                    policy_finding_id=values.get("policy_finding_id"), created_by_id=auth.user.id,
                    item_type="mitigation" if kind == "mitigations" else "vulnerability" if values.get("finding_id") else "missing_requirement",
                    title=values["title"], description=values["description"], remediation=values["remediation"],
                    due_date=date, ticket=values.get("ticket") or None, status="pending_approval")
                db.add(entry)
                db.flush()
                add_poam_request(db, auth, entry, batch, row["row"])
        record_audit(db, auth, "review_import.requested", "workflow_import_batch", batch.id,
            service_id=service.id, kind=kind, rows=len(rows))
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"batch_id": batch.id, "count": len(rows), "href": f"/requests?batch_id={batch.id}"}


def batch_summaries(workflows):
    batches = {}
    for workflow in workflows:
        if not workflow.import_batch_id:
            continue
        batch = workflow.import_batch
        summary = batches.setdefault(batch.id, {"id": batch.id, "name": f"{workflow.service.name} · {batch.filename}",
            "kind": batch.kind, "count": 0, "pending": 0, "approved": 0, "rejected": 0,
            "href": f"/requests?batch_id={batch.id}"})
        summary["count"] += 1
        if workflow.status in {"pending", "approved", "rejected"}:
            summary[workflow.status] += 1
    return list(batches.values())
