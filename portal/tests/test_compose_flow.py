"""Compose discovery through the real durable worker protocol and ingestion."""
import hashlib
import io
import json
import tarfile
import threading
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from test_scan_worker_coordination import database
from app import main, scan_control, scan_protocol, scan_worker
from app.models import Execution, Finding


APP_IMAGE = "harbor.example.test/team/app:2"
SECRETS = ("compose-env-secret-value", "compose-command-secret-value", "service-env-secret-value")


def compose_bundle():
    """An ordered override, shared image, profile, env files and build-only service."""
    files = {
        "compose.yaml": """services:
  web:
    image: harbor.example.test/team/app:1
    ports: ["8080:80"]
    privileged: true
    environment:
      API_PASSWORD: compose-env-secret-value
    command: ["serve", "--password=compose-command-secret-value"]
    env_file: [service.env]
  worker:
    image: harbor.example.test/team/app:1
  db:
    image: postgres:16
    profiles: [data]
  debug:
    image: redis:7
    profiles: [debug]
  local:
    build: .
    profiles: [local]
""",
        "override.yaml": """services:
  web:
    image: harbor.example.test/team/app:${TAG}
  worker:
    image: harbor.example.test/team/app:${TAG}
""",
        ".env": "TAG=2\n",
        "service.env": "DB_PASSWORD=service-env-secret-value\n",
    }
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return stream.getvalue(), "application.zip"


@pytest.fixture
def compose_http(database, monkeypatch, tmp_path):
    monkeypatch.setattr(main, "SessionLocal", database)
    monkeypatch.setattr(scan_protocol, "SessionLocal", database)
    monkeypatch.setattr(main, "PUBLIC_JOB_ROOT", tmp_path / "portal")
    main.PUBLIC_JOB_ROOT.mkdir()
    monkeypatch.setattr(scan_worker, "ROOT", tmp_path / "worker")
    scan_worker.ROOT.mkdir()
    monkeypatch.setattr(scan_worker, "STOP", threading.Event())
    monkeypatch.setattr(scan_worker.shutil, "which", lambda _: None)
    token = "compose-integration-" + "7" * 40
    monkeypatch.setenv("CATS_SCAN_WORKER_TOKEN", token)
    monkeypatch.setenv("CATS_SCAN_WORKER_ID", "compose-integration-worker")
    monkeypatch.delenv("CATS_SCAN_WORKER_CREDENTIALS", raising=False)
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    # This value must not affect Compose interpolation.
    monkeypatch.setenv("TAG", "process-must-not-win")
    registry = tmp_path / "registry-auth"
    registry.mkdir()
    (registry / "config.json").write_text(json.dumps({"auths": {
        "harbor.example.test": {"auth": "harbor-registry-credential"},
        "unrelated.example.test": {"auth": "unrelated-registry-credential"},
    }}), encoding="utf-8")
    monkeypatch.setenv("DOCKER_CONFIG", str(registry))
    headers = {"Authorization": "Bearer " + token, "X-Worker-ID": "compose-integration-worker"}
    control = TestClient(scan_control.app)
    browser = TestClient(main.app)
    main.app.dependency_overrides[main.optional_user] = lambda: SimpleNamespace(
        user=SimpleNamespace(id=17), has=lambda permission, service: False)
    calls = []

    class ScannerStandIn:
        returncode = 0
        def __init__(self, command, **kwargs):
            inputs, output = Path(command[-2]), Path(command[-1])
            artifacts = {str(p.relative_to(inputs)): p.read_bytes() for p in inputs.rglob("*") if p.is_file()}
            auth_path = Path(kwargs["env"]["DOCKER_CONFIG"]) / "config.json"
            calls.append({"command": command, "inputs": artifacts,
                          "auth": json.loads(auth_path.read_text()) if auth_path.exists() else {}})
            (output / "scan-summary.json").write_text(json.dumps({"status": "complete", "images": 2, "findings": 1}), encoding="utf-8")
            result = {"schema_version": "1.0", "execution_id": "scanner-placeholder",
                "scanned_at": datetime.now(timezone.utc).isoformat(), "complete": True,
                "sbom_images": [APP_IMAGE, "postgres:16"],
                "service": {"id": "service", "name": "Coordination test service", "version": "v1"},
                "findings": [{"cve": "CVE-2026-1212", "severity": "High", "image": APP_IMAGE,
                              "package": "openssl", "fixed_version": "9.9",
                              "evidence": {"description": "Scanner stand-in evidence"}}]}
            (output / "portal-result.json").write_text(json.dumps(result), encoding="utf-8")
        def poll(self): return 0

    monkeypatch.setattr(scan_worker.subprocess, "Popen", ScannerStandIn)
    # No extra subprocess is needed to interpret submitted Compose files.
    monkeypatch.setattr(scan_worker.subprocess, "run", lambda *a, **kw: pytest.fail("Unexpected subprocess during Compose discovery"))

    class ResponseAdapter:
        def __init__(self, response): self.response = response
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def json(self): return self.response.json()
        def iter_content(self, size): yield self.response.content

    worker = scan_worker.Worker.__new__(scan_worker.Worker)
    worker.worker_id = "compose-integration-worker"
    def request(method, suffix, **kwargs):
        kwargs.pop("timeout", None)
        kwargs.pop("stream", None)
        request_headers = {**headers, **kwargs.pop("headers", {})}
        if "data" in kwargs:
            kwargs["content"] = kwargs.pop("data").read()
        response = control.request(method, "/internal/scan-worker" + suffix, headers=request_headers, **kwargs)
        response.raise_for_status()
        return ResponseAdapter(response)
    worker.request = request

    def run(*, preview=False):
        job_id = main._start_public_scan("", owner_user_id=17,
            ingest_service_id=None if preview else "service", ingest_service_version="v1",
            compose_uploads=[compose_bundle()], compose_files=["compose.yaml", "override.yaml"],
            compose_profiles=["data", "local"], compose_preview=preview)
        claim = worker.claim()
        assert claim["job_id"] == job_id
        worker._execute(claim, None)
        scan_protocol.ingest_one()
        output = main.PUBLIC_JOB_ROOT / job_id / "output"
        result = json.loads((output / "portal-result.json").read_text(encoding="utf-8"))
        summary = json.loads((output / "scan-summary.json").read_text(encoding="utf-8"))
        return job_id, output, result, summary

    try:
        yield SimpleNamespace(run=run, calls=calls, browser=browser, database=database)
    finally:
        main.app.dependency_overrides.clear()
        browser.close()
        control.close()


def assert_sanitized(output, result, summary):
    encoded = json.dumps({"result": result, "summary": summary}).encode()
    returned = b"\n".join(path.read_bytes() for path in output.rglob("*") if path.is_file())
    for secret in (*SECRETS, "harbor-registry-credential", "unrelated-registry-credential", "process-must-not-win"):
        assert secret.encode() not in encoded
        assert secret.encode() not in returned
    assert not any(path.name in {"compose.yaml", "override.yaml", ".env", "service.env"} for path in output.rglob("*"))


def test_compose_http_worker_ingests_shared_image_and_missing_evidence(compose_http):
    job_id, output, result, summary = compose_http.run()
    assert len(compose_http.calls) == 1
    call = compose_http.calls[0]
    images = call["inputs"]["images.txt"].decode().splitlines()
    assert sorted(images) == sorted([APP_IMAGE, "postgres:16"])
    assert set(call["auth"]["auths"]) == {"harbor.example.test"}
    assert "compose-discovery.json" in call["inputs"]
    assert not any(name.endswith(("compose.yaml", "override.yaml", ".env", "service.env")) for name in call["inputs"])
    for secret in SECRETS:
        assert all(secret.encode() not in content for content in call["inputs"].values())
    assert_sanitized(output, result, summary)
    model = result["service_overview"]["compose"]
    services = {service["name"]: service for service in model["services"]}
    assert set(services) == {"web", "worker", "db", "local"}
    assert services["web"]["image"] == services["worker"]["image"] == APP_IMAGE
    assert services["web"]["source_files"] == ["compose.yaml", "override.yaml"]
    assert str(services["web"]["configuration"]["ports"][0]["target"]) == "80"
    assert services["local"]["image"] is None
    assert any(issue["code"] == "build_only" for issue in model["issues"])
    assert any(finding.get("finding") == "COMPOSE-PRIVILEGED" for finding in model["findings"])
    assert summary["compose"]
    with tarfile.open(output / "results-export.tar.gz", "r:gz") as archive:
        exported = {member.name: archive.extractfile(member).read()
                    for member in archive.getmembers() if member.isfile()}
    manifest = json.loads(exported["results/manifest.json"])
    assert manifest["metadata"]["service-overview.json"]["compose"] == model
    assert manifest["metadata"]["scan-summary.json"]["compose"] == summary["compose"]
    report = json.loads(exported["results/cats/compose-findings.json"])
    assert {item["unique_id_from_tool"] for item in report["findings"]} == {
        item["finding"] + ":" + item["target"] for item in model["findings"]}
    entry = next(item for item in manifest["reports"] if item["path"] == "cats/compose-findings.json")
    assert entry["sha256"] == hashlib.sha256(exported["results/" + entry["path"]]).hexdigest()
    assert all(secret.encode() not in content for secret in SECRETS for content in exported.values())
    assert result["complete"] is False
    assert main.PUBLIC_JOBS[job_id]["status"] == "incomplete"
    assert main.PUBLIC_JOBS[job_id]["ingested"] is True
    with compose_http.database() as db:
        execution = db.scalar(select(Execution).where(Execution.execution_key == "public:" + job_id))
        assert execution is not None
        assert db.scalar(select(Finding).where(Finding.service_id == 1, Finding.cve == "CVE-2026-1212")) is not None
    status = compose_http.browser.get("/api/public/jobs/" + job_id)
    assert status.status_code == 200
    assert status.json()["status"] == "incomplete"
    assert all(secret not in status.text for secret in SECRETS)


def test_compose_preview_returns_discovery_without_scanning_or_ingesting(compose_http):
    job_id, output, result, summary = compose_http.run(preview=True)
    assert compose_http.calls == []
    model = result["service_overview"]["compose"]
    services = {service["name"]: service for service in model["services"]}
    assert set(services) == {"web", "worker", "db", "local"}
    assert services["web"]["image"] == services["worker"]["image"] == APP_IMAGE
    assert services["web"]["source_files"] == ["compose.yaml", "override.yaml"]
    assert str(services["web"]["configuration"]["ports"][0]["target"]) == "80"
    assert services["local"]["image"] is None
    assert any(issue["code"] == "build_only" for issue in model["issues"])
    assert any(finding.get("finding") == "COMPOSE-PRIVILEGED" for finding in model["findings"])
    assert summary["compose"]
    assert_sanitized(output, result, summary)
    assert main.PUBLIC_JOBS[job_id]["ingested"] is False
    with compose_http.database() as db:
        assert db.scalar(select(Execution).where(Execution.execution_key == "public:" + job_id)) is None


@pytest.mark.parametrize("preview", [False, True])
def test_compose_form_preserves_upload_order_and_preview_scope(monkeypatch, preview):
    captured = []
    def start(*args, **kwargs):
        uploads = [(source.read() if hasattr(source, "read") else source, name)
                   for source, name in kwargs["compose_uploads"]]
        captured.append((args, kwargs, uploads))
        return "c" * 32
    monkeypatch.setattr(main, "_start_public_scan", start)
    monkeypatch.setattr(main, "get_global_configuration", lambda db: {})
    main.app.dependency_overrides[main.get_db] = lambda: None
    main.app.dependency_overrides[main.optional_user] = lambda: SimpleNamespace(
        user=SimpleNamespace(id=17), has=lambda permission, service: False,
        accessible_service_ids=lambda permission: set())
    try:
        with TestClient(main.app) as browser:
            response = browser.post("/scan", data={
                "compose_files": "base.yaml\noverride.yaml",
                "compose_profiles": "data,local\nmetrics",
                "compose_preview": "true" if preview else "false",
                "ingest_service_id": "must-not-ingest-preview" if preview else "",
                "ingest_service_version": "v1" if preview else "",
            }, files=[
                ("compose_uploads", ("base.yaml", b"services: {}", "application/yaml")),
                ("compose_uploads", ("override.yaml", b"services: {}", "application/yaml")),
                ("compose_uploads", (".env", b"TAG=2", "text/plain")),
            ], follow_redirects=False)
        assert response.status_code == 303
        assert response.headers["location"] == "/scan?job_id=" + "c" * 32
        args, kwargs, uploads = captured[0]
        assert args[4] is None
        assert kwargs["owner_user_id"] == 17
        assert kwargs["compose_files"] == ["base.yaml", "override.yaml"]
        assert kwargs["compose_profiles"] == ["data", "local", "metrics"]
        assert kwargs["compose_preview"] is preview
        assert kwargs["ingest_service_version"] == ""
        assert uploads == [(b"services: {}", "base.yaml"), (b"services: {}", "override.yaml"), (b"TAG=2", ".env")]
    finally:
        main.app.dependency_overrides.clear()


@pytest.mark.parametrize("skipped", [False, True])
def test_missing_or_failed_image_cannot_count_as_compose_coverage(tmp_path, skipped):
    from app.compose_ingestion import discover
    from app.compose_scan import finalize_compose
    (tmp_path / "compose.yaml").write_text("""services:
  web: {image: 'app:2'}
  worker: {image: 'app:2'}
  db: {image: 'postgres:16'}
""", encoding="utf-8")
    model = discover(tmp_path, ["compose.yaml"], [])
    output = tmp_path / "output"
    output.mkdir()
    scanned = ["app:2", "postgres:16"] if skipped else ["app:2"]
    (output / "portal-result.json").write_text(json.dumps({"complete": True,
        "sbom_images": scanned, "skipped_images": ["postgres:16 :: pull failed"] if skipped else []}), encoding="utf-8")
    (output / "scan-summary.json").write_text(json.dumps({"status": "complete"}), encoding="utf-8")
    finalize_compose(output, model)
    result = json.loads((output / "portal-result.json").read_text())
    summary = json.loads((output / "scan-summary.json").read_text())
    model = result["service_overview"]["compose"]
    services = {item["name"]: item for item in model["services"]}
    assert services["web"]["image_scan"] == services["worker"]["image_scan"] == "complete"
    assert services["db"]["image_scan"] == "incomplete"
    assert model["coverage"]["scanned_images"] == 1
    assert any(issue["code"] == "image_scan_incomplete" and issue["service"] == "db" for issue in model["issues"])
    assert result["service_overview"]["missing_evidence"]
    assert result["complete"] is False
    assert summary["status"] == "incomplete"


@pytest.mark.parametrize("source", ["services: {}", "services:\n  app:\n    build: .\n"])
def test_empty_or_build_only_compose_remains_incomplete_without_scanner(tmp_path, source):
    from app.compose_ingestion import discover
    from app.compose_scan import finalize_compose
    (tmp_path / "compose.yaml").write_text(source, encoding="utf-8")
    model = discover(tmp_path, ["compose.yaml"], [])
    assert model["images"] == []
    finalize_compose(tmp_path, model)
    result = json.loads((tmp_path / "portal-result.json").read_text())
    summary = json.loads((tmp_path / "scan-summary.json").read_text())
    assert result["complete"] is False
    assert summary["status"] == "incomplete"
    assert summary["compose"]["coverage"]["scanned_images"] == 0
    assert result["service_overview"]["missing_evidence"]


def test_resolved_compose_preview_is_not_completed_image_evidence(tmp_path):
    from app.compose_ingestion import discover
    from app.compose_scan import finalize_compose
    (tmp_path / "compose.yaml").write_text("services:\n  app: {image: 'app:2'}\n", encoding="utf-8")
    model = discover(tmp_path, ["compose.yaml"], [])
    finalize_compose(tmp_path, model, preview=True)
    result = json.loads((tmp_path / "portal-result.json").read_text())
    summary = json.loads((tmp_path / "scan-summary.json").read_text())
    assert result["complete"] is False
    assert summary["status"] == "complete"
    assert summary["compose"]["coverage"] == {"services": 1, "unique_images": 1, "scanned_images": 0, "preview": True}
    assert result["service_overview"]["compose"]["services"][0]["image_scan"] == "not_scanned"


def test_service_workbook_preserves_compose_evidence_and_escapes_formulas(tmp_path):
    from openpyxl import load_workbook
    from app.service_export import build_service_workbook
    model = {
        "services": [{"name": "=web", "image": "+image", "status": "-blocked",
                      "image_scan": "incomplete", "source_files": ["@compose.yaml"],
                      "configuration": {"environment": {"TOKEN": "[REDACTED]"}}}],
        "coverage": {"services": 1, "unique_images": 1, "scanned_images": 0,
                     "preview": False, "=untrusted": "+value"},
        "issues": [{"code": "=issue", "message": "+message", "service": "-web",
                    "source": "@compose.yaml", "blocking": True}],
    }
    service = SimpleNamespace(id=1, name="Compose service", service_key="compose-service",
        owner="Test owner", created_at=datetime.now(timezone.utc), groups=[], findings=[],
        policy_findings=[], images=[], poam_entries=[], archive_events=[])
    execution = SimpleNamespace(raw_payload={"service_overview": {"compose": model}},
                                scanned_at=datetime.now(timezone.utc))
    workbook = build_service_workbook(service, {}, {}, datetime.now(timezone.utc), execution)
    target = tmp_path / "service.xlsx"
    workbook.save(target)
    workbook = load_workbook(target)
    services = workbook["Compose Services"]
    assert list(services.values)[1][:5] == ("'=web", "'+image", "'-blocked", "incomplete", "'@compose.yaml")
    assert json.loads(services["F2"].value) == model["services"][0]["configuration"]
    coverage = dict(workbook["Compose Coverage"].iter_rows(min_row=2, values_only=True))
    assert coverage == {"services": 1, "unique_images": 1, "scanned_images": 0,
                        "preview": False, "'=untrusted": "'+value"}
    assert list(workbook["Compose Input Issues"].values)[1] == (
        "'=issue", "'-web", "'@compose.yaml", "'+message", True)
    assert all(cell.data_type != "f" for name in ("Compose Services", "Compose Coverage", "Compose Input Issues")
               for row in workbook[name] for cell in row)
    workbook.close()


def test_compose_keeps_mixed_helm_evidence_applicable(tmp_path):
    from app.compose_scan import finalize_compose
    (tmp_path / "portal-result.json").write_text(json.dumps({
        "complete": True, "sbom_images": [], "service_overview": {"charts": [{"name": "existing-chart"}]}
    }), encoding="utf-8")
    (tmp_path / "scan-summary.json").write_text(json.dumps({"status": "complete"}), encoding="utf-8")
    finalize_compose(tmp_path, {"services": [], "images": [], "issues": [], "findings": []})
    result = json.loads((tmp_path / "portal-result.json").read_text(encoding="utf-8"))
    assert result["artifact_type"] == "helm"
    assert result["service_overview"]["charts"] == [{"name": "existing-chart"}]
    assert "compose" in result["service_overview"]
