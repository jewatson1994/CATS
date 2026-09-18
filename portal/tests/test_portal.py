from io import BytesIO
import os
import re
import json
import urllib.parse
import urllib.error
import ssl
import tarfile
from zipfile import ZipFile
from datetime import datetime, timedelta, timezone
from pathlib import Path

os.environ["DATABASE_URL"] = "sqlite://"
os.environ["PIPELINE_API_TOKEN"] = "test-token"
os.environ["CATS_BOOTSTRAP_USERNAME"] = "admin"
os.environ["CATS_BOOTSTRAP_PASSWORD"] = "test-password-long"
os.environ["SESSION_COOKIE_SECURE"] = "false"
os.environ["CATS_DEPLOYMENT_VALIDATION_ENABLED"] = "false"

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event, select

from app.auth import AuthContext, hash_password, seed_auth, token_hash
from app.database import Base, SessionLocal, engine
from app.main import app, configuration_for_service
from app.models import AuditEvent, DeploymentValidationRun, ExceptionRecord, Execution, Finding, Group, PoamEntry, PolicyExceptionRecord, PolicyFinding, PortalSetting, RemediationExecution, Role, Service, ServiceArchiveEvent, ServiceArtifact, ServiceArtifactRevision, ServiceImage, User, UserRoleAssignment, UserSession, WorkflowRequest

pipeline_headers = {"Authorization": "Bearer test-token"}


def setup_function():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    seed_auth()


def new_client(username="admin", password="test-password-long"):
    client = TestClient(app)
    response = client.post("/login", data={"username": username, "password": password}, follow_redirects=False)
    assert response.status_code == 303
    return client


def csrf(client):
    raw = client.cookies.get("cats_session")
    with SessionLocal() as db:
        return db.scalar(select(UserSession.csrf_token).where(UserSession.token_hash == token_hash(raw)))


def payload(execution_id, scanned_at, cves, service_id="payments-service", complete=True, skipped_images=None):
    return {
        "schema_version": "1.0", "execution_id": execution_id,
        "scanned_at": scanned_at.isoformat(), "complete": complete, "skipped_images": skipped_images or [], "fixable_only": True,
        "service": {"id": service_id, "name": service_id.replace("-", " ").title(), "version": "2.4.1", "poc": "cats-test-poc@example.invalid"},
        "findings": [{"cve": cve, "severity": "High", "image": f"registry/payments:{execution_id}",
                      "package": "openssl", "fixed_version": "9.9", "evidence": {"description": "Test vulnerability"}}
                     for cve in cves],
    }


def ingest(client, execution="run-1", cves=None, service_id="payments-service", complete=True, when=None, skipped_images=None):
    when = when or datetime.now(timezone.utc)
    return client.post("/api/v1/pipeline-results", json=payload(execution, when, cves or [], service_id, complete, skipped_images), headers=pipeline_headers)


def helm_payload(execution="helm-validation", service_id="payments-service"):
    body = payload(execution, datetime.now(timezone.utc), [], service_id)
    body.update({"artifact_type": "helm", "helm_source_files": {"Chart.yaml": "apiVersion: v2\nname: demo\nversion: 1.0.0\n"}})
    body["service_overview"] = {"rendered_resources": [{"apiVersion": "v1", "kind": "Service", "metadata": {"name": "demo"}}]}
    return body


def add_user(username, role_name, service_id=None):
    with SessionLocal() as db:
        role = db.scalar(select(Role).where(Role.name == role_name))
        user = User(username=username, display_name=username.title(), password_hash=hash_password("test-password-long"), must_change_password=False)
        db.add(user); db.flush()
        db.add(UserRoleAssignment(user_id=user.id, role_id=role.id, service_id=service_id)); db.commit()


def test_login_and_http_cookie_mode():
    client = new_client()
    assert client.get("/").status_code == 200
    assert client.cookies.get("cats_session")
    public = TestClient(app).get("/", follow_redirects=False)
    assert public.status_code == 200
    assert "Start Scan" in public.text and "Generate SBOM" in public.text and "Sign In" in public.text


def test_public_home_scan_and_patch_workspaces():
    client = TestClient(app)
    home = client.get("/")
    assert home.status_code == 200
    assert all(label in home.text for label in ("CATS", "Production Services", "Scan Images &amp; Charts", "Generate SBOM", "Patch Images", "Sign In"))
    assert "Understand what is deployed" not in home.text
    assert "SECURITY LIFECYCLE" not in home.text
    assert "Discover" not in home.text and "Monitor" not in home.text
    assert "Authenticated service operations" not in home.text
    assert client.get("/scan").status_code == 200
    patch = client.get("/patch")
    assert patch.status_code == 200
    assert "OCI Registry" in patch.text and "Upload Image" in patch.text
    invalid = client.post("/patch", data={"source_mode": "oci", "output_mode": "download"})
    assert invalid.status_code == 200
    assert "Image URI is required" in invalid.text


def test_scan_and_sbom_are_separate_workspaces_with_multi_format_generation(monkeypatch, tmp_path):
    from app import main as portal_main

    monkeypatch.setattr(portal_main, "PUBLIC_JOB_ROOT", tmp_path)
    submitted = []
    monkeypatch.setattr(portal_main.PUBLIC_WORKERS, "submit", lambda function, *args: submitted.append((function, args)))
    client = TestClient(app)

    scan = client.get("/scan")
    assert scan.status_code == 200
    assert "<h1>Scan</h1>" in scan.text
    assert 'href="/scan">Scan</a>' in scan.text
    assert 'href="/sbom">SBOM</a>' in scan.text

    sbom = client.get("/sbom")
    assert sbom.status_code == 200
    assert "<h1>SBOM</h1>" in sbom.text and "SBOM generator" in sbom.text
    assert sbom.text.index('href="/scan"') < sbom.text.index('href="/sbom"') < sbom.text.index('href="/patch"')
    for value in ("syft-json", "cyclonedx-json", "cyclonedx-xml", "spdx-json"):
        assert f'value="{value}"' in sbom.text
    assert 'name="sbom_formats"' in sbom.text
    assert 'name="cyclonedx_spec_version"' in sbom.text

    response = client.post("/sbom", data={
        "image_list": "docker.io/library/alpine:3.19",
        "sbom_formats": ["cyclonedx-json", "cyclonedx-xml", "spdx-json"],
        "cyclonedx_spec_version": "1.6",
    }, follow_redirects=False)
    assert response.status_code == 303
    job_id = response.headers["location"].split("job_id=", 1)[1]
    assert submitted and submitted[0][1][0] == job_id
    job = portal_main.PUBLIC_JOBS[job_id]
    assert job["job_kind"] == "sbom"
    assert job["sbom_formats"] == ["cyclonedx-json", "cyclonedx-xml", "spdx-json"]
    assert job["cyclonedx_spec_version"] == "1.6"

    missing_format = client.post("/sbom", data={
        "image_list": "docker.io/library/alpine:3.19",
        "cyclonedx_spec_version": "1.5",
    })
    assert missing_format.status_code == 200
    assert "Select at least one SBOM format" in missing_format.text


def test_sbom_job_passes_one_inventory_configuration_to_runner(monkeypatch, tmp_path):
    from app import main as portal_main

    monkeypatch.setattr(portal_main, "PUBLIC_JOB_ROOT", tmp_path)
    monkeypatch.setenv("CATS_SCANNER_RUNNER", "test-cats-scan")
    captured = {}

    class CompletedProcess:
        returncode = 0

        def __init__(self, arguments, **kwargs):
            captured["arguments"] = arguments
            captured["environment"] = kwargs["env"]
            output = Path(arguments[2])
            output.mkdir(parents=True, exist_ok=True)
            (output / "scan-summary.json").write_text(json.dumps({
                "status": "complete", "sboms": 1, "reports": 3,
                "formats": ["cyclonedx-json", "cyclonedx-xml", "spdx-json"],
            }), encoding="utf-8")

        def poll(self):
            return 0

    monkeypatch.setattr(portal_main.subprocess, "Popen", CompletedProcess)
    job_id = "sbom-runner-test"
    portal_main.PUBLIC_JOBS[job_id] = {
        "job_id": job_id, "job_kind": "sbom", "status": "queued", "phase": "queued",
        "sbom_formats": ["cyclonedx-json", "cyclonedx-xml", "spdx-json"],
        "cyclonedx_spec_version": "1.6",
    }
    portal_main._run_public_scan(job_id, "docker.io/library/alpine:3.19\n")

    assert captured["environment"]["CATS_JOB_MODE"] == "sbom"
    assert captured["environment"]["SBOM_FORMATS"] == "cyclonedx-json,cyclonedx-xml,spdx-json"
    assert captured["environment"]["SBOM_CYCLONEDX_SPEC_VERSION"] == "1.6"
    assert portal_main.PUBLIC_JOBS[job_id]["status"] == "complete"
    assert portal_main.PUBLIC_JOBS[job_id]["phase"] == "generate_sboms"


def test_sbom_download_contains_only_manifest_reports(monkeypatch, tmp_path):
    from app import main as portal_main

    job_id = "sbom-download-test"
    output = tmp_path / job_id / "output"
    formats_dir = output / "sboms" / "formats"
    formats_dir.mkdir(parents=True)
    raw = output / "sboms" / "alpine.json"
    cyclonedx = formats_dir / "alpine.cyclonedx.xml"
    raw.write_text('{"artifacts": []}', encoding="utf-8")
    cyclonedx.write_text("<bom/>", encoding="utf-8")
    (output / "worker.log").write_text("not part of the SBOM bundle", encoding="utf-8")
    (formats_dir / "manifest.json").write_text(json.dumps({"reports": [
        {"format": "syft-json", "path": "sboms/alpine.json"},
        {"format": "cyclonedx-xml", "path": "sboms/formats/alpine.cyclonedx.xml"},
    ]}), encoding="utf-8")
    monkeypatch.setattr(portal_main, "PUBLIC_JOB_ROOT", tmp_path)
    portal_main.PUBLIC_JOBS[job_id] = {"job_id": job_id, "job_kind": "sbom", "status": "complete"}

    response = TestClient(app).get(f"/api/public/jobs/{job_id}/sboms")
    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as bundle:
        assert set(bundle.namelist()) == {
            "manifest.json", "sboms/alpine.json", "sboms/formats/alpine.cyclonedx.xml",
        }


def test_public_results_view_handles_recursive_partial_result(monkeypatch, tmp_path):
    from app import main as portal_main

    job_id = "recursive-partial"
    output = tmp_path / job_id / "output"
    output.mkdir(parents=True)
    (output / "portal-result.json").write_text(json.dumps({
        "service": {"name": "torture-test"},
        "findings": [{"cve": "CVE-2026-0001", "image": "registry.example/api:1", "evidence": None}],
        "policy_findings": ["malformed-policy-row"],
        "skipped_charts": [{"item": "backend", "reason": "helm template failed at templates/deployment.yaml: rendered error"}],
    }), encoding="utf-8")
    (output / "results-export.tar.gz").write_bytes(b"test archive")
    trivy_result = output / "trivy-results" / "nginx.json"
    trivy_result.parent.mkdir()
    trivy_result.write_text('{"Results": []}', encoding="utf-8")
    monkeypatch.setattr(portal_main, "PUBLIC_JOB_ROOT", tmp_path)
    monkeypatch.setitem(portal_main.PUBLIC_JOBS, job_id, {"job_id": job_id, "status": "complete", "chart_names": []})
    response = TestClient(app).get(f"/api/public/jobs/{job_id}/results/view")
    assert response.status_code == 200
    assert "torture-test" in response.text
    assert "backend" in response.text
    html_overview = TestClient(app).get(f"/api/public/jobs/{job_id}/overview.html")
    assert html_overview.status_code == 200
    assert "INTERACTIVE SCAN OVERVIEW" in html_overview.text
    assert "data-type-filter=\"Vulnerability\"" in html_overview.text
    assert "CVE-2026-0001" in html_overview.text
    assert "https://nvd.nist.gov/vuln/detail/CVE-2026-0001" in html_overview.text
    assert 'id="search"' in html_overview.text and 'id="detail"' in html_overview.text
    assert "Open this file after extracting the ZIP" in html_overview.text
    assert 'href="portal-result.json"' in html_overview.text
    assert 'href="trivy-results/nginx.json"' in html_overview.text
    assert 'href="#artifacts"' in html_overview.text
    artifacts = TestClient(app).get(f"/api/public/jobs/{job_id}/artifacts")
    assert artifacts.status_code == 200
    with ZipFile(BytesIO(artifacts.content)) as bundle:
        assert "scan-overview.html" in bundle.namelist()
        exported_html = bundle.read("scan-overview.html").decode()
        assert "CVE-2026-0001" in exported_html and "const findings=" in exported_html
        assert 'href="scan-results.xlsx"' in exported_html
        assert 'href="results-export.tar.gz"' in exported_html
    archive = TestClient(app).get(f"/api/public/jobs/{job_id}/results-export")
    assert archive.status_code == 200
    assert archive.headers["content-type"] == "application/gzip"


def test_public_ingest_retains_large_helm_source_set_and_is_idempotent(monkeypatch, tmp_path):
    from app import main as portal_main

    client = new_client()
    monkeypatch.setattr(AuthContext, "accessible_service_ids", lambda self, permission: {1})
    monkeypatch.setattr(AuthContext, "has", lambda self, permission, service_id=None: True)
    assert ingest(client, execution="large-source-baseline").status_code == 201
    job_id = "large-source-ingest"
    output = tmp_path / job_id / "output"
    charts = tmp_path / job_id / "input" / "charts"
    output.mkdir(parents=True); charts.mkdir(parents=True)
    body = payload("replaced-by-public-route", datetime.now(timezone.utc), ["CVE-2026-1111"] * 20)
    body["policy_findings"] = [{"finding": f"KSV-{index}", "target": f"Deployment/item-{index}"} for index in range(25)]
    body["service_overview"] = {"rendered_resources": [{"apiVersion": "v1", "kind": "ConfigMap", "metadata": {"name": "large"}}]}
    (output / "portal-result.json").write_text(json.dumps(body), encoding="utf-8")
    for index in range(734):
        directory = charts / f"chart-{index:04d}"
        directory.mkdir()
        (directory / "values.yaml").write_text(f"index: {index}\n", encoding="utf-8")
    monkeypatch.setattr(portal_main, "PUBLIC_JOB_ROOT", tmp_path)
    monkeypatch.setitem(portal_main.PUBLIC_JOBS, job_id, {"job_id": job_id, "status": "complete", "image_list": ""})

    first = client.post(f"/api/public/jobs/{job_id}/ingest?service_id=payments-service")
    assert first.status_code == 200
    assert first.json()["accepted"] is True and first.json()["duplicate"] is False
    assert set(first.json()["ingest_timings_ms"]) >= {"read_result", "validate_payload", "database_ingest", "total"}
    with SessionLocal() as db:
        execution = db.scalar(select(Execution).where(Execution.execution_key == f"public:{job_id}"))
        assert len(execution.raw_payload["helm_source_files"]) == 734
        assert len(execution.raw_payload["findings"]) == 20
        assert len(execution.raw_payload["policy_findings"]) == 25

    second = client.post(f"/api/public/jobs/{job_id}/ingest?service_id=payments-service")
    assert second.status_code == 200 and second.json()["duplicate"] is True
    with SessionLocal() as db:
        assert len(db.scalars(select(Execution).where(Execution.execution_key == f"public:{job_id}")).all()) == 1


def test_duplicate_execution_ignores_mutable_service_presentation_but_not_evidence():
    client = new_client()
    body = payload("stable-evidence-id", datetime.now(timezone.utc), ["CVE-2026-1010"])
    assert client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers).status_code == 201
    presentation_change = json.loads(json.dumps(body))
    presentation_change["service"].update({"name": "Renamed Service", "owner": "New Owner", "groups": ["Updated"]})
    duplicate = client.post("/api/v1/pipeline-results", json=presentation_change, headers=pipeline_headers)
    assert duplicate.status_code == 201 and duplicate.json()["duplicate"] is True
    changed_evidence = json.loads(json.dumps(presentation_change))
    changed_evidence["findings"][0]["fixed_version"] = "10.0"
    conflict = client.post("/api/v1/pipeline-results", json=changed_evidence, headers=pipeline_headers)
    assert conflict.status_code == 409


def test_public_ingest_validation_error_is_safe_structured_and_atomic(monkeypatch, tmp_path):
    from app import main as portal_main

    client = new_client()
    monkeypatch.setattr(AuthContext, "accessible_service_ids", lambda self, permission: {1})
    monkeypatch.setattr(AuthContext, "has", lambda self, permission, service_id=None: True)
    assert ingest(client, execution="atomic-baseline").status_code == 201
    job_id = "invalid-public-ingest"
    output = tmp_path / job_id / "output"; output.mkdir(parents=True)
    body = payload("ignored", datetime.now(timezone.utc), [])
    body["policy_findings"] = [{"finding": ""}]
    (output / "portal-result.json").write_text(json.dumps(body), encoding="utf-8")
    monkeypatch.setattr(portal_main, "PUBLIC_JOB_ROOT", tmp_path)
    monkeypatch.setitem(portal_main.PUBLIC_JOBS, job_id, {"job_id": job_id, "status": "complete", "image_list": ""})
    before_name = "Payments Service"
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        service.name = before_name; db.commit()

    response = client.post(f"/api/public/jobs/{job_id}/ingest?service_id=payments-service")
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert detail["code"] == "PUBLIC_INGEST_VALIDATION_ERROR"
    assert detail["stage"] == "validate_payload" and detail["record"].startswith("policy_findings.0.finding")
    assert "input" not in json.dumps(detail).lower()
    with SessionLocal() as db:
        assert db.scalar(select(Execution).where(Execution.execution_key == f"public:{job_id}")) is None
        assert db.scalar(select(Service).where(Service.service_key == "payments-service")).name == before_name


def test_public_ingest_database_failure_rolls_back_all_partial_state(monkeypatch, tmp_path):
    from app import main as portal_main

    client = new_client()
    monkeypatch.setattr(AuthContext, "accessible_service_ids", lambda self, permission: {1})
    monkeypatch.setattr(AuthContext, "has", lambda self, permission, service_id=None: True)
    assert ingest(client, execution="rollback-baseline").status_code == 201
    job_id = "database-failure-ingest"
    output = tmp_path / job_id / "output"; output.mkdir(parents=True)
    body = payload("ignored", datetime.now(timezone.utc), ["CVE-2026-9898"])
    body["service"]["name"] = "Should Roll Back"
    body["policy_findings"] = [{"finding": "KSV-rollback"}]
    (output / "portal-result.json").write_text(json.dumps(body), encoding="utf-8")
    monkeypatch.setattr(portal_main, "PUBLIC_JOB_ROOT", tmp_path)
    monkeypatch.setitem(portal_main.PUBLIC_JOBS, job_id, {"job_id": job_id, "status": "complete", "image_list": ""})
    monkeypatch.setattr(portal_main, "sync_policy_findings", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("forced transaction failure")))

    failing_client = TestClient(app, raise_server_exceptions=False)
    failing_client.cookies.update(client.cookies)
    response = failing_client.post(f"/api/public/jobs/{job_id}/ingest?service_id=payments-service")
    assert response.status_code == 500
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        assert service.name == "Payments Service"
        assert db.scalar(select(Execution).where(Execution.execution_key == f"public:{job_id}")) is None
        assert db.scalar(select(Finding).where(Finding.service_id == service.id, Finding.cve == "CVE-2026-9898")) is None


def test_service_remediation_creates_auditable_review_candidate(monkeypatch):
    from app import main as portal_main

    client = new_client()
    data = payload("remediation-source", datetime.now(timezone.utc), [])
    data["policy_findings"] = [{
        "finding": "KSV-allowPrivilegeEscalation", "severity": "High", "target": "Deployment/payments",
        "title": "Container allows privilege escalation", "description": "allowPrivilegeEscalation should be false",
    }]
    data["service_overview"] = {"rendered_resources": [{
        "apiVersion": "apps/v1", "kind": "Deployment", "metadata": {"name": "payments"},
        "_cats_source_file": "templates/deployment.yaml",
        "spec": {"template": {"spec": {"containers": [{"name": "payments", "image": "registry/payments:1"}]}}},
    }]}
    assert client.post("/api/v1/pipeline-results", json=data, headers=pipeline_headers).status_code == 201
    monkeypatch.setattr(portal_main.REMEDIATION_WORKERS, "submit", lambda function, *args: function(*args))
    response = client.post("/services/payments-service/remediate", data={"csrf_token": csrf(client)}, follow_redirects=False)
    assert response.status_code == 303
    with SessionLocal() as db:
        job = db.scalar(select(RemediationExecution))
        assert job.status == "review_required"
        assert job.rollback_reference == "remediation-source"
        assert job.configuration_changes[0]["classification"] == "REVIEW REQUIRED"
        assert Path(job.artifact_path).is_file()
        location = f"/services/payments-service/remediations/{job.job_key}"
    report = client.get(location)
    assert report.status_code == 200
    assert "Before / After" in report.text and "Deployment validation" in report.text and "Review Required" in report.text


def test_administrator_can_save_oci_registry():
    client = new_client()
    response = client.post("/admin/configuration/registries", data={
        "csrf_token": csrf(client), "display_name": "Production Harbor",
        "endpoint": "https://harbor.example.invalid/", "namespace": "/cyber-approved/",
        "auth_mode": "none", "username": "", "password": "", "action": "save",
    }, follow_redirects=False)
    assert response.status_code == 303
    assert response.headers["location"] == "/admin/configuration?saved=1"


def test_internal_login_navigation_stays_in_current_tab():
    """Internal authentication links must not open a second browser tab."""
    template_root = Path(__file__).parents[1] / "app" / "templates"
    static_root = Path(__file__).parents[1] / "app" / "static"
    template_paths = sorted(template_root.glob("*.html"))
    assert template_paths

    for path in template_paths:
        text = path.read_text(encoding="utf-8")
        for tag in re.findall(r"<a\b[^>]*>", text, flags=re.IGNORECASE | re.DOTALL):
            href_match = re.search(r"\bhref\s*=\s*['\"]([^'\"]+)['\"]", tag, flags=re.IGNORECASE)
            if not href_match:
                continue
            href = href_match.group(1)
            if href.startswith("/login") or href.startswith("/auth/oidc/login"):
                assert not re.search(r"\btarget\s*=\s*['\"]_blank['\"]", tag, flags=re.IGNORECASE), (
                    f"internal login link opens a new tab in {path.name}: {tag}"
                )

    for path in static_root.rglob("*.js"):
        assert "window.open(" not in path.read_text(encoding="utf-8"), f"new-window behavior found in {path.name}"


def test_authenticated_primary_navigation_and_gear_cleanup():
    client = new_client()
    page = client.get("/")
    assert page.status_code == 200
    assert page.text.index('href="/patch"') < page.text.index('href="/remediations"')
    assert "Remediations workspace" not in page.text
    assert "Service staging" not in page.text.split("admin-menu", 1)[-1].split("user-menu", 1)[0]
    assert "Audit Policy" not in page.text.split("admin-menu", 1)[-1].split("user-menu", 1)[0]
    remediations = client.get("/remediations")
    assert remediations.status_code == 200 and 'href="/remediations"' in remediations.text
    assert "class=\"active\"" in remediations.text


def test_active_services_are_alphabetical_without_compliance_explainer():
    client = new_client()
    ingest(client, execution="zeta", service_id="zeta-service")
    ingest(client, execution="alpha", service_id="alpha-service")
    page = client.get("/")
    assert page.status_code == 200
    assert page.text.index("Alpha Service") < page.text.index("Zeta Service")
    assert "Non-compliant means at least one unexcepted fixable CVE is older than 90 days." not in page.text


def test_services_overview_query_count_is_constant_as_services_grow():
    client = new_client()
    for index in range(20):
        assert ingest(client, execution=f"overview-scale-{index}", service_id=f"overview-service-{index}").status_code == 201
    calls = []
    listener = lambda *args: calls.append(args[2])
    event.listen(engine, "before_cursor_execute", listener)
    try:
        response = client.get("/")
    finally:
        event.remove(engine, "before_cursor_execute", listener)
    assert response.status_code == 200
    assert len(calls) <= 20


def test_configuration_policy_findings_render_as_generic_active_findings():
    client = new_client()
    policy_payload = payload("policy-run", datetime.now(timezone.utc), [], service_id="payments-service")
    policy_payload["policy_findings"] = [{
        "type": "Configuration", "finding": "KSV014", "severity": "High",
        "scanner": "Trivy", "framework": "CIS Kubernetes",
        "target": "Deployment/payments-api", "title": "Privilege escalation enabled",
        "remediation": "Set allowPrivilegeEscalation to false",
    }]
    response = client.post("/api/v1/pipeline-results", json=policy_payload, headers=pipeline_headers)
    assert response.status_code == 201
    page = client.get("/services/payments-service?finding_state=active")
    assert page.status_code == 200
    assert "Active Fixable Findings" in page.text
    assert "KSV014" in page.text and "Configuration" in page.text
    assert "CIS Kubernetes" in page.text and "Deployment/payments-api" in page.text
    overview = client.get("/services/payments-service?overview=true")
    assert "Active Fixable</span><strong>1</strong>" in overview.text
    dashboard = client.get("/")
    assert "Active Findings" in dashboard.text and "Overdue Findings" in dashboard.text
    assert "Active CVEs" not in dashboard.text and "Over 90d" not in dashboard.text


def test_overview_navigation_never_runs_digest_resolution(monkeypatch):
    """Overview rendering must use persisted scan data, not Docker/registry I/O."""
    client = new_client()
    response = ingest(client, execution="overview-fast", cves=["CVE-2025-0001"])
    assert response.status_code == 201

    def fail_if_called(*args, **kwargs):
        raise AssertionError("digest resolution must not run while rendering an overview")

    monkeypatch.setattr("app.main._resolve_manifest_digest", fail_if_called)
    overview = client.get("/services/payments-service?overview=true")
    assert overview.status_code == 200
    assert "Overview" in overview.text


def test_service_findings_can_be_filtered_by_type():
    client = new_client()
    policy_payload = payload("filter-run", datetime.now(timezone.utc), ["CVE-2026-0001"])
    policy_payload["policy_findings"] = [{
        "type": "Configuration", "finding": "KSV014", "severity": "High",
        "framework": "CIS Kubernetes", "target": "Deployment/payments-api",
    }]
    response = client.post("/api/v1/pipeline-results", json=policy_payload, headers=pipeline_headers)
    assert response.status_code == 201
    configuration_page = client.get("/services/payments-service?finding_type=configuration")
    assert configuration_page.status_code == 200
    assert "KSV014" in configuration_page.text
    assert "CVE-2026-0001" not in configuration_page.text
    assert "&ndash;" in configuration_page.text
    vulnerability_page = client.get("/services/payments-service?finding_type=vulnerability")
    assert vulnerability_page.status_code == 200
    assert "CVE-2026-0001" in vulnerability_page.text
    assert "KSV014" not in vulnerability_page.text
    assert ">Configuration</option>" in vulnerability_page.text


def test_warning_filter_renders_warning_records_instead_of_active_findings():
    client = new_client()
    ingest(client, execution="warning-seed")
    when = datetime.now(timezone.utc)
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        finding = Finding(service_id=service.id, cve="CVE-2026-WARN", severity="High", first_seen=when - timedelta(days=80),
                          episode_started=when - timedelta(days=80), last_seen=when, active=True)
        db.add(finding)
        db.commit()
    with SessionLocal() as db:
        setting = PortalSetting(key="raw_due_rules", value='[{"severity":"High","days":90}]')
        db.add(setting); db.commit()
    response = client.get("/services/payments-service?finding_state=warnings")
    assert response.status_code == 200
    assert "Warnings" in response.text and "CVE-2026-WARN" in response.text
    assert "No warnings for this service" not in response.text


def test_remediations_workspace_uses_existing_records_and_tabs():
    client = new_client()
    ingest(client, cves=["CVE-2026-REMEDIATE"])
    page = client.get("/remediations?tab=poams")
    assert page.status_code == 200
    assert all(label in page.text for label in ("POA&amp;Ms", "Exceptions", "Mitigations"))
    assert client.get("/remediations?tab=exceptions").status_code == 200
    assert client.get("/remediations?tab=mitigations").status_code == 200


def test_service_remediations_replaces_service_poam_tab_without_breaking_legacy_url():
    client = new_client()
    ingest(client, cves=["CVE-2026-SERVICE-REMEDIATION"])
    page = client.get("/services/payments-service?remediations=true&tab=poams")
    assert page.status_code == 200
    assert "Payments Service Remediations" in page.text
    assert all(label in page.text for label in ("POA&amp;Ms", "Exceptions", "Mitigations"))
    assert 'href="/services/payments-service?remediations=true&amp;tab=poams"' in page.text
    assert client.get("/services/payments-service?remediations=true&tab=exceptions").status_code == 200
    assert client.get("/services/payments-service?remediations=true&tab=mitigations").status_code == 200
    legacy = client.get("/services/payments-service?poam=true")
    assert legacy.status_code == 200 and "Payments Service POA&amp;M" in legacy.text


def test_account_last_login_delete_protection_and_audit_snapshot():
    client = new_client()
    assert client.post("/admin/users", data={"csrf_token": csrf(client), "username": "remove-me", "display_name": "Remove Me", "password": "temporary-password"}, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == "remove-me"))
        user_id = user.id
        assert user.last_login_at is None
        service = Service(service_key="deleted-requester-service", name="Deleted requester service")
        db.add(service); db.flush()
        execution = Execution(execution_key="deleted-requester-execution", service_id=service.id, scanned_at=datetime.now(timezone.utc), complete=True, raw_payload={})
        db.add(execution); db.flush()
        validation_run = DeploymentValidationRun(run_key="DV-DELETED-REQUESTER", service_id=service.id, execution_id=execution.id, requested_by_id=user.id)
        db.add(validation_run); db.commit(); validation_run_id = validation_run.id
    failed_login = TestClient(app).post("/login", data={"username": "remove-me", "password": "wrong-password"})
    assert failed_login.status_code == 401
    with SessionLocal() as db:
        assert db.get(User, user_id).last_login_at is None
    new_client("remove-me", "temporary-password")
    with SessionLocal() as db:
        assert db.get(User, user_id).last_login_at is not None
    assert client.post("/admin/users/{}/delete".format(user_id), data={"csrf_token": csrf(client)}, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        assert db.get(User, user_id) is None
        assert db.get(DeploymentValidationRun, validation_run_id).requested_by_id is None
        deleted = db.scalar(select(AuditEvent).where(AuditEvent.action == "user.deleted", AuditEvent.detail["username"].as_string() == "remove-me"))
        assert deleted is not None
    with SessionLocal() as db:
        admin = db.scalar(select(User).where(User.username == "admin"))
    assert client.post(f"/admin/users/{admin.id}/delete", data={"csrf_token": csrf(client)}).status_code == 409


def test_configuration_finding_exception_and_poam_use_first_class_workflows():
    admin = new_client()
    policy_payload = payload("policy-workflow", datetime.now(timezone.utc), [])
    policy_payload["policy_findings"] = [{
        "type": "Hardening", "finding": "KSV014", "severity": "High",
        "scanner": "Trivy", "framework": "CIS Kubernetes",
        "target": "Deployment/payments-api", "title": "Privilege escalation enabled",
        "description": "The workload permits privilege escalation.",
        "remediation": "Set allowPrivilegeEscalation to false", "fingerprint": "ksv014-payments",
    }]
    assert admin.post("/api/v1/pipeline-results", json=policy_payload, headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        finding = db.scalar(select(PolicyFinding).where(PolicyFinding.service_id == service.id))
        service_id, policy_finding_id = service.id, finding.id
        assert finding.finding == "KSV014" and finding.type == "Configuration"
    other_payload = payload("other-policy-workflow", datetime.now(timezone.utc), [], service_id="other-service")
    other_payload["policy_findings"] = [{"finding": "KSV999", "fingerprint": "other-config"}]
    assert admin.post("/api/v1/pipeline-results", json=other_payload, headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        other_policy_finding_id = db.scalar(select(PolicyFinding.id).join(Service).where(Service.service_key == "other-service"))
    add_user("policy-manager", "Service Manager", service_id)
    manager = new_client("policy-manager")
    page = manager.get("/services/payments-service?finding_type=configuration")
    assert page.status_code == 200
    assert f'data-exception-url="/policy-findings/{policy_finding_id}/exceptions"' in page.text
    assert f'data-poam-url="/policy-findings/{policy_finding_id}/poams"' in page.text
    assert 'data-label="KSV014"' in page.text
    expiry = datetime.now(timezone.utc) + timedelta(days=30)
    assert manager.post(f"/policy-findings/{other_policy_finding_id}/exceptions", data={
        "csrf_token": csrf(manager), "justification": "Out of scope request",
        "expires_at": expiry.isoformat(),
    }).status_code == 403
    assert manager.post(f"/policy-findings/{other_policy_finding_id}/poams", data={
        "csrf_token": csrf(manager), "title": "Out of scope configuration",
        "remediation": "Must not be accepted",
    }).status_code == 403
    exception_response = manager.post(f"/policy-findings/{policy_finding_id}/exceptions", data={
        "csrf_token": csrf(manager), "justification": "Compensating admission policy is active",
        "expires_at": expiry.isoformat(), "ticket": "RISK-CONFIG-14",
    }, follow_redirects=False)
    assert exception_response.status_code == 303
    poam_response = manager.post(f"/policy-findings/{policy_finding_id}/poams", data={
        "csrf_token": csrf(manager), "title": "Harden payments deployment",
        "remediation": "Disable privilege escalation and redeploy", "ticket": "POAM-CONFIG-14",
    }, follow_redirects=False)
    assert poam_response.status_code == 303
    with SessionLocal() as db:
        exception_workflow = db.scalar(select(WorkflowRequest).where(
            WorkflowRequest.request_type == "exception", WorkflowRequest.policy_finding_id == policy_finding_id
        ))
        poam_workflow = db.scalar(select(WorkflowRequest).where(
            WorkflowRequest.request_type == "poam", WorkflowRequest.policy_finding_id == policy_finding_id
        ))
        poam_id = poam_workflow.poam_id
    add_user("policy-cyber", "Cybersecurity")
    cyber = new_client("policy-cyber")
    assert "KSV014" in cyber.get("/requests").text
    assert cyber.post(f"/requests/{exception_workflow.id}/review", data={
        "csrf_token": csrf(cyber), "decision": "approved",
    }, follow_redirects=False).status_code == 303
    assert cyber.post(f"/requests/{poam_workflow.id}/review", data={
        "csrf_token": csrf(cyber), "decision": "approved",
    }, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        policy_exception = db.scalar(select(PolicyExceptionRecord).where(
            PolicyExceptionRecord.policy_finding_id == policy_finding_id
        ))
        entry = db.get(PoamEntry, poam_id)
        assert policy_exception is not None
        assert entry.policy_finding_id == policy_finding_id and entry.status == "active"
        policy_exception_id = policy_exception.id
    finding_export = load_workbook(BytesIO(cyber.get("/services/payments-service/export.xlsx").content))
    assert "Configuration Findings" in finding_export.sheetnames
    assert "KSV014" in [cell.value for cell in finding_export["Configuration Findings"]["A"]]
    poam_export = load_workbook(BytesIO(cyber.get("/poam/services/payments-service/export.xlsx").content))["POA&M"]
    assert "KSV014" in [cell.value for cell in poam_export["F"]]
    poam_service_page = cyber.get("/poam/services/payments-service")
    assert "<code>KSV014</code>" in poam_service_page.text
    poam_detail_page = cyber.get(f"/poam/entries/{poam_id}")
    assert "<dt>Finding</dt><dd>KSV014</dd>" in poam_detail_page.text
    exceptions_page = cyber.get("/services/payments-service?finding_state=exceptions&finding_type=configuration")
    assert "KSV014" in exceptions_page.text
    assert f'action="/policy-exceptions/{policy_exception_id}/revoke"' in exceptions_page.text
    assert cyber.post(f"/policy-exceptions/{policy_exception_id}/revoke", data={
        "csrf_token": csrf(cyber),
    }, follow_redirects=False).status_code == 303


def test_configuration_findings_age_resolve_and_recur_by_stable_identity():
    client = new_client()
    old = datetime.now(timezone.utc) - timedelta(days=100)
    first = payload("policy-old", old, [])
    first["policy_findings"] = [{
        "type": "Compliance", "finding": "AVD-KSV-0014", "severity": "High",
        "scanner": "Trivy", "target": "Deployment/api", "fingerprint": "stable-config-14",
    }]
    assert client.post("/api/v1/pipeline-results", json=first, headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        overdue_policy_finding_id = db.scalar(select(PolicyFinding.id))
    overdue = client.get("/services/payments-service?finding_state=noncompliant&finding_type=configuration")
    assert "AVD-KSV-0014" in overdue.text and "Configuration" in overdue.text
    overview = client.get("/services/payments-service?overview=true")
    assert "Active Non-Compliance</span><strong>1</strong>" in overview.text
    assert f'data-exception-url="/policy-findings/{overdue_policy_finding_id}/exceptions"' in overdue.text
    assert f'data-poam-url="/policy-findings/{overdue_policy_finding_id}/poams"' in overdue.text
    second = payload("policy-clear", datetime.now(timezone.utc) - timedelta(days=1), [])
    second["policy_findings"] = []
    assert client.post("/api/v1/pipeline-results", json=second, headers=pipeline_headers).status_code == 201
    assert "AVD-KSV-0014" in client.get("/services/payments-service?finding_state=resolved&finding_type=configuration").text
    third = payload("policy-return", datetime.now(timezone.utc), [])
    third["policy_findings"] = first["policy_findings"]
    assert client.post("/api/v1/pipeline-results", json=third, headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        finding = db.scalar(select(PolicyFinding))
        assert finding.active and finding.recurrence_count == 1


def test_group_scoped_manager_can_act_when_service_has_multiple_groups():
    admin = new_client()
    policy_payload = payload("group-policy", datetime.now(timezone.utc), [])
    policy_payload["policy_findings"] = [{"finding": "KSV-GROUP", "fingerprint": "group-config"}]
    assert admin.post("/api/v1/pipeline-results", json=policy_payload, headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        first_group = Group(name="First Group")
        assigned_group = Group(name="Assigned Group")
        service.groups.extend([first_group, assigned_group])
        db.flush()
        role = db.scalar(select(Role).where(Role.name == "Service Manager"))
        user = User(username="group-manager", display_name="Group Manager",
                    password_hash=hash_password("test-password-long"), must_change_password=False)
        db.add(user); db.flush()
        db.add(UserRoleAssignment(user_id=user.id, role_id=role.id, group_id=assigned_group.id))
        policy_finding_id = db.scalar(select(PolicyFinding.id).where(PolicyFinding.service_id == service.id))
        db.commit()
    manager = new_client("group-manager")
    assert "Payments Service" in manager.get("/").text
    response = manager.post(f"/policy-findings/{policy_finding_id}/poams", data={
        "csrf_token": csrf(manager), "title": "Group-scoped remediation",
        "remediation": "Apply hardened configuration",
    }, follow_redirects=False)
    assert response.status_code == 303


def test_complete_scan_resolves_finding():
    client = new_client()
    now = datetime.now(timezone.utc)
    assert ingest(client, cves=["CVE-2026-0001"], when=now).status_code == 201
    ingest(client, "run-2", [], complete=True, when=now + timedelta(days=1))
    assert "CVE-2026-0001" in client.get("/services/payments-service?finding_state=resolved").text


def test_group_scoped_access_is_not_global():
    client = new_client()
    assert ingest(client, execution="group-seed", cves=[]).status_code == 201
    with SessionLocal() as db:
        group = Group(name="Payments Group")
        db.add(group)
        db.flush()
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        if not service:
            service = Service(service_key="payments-service", name="Payments", owner="Platform")
            db.add(service)
            db.flush()
        service.groups.append(group)
        role = db.scalar(select(Role).where(Role.name == "Service Manager"))
        user = User(username="scoped", display_name="Scoped", password_hash=hash_password("test-password-long"), must_change_password=False)
        db.add(user)
        db.flush()
        db.add(UserRoleAssignment(user_id=user.id, role_id=role.id, group_id=group.id))
        db.commit()
    scoped = new_client("scoped")
    assert scoped.get("/services/payments-service").status_code == 200
    assert scoped.get("/services/unknown-service").status_code == 403


def test_skipped_images_are_visible_on_service():
    client = new_client()
    ingest(client, skipped_images=["registry/legacy-centos:1.0", "registry/vendor:2.0"], complete=False)
    page = client.get("/services/payments-service").text
    assert "Skipped images (2)" in page
    assert "registry/legacy-centos:1.0" in page
    assert "Incomplete · 2 skipped" in client.get("/").text


def test_incomplete_execution_without_skipped_images_is_noncompliant():
    client = new_client()
    ingest(client, complete=False, skipped_images=[])
    page = client.get("/services/payments-service?finding_state=noncompliant")
    assert page.status_code == 200
    assert "Latest assessment reported incomplete evidence" in page.text
    assert "No image details reported" not in page.text
    assert "<th>Details</th>" in page.text
    overview = client.get("/services/payments-service?overview=true")
    assert "Active Non-Compliance</span><strong>1</strong>" in overview.text
    assert "overdue-card" in overview.text


def test_service_manager_can_add_incomplete_evidence_to_poam():
    admin = new_client()
    ingest(admin, skipped_images=["registry/unavailable:demo"], complete=False)
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        service_id = service.id
    add_user("evidence-manager", "Service Manager", service_id)
    manager = new_client("evidence-manager")
    page = manager.get("/services/payments-service?finding_state=noncompliant")
    assert page.status_code == 200
    assert "Add missing evidence to POA&amp;M" in page.text
    assert "registry/unavailable:demo" in page.text
    assert 'class="secondary-button evidence-poam-action"' in page.text
    assert 'data-image="registry/unavailable:demo"' in page.text
    created = manager.post("/poam", data={
        "csrf_token": csrf(manager), "service_id": service_id, "item_type": "missing_evidence",
        "finding_id": "", "title": "Missing evidence for registry/unavailable:demo",
        "description": "Unavailable for assessment: registry/unavailable:demo",
        "remediation": "Restore scanner access and submit a complete report",
    }, follow_redirects=False)
    assert created.status_code == 303
    with SessionLocal() as db:
        entry = db.scalar(select(PoamEntry))
        assert entry.item_type == "missing_evidence" and entry.status == "pending_approval"


def test_service_findings_are_paginated_and_gzip_compressed():
    client = new_client()
    ingest(client)
    now = datetime.now(timezone.utc)
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        for number in range(60):
            db.add(Finding(
                service_id=service.id, cve=f"CVE-PAGE-{number:04d}", severity="High",
                first_seen=now, episode_started=now, last_seen=now, active=True,
            ))
        db.commit()
    first = client.get("/services/payments-service?page_size=50")
    assert first.status_code == 200
    assert first.headers.get("content-encoding") == "gzip"
    assert first.text.count("<strong>CVE-PAGE-") == 50
    assert "Page 1 of 2" in first.text and "Showing 1" in first.text and "of 60" in first.text
    second = client.get("/services/payments-service?page=2&page_size=50")
    assert second.text.count("<strong>CVE-PAGE-") == 10
    assert "Page 2 of 2" in second.text


def test_service_finding_filters_combine_and_reset_pagination():
    client = new_client()
    ingest(client)
    now = datetime.now(timezone.utc)
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        db.add_all([
            Finding(service_id=service.id, cve="CVE-FILTER-CRITICAL", severity="Critical", first_seen=now, episode_started=now, last_seen=now, active=True),
            Finding(service_id=service.id, cve="CVE-FILTER-HIGH", severity="High", first_seen=now, episode_started=now, last_seen=now, active=True),
            Finding(service_id=service.id, cve="CVE-FILTER-LOW", severity="Low", first_seen=now, episode_started=now, last_seen=now, active=True),
            *[Finding(service_id=service.id, cve=f"CVE-BULK-CRIT-{number:03d}", severity="Critical", first_seen=now, episode_started=now, last_seen=now, active=True) for number in range(55)],
        ])
        db.commit()
    response = client.get("/services/payments-service?page=4&page_size=50&severity=Critical&severity=High&q=FILTER")
    assert response.status_code == 200
    assert "CVE-FILTER-CRITICAL" in response.text and "CVE-FILTER-HIGH" in response.text
    assert "CVE-FILTER-LOW" not in response.text
    assert "Showing 1&ndash;2 of 2" in response.text
    assert "Filters active" in response.text
    assert response.text.count('aria-label="Finding pages"') == 0  # one page keeps the compact layout
    response = client.get("/services/payments-service?page_size=50&severity=Critical")
    assert response.status_code == 200
    assert response.text.count('aria-label="Finding pages"') == 2
    assert 'class="pagination pagination-top"' in response.text and 'class="pagination pagination-bottom"' in response.text
    assert 'href="/services/payments-service?overview=false&amp;finding_state=active&amp;finding_type=all&amp;page_size=50&amp;severity=Critical&amp;page=2"' in response.text


def test_risk_overlay_filters_active_findings_before_age_drives_noncompliance():
    client = new_client()
    ingest(client, cves=["CVE-KEV-MATCH", "CVE-EPSS-MATCH", "CVE-NO-MATCH"])
    now = datetime.now(timezone.utc)
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        group = Group(name="Risk Overlay Group")
        service.groups.append(group)
        db.add(group)
        db.flush()
        findings = {finding.cve: finding for finding in service.findings}
        kev_observation = findings["CVE-KEV-MATCH"].observations[-1]
        epss_observation = findings["CVE-EPSS-MATCH"].observations[-1]
        kev_observation.evidence = {**kev_observation.evidence, "kev": True}
        epss_observation.evidence = {**epss_observation.evidence, "epss": 0.95}
        for key, value in {
            "compliance_mode": "risk_based",
            "overdue_days": "90",
            "minimum_severity": "None",
            "kev_enabled": "true",
            "kev_noncompliant": "true",
            "epss_enabled": "true",
            "epss_rules": '[{"severity":"Any","threshold":0.9,"noncompliant":true}]',
        }.items():
            db.add(PortalSetting(key=f"group:{group.id}:{key}", group_id=group.id, value=value))
        db.commit()

    active = client.get("/services/payments-service?finding_state=active")
    assert "CVE-KEV-MATCH" in active.text
    assert "CVE-EPSS-MATCH" in active.text
    assert "CVE-NO-MATCH" not in active.text
    overview = client.get("/services/payments-service?overview=true")
    assert "Active Fixable</span><strong>2</strong>" in overview.text
    assert "Active Non-Compliance</span><strong>0</strong>" in overview.text

    with SessionLocal() as db:
        for finding in db.scalars(select(Finding)).all():
            finding.episode_started = now - timedelta(days=100)
        db.commit()
    noncompliant = client.get("/services/payments-service?finding_state=noncompliant")
    assert "CVE-KEV-MATCH" in noncompliant.text
    assert "CVE-EPSS-MATCH" in noncompliant.text
    assert "CVE-NO-MATCH" not in noncompliant.text
    overview = client.get("/services/payments-service?overview=true")
    assert "Active Fixable</span><strong>0</strong>" in overview.text
    assert "Active Non-Compliance</span><strong>2</strong>" in overview.text
    dashboard = client.get("/").text
    service_row = dashboard[dashboard.index("Payments Service"):dashboard.index("</tr>", dashboard.index("Payments Service"))]
    assert ">2</td>" in service_row
    assert ">3</td>" not in service_row


def test_administrator_can_edit_service_metadata():
    client = new_client(); ingest(client)
    detail = client.get("/services/payments-service").text
    assert "Edit service information" in detail
    assert "Actions" in detail and ">Edit</button>" in detail and ">Archive</button>" in detail
    overview = client.get("/services/payments-service?overview=true").text
    assert "Edit service information" in overview
    assert '<details class="actions-menu">' in overview
    assert ">Edit</button>" in overview and ">Archive</button>" in overview
    response = client.post("/admin/services/payments-service", data={
        "csrf_token": csrf(client), "name": "Payments Platform", "owner": "Cyber Team",
        "description": "Payments service metadata", "poc": "owner@example.invalid", "manual_version": "3.0",
    }, follow_redirects=False)
    assert response.status_code == 303
    page = client.get("/services/payments-service").text
    assert "Payments Platform" in page and "Version 3.0" in page and "Payments service metadata" in page


def test_edit_service_preserves_service_tab_query_rename_and_confirmation():
    client = new_client(); ingest(client)
    destinations = [
        "/services/payments-service?overview=true",
        "/services/payments-service?architecture=true",
        "/services/payments-service?artifacts=true",
        "/services/payments-service?validation=true&validation_run=run-7",
        "/services/payments-service?findings=true&findings_view=raw&q=openssl&severity=HIGH&page=3&page_size=25",
        "/services/payments-service?remediations=true&tab=pipeline",
        "/services/payments-service?activity=true",
    ]
    for index, destination in enumerate(destinations):
        response = client.post("/admin/services/payments-service", data={
            "csrf_token": csrf(client),
            "name": f"Payments Platform {index}",
            "return_to": destination,
        }, follow_redirects=False)
        assert response.status_code == 303
        location = response.headers["location"]
        parsed = urllib.parse.urlsplit(location)
        assert parsed.path == "/services/payments-service"
        expected_query = urllib.parse.parse_qsl(urllib.parse.urlsplit(destination).query, keep_blank_values=True)
        assert urllib.parse.parse_qsl(parsed.query, keep_blank_values=True) == expected_query + [("saved", "1")]

    # The display-name rename does not alter the stable service route, and the
    # shared service shell renders confirmation on the preserved destination.
    page = client.get(response.headers["location"])
    assert page.status_code == 200
    assert "Payments Platform 6" in page.text
    assert "Service information saved successfully." in page.text


def test_findings_filter_controls_and_view_selector_keep_expected_contract():
    client = new_client(); ingest(client)
    raw = client.get(
        "/services/payments-service?findings=true&findings_view=raw&q=CVE&severity=Critical&resource=registry&page=1&page_size=50"
    )
    assert raw.status_code == 200
    assert 'class="finding-filter-field">Search findings' in raw.text
    assert '<span>Severity</span><details class="multi-select-filter">' in raw.text
    assert 'class="finding-filter-field">Resource' in raw.text
    assert 'name="q" value="CVE"' in raw.text
    assert "1 selected" in raw.text
    assert 'name="resource" value="registry"' in raw.text
    assert '<nav class="finding-view-selector"' in raw.text
    assert ">Simplified</a>" in raw.text and ">Raw</a>" in raw.text
    assert 'class="secondary-button active"' in raw.text

    simplified = client.get("/services/payments-service?findings=true&findings_view=simplified&page_size=50")
    assert simplified.status_code == 200
    assert "Simplified Findings" in simplified.text
    assert 'class="finding-filter-field">Search findings' in simplified.text

    css = (Path(__file__).parents[1] / "app" / "static" / "app.css").read_text(encoding="utf-8")
    assert "--finding-filter-height:42px" in css
    assert ".finding-view-selector{display:inline-flex" in css
    assert "margin:0 0 .75rem" in css


def _helm_chart_archive(name: str, extra_name: str = "templates/deployment.yaml", extra_content: str = "apiVersion: apps/v1\nkind: Deployment\n") -> bytes:
    output = BytesIO()
    with tarfile.open(fileobj=output, mode="w:gz") as bundle:
        for path, content in {
            f"{name}/Chart.yaml": f"apiVersion: v2\nname: {name}\nversion: 1.0.0\n",
            f"{name}/{extra_name}": extra_content,
        }.items():
            data = content.encode()
            member = tarfile.TarInfo(path)
            member.size = len(data)
            bundle.addfile(member, BytesIO(data))
    return output.getvalue()


def test_artifacts_page_exposes_first_class_helm_and_kubernetes_workflow():
    client = new_client(); ingest(client)
    page = client.get("/services/payments-service?artifacts=true")
    assert page.status_code == 200
    assert 'data-open-artifact-dialog>+ Add Artifact</button>' in page.text
    assert 'data-artifact-type="helm_repository">+ Add Repository</button>' in page.text
    assert 'data-artifact-type="helm_chart">+ Add Helm Chart</button>' in page.text
    assert 'id="add-artifact-dialog"' in page.text
    assert 'value="helm_repository" checked>' in page.text and "Helm Repository" in page.text
    assert 'value="helm_chart">' in page.text and "Helm Chart" in page.text
    assert 'value="kubernetes">' in page.text and "Kubernetes Manifest" in page.text
    assert 'value="repository"' in page.text
    assert 'value="oci">' in page.text and "OCI Registry" in page.text
    assert 'value="upload" checked>' in page.text and "Upload Chart" in page.text
    assert "catalog is discovered without treating the repository as a deployable chart" in page.text
    assert 'src="/static/artifacts.js"' in page.text and 'data-artifact-table="charts"' not in page.text
    assert "No Kubernetes manifests added." in page.text


def test_service_artifact_acquisition_reuses_multi_chart_repository_pipeline(monkeypatch):
    from app import main as portal_main
    client = new_client(); ingest(client)
    catalog = {"repository_url": "https://charts.example.invalid/helm-charts", "index_url": "https://charts.example.invalid/helm-charts/index.yaml", "api_version": "v1", "generated": "now", "charts": [
        {"name": "alpha", "latest": {"version": "1.0.0", "url": "https://charts.example.invalid/alpha.tgz"}, "versions": [{"version": "1.0.0", "url": "https://charts.example.invalid/alpha.tgz"}]},
        {"name": "beta", "latest": {"version": "2.0.0", "url": "https://charts.example.invalid/beta.tgz"}, "versions": [{"version": "2.0.0", "url": "https://charts.example.invalid/beta.tgz"}]},
    ]}
    monkeypatch.setattr(portal_main, "_discover_helm_repository", lambda reference, certificates: catalog)
    response = client.post("/services/payments-service/artifacts/acquire", data={
        "csrf_token": csrf(client), "artifact_type": "helm_repository", "source_method": "repository",
        "source_reference": "https://charts.example.invalid/helm-charts",
    }, follow_redirects=False)
    assert response.status_code == 303 and "artifact_added=repository" in response.headers["location"]
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        repository = db.scalar(select(ServiceArtifact).where(ServiceArtifact.service_id == service.id, ServiceArtifact.artifact_type == "helm_repository"))
        charts = db.scalars(select(ServiceArtifact).where(ServiceArtifact.parent_repository_id == repository.id).order_by(ServiceArtifact.chart_name)).all()
        assert repository.source_reference == catalog["repository_url"] and repository.source_metadata["chart_count"] == 2
        assert [(chart.chart_name, chart.chart_version) for chart in charts] == [("alpha", "1.0.0"), ("beta", "2.0.0")]
        assert all(not chart.revisions for chart in charts)
        alpha_id = charts[0].id
    page = client.get("/services/payments-service?artifacts=true")
    assert "Helm Repositories" in page.text and "Helm Charts" in page.text
    assert "alpha" in page.text and "beta" in page.text and ">2</strong> charts" in page.text
    assert "Discovered" in page.text and "Not Validated" in page.text
    assert "https://charts.example.invalid/helm-charts" in page.text
    monkeypatch.setattr(portal_main, "_download_public_chart", lambda reference, certificates: [(_helm_chart_archive("alpha"), "alpha-1.0.0.tgz")])
    materialized = client.post(f"/services/payments-service/artifacts/charts/{alpha_id}/materialize", data={"csrf_token": csrf(client), "version": "1.0.0"}, follow_redirects=False)
    assert materialized.status_code == 303
    with SessionLocal() as db:
        revisions = db.scalars(select(ServiceArtifactRevision).where(ServiceArtifactRevision.artifact_id == alpha_id)).all()
        assert len(revisions) == 1 and revisions[0].source_metadata["repository_id"] == repository.id
        assert revisions[0].source_metadata["chart_version"] == "1.0.0"


def test_packaged_helm_and_existing_kubernetes_uploads_are_retained_and_rbac_enforced():
    client = new_client(); ingest(client)
    response = client.post("/services/payments-service/artifacts/acquire", data={
        "csrf_token": csrf(client), "artifact_type": "helm", "source_method": "upload",
    }, files={"files": ("application-1.0.0.tgz", _helm_chart_archive("application"), "application/gzip")}, follow_redirects=False)
    assert response.status_code == 303
    manifests = client.post("/services/payments-service/artifacts/upload", data={
        "csrf_token": csrf(client), "artifact_type": "kubernetes",
    }, files={"files": ("deployment.yaml", b"apiVersion: apps/v1\nkind: Deployment\n", "text/yaml")}, follow_redirects=False)
    assert manifests.status_code == 303
    with SessionLocal() as db:
        service_id = db.scalar(select(Service.id).where(Service.service_key == "payments-service"))
        artifacts = db.scalars(select(ServiceArtifact).where(ServiceArtifact.service_id == service_id)).all()
        assert {artifact.artifact_type for artifact in artifacts} == {"helm_chart", "kubernetes"}
    add_user("artifact-viewer", "Assessor", service_id=service_id)
    viewer = new_client("artifact-viewer")
    denied = viewer.post("/services/payments-service/artifacts/acquire", data={
        "csrf_token": csrf(viewer), "artifact_type": "helm", "source_method": "upload",
    }, files={"files": ("blocked.tgz", _helm_chart_archive("blocked"), "application/gzip")})
    assert denied.status_code == 403


def test_helm_archive_traversal_and_links_are_rejected():
    client = new_client(); ingest(client)
    output = BytesIO()
    with tarfile.open(fileobj=output, mode="w:gz") as bundle:
        member = tarfile.TarInfo("../../outside/Chart.yaml")
        data = b"name: unsafe\nversion: 1\n"
        member.size = len(data)
        bundle.addfile(member, BytesIO(data))
    response = client.post("/services/payments-service/artifacts/acquire", data={
        "csrf_token": csrf(client), "artifact_type": "helm", "source_method": "upload",
    }, files={"files": ("unsafe.tgz", output.getvalue(), "application/gzip")})
    assert response.status_code == 400
    assert "unsafe path" in response.json()["detail"]

    linked = BytesIO()
    with tarfile.open(fileobj=linked, mode="w:gz") as bundle:
        chart = tarfile.TarInfo("unsafe/Chart.yaml")
        chart_data = b"name: unsafe\nversion: 1\n"
        chart.size = len(chart_data)
        bundle.addfile(chart, BytesIO(chart_data))
        symlink = tarfile.TarInfo("unsafe/templates/escape.yaml")
        symlink.type = tarfile.SYMTYPE
        symlink.linkname = "../../outside.yaml"
        bundle.addfile(symlink)
    response = client.post("/services/payments-service/artifacts/acquire", data={
        "csrf_token": csrf(client), "artifact_type": "helm", "source_method": "upload",
    }, files={"files": ("linked.tgz", linked.getvalue(), "application/gzip")})
    assert response.status_code == 400
    assert "unsafe link" in response.json()["detail"]


def test_helm_repository_and_tls_failures_are_specific_and_secure(monkeypatch):
    from app import main as portal_main
    client = new_client(); ingest(client)
    real_fetch = portal_main._fetch_public_url
    monkeypatch.setattr(portal_main, "_fetch_public_url", lambda url, certificates=None: (b"not a helm index", url))
    invalid = client.post("/services/payments-service/artifacts/acquire", data={
        "csrf_token": csrf(client), "artifact_type": "helm", "source_method": "repository",
        "source_reference": "https://charts.example.invalid",
    })
    assert invalid.status_code == 400
    assert "valid index.yaml" in invalid.json()["detail"]

    monkeypatch.setattr(portal_main, "_fetch_public_url", real_fetch)
    def tls_failure(*_args, **_kwargs):
        raise urllib.error.URLError(ssl.SSLCertVerificationError("untrusted issuer"))
    monkeypatch.setattr(portal_main.urllib.request, "urlopen", tls_failure)
    with __import__("pytest").raises(__import__("fastapi").HTTPException) as raised:
        portal_main._fetch_public_url("https://private.example.invalid/chart.tgz", [])
    assert "TLS certificate verification failed" in raised.value.detail
    source = Path(portal_main.__file__).read_text(encoding="utf-8")
    assert "--insecure-skip-tls-verify" not in source


def test_uploaded_helm_revision_is_directly_available_to_deployment_validation(monkeypatch):
    from app import main as portal_main
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    submitted = []
    monkeypatch.setattr(portal_main, "_submit_validation_run", lambda run_id: submitted.append(run_id))
    client = new_client()
    with SessionLocal() as db:
        db.add(Service(service_key="artifact-only", name="Artifact Only"))
        db.commit()
    added = client.post("/services/artifact-only/artifacts/acquire", data={
        "csrf_token": csrf(client), "artifact_type": "helm", "source_method": "upload",
    }, files={"files": ("application.tgz", _helm_chart_archive("application"), "application/gzip")}, follow_redirects=False)
    assert added.status_code == 303
    with SessionLocal() as db:
        artifact = db.scalar(select(ServiceArtifact).join(Service).where(Service.service_key == "artifact-only"))
        revision = db.scalar(select(ServiceArtifactRevision).where(ServiceArtifactRevision.artifact_id == artifact.id))
        artifact_id, revision_id = artifact.id, revision.id
    validation = client.post("/services/artifact-only/deployment-validations", data={
        "csrf_token": csrf(client), "artifact_id": str(artifact_id),
    }, follow_redirects=False)
    assert validation.status_code == 303
    with SessionLocal() as db:
        run = db.scalar(select(DeploymentValidationRun).where(DeploymentValidationRun.service_id == select(Service.id).where(Service.service_key == "artifact-only").scalar_subquery()))
        assert run.artifact_revision_id == revision_id and run.execution_id is None
        assert run.status == "QUEUED" and submitted == [run.id]


def test_artifact_validation_badge_is_exact_revision_safe(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    client = new_client()
    with SessionLocal() as db:
        service = Service(service_key="revision-safe", name="Revision Safe")
        db.add(service); db.flush()
        artifact = ServiceArtifact(service_id=service.id, artifact_type="helm_chart", artifact_name="application",
                                   chart_name="application", chart_version="1.2.0", source_type="upload")
        db.add(artifact); db.flush()
        original = ServiceArtifactRevision(artifact_id=artifact.id, revision_number=1, revision_label="ORIGINAL",
            files={"application/Chart.yaml": "name: application\nversion: 1.2.0\n"}, checksum="a" * 64)
        db.add(original); db.flush()
        db.add(DeploymentValidationRun(run_key="validated-original", service_id=service.id,
            artifact_revision_id=original.id, artifact_type="WORKING", artifact_reference=f"artifact:{artifact.id}:r1",
            status="VERIFIED", phase="COMPLETE", cleanup_status="COMPLETE"))
        db.commit(); artifact_id = artifact.id
    page = client.get("/services/revision-safe?artifacts=true")
    assert "✓</span> Validated" in page.text and "Revalidate" in page.text
    with SessionLocal() as db:
        db.add(ServiceArtifactRevision(artifact_id=artifact_id, revision_number=2, revision_label="WORKING",
            files={"application/Chart.yaml": "name: application\nversion: 1.2.0\n", "application/values.yaml": "replicas: 2\n"}, checksum="b" * 64))
        db.commit()
    page = client.get("/services/revision-safe?artifacts=true")
    assert "Not Validated" in page.text and "✓</span> Validated" not in page.text
    assert "WORKING · Revision 2" in page.text and "validated-original" not in page.text


def test_service_snapshot_separates_staged_services_from_active_and_archived():
    client = new_client()
    with SessionLocal() as db:
        group = Group(name="Staged Services")
        db.add(group)
        db.commit()
        group_id = group.id
    response = client.post("/admin/services/stage", data={
        "csrf_token": csrf(client), "service_id": "helm-test", "group_id": str(group_id), "next_path": "/",
    }, follow_redirects=False)
    assert response.status_code == 303
    default = client.get("/")
    assert "helm-test" not in default.text
    assert "View Staged (1)" in default.text
    staged = client.get("/?lifecycle=staged&q=helm-test&sort=name")
    assert staged.status_code == 200
    assert "helm-test" in staged.text
    assert "View Staged (1)" in staged.text
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "helm-test"))
        assert service.lifecycle_status == "staged"
    ingested = ingest(client, execution="helm-test-run", service_id="helm-test")
    assert ingested.status_code == 201
    assert "helm-test" in client.get("/").text
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "helm-test"))
        assert service.lifecycle_status == "active"


def test_staged_service_promotes_on_any_ingested_finding_or_evidence():
    client = new_client()
    with SessionLocal() as db:
        group = Group(name="Ingest Promotion")
        db.add(group)
        db.commit()
        group_id = group.id

    response = client.post("/admin/services/stage", data={
        "csrf_token": csrf(client), "service_id": "staged-evidence", "group_id": str(group_id), "next_path": "/",
    }, follow_redirects=False)
    assert response.status_code == 303

    data = payload("staged-evidence-run", datetime.now(timezone.utc), ["CVE-2026-9999"], service_id="staged-evidence")
    data["policy_findings"] = [{
        "finding": "KSV-999", "severity": "High", "target": "Deployment/staged-evidence",
    }]
    data["service_overview"] = {
        "images": [{"image": "registry.example/staged-evidence:1", "digest": "sha256:" + "a" * 64}],
        "artifacts": [{"type": "chart", "name": "staged-evidence", "version": "1.0.0"}],
    }
    ingested = client.post("/api/v1/pipeline-results", json=data, headers=pipeline_headers)
    assert ingested.status_code == 201
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "staged-evidence"))
        assert service.lifecycle_status == "active"
        promotion = db.scalar(select(AuditEvent).where(
            AuditEvent.action == "service.promoted", AuditEvent.target_id == str(service.id),
        ))
        assert promotion is not None
        assert promotion.detail["reason"] == "ingested_evidence"


def test_generated_staged_name_is_restored_idempotently_without_changing_identity_or_evidence():
    client = new_client()
    with SessionLocal() as db:
        group = Group(name="Lifecycle Identity")
        db.add(group); db.commit()
        group_id = group.id
    staged = client.post("/admin/services/stage", data={
        "csrf_token": csrf(client), "service_id": "torture-test", "group_id": str(group_id), "next_path": "/",
    }, follow_redirects=False)
    assert staged.status_code == 303
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "torture-test"))
        service_id = service.id
        assert service.name == "Staged — torture-test"
        assert service.staging_name_generated is True and service.staging_original_name == "torture-test"

    body = helm_payload("torture-lifecycle-run", "torture-test")
    # Reproduce the reported client behavior: the submitted display name still
    # contains CATS' generated staging label.
    body["service"]["name"] = "Staged — torture-test"
    body["findings"] = [{"cve": "CVE-2026-4242", "severity": "High", "image": "registry/torture:1"}]
    response = client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers)
    assert response.status_code == 201

    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "torture-test"))
        execution = db.scalar(select(Execution).where(Execution.service_id == service.id))
        finding = db.scalar(select(Finding).where(Finding.service_id == service.id))
        assert service.id == service_id and service.lifecycle_status == "active"
        assert service.name == "torture-test"
        assert service.staging_name_generated is False and service.staging_original_name is None
        assert execution.service_id == service_id and finding.service_id == service_id
        db.add(DeploymentValidationRun(
            run_key="DV-LIFECYCLE-VERIFIED", service_id=service.id, execution_id=execution.id,
            artifact_type="ORIGINAL", artifact_reference=execution.execution_key,
            status="VERIFIED", phase="COMPLETE", engine="kind", completed_at=datetime.now(timezone.utc),
            diagnostics={"classification_summary": {"expected_resources": 1, "observed_expected": 1,
                                                       "expected_only": 0, "failed": 0}},
        ))
        db.commit()

    evidence = client.get("/api/v1/services/torture-test/architecture-evidence")
    assert evidence.status_code == 200 and evidence.json()["architecture"]["state"] == "VERIFIED"
    for url in (
        "/services/torture-test?overview=true", "/services/torture-test?architecture=true",
        "/services/torture-test?artifacts=true", "/services/torture-test?validation=true",
        "/services/torture-test?remediations=true",
    ):
        page = client.get(url)
        assert page.status_code == 200 and "Staged — torture-test" not in page.text
    activity = client.get("/services/torture-test?activity=true")
    assert activity.status_code == 200 and "torture-test" in activity.text
    assert "Staged — torture-test" in activity.text  # immutable name-at-event audit context

    # A retry is idempotent and cannot modify the restored name.
    duplicate = client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers)
    assert duplicate.status_code == 201 and duplicate.json()["duplicate"] is True
    with SessionLocal() as db:
        service = db.get(Service, service_id)
        assert service.name == "torture-test" and service.lifecycle_status == "active"


def test_user_owned_staged_prefix_is_never_stripped_without_generation_metadata():
    client = new_client()
    with SessionLocal() as db:
        service = Service(service_key="production-app", name="Staged — Production Application",
                          lifecycle_status="staged", staging_name_generated=False)
        db.add(service); db.commit(); service_id = service.id
    body = payload("production-app-run", datetime.now(timezone.utc), [], service_id="production-app")
    body["service"]["name"] = "Staged — Production Application"
    assert client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        service = db.get(Service, service_id)
        assert service.lifecycle_status == "active"
        assert service.name == "Staged — Production Application"


def test_startup_reconciles_preexisting_staged_evidence():
    from app import main as portal_main

    with SessionLocal() as db:
        service = Service(service_key="legacy-staged", name="Staged — legacy-staged", lifecycle_status="staged")
        db.add(service)
        db.flush()
        db.add(Execution(
            execution_key="legacy-staged-run", service_id=service.id,
            scanned_at=datetime.now(timezone.utc), complete=True,
            raw_payload={"service": {"id": "legacy-staged"}, "findings": [], "service_overview": {"artifacts": [{"type": "chart"}]}},
        ))
        db.commit()

    assert portal_main.promote_staged_services_with_evidence() == 1
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "legacy-staged"))
        assert service.lifecycle_status == "active"
        promotion = db.scalar(select(AuditEvent).where(
            AuditEvent.action == "service.promoted", AuditEvent.target_id == str(service.id),
        ))
        assert promotion is not None
        assert promotion.detail["reason"] == "existing_ingested_evidence"


def test_service_manager_scope_and_exception_approval_separation():
    admin = new_client()
    ingest(admin, cves=["CVE-2026-0001"])
    ingest(admin, "run-other", [], service_id="other-service")
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        finding = db.scalar(select(Finding).where(Finding.service_id == service.id))
        service_id, finding_id = service.id, finding.id
    add_user("manager", "Service Manager", service_id)
    manager = new_client("manager")
    page = manager.get("/").text
    assert "Payments Service" in page and "Other Service" not in page
    expiry = datetime.now(timezone.utc) + timedelta(days=30)
    requested = manager.post(f"/findings/{finding_id}/exceptions", data={
        "csrf_token": csrf(manager), "justification": "Vendor remediation scheduled",
        "expires_at": expiry.isoformat(), "ticket": "RISK-42",
    }, follow_redirects=False)
    assert requested.status_code == 303
    with SessionLocal() as db:
        workflow_id = db.scalar(select(WorkflowRequest.id))
    assert manager.post(f"/requests/{workflow_id}/review", data={"csrf_token": csrf(manager), "decision": "approved"}).status_code == 403
    add_user("cyber", "Cybersecurity")
    cyber = new_client("cyber")
    assert cyber.post(f"/requests/{workflow_id}/review", data={"csrf_token": csrf(cyber), "decision": "approved"}, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        exception_id = db.scalar(select(ExceptionRecord.id))
    assert cyber.post(f"/exceptions/{exception_id}/revoke", data={"csrf_token": csrf(cyber)}, follow_redirects=False).status_code == 303
    assert "Revoked" in cyber.get("/requests").text


def test_service_workspace_exposes_embedded_poam_and_activity_tabs():
    client = new_client()
    ingest(client, cves=["CVE-2026-0099"])
    poam = client.get("/services/payments-service?poam=true")
    assert poam.status_code == 200
    assert "Payments Service POA&amp;M" in poam.text
    assert 'href="/services/payments-service?activity=true"' in poam.text
    activity = client.get("/services/payments-service?activity=true")
    assert activity.status_code == 200
    assert "Service activity" in activity.text
    assert "Scan" in activity.text


def test_poam_is_scoped_and_requires_cybersecurity_approval():
    admin = new_client()
    ingest(admin, cves=["CVE-2026-0002"])
    ingest(admin, "other-run", [], service_id="other-service")
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        other = db.scalar(select(Service).where(Service.service_key == "other-service"))
        finding = db.scalar(select(Finding).where(Finding.service_id == service.id))
        finding.episode_started = datetime.now(timezone.utc) - timedelta(days=100)
        service_id, other_id, finding_id = service.id, other.id, finding.id
        db.commit()
    add_user("poam-manager", "Service Manager", service_id)
    manager = new_client("poam-manager")
    poam_index = manager.get("/poam").text
    assert "Create POA&amp;M entry" in poam_index
    assert "Payments Service" in poam_index and "Other Service" not in poam_index
    assert manager.get("/poam/services/other-service").status_code == 403
    service_poam = manager.get("/poam/services/payments-service")
    assert service_poam.status_code == 200
    assert "Payments Service POA&amp;M" in service_poam.text
    noncompliant_page = manager.get("/services/payments-service?finding_state=noncompliant").text
    assert 'class="secondary-button poam-action"' in noncompliant_page
    assert 'data-title="Remediate CVE-2026-0002"' in noncompliant_page
    assert 'id="shared-poam-dialog"' in noncompliant_page
    created = manager.post("/poam", data={
        "csrf_token": csrf(manager), "service_id": service_id, "item_type": "missing_evidence",
        "title": "Missing authorization evidence", "description": "Authorization package is absent",
        "remediation": "Upload signed authorization package", "ticket": "POAM-7", "finding_id": "",
    }, follow_redirects=False)
    assert created.status_code == 303
    forbidden = manager.post("/poam", data={
        "csrf_token": csrf(manager), "service_id": other_id, "item_type": "missing_requirement",
        "title": "Out of scope", "description": "Should not be accepted", "remediation": "None",
    })
    assert forbidden.status_code == 403
    vulnerability = manager.post(f"/findings/{finding_id}/poams", data={
        "csrf_token": csrf(manager), "title": "Patch vulnerability", "remediation": "Deploy fixed image",
    }, follow_redirects=False)
    assert vulnerability.status_code == 303
    with SessionLocal() as db:
        entries = db.scalars(select(PoamEntry).order_by(PoamEntry.id)).all()
        assert [entry.status for entry in entries] == ["pending_approval", "pending_approval"]
        workflow_id = db.scalar(select(WorkflowRequest.id).where(WorkflowRequest.poam_id == entries[0].id))
    assert manager.post(f"/requests/{workflow_id}/review", data={
        "csrf_token": csrf(manager), "decision": "approved",
    }).status_code == 403
    add_user("poam-cyber", "Cybersecurity")
    cyber = new_client("poam-cyber")
    assert cyber.post(f"/requests/{workflow_id}/review", data={
        "csrf_token": csrf(cyber), "decision": "approved",
    }, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        entry = db.get(PoamEntry, entries[0].id)
        assert entry.status == "active" and entry.approved_by_id is not None
    detail = cyber.get("/poam/services/payments-service").text
    assert "Missing authorization evidence" in detail and "Active" in detail
    updated = manager.post(f"/poam/entries/{entries[0].id}/update", data={
        "csrf_token": csrf(manager), "title": "Updated authorization evidence",
        "description": "The authorization package remains incomplete",
        "remediation": "Upload and validate the signed authorization package", "due_date": "",
        "ticket": "POAM-8", "reason": "Milestones changed",
    }, follow_redirects=False)
    assert updated.status_code == 303
    with SessionLocal() as db:
        update_workflow = db.scalar(select(WorkflowRequest).where(WorkflowRequest.request_type == "poam_update"))
    assert cyber.post(f"/requests/{update_workflow.id}/review", data={
        "csrf_token": csrf(cyber), "decision": "approved", "review_reason": "Update validated",
    }, follow_redirects=False).status_code == 303
    assert "Updated authorization evidence" in cyber.get(f"/poam/entries/{entries[0].id}").text
    completed = manager.post(f"/poam/entries/{entries[0].id}/complete", data={
        "csrf_token": csrf(manager), "closure_note": "Authorization package supplied and validated",
        "evidence_reference": "EVIDENCE-2026-42",
    }, follow_redirects=False)
    assert completed.status_code == 303
    with SessionLocal() as db:
        complete_workflow = db.scalar(select(WorkflowRequest).where(WorkflowRequest.request_type == "poam_complete"))
    assert cyber.post(f"/requests/{complete_workflow.id}/review", data={
        "csrf_token": csrf(cyber), "decision": "approved", "review_reason": "Evidence verified",
    }, follow_redirects=False).status_code == 303
    completed_page = cyber.get(f"/poam/entries/{entries[0].id}").text
    assert "Completed" in completed_page and "EVIDENCE-2026-42" in completed_page
    assert manager.post(f"/poam/entries/{entries[0].id}/reopen", data={
        "csrf_token": csrf(manager), "reason": "New evidence gap was identified",
    }, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        reopen_workflow = db.scalar(select(WorkflowRequest).where(WorkflowRequest.request_type == "poam_reopen"))
    assert cyber.post(f"/requests/{reopen_workflow.id}/review", data={
        "csrf_token": csrf(cyber), "decision": "approved", "review_reason": "Reopen approved",
    }, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        assert db.get(PoamEntry, entries[0].id).status == "active"
    export = cyber.get("/poam/services/payments-service/export.xlsx")
    assert export.status_code == 200
    export_sheet = load_workbook(BytesIO(export.content))["POA&M"]
    assert "Updated authorization evidence" in [cell.value for cell in export_sheet["D"]]
    add_user("poam-assessor", "Assessor")
    assessor = new_client("poam-assessor")
    assessor_index = assessor.get("/poam").text
    assert "Payments Service" in assessor_index and "Other Service" in assessor_index
    assert "Create POA&amp;M entry" not in assessor_index


def test_audit_logs_default_to_ten_and_expand_within_retention():
    client = new_client()
    with SessionLocal() as db:
        for number in range(15):
            db.add(AuditEvent(action=f"test.event.{number}", target_type="test", target_id=str(number)))
        db.commit()
    default_page = client.get("/admin/audit")
    assert default_page.status_code == 200
    assert "Showing 10 of" in default_page.text
    assert default_page.text.count("test.event.") == 10
    expanded = client.get("/admin/audit?show=60")
    assert expanded.text.count("test.event.") == 15
    assert "Collapse to latest 10" in expanded.text


def test_archive_requires_request_and_separate_approval():
    admin = new_client(); ingest(admin)
    requested = admin.post("/services/payments-service/archive", data={"csrf_token": csrf(admin), "reason": "Retired"}, follow_redirects=False)
    assert requested.status_code == 303
    with SessionLocal() as db:
        workflow_id = db.scalar(select(WorkflowRequest.id))
    assert admin.post(f"/requests/{workflow_id}/review", data={"csrf_token": csrf(admin), "decision": "approved"}).status_code == 409
    add_user("cyber", "Cybersecurity")
    cyber = new_client("cyber")
    assert cyber.post(f"/requests/{workflow_id}/review", data={"csrf_token": csrf(cyber), "decision": "approved"}, follow_redirects=False).status_code == 303
    assert "Payments Service" in cyber.get("/?archived=true").text


def test_admin_can_create_account_custom_role_and_assignment():
    client = new_client(); ingest(client)
    assert client.post("/admin/users", data={"csrf_token": csrf(client), "username": "assessor1", "display_name": "Assessor One", "password": "temporary-password"}, follow_redirects=False).status_code == 303
    assert client.post("/admin/roles", data={"csrf_token": csrf(client), "name": "Evidence Exporter", "description": "Exports evidence", "permissions": ["service.view", "service.export"]}, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == "assessor1")); role = db.scalar(select(Role).where(Role.name == "Evidence Exporter"))
    assert client.post("/admin/assignments", data={"csrf_token": csrf(client), "user_id": user.id, "role_id": role.id, "service_id": ""}, follow_redirects=False).status_code == 303


def test_reset_password_dialog_is_not_constrained_as_a_table_action():
    css = (Path(__file__).parents[1] / "app" / "static" / "app.css").read_text(encoding="utf-8")
    assert ".account-actions > form,.account-actions > button" in css
    assert ".account-actions form,.account-actions > button" not in css
    assert ".account-actions form button" not in css


def test_excel_export_remains_valid():
    client = new_client(); ingest(client, cves=["CVE-2026-0001"])
    service_page = client.get("/services/payments-service").text
    assert "Export to Excel" in service_page
    overview_page = client.get("/services/payments-service?overview=true").text
    assert '>Export</a>' in overview_page
    assert 'class="secondary-button" href="/services/payments-service?architecture=true">Architecture</a>' not in overview_page
    assert 'class="secondary-button" href="/services/payments-service/helm-diagram.svg">Helm Diagram</a>' not in overview_page
    diagram = client.get("/services/payments-service/helm-diagram.svg")
    assert diagram.status_code == 200
    assert diagram.headers["content-type"].startswith("image/svg+xml")
    assert "Helm rendering" in diagram.text
    assert "Ports / protocols" in diagram.text
    assert "Latest Evidence:" in service_page and "None received" not in service_page
    # Findings now prioritizes the table; service-level metric cards live on Overview.
    assert "metric-card selected" not in service_page
    response = client.get("/services/payments-service/export.xlsx")
    assert response.status_code == 200
    assert load_workbook(BytesIO(response.content))["Findings"]["A2"].value == "CVE-2026-0001"
    package = client.get("/services/payments-service/export.xlsx?include_diagrams=true")
    assert package.status_code == 200
    assert package.headers["content-type"].startswith("application/zip")
    with ZipFile(BytesIO(package.content)) as archive:
        names = set(archive.namelist())
        assert {"service-export.xlsx", "legacy-helm-diagram.svg"} <= names
        assert {f"architecture-{view}.svg" for view in ("all", "configuration", "containers", "flow", "network", "storage")} <= names


def test_architecture_reflow_endpoint_and_canonical_legacy_export():
    from test_architecture_packing import flows

    client = new_client()
    body = payload("responsive", datetime.now(timezone.utc), [])
    body["service_overview"] = flows(8)
    assert client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers).status_code == 201
    url = "/services/payments-service?architecture=true"
    narrow = client.get(url + "&layout_width=1000")
    wide = client.get(url + "&layout_width=1920")
    assert narrow.status_code == wide.status_code == 200
    assert len(wide.json()["flow"]["components"]) == 8
    assert narrow.json()["flow"]["bounds"]["height"] > wide.json()["flow"]["bounds"]["height"]
    assert client.get(url + "&layout_width=0").status_code == 422
    assert client.get(url + "&layout_width=10001").status_code == 422
    assert TestClient(app).get(url + "&layout_width=1000", follow_redirects=False).status_code in {303, 401, 403}
    diagram = client.get("/services/payments-service/helm-diagram.svg")
    assert "CATS Architecture" in diagram.text
    assert "application-7" in diagram.text


def test_architecture_and_overview_use_exact_persisted_validation_evidence():
    client = new_client()
    assert client.post("/api/v1/pipeline-results", json=helm_payload("architecture-verified"), headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        run = db.scalar(select(DeploymentValidationRun))
        run.status = "VERIFIED"; run.phase = "COMPLETE"; run.cleanup_status = "COMPLETE"
        run.completed_at = datetime.now(timezone.utc)
        run.comparison = {"matched": ["Service/validation/demo"], "declared_only": [], "defaulted": [], "observed_only": [],
                          "expected_evidence": [{"apiVersion": "v1", "kind": "Service", "namespace": "validation", "name": "demo", "matched": True}], "observed_evidence": []}
        run.diagnostics = {"classification_summary": {"expected_resources": 1, "observed_expected": 1, "expected_only": 0, "runtime_generated": 0, "observed_only": 0, "failed": 0}}
        db.commit()
    architecture = client.get("/services/payments-service?architecture=true")
    overview = client.get("/services/payments-service?overview=true")
    evidence = client.get("/api/v1/services/payments-service/architecture-evidence").json()
    assert 'data-architecture-badge>✓ Verified<' in architecture.text
    assert 'data-architecture-summary-badge>Verified<' in overview.text
    assert evidence["architecture"]["state"] == "VERIFIED"
    assert evidence["graph"]["summary"]["runtime_verified"] == 1


def test_new_working_revision_is_declared_until_that_exact_revision_is_verified():
    client = new_client()
    assert client.post("/api/v1/pipeline-results", json=helm_payload("architecture-revision"), headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        execution = db.scalar(select(Execution).where(Execution.service_id == service.id))
        original = db.scalar(select(DeploymentValidationRun).where(DeploymentValidationRun.execution_id == execution.id))
        original.status = "VERIFIED"; original.phase = "COMPLETE"; original.completed_at = datetime.now(timezone.utc)
        original.diagnostics = {"classification_summary": {"expected_resources": 1, "observed_expected": 1,
                                                               "expected_only": 0, "failed": 0}}
        db.commit(); execution_id = execution.id; service_id = service.id
    assert client.get("/api/v1/services/payments-service/architecture-evidence").json()["architecture"]["state"] == "VERIFIED"

    assert client.post("/services/payments-service/artifacts/from-original", data={
        "csrf_token": csrf(client), "execution_id": str(execution_id),
    }, follow_redirects=False).status_code == 303
    with SessionLocal() as db:
        artifact = db.scalar(select(ServiceArtifact).where(ServiceArtifact.service_id == service_id))
        artifact_id = artifact.id
    assert client.post(f"/services/payments-service/artifacts/{artifact_id}/files", data={
        "csrf_token": csrf(client), "path": "values.yaml", "content": "replicaCount: 2\n",
    }, follow_redirects=False).status_code == 303

    unvalidated = client.get("/api/v1/services/payments-service/architecture-evidence").json()
    assert unvalidated["architecture"]["state"] == "DECLARED"
    assert unvalidated["graph"]["nodes"]  # declared/static graph remains available
    with SessionLocal() as db:
        revision = db.scalar(select(ServiceArtifactRevision).where(
            ServiceArtifactRevision.artifact_id == artifact_id,
            ServiceArtifactRevision.revision_label == "WORKING",
        ))
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        db.add(DeploymentValidationRun(
            run_key="DV-EXACT-WORKING", service_id=service.id, execution_id=execution_id,
            artifact_revision_id=revision.id, artifact_type="WORKING",
            artifact_reference=f"artifact:{artifact_id}:r{revision.revision_number}", engine="kind",
            status="VERIFIED", phase="COMPLETE", completed_at=datetime.now(timezone.utc),
            diagnostics={"classification_summary": {"expected_resources": 1, "observed_expected": 1,
                                                       "expected_only": 0, "failed": 0}},
        ))
        db.commit()
    verified = client.get("/api/v1/services/payments-service/architecture-evidence").json()
    assert verified["architecture"]["state"] == "VERIFIED"


def test_service_export_contains_authoritative_overview_sheets_and_provenance():
    client = new_client()
    body = payload("export-complete", datetime.now(timezone.utc), ["CVE-2026-0100"], service_id="export-service")
    body["service_overview"] = {
        "images": [{"image": "registry.example/team/api:2.0", "source_file": "charts/api/templates/deployment.yaml", "discovered_from": "Deployment/api"}],
        "helm_components": [{"chart": "api", "path": "charts/api", "declared_by": "charts/root/values.yaml", "enabled": False}],
        "missing_evidence": [{"type": "Chart", "item": "worker", "reason": "chart not found", "source_file": "charts/root/values.yaml"}],
    }
    assert client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers).status_code == 201
    response = client.get("/services/export-service/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    required = {"Summary", "Containers", "Helm Charts", "Artifacts", "Missing Evidence", "Ports & Protocols", "Accounts", "Dependencies", "Render Warnings", "Raw Findings", "Vulnerabilities", "Simplified Findings", "Configuration Findings", "POA&Ms", "Exceptions", "Mitigations", "Activity"}
    assert required.issubset(set(workbook.sheetnames))
    containers = list(workbook["Containers"].values)
    assert any("charts/api/templates/deployment.yaml" in str(row) for row in containers)
    missing = list(workbook["Missing Evidence"].values)
    assert any("charts/root/values.yaml" in str(row) for row in missing)


def test_service_export_accepts_singleton_recursive_helm_evidence():
    client = new_client()
    body = payload("export-singletons", datetime.now(timezone.utc), [], service_id="singleton-export")
    body["service_overview"] = {
        "images": {"image": "registry.example/team/api:2.0", "source_file": "charts/api/templates/deployment.yaml"},
        "helm_components": {"chart": "api", "path": "charts/api", "declared_by": "values.yml"},
        "ports": {"port": 8080, "protocol": "TCP", "service": "api"},
        "accounts": {"name": "api", "kind": "ServiceAccount", "namespace": "default"},
        "rendered_resources": {"items": [{
            "apiVersion": "v1", "kind": "Service", "metadata": {"name": "api"},
            "spec": {"ports": [{"port": 8080}]},
        }]},
    }
    assert client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers).status_code == 201
    response = client.get("/services/singleton-export/export.xlsx?include_diagrams=true")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/zip")


def test_service_poc_is_displayed_and_exported():
    client = new_client(); ingest(client)
    page = client.get("/services/payments-service")
    assert "POC: cats-test-poc@example.invalid" in page.text
    workbook = load_workbook(BytesIO(client.get("/exports/services.xlsx").content))
    assert "POC" in [cell.value for cell in workbook["Service Snapshot"][1]]


def test_administrator_can_save_configuration():
    client = new_client()
    page = client.get("/admin/configuration")
    assert page.status_code == 200
    response = client.post("/admin/configuration", data={
        "csrf_token": csrf(client), "display_timezone": "America/New_York",
        "date_format": "%Y-%m-%d", "time_format": "%H:%M UTC",
        "log_level": "WARNING", "audit_retention_days": "730", "identity_mode": "local",
    }, follow_redirects=False)
    assert response.status_code == 303
    policy = client.post("/admin/compliance", data={
        "csrf_token": csrf(client), "overdue_days": "120", "exception_max_days": "180", "skipped_images_incomplete": "true",
        "minimum_severity": "Critical", "epss_severity": ["Critical", "High"],
        "epss_rule_threshold": ["0.90", "0.80"], "epss_rule_noncompliant": ["true", "true"],
    }, follow_redirects=False)
    assert policy.status_code == 303
    assert "Overdue Findings" in client.get("/").text


def test_configuration_is_global_and_migrates_legacy_group_preference():
    client = new_client()
    with SessionLocal() as db:
        legacy_group = Group(name="Legacy configuration group")
        db.add(legacy_group)
        db.flush()
        db.add(PortalSetting(
            key=f"group:{legacy_group.id}:display_timezone",
            group_id=legacy_group.id,
            value="America/Los_Angeles",
        ))
        db.add(PortalSetting(
            key=f"group:{legacy_group.id}:log_level",
            group_id=legacy_group.id,
            value="DEBUG",
        ))
        db.commit()

    page = client.get("/admin/configuration")
    assert page.status_code == 200
    assert "Global runtime preferences" in page.text
    assert "Group Scope" not in page.text
    assert "America/Los_Angeles" in page.text
    with SessionLocal() as db:
        global_setting = db.scalar(select(PortalSetting).where(
            PortalSetting.key == "display_timezone", PortalSetting.group_id.is_(None)
        ))
        assert global_setting is not None
        assert global_setting.value == "America/Los_Angeles"
        global_log_level = db.scalar(select(PortalSetting).where(
            PortalSetting.key == "log_level", PortalSetting.group_id.is_(None)
        ))
        assert global_log_level is not None
        assert global_log_level.value == "DEBUG"


def test_service_configuration_does_not_reapply_legacy_global_group_settings():
    with SessionLocal() as db:
        group = Group(name="Scoped group")
        service = Service(service_key="global-config-service", name="Global config service", manual_version="1")
        db.add_all([group, service]); db.flush()
        service.groups.append(group)
        db.add(PortalSetting(key="display_timezone", group_id=None, value="UTC"))
        db.add(PortalSetting(key=f"group:{group.id}:display_timezone", group_id=group.id, value="America/New_York"))
        db.commit(); db.refresh(service)
        assert configuration_for_service(db, service)["display_timezone"] == "UTC"


def test_configuration_lists_all_builtin_os_repositories_and_image_defined_default():
    client = new_client()
    page = client.get("/admin/configuration")
    assert page.status_code == 200
    for os_id in ("ubuntu", "debian", "rhel", "rocky", "almalinux", "centos", "fedora", "alpine"):
        assert os_id in page.text
    assert "Operating system" in page.text
    assert "OS ID" in page.text
    assert "Package manager" in page.text
    assert "Image-defined" in page.text


def test_custom_os_repository_can_be_added_and_builtins_are_protected():
    client = new_client()
    response = client.post("/admin/configuration/repositories", data={
        "csrf_token": csrf(client), "action": "add", "os_id": "test-linux",
        "display_name": "Test Linux", "package_manager": "dnf",
        "mode": "custom", "url": "https://mirror.example.invalid/test-linux",
    }, follow_redirects=False)
    assert response.status_code == 303
    assert "test-linux" in client.get("/admin/configuration").text
    assert "https://mirror.example.invalid/test-linux" in client.get("/admin/configuration").text
    protected = client.post("/admin/configuration/repositories", data={
        "csrf_token": csrf(client), "action": "delete", "remove_os_id": "ubuntu",
    })
    assert protected.status_code == 422


def test_audit_logging_level_is_saved_globally():
    client = new_client()
    response = client.post("/admin/audit-policy", data={
        "csrf_token": csrf(client), "audit_retention_days": "365",
        "audit_tool_integration": "planned", "log_level": "DEBUG", "group_id": "",
    }, follow_redirects=False)
    assert response.status_code == 303
    with SessionLocal() as db:
        setting = db.scalar(select(PortalSetting).where(
            PortalSetting.key == "log_level", PortalSetting.group_id.is_(None)
        ))
        assert setting is not None and setting.value == "DEBUG"


def test_account_can_save_personal_theme():
    client = new_client()
    assert client.get("/account/appearance").status_code == 200
    response = client.post("/account/appearance", data={
        "csrf_token": csrf(client), "theme": "blue",
    }, follow_redirects=False)
    assert response.status_code == 303
    assert 'data-theme="blue"' in client.get("/").text
    assert "Saved successfully." in client.get("/account/appearance?saved=1").text


def test_image_scoped_ingest_only_reconciles_the_scanned_image():
    client = new_client()
    first = payload("multi-image-1", datetime.now(timezone.utc), ["CVE-A"], service_id="multi-image")
    first["findings"].append({"cve": "CVE-B", "severity": "High", "image": "registry.example/image-b:2.3", "package": "openssl", "fixed_version": "9.9", "evidence": {}})
    first["findings"][0]["image"] = "registry.example/image-a:2.3"
    assert client.post("/api/v1/pipeline-results", json=first, headers=pipeline_headers).status_code == 201
    second = payload("image-only-2", datetime.now(timezone.utc), [], service_id="multi-image")
    second["scan_scope"] = "image"
    second["scope_image"] = "registry.example/image-a:2.3"
    second["findings"] = []
    assert client.post("/api/v1/pipeline-results", json=second, headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "multi-image"))
        findings = {finding.cve: finding.active for finding in db.scalars(select(Finding).where(Finding.service_id == service.id)).all()}
        images = {image.image_reference: image.lifecycle_status for image in db.scalars(select(ServiceImage).where(ServiceImage.service_id == service.id)).all()}
    assert findings == {"CVE-A": False, "CVE-B": True}
    assert images["registry.example/image-a:2.3"] == "active"
    assert images["registry.example/image-b:2.3"] == "active"


def test_missing_evidence_source_file_is_preserved_in_service_overview():
    client = new_client()
    body = payload("provenance-1", datetime.now(timezone.utc), [], service_id="provenance-service")
    body["service_overview"] = {"missing_evidence": [{"type": "Image", "item": "registry.example/missing:1", "reason": "repository unavailable", "source_file": "charts/app/values.yaml"}]}
    assert client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers).status_code == 201
    page = client.get("/services/provenance-service?overview=true")
    assert page.status_code == 200
    assert "charts/app/values.yaml" in page.text


def test_helm_ingest_queues_deployment_validation_without_changing_static_result(monkeypatch):
    from app import main as portal_main
    submitted = []
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    monkeypatch.setattr(portal_main, "_submit_validation_run", lambda run_id: submitted.append(run_id))
    response = new_client().post("/api/v1/pipeline-results", json=helm_payload(), headers=pipeline_headers)
    assert response.status_code == 201 and response.json()["accepted"] is True
    with SessionLocal() as db:
        run = db.scalar(select(DeploymentValidationRun))
        assert run.status == "QUEUED" and submitted == [run.id]


def test_validation_persistence_failure_cannot_roll_back_static_ingest(monkeypatch):
    from app import main as portal_main
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    monkeypatch.setattr(portal_main, "_new_validation_run", lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("optional store unavailable")))
    response = new_client().post("/api/v1/pipeline-results", json=helm_payload("static-survives-validation-store"), headers=pipeline_headers)
    assert response.status_code == 201
    assert response.json()["deployment_validation_run_id"] is None
    with SessionLocal() as db:
        execution = db.scalar(select(Execution).where(Execution.execution_key == "static-survives-validation-store"))
        assert execution is not None and execution.complete is True
        assert db.scalar(select(DeploymentValidationRun)) is None


def test_helm_source_limits_reject_before_persistence(monkeypatch):
    monkeypatch.setenv("CATS_INGEST_MAX_SOURCE_BYTES", "32")
    body = helm_payload("oversized-source")
    body["helm_source_files"] = {"Chart.yaml": "x" * 64}
    response = new_client().post("/api/v1/pipeline-results", json=body, headers=pipeline_headers)
    assert response.status_code == 422
    with SessionLocal() as db:
        assert db.scalar(select(Execution)) is None


def test_pipeline_content_length_limit_rejects_early():
    from app import main as portal_main
    response = new_client().post("/api/v1/pipeline-results", content=b"{}", headers={**pipeline_headers, "content-length": str(portal_main.PIPELINE_MAX_REQUEST_BYTES + 1)})
    assert response.status_code == 413


def test_pipeline_size_limit_counts_actual_body_when_header_is_misleading(monkeypatch):
    from app import main as portal_main
    monkeypatch.setattr(portal_main, "PIPELINE_MAX_REQUEST_BYTES", 64)
    body = json.dumps({"padding": "x" * 200}).encode()
    response = new_client().post("/api/v1/pipeline-results", content=body, headers={**pipeline_headers, "content-type": "application/json", "content-length": "1"})
    assert response.status_code == 413


def test_disabled_deployment_validation_is_not_attempted_and_static_scan_remains_visible(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    client = new_client()
    assert client.post("/api/v1/pipeline-results", json=helm_payload("disabled-validation"), headers=pipeline_headers).status_code == 201
    with SessionLocal() as db:
        assert db.scalar(select(DeploymentValidationRun)).status == "NOT_ATTEMPTED"
    overview = client.get("/services/payments-service?overview=true")
    details = client.get("/services/payments-service?validation=true")
    assert overview.status_code == details.status_code == 200
    assert "Deployment Validation" in overview.text and "Not Attempted" in overview.text
    assert "Static Scan:" in overview.text
    assert "authoritative static CATS scan" in details.text
    assert "Validation history" in details.text and "Not Attempted" in details.text


def test_helm_ingest_without_retained_sources_records_not_attempted(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    body = helm_payload("missing-validation-sources")
    body["helm_source_files"] = {}
    response = new_client().post("/api/v1/pipeline-results", json=body, headers=pipeline_headers)
    assert response.status_code == 201
    with SessionLocal() as db:
        run = db.scalar(select(DeploymentValidationRun))
        assert run.status == "NOT_ATTEMPTED"
        assert "did not retain Helm source files" in run.reason


def test_validation_ui_preserves_helm_provenance_after_a_later_non_helm_scan(monkeypatch):
    from app import main as portal_main
    submitted = []
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    monkeypatch.setattr(portal_main, "_submit_validation_run", lambda run_id: submitted.append(run_id))
    client = new_client()
    assert client.post("/api/v1/pipeline-results", json=helm_payload("helm-provenance"), headers=pipeline_headers).status_code == 201
    later = payload("image-evidence-later", datetime.now(timezone.utc), [], service_id="payments-service")
    later["artifact_type"] = "image"
    assert client.post("/api/v1/pipeline-results", json=later, headers=pipeline_headers).status_code == 201
    response = client.post("/services/payments-service/deployment-validations", data={"csrf_token": csrf(client)}, follow_redirects=False)
    assert response.status_code == 303
    with SessionLocal() as db:
        runs = db.scalars(select(DeploymentValidationRun).order_by(DeploymentValidationRun.id)).all()
        assert len(runs) == 2
        assert all(run.execution.execution_key == "helm-provenance" for run in runs)
    overview = client.get("/services/payments-service?overview=true")
    assert "Helm evidence:" in overview.text and "helm-provenance" in overview.text


def test_validation_ui_reports_incomplete_linked_static_scan_and_hides_disabled_rerun(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    body = helm_payload("incomplete-static")
    body["complete"] = False
    client = new_client()
    assert client.post("/api/v1/pipeline-results", json=body, headers=pipeline_headers).status_code == 201
    details = client.get("/services/payments-service?validation=true")
    assert "Static Scan: Incomplete" in details.text
    assert "Deployment Validation is disabled by configuration." in details.text
    assert "Re-run Validation" not in details.text


def test_validation_history_links_open_the_selected_immutable_run(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("history-selection"), headers=pipeline_headers)
    with SessionLocal() as db:
        first = db.scalar(select(DeploymentValidationRun))
        first.reason = "Older run explanation"
        second = DeploymentValidationRun(run_key="DV-NEWER-HISTORY", service_id=first.service_id, execution_id=first.execution_id,
                                         status="COULD_NOT_VALIDATE", phase="COMPLETE", reason="Newer run explanation", artifact_reference="history-selection")
        db.add(second); db.commit(); first_key = first.run_key
    latest = client.get("/services/payments-service?validation=true")
    selected = client.get(f"/services/payments-service?validation=true&validation_run={first_key}")
    assert "Newer run explanation" in latest.text
    assert "Older run explanation" in selected.text
    assert f"validation_run={first_key}" in selected.text


def test_overview_collapses_worker_phase_to_in_progress(monkeypatch):
    from app import main as portal_main
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    monkeypatch.setattr(portal_main, "_submit_validation_run", lambda _run_id: None)
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("queued-overview"), headers=pipeline_headers)
    overview = client.get("/services/payments-service?overview=true")
    assert "In Progress" in overview.text
    assert ">Queued<" not in overview.text


def test_stale_never_started_validation_recovers_as_not_attempted(monkeypatch):
    from app import main as portal_main
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("stale-queued"), headers=pipeline_headers)
    with SessionLocal() as db:
        run = db.scalar(select(DeploymentValidationRun))
        run.status = "QUEUED"; run.started_at = None; run.cluster_name = None
        run.created_at = datetime.now(timezone.utc) - timedelta(days=1)
        db.commit(); run_id = run.id
    portal_main.recover_stale_validation_runs()
    with SessionLocal() as db:
        recovered = db.get(DeploymentValidationRun, run_id)
        assert recovered.status == "NOT_ATTEMPTED"
        assert recovered.cleanup_status == "NOT_REQUIRED"


def test_failed_terminal_cleanup_is_retried_by_stale_recovery(monkeypatch):
    from app import main as portal_main
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("cleanup-retry"), headers=pipeline_headers)
    with SessionLocal() as db:
        run = db.scalar(select(DeploymentValidationRun))
        run.status = "COULD_NOT_VALIDATE"; run.cleanup_status = "FAILED"; run.cluster_name = "cats-validation-cleanup-retry"
        run.created_at = datetime.now(timezone.utc) - timedelta(days=1)
        db.commit(); run_id = run.id
    monkeypatch.setattr(portal_main, "cleanup_stale_clusters", lambda names, config: {"deleted": list(names), "failed": []})
    portal_main.recover_stale_validation_runs()
    with SessionLocal() as db:
        recovered = db.get(DeploymentValidationRun, run_id)
        assert recovered.status == "COULD_NOT_VALIDATE"
        assert recovered.cleanup_status == "COMPLETE"


def test_rerun_creates_historical_validation_row_and_schedules_it(monkeypatch):
    from app import main as portal_main
    submitted = []
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    monkeypatch.setattr(portal_main, "_submit_validation_run", lambda run_id: submitted.append(run_id))
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("rerun-original"), headers=pipeline_headers)
    response = client.post("/services/payments-service/deployment-validations", data={"csrf_token": csrf(client)}, follow_redirects=False)
    assert response.status_code == 303 and "validation=true&validation_run=DV-" in response.headers["location"]
    with SessionLocal() as db:
        runs = db.scalars(select(DeploymentValidationRun).order_by(DeploymentValidationRun.id)).all()
        assert len(runs) == 2 and [run.id for run in runs] == submitted


def test_deployment_validation_json_history_and_detail_apis(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("api-validation"), headers=pipeline_headers)
    with SessionLocal() as db:
        stored = db.scalar(select(DeploymentValidationRun))
        stored.status = "PARTIALLY_VERIFIED"
        stored.reason_category = "EXPECTED_RESOURCE_NOT_OBSERVED"
        stored.classification_reasons = [{
            "code": "EXPECTED_RESOURCE_NOT_OBSERVED", "resource": {"kind": "CronJob", "namespace": "demo", "name": "maintenance"},
            "expected_state": "Observed", "observed_state": "Missing", "explanation": "Expected object was not observed.",
        }]
        stored.capability_preflight = [{"capability": "Configuration dependencies", "required": True, "status": "AVAILABLE"}]
        stored.diagnostics = {"classification_summary": {"expected_resources": 2, "observed_expected": 1, "expected_only": 1, "runtime_generated": 1, "observed_only": 0, "failed": 0}}
        db.commit()
    history = client.get("/api/v1/services/payments-service/deployment-validations")
    assert history.status_code == 200 and len(history.json()["runs"]) == 1
    run = history.json()["runs"][0]
    result = client.get(f"/api/v1/services/payments-service/deployment-validations/{run['run_key']}")
    assert result.status_code == 200 and result.json()["status"] == "PARTIALLY_VERIFIED"
    # The detail endpoint exposes persisted reasons and grouped capability data.
    assert result.json()["classification_reasons"][0]["code"] == "EXPECTED_RESOURCE_NOT_OBSERVED"
    assert result.json()["classification_summary"]["expected_only"] == 1
    assert result.json()["capability_assessment"][0]["capability"] == "Configuration dependencies"
    assert result.json()["checks"]["helm_template"] == "NOT_ATTEMPTED"
    assert result.json()["terminal"] is True and result.json()["cleanup_terminal"] is True


def test_rerun_json_response_returns_authoritative_new_run(monkeypatch):
    from app import main as portal_main
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "true")
    monkeypatch.setattr(portal_main, "_submit_validation_run", lambda _run_id: None)
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("json-rerun"), headers=pipeline_headers)
    response = client.post("/services/payments-service/deployment-validations", data={"csrf_token": csrf(client)}, headers={"Accept": "application/json"})
    assert response.status_code == 200
    payload = response.json()
    assert payload["run_id"].startswith("DV-")
    assert payload["run"]["run_key"] == payload["run_id"]
    assert payload["run"]["status"] == "QUEUED"
    assert payload["run"]["terminal"] is False


def test_deployment_validation_requires_permission_and_csrf(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    admin = new_client()
    admin.post("/api/v1/pipeline-results", json=helm_payload("permission-validation"), headers=pipeline_headers)
    assert admin.post("/services/payments-service/deployment-validations", data={}, follow_redirects=False).status_code == 422
    with SessionLocal() as db:
        service_id = db.scalar(select(Service.id).where(Service.service_key == "payments-service"))
    add_user("validation-viewer", "Assessor", service_id=service_id)
    viewer = new_client("validation-viewer")
    assert viewer.get("/api/v1/services/payments-service/deployment-validations").status_code == 200
    assert viewer.post("/services/payments-service/deployment-validations", data={"csrf_token": csrf(viewer)}, follow_redirects=False).status_code == 403


def test_service_deletion_removes_deployment_validation_rows(monkeypatch):
    monkeypatch.setenv("CATS_DEPLOYMENT_VALIDATION_ENABLED", "false")
    client = new_client()
    client.post("/api/v1/pipeline-results", json=helm_payload("delete-validation"), headers=pipeline_headers)
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "payments-service"))
        db.add(ServiceArchiveEvent(service_id=service.id, action="archive", reason="test", performed_by="admin"))
        db.commit()
    monkeypatch.setenv("ALLOW_SERVICE_DELETE", "true")
    response = client.post("/services/payments-service/delete", data={"confirmation": "payments-service", "reason": "Remove test service", "csrf_token": csrf(client)}, follow_redirects=False)
    assert response.status_code == 303
    with SessionLocal() as db:
        assert db.scalar(select(DeploymentValidationRun)) is None
