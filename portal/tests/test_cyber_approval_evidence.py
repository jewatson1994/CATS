from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
import json
import pytest
from sqlalchemy import select
from app import cyber_approval_evidence as evidence
from app.models import Service, Execution
from app.execution_summaries import payload_digest
from test_portal import setup_function, new_client, payload, pipeline_headers, SessionLocal

D = 'sha256:' + 'a' * 64


def seeded():
    client = new_client()
    body = payload('approval-source', datetime.now(timezone.utc), ['CVE-1'])
    body['helm_source_files'] = {'app/Chart.yaml': 'apiVersion: v2\nname: app\nversion: 1.0.0\n',
        'app/values.yaml': 'replicaCount: 2\n', 'app/templates/deployment.yaml': 'retained template bytes'}
    for finding in body['findings']:
        finding['image'] = 'registry.example/app:1'
    body['sbom_components'] = [{'name': 'pkg', 'version': '1', 'image': 'registry.example/app:1'}]
    body['helm_values_files'] = ['app/values.yaml']
    body['service_overview'] = {'images': [{'image': 'registry.example/app:1', 'digest': D}],
        'charts': [{'name': 'app', 'version': '1.0.0'}],
        'rendered_resources': [{'apiVersion': 'v1', 'kind': 'Pod', 'metadata': {'name': 'app'},
            'spec': {'containers': [{'name': 'app', 'image': 'registry.example/app:1'}]}}]}
    assert client.post('/api/v1/pipeline-results', json=body, headers=pipeline_headers).status_code == 201


def fake_export(output, service, db, auth):
    with ZipFile(output, 'w') as archive:
        archive.writestr('findings.xlsx', b'workbook fixture')
        archive.writestr('mitigations.xlsx', b'mitigation fixture')
    return {'entries': [{'name': 'findings.xlsx', 'status': 'included'}]}


def test_ready_uses_exact_scan_and_preserves_chart_values_and_inventory(tmp_path, monkeypatch):
    seeded()
    monkeypatch.setattr('app.service_exports.write_all', fake_export)
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        ready = evidence.readiness(db, service)
        assert ready['ready'], ready['blockers']
        assert not any('Image identity was not retained' in note for note in ready['review_notes'])
        snapshot = evidence.freeze(db, service, SimpleNamespace(), ready, tmp_path)
        assert snapshot['artifacts'][0]['reference'] == 'registry.example/app@' + D
        assert snapshot['evidence']['digest'] == evidence.file_digest(snapshot['evidence']['path'])
        with ZipFile(snapshot['evidence']['path']) as archive:
            assert archive.read('findings.xlsx') == b'workbook fixture'
            assert archive.read('chart-sources/app/values.yaml') == b'replicaCount: 2\n'
            assert json.loads(archive.read('scan-values-files.json')) == ['app/values.yaml']
        import tarfile
        with tarfile.open(snapshot['artifacts'][1]['path']) as archive:
            assert archive.extractfile('app/values.yaml').read() == b'replicaCount: 2\n'
        assert snapshot['findings']['vulnerabilities'][0]['cve'] == 'CVE-1'


@pytest.mark.parametrize('mutation,expected', [
    ('digest', 'digest'), ('skip', 'Missing'),
    ('dependency', 'dependency'), ('chart', 'matched'),
])
def test_missing_or_mismatched_evidence_blocks(mutation, expected):
    seeded()
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        execution = db.scalar(select(Execution))
        body = json.loads(json.dumps(execution.raw_payload))
        if mutation == 'skip': body['skipped_charts'] = [{'chart': 'lost', 'reason': 'download failed'}]
        if mutation == 'dependency': body['helm_source_files']['app/Chart.yaml'] += 'dependencies:\n- name: lost\n  version: 1.0.0\n  repository: https://example.test\n'
        if mutation == 'chart': body['service_overview']['charts'].append({'name': 'missing', 'version': '2'})
        execution.raw_payload = body
        db.flush()
        if mutation == 'digest':
            from sqlalchemy import text
            db.execute(text("UPDATE executions SET payload_digest = :digest WHERE id = :id"), {'digest': '0' * 64, 'id': execution.id})
        db.flush()
        result = evidence.readiness(db, service)
        assert not result['ready']
        assert expected.lower() in ' '.join(result['blockers']).lower(), result['blockers']


def test_newer_incomplete_scan_does_not_fall_back_to_older_complete():
    seeded()
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        first = db.scalar(select(Execution))
        db.add(Execution(execution_key='incomplete', service_id=service.id, service_version_id=service.current_version_id,
            scanned_at=datetime(2099, 1, 1, tzinfo=timezone.utc), complete=False, scan_scope='service', raw_payload=dict(first.raw_payload)))
        db.commit()
        result = evidence.readiness(db, service)
        assert not result['ready'] and result['execution'].execution_key == 'incomplete'


def test_unavailable_export_all_is_not_silently_approved(tmp_path, monkeypatch):
    seeded()
    monkeypatch.setattr('app.service_exports.write_all', lambda *args: {'entries': [{'name': 'sbom', 'status': 'unavailable'}]})
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        with pytest.raises(ValueError, match='Export All evidence unavailable'):
            evidence.freeze(db, service, SimpleNamespace(), evidence.readiness(db, service), tmp_path)


def test_candidate_requires_matching_retained_verified_evidence():
    seeded()
    with SessionLocal() as db:
        result = evidence.readiness(db, db.scalar(select(Service)), 'cats_remediated', 999999)
        assert not result['ready']
        assert any('Candidate must belong' in item for item in result['blockers'])


@pytest.mark.parametrize('corrupt', [False, True])
@pytest.mark.parametrize('unverified', [False, True])
def test_candidate_freeze_binds_verified_sources_and_labels_baseline(tmp_path, monkeypatch, corrupt, unverified):
    from app import main
    from app.models import RemediationExecution
    import hashlib
    import yaml
    seeded()
    monkeypatch.setattr(main, 'REMEDIATION_JOB_ROOT', tmp_path)
    monkeypatch.setattr('app.service_exports.write_all', fake_export)
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        execution = db.scalar(select(Execution))
        source = json.loads(json.dumps(execution.raw_payload))
        files = {'candidate/' + name: content.encode() for name, content in source['helm_source_files'].items()}
        files['scans/rendered-after.yaml'] = yaml.safe_dump_all(source['service_overview']['rendered_resources']).encode()
        if not unverified:
            files['scans/trivy-config-after.json'] = b'{"Results": []}'
        manifest = {'source_execution_id': execution.id, 'source_version_id': service.current_version_id,
            'values_files': ['app/values.yaml'], 'images': [],
            'deployment_manifest': {'deployment': {'chartPath': 'candidate/app'},
            'files': {name: {'bytes': len(content), 'sha256': 'sha256:' + hashlib.sha256(content).hexdigest()} for name, content in files.items()}}}
        if corrupt:
            manifest['deployment_manifest']['files']['candidate/app/values.yaml']['sha256'] = 'sha256:' + '0' * 64
        candidate_dir = tmp_path / 'candidatejob'
        candidate_dir.mkdir()
        candidate = candidate_dir / 'candidate.zip'
        with ZipFile(candidate, 'w') as archive:
            for name, content in files.items(): archive.writestr(name, content)
            archive.writestr('manifest.json', json.dumps(manifest))
        record = RemediationExecution(job_key='candidatejob', service_id=service.id, requested_by_id=1,
            source_execution_id=execution.id, source_version_id=service.current_version_id,
            status='validated', remediation_status='complete', verification_status='failed' if unverified else 'verified',
            artifact_path=str(candidate), artifact_digest=evidence.file_digest(candidate),
            scan_results={'configuration': {'status': 'SKIP' if unverified else 'PASS'}, 'vulnerabilities': {'status': 'FAIL' if unverified else 'PASS'}},
            after_snapshot={'findings': [{'cve': 'CVE-CANDIDATE', 'status': 'remaining'}]})
        db.add(record); db.commit()
        ready = evidence.readiness(db, service, 'cats_remediated', record.id)
        if corrupt:
            assert not ready['ready'] and any('digest mismatch' in item for item in ready['blockers'])
            return
        assert ready['ready'], ready['blockers']
        snapshot = evidence.freeze(db, service, SimpleNamespace(), ready, tmp_path / 'frozen')
        assert snapshot['evidence_scope'] == 'remediation-candidate'
        assert snapshot['overview']['source'] == 'CATS Remediated candidate'
        assert snapshot['findings']['summary'] == record.after_snapshot
        assert snapshot['findings']['configuration'] == (None if unverified else {'Results': []})
        if unverified:
            assert any('verification: failed' in note for note in snapshot['review_notes'])
        with ZipFile(snapshot['evidence']['path']) as archive:
            assert 'current-service-baseline/all.zip' in archive.namelist()
            assert 'candidate/evidence.xlsx' in archive.namelist()
            assert archive.read('candidate/candidate.zip') == candidate.read_bytes()


def test_freeze_detects_new_scan_arriving_during_export(tmp_path, monkeypatch):
    seeded()
    def changing_export(output, service, db, auth):
        result = fake_export(output, service, db, auth)
        first = db.scalar(select(Execution))
        db.add(Execution(execution_key='arrived-during-freeze', service_id=service.id, service_version_id=service.current_version_id,
            scanned_at=datetime(2099, 1, 1, tzinfo=timezone.utc), complete=True, scan_scope='service', raw_payload=dict(first.raw_payload)))
        db.flush()
        return result
    monkeypatch.setattr('app.service_exports.write_all', changing_export)
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        with pytest.raises(ValueError, match='changed while freezing'):
            evidence.freeze(db, service, SimpleNamespace(), evidence.readiness(db, service), tmp_path)


def test_full_service_image_only_freezes_actual_export_all(tmp_path):
    seeded()
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        execution = db.scalar(select(Execution))
        body = json.loads(json.dumps(execution.raw_payload))
        body.pop('helm_source_files')
        body.pop('helm_values_files')
        body['service_overview'].pop('charts')
        body['service_overview'].pop('rendered_resources')
        execution.raw_payload = body
        db.flush()
        ready = evidence.readiness(db, service)
        assert ready['ready'], ready['blockers']
        assert [item['kind'] for item in ready['artifacts']] == ['image']
        auth = SimpleNamespace(has=lambda permission, service_id: permission == 'service.export')
        snapshot = evidence.freeze(db, service, auth, ready, tmp_path)
        with ZipFile(snapshot['evidence']['path']) as archive:
            assert {'ppsm.xlsx', 'poam.xlsx', 'asset_list.xlsx', 'mitigations.xlsx', 'findings.xlsx',
                    'diagrams.zip', 'sbom-components.json', 'frozen-snapshot.json'}.issubset(archive.namelist())
            from openpyxl import load_workbook
            with archive.open('findings.xlsx') as stream:
                book = load_workbook(stream)
                assert book.active.max_row >= 2
                book.close()


def test_chart_root_under_charts_directory_is_retained():
    charts = evidence._charts({'charts/package-123/app/Chart.yaml': 'apiVersion: v2\nname: app\nversion: 1.0.0\n',
                              'charts/package-123/app/values.yaml': 'replicas: 1\n'})
    assert len(charts) == 1
    assert charts[0]['source_root'] == 'charts/package-123/app'


def test_candidate_without_selection_returns_blockers():
    seeded()
    with SessionLocal() as db:
        result = evidence.readiness(db, db.scalar(select(Service)), 'cats_remediated')
        assert not result['ready']
        assert any('Candidate must belong' in item for item in result['blockers'])


def test_large_vendored_chart_inventory_is_complete_and_packaged(tmp_path, monkeypatch):
    import tarfile
    seeded()
    monkeypatch.setattr('app.service_exports.write_all', fake_export)
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        execution = db.scalar(select(Execution))
        body = json.loads(json.dumps(execution.raw_payload))
        declarations = []
        for index in range(90):
            name = f'component-{index}'
            declarations.append(f'- name: {name}\n  version: 1.0.0\n')
            body['helm_source_files'][f'app/charts/{name}/Chart.yaml'] = f'apiVersion: v2\nname: {name}\nversion: 1.0.0\n'
            body['service_overview']['charts'].append({'name': name, 'version': '1.0.0'})
        body['helm_source_files']['app/Chart.yaml'] += 'dependencies:\n' + ''.join(declarations)
        execution.raw_payload = body
        db.flush()
        ready = evidence.readiness(db, service)
        assert ready['ready'], ready['blockers']
        snapshot = evidence.freeze(db, service, SimpleNamespace(), ready, tmp_path)
        charts = [item for item in snapshot['artifacts'] if item['kind'] == 'chart']
        assert len(charts) == 1
        assert len(charts[0]['dependencies']) == 90
        with tarfile.open(charts[0]['path']) as archive:
            assert len([name for name in archive.getnames() if name.endswith('/Chart.yaml')]) == 91


@pytest.mark.parametrize('retained_digest', [D, 'registry.example/app@' + D])
def test_repeated_finding_image_digests_produce_one_artifact(retained_digest):
    seeded()
    with SessionLocal() as db:
        execution = db.scalar(select(Execution))
        body = json.loads(json.dumps(execution.raw_payload))
        body['service_overview']['images'][0]['digest'] = retained_digest
        body['findings'] = [dict(body['findings'][0], image='registry.example/app:1',
            image_digest=retained_digest, cve=f'CVE-{i}') for i in range(250)]
        execution.raw_payload = body
        db.flush()
        result = evidence.readiness(db, db.scalar(select(Service)))
        assert result['ready'], result['blockers']
        images = [item for item in result['artifacts'] if item['kind'] == 'image']
        assert len(images) == 1
        assert images[0]['reference'] == 'registry.example/app@' + D


@pytest.mark.parametrize('digest,message', [
    ('bad-digest', 'Invalid retained image digest'),
    ('sha256:' + 'b' * 64, 'Conflicting retained image digests'),
    ('other.example/app@' + D, 'Retained image digest repository does not match'),
])
def test_repeated_image_errors_are_unique_and_still_block_approval(digest, message):
    seeded()
    with SessionLocal() as db:
        execution = db.scalar(select(Execution))
        body = json.loads(json.dumps(execution.raw_payload))
        body['findings'] = [dict(body['findings'][0], image='registry.example/app:1',
            image_digest=digest, cve=f'CVE-{i}') for i in range(250)]
        # A separate affected image must retain its own diagnostic.
        body['findings'].append(dict(body['findings'][0], image='registry.example/other:1',
            image_digest='bad-digest'))
        execution.raw_payload = body
        db.flush()
        result = evidence.readiness(db, db.scalar(select(Service)))
        assert not result['ready']
        assert result['blockers'].count(message + ': registry.example/app:1') == 1
        assert 'Invalid retained image digest: registry.example/other:1' in result['blockers']
        assert len(result['blockers']) == len(set(result['blockers']))


@pytest.mark.parametrize('reference,retained', [
    ('tsuru/nginx-operator:0.17.0', 'docker.io/tsuru/nginx-operator'),
    ('nginx:1', 'docker.io/library/nginx'),
    ('docker.io/nginx:1', 'index.docker.io/library/nginx'),
    ('localhost:5000/team/app:1', 'localhost:5000/team/app'),
])
def test_repository_digest_accepts_same_image_and_docker_hub_aliases(reference, retained):
    pinned, digest = evidence._identity(reference, retained + '@' + D)
    assert digest == D
    assert pinned.endswith('@' + D)


@pytest.mark.parametrize('retained', [
    'registry.example/app@sha256:short',
    'registry.example/app@sha256:' + 'b' * 64,
    'registry.example/other@' + D,
    'other.example/app@' + D,
    'registry.example/app@@' + D,
])
def test_repository_digest_rejects_invalid_or_inconsistent_identity(retained):
    with pytest.raises(ValueError):
        evidence._identity('registry.example/app@' + D, retained)


@pytest.mark.parametrize('legacy_digest', [False, True])
def test_tag_only_vulnerable_image_can_be_submitted_without_helm_or_config(tmp_path, monkeypatch, legacy_digest):
    seeded()
    monkeypatch.setattr('app.service_exports.write_all', fake_export)
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        execution = db.scalar(select(Execution))
        body = json.loads(json.dumps(execution.raw_payload))
        body.pop('helm_source_files')
        body.pop('helm_values_files')
        reference = 'tsuru/nginx-operator:0.17.0'
        body['service_overview'] = {'images': [{'image': reference}]}
        body['findings'] = [dict(body['findings'][0], image=reference,
            image_digest=reference if legacy_digest else None, cve=f'CVE-{i}') for i in range(250)]
        execution.raw_payload = body
        db.flush()
        ready = evidence.readiness(db, service)
        assert ready['ready'], ready['blockers']
        assert len(ready['artifacts']) == 1
        assert ready['artifacts'][0]['resolve_at_publication']
        assert sum('Image identity' in note for note in ready['review_notes']) == 1
        assert any('No Helm charts' in note for note in ready['review_notes'])
        assert any('Configuration coverage' in note for note in ready['review_notes'])
        snapshot = evidence.freeze(db, service, SimpleNamespace(), ready, tmp_path)
        assert len(snapshot['findings']['vulnerabilities']) == 250
        assert snapshot['review_notes'] == ready['review_notes']


@pytest.mark.parametrize('missing', [None, 'sbom/image1.json', 'scans/image1-grype-after.json', 'images/image1.tar'])
def test_image_only_candidate_retains_archive_and_findings_without_runtime_pass(tmp_path, monkeypatch, missing):
    from app import main
    from app.models import RemediationExecution
    import hashlib
    seeded()
    monkeypatch.setattr(main, 'REMEDIATION_JOB_ROOT', tmp_path)
    monkeypatch.setattr('app.service_exports.write_all', fake_export)
    with SessionLocal() as db:
        service = db.scalar(select(Service))
        execution = db.scalar(select(Execution))
        archive_bytes = b'retained candidate image'
        files = {'images/image1.tar': archive_bytes, 'sbom/image1.json': b'{"artifacts": []}',
            'scans/image1-grype-after.json': b'{"matches": [{"vulnerability": {"id": "CVE-REMAINING"}}]}'}
        manifest = {'source_execution_id': execution.id, 'source_version_id': service.current_version_id,
            'images': [{'original': 'registry.example/app:1', 'remediated': 'local/candidate:patched',
                'archive_path': 'images/image1.tar', 'artifact_sha256': hashlib.sha256(archive_bytes).hexdigest(),
                'evidence': {'grype': 'complete', 'sbom': 'complete'}}]}
        folder = tmp_path / 'imagecandidate'
        folder.mkdir()
        candidate = folder / 'candidate.zip'
        with ZipFile(candidate, 'w') as archive:
            archive.writestr('manifest.json', json.dumps(manifest))
            for name, content in files.items():
                if name != missing:
                    archive.writestr(name, content)
        record = RemediationExecution(job_key='imagecandidate', service_id=service.id, requested_by_id=1,
            source_execution_id=execution.id, source_version_id=service.current_version_id,
            status='failed', remediation_status='complete', verification_status='failed',
            artifact_path=str(candidate), artifact_digest=evidence.file_digest(candidate),
            scan_results={'images': {'image1': {}}, 'vulnerabilities': {'status': 'FAIL'}},
            after_snapshot={'findings': [{'cve': 'CVE-REMAINING'}]})
        db.add(record); db.commit()
        ready = evidence.readiness(db, service, 'cats_remediated', record.id)
        if missing:
            assert not ready['ready']
            return
        assert ready['ready'], ready['blockers']
        snapshot = evidence.freeze(db, service, SimpleNamespace(), ready, tmp_path / 'frozen')
        item = next(item for item in snapshot['artifacts'] if item.get('format') == 'docker-archive')
        assert Path(item['path']).read_bytes() == archive_bytes
        assert snapshot['findings']['vulnerability_scans']['image1']['matches'][0]['vulnerability']['id'] == 'CVE-REMAINING'
        assert any('No Helm charts' in note for note in snapshot['review_notes'])


def test_absent_rendered_output_is_review_note_when_chart_sources_are_retained():
    seeded()
    with SessionLocal() as db:
        execution = db.scalar(select(Execution))
        body = json.loads(json.dumps(execution.raw_payload))
        body['service_overview']['rendered_resources'] = []
        execution.raw_payload = body
        db.commit()
        ready = evidence.readiness(db, db.scalar(select(Service)))
        assert ready['ready'], ready['blockers']
        assert 'Rendered deployment evidence is unavailable.' in ready['review_notes']
        assert any(item['kind'] == 'chart' for item in ready['artifacts'])
