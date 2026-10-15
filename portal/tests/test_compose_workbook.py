"""Standalone Compose inventory workbook regression coverage."""
import json

from openpyxl import load_workbook

from app.main import _write_public_scan_workbook


def test_compose_workbook_preserves_coverage_and_escapes_formulas(tmp_path):
    model = {
        "complete": False,
        "images": ["nginx:stable"],
        "services": [{"name": "=web", "image": "+image", "status": "-blocked", "image_scan": "incomplete",
                      "source_files": ["@compose.yaml"], "configuration": {"environment": {"TOKEN": "[REDACTED]"}}}],
        "coverage": {"services": 1, "unique_images": 1, "scanned_images": 0, "preview": False},
        "issues": [{"code": "=issue", "message": "+message", "service": "-web", "source": "@compose.yaml", "blocking": True}],
    }
    (tmp_path / "portal-result.json").write_text(json.dumps({"service_overview": {"compose": model}}), encoding="utf-8")
    target = tmp_path / "results.xlsx"
    _write_public_scan_workbook(tmp_path, target)
    workbook = load_workbook(target)
    services = workbook["Compose Services"]
    assert [cell.value for cell in services[2]][:5] == ["'=web", "'+image", "'-blocked", "incomplete", "'@compose.yaml"]
    assert json.loads(services["F2"].value) == model["services"][0]["configuration"]
    coverage = dict(workbook["Compose Coverage"].iter_rows(min_row=2, values_only=True))
    assert coverage["scanned_images"] == 0
    assert coverage["Discovery complete"] is False
    assert coverage["preview"] is False
    issues = workbook["Compose Issues"]
    assert list(issues.values)[1] == ("'=issue", "'+message", "'-web", "'@compose.yaml", True)
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
