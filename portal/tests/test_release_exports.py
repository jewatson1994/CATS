"""Release documentation is the permission-scoped Export All archive."""
from datetime import datetime, timezone
from io import BytesIO
import json
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.models import Execution, PortalSetting, Service, ServiceImage, ServiceVersion
from app.purpose_exports import default_template
from app.release_exports import EvidenceUnavailable, write_documentation, write_export_archive


@pytest.fixture
def source():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        service = Service(service_key="sample", name="New mutable service name")
        db.add(service)
        db.flush()
        version = ServiceVersion(service_id=service.id, version="1.0")
        db.add(version)
        db.flush()
        execution = Execution(service_id=service.id, service_version_id=version.id,
            execution_key="old-scan", scanned_at=datetime(2025, 1, 1, tzinfo=timezone.utc), complete=True,
            raw_payload={"service": {"id": "sample", "name": "Original name", "version": "1.0"},
                "findings": [{"cve": "CVE-OLD", "severity": "High", "image": "example/old:1",
                              "installed_version": "1", "fixed_version": "2", "package": "libold"}],
                "policy_findings": [{"finding": "POLICY-OLD", "severity": "Low",
                                     "title": "Original policy", "remediation": "Set read-only filesystem"}],
                "sbom_components": [{"name": "old-component", "version": "1"}],
                "helm_source_files": {"secret.yaml": "VERY_SECRET_SOURCE"},
                "service_overview": {"ports": [{"port": 443, "protocol": "TCP", "service": "HTTPS"}]},
                "access_token": "VERY_SECRET_TOKEN", "description": "Bearer VERY_SECRET_BEARER"})
        db.add(execution)
        db.flush()
        newer = ServiceVersion(service_id=service.id, version="2.0")
        db.add(newer)
        db.flush()
        db.add(Execution(service_id=service.id, service_version_id=newer.id, execution_key="new-scan",
            scanned_at=datetime(2026, 1, 1, tzinfo=timezone.utc), complete=True,
            raw_payload={"findings": [{"cve": "CVE-NEW", "severity": "Critical"}],
                         "sbom_components": [{"name": "new-component"}]}))
        db.add(ServiceImage(service_id=service.id, image_reference="example/new-global:2"))
        columns = default_template("poam")
        columns[0]["heading"] = "Configured Service Heading"
        db.add(PortalSetting(key="purpose_export_template:poam", value=json.dumps(columns)))
        db.flush()
        record = SimpleNamespace(id=7, job_key="remediation-7", service_id=service.id,
            source_execution_id=execution.id, source_version_id=version.id,
            validation_results={"deployment": {"run_key": "candidate-7", "status": "PASSED",
                "logs": "VERY_SECRET_LOG", "script": "VERY_SECRET_SCRIPT"}})
        yield db, record
    engine.dispose()


def _auth(allowed=True):
    return SimpleNamespace(has=lambda permission, service_id: allowed and permission == "service.export")


def test_documentation_uses_all_export_data_and_configured_templates(source, tmp_path):
    db, record = source
    path = tmp_path / "docs.zip"
    manifest = write_documentation(db, record, _auth(), path)
    with ZipFile(path) as archive:
        assert set(archive.namelist()) == {"ppsm.xlsx", "poam.xlsx", "asset_list.xlsx", "findings.xlsx",
            "mitigations.xlsx", "sbom-components.json", "diagrams.zip"}
        assert json.loads(archive.read("sbom-components.json"))["components"] == [{"name": "new-component"}]
        all_cells = []
        for name in ("ppsm.xlsx", "poam.xlsx", "asset_list.xlsx", "findings.xlsx", "mitigations.xlsx"):
            workbook = load_workbook(BytesIO(archive.read(name)))
            cells = [cell.value for row in workbook.active for cell in row]
            all_cells.extend(cells)
            if name == "poam.xlsx":
                assert cells[0] == "Configured Service Heading"
        assert "example/new-global:2" in all_cells
    assert manifest["source_execution_id"] == record.source_execution_id
    assert manifest["evidence_scope"] == "current-service-at-export"


def test_rejects_cross_service_or_version_scope_before_writing(source, tmp_path):
    db, record = source
    record.source_version_id += 1
    with pytest.raises(ValueError, match="consistent source"):
        write_documentation(db, record, _auth(), tmp_path / "docs.zip")
    assert not (tmp_path / "docs.zip").exists()


def test_missing_source_does_not_fall_back_to_latest(source, tmp_path):
    db, record = source
    record.source_execution_id = 99999
    with pytest.raises(ValueError, match="consistent source"):
        write_documentation(db, record, _auth(), tmp_path / "docs.zip")


def test_export_permission_is_required(source, tmp_path):
    db, record = source
    with pytest.raises(PermissionError):
        write_documentation(db, record, _auth(False), tmp_path / "docs.zip")


def test_missing_evidence_is_explicit(source, tmp_path):
    db, record = source
    for execution in db.query(Execution).all():
        execution.raw_payload = {}
    manifest = write_documentation(db, record, _auth(), tmp_path / "docs.zip")
    unavailable = {entry["name"] for entry in manifest["entries"] if entry["status"] == "unavailable"}
    assert unavailable == {"diagrams.zip", "sbom-components.json"}


def test_byte_limit_preserves_existing_file_and_cleans_temporary(source, tmp_path, monkeypatch):
    db, record = source
    path = tmp_path / "docs.zip"
    path.write_bytes(b"existing file")
    monkeypatch.setenv("CATS_RELEASE_DOC_MAX_BYTES", "4096")
    with pytest.raises(ValueError, match="byte limit"):
        write_documentation(db, record, _auth(), path)
    assert path.read_bytes() == b"existing file"
    assert list(tmp_path.iterdir()) == [path]


def test_common_export_writer_preserves_unavailable_contract():
    output = BytesIO()
    def absent():
        raise EvidenceUnavailable("No source evidence")
    manifest = write_export_archive(output, workbook_builders={}, evidence_builders={"data.json": absent},
        supplemental_entries={"extra.json": {"value": 1}})
    with ZipFile(output) as archive:
        assert archive.read("data.json.unavailable.txt") == b"No source evidence\n"
        assert json.loads(archive.read("extra.json")) == {"value": 1}
    assert manifest["entries"][0]["status"] == "unavailable"


def test_formula_cells_are_escaped(source, tmp_path):
    db, record = source
    from app.models import PoamEntry, User
    user = User(username="planner", display_name="Planner")
    db.add(user)
    db.flush()
    db.add(PoamEntry(service_id=record.service_id, item_type="manual", title="=DANGEROUS()",
                     description="Manual finding", remediation="=DANGEROUS()", status="open",
                     created_by_id=user.id))
    db.flush()
    path = tmp_path / "docs.zip"
    write_documentation(db, record, _auth(), path)
    with ZipFile(path) as archive:
        for name in ("poam.xlsx", "findings.xlsx", "mitigations.xlsx"):
            workbook = load_workbook(BytesIO(archive.read(name)))
            assert all(cell.data_type != "f" for row in workbook.active for cell in row)
