from io import BytesIO

import pytest

from openpyxl import Workbook

from app.exchange import builtins, missing_fields, parse_workbook, workbook
from app.exchange_validation import validate_values


@pytest.mark.parametrize("field,value", [
    ("asset.ip", "999.1.2.3"), ("asset.ip", 123), ("asset.public_ip", "10.0.0.1/24"),
    ("asset.fqdn", "not a name"), ("network.fqdn", "-bad.example.org"),
    ("asset.public_url", "javascript:alert(1)"), ("asset.public_url", "https://user:secret@example.org"),
    ("asset.public_url", "https://example.org:99999"), ("asset.virtual", "maybe"),
    ("asset.public_facing", 2), ("finding.severity", "extreme"),
    ("asset.critical_information", "maybe"),
])
def test_invalid_semantic_fields(field, value):
    assert validate_values({field: value})


def test_valid_semantic_fields_and_normalization():
    values = {"asset.ip": "2001:db8::1", "asset.public_ip": "192.0.2.1", "asset.fqdn": "host.example.org.", "network.fqdn": "api.example.org", "asset.public_url": "https://example.org/path\nhttp://[2001:db8::1]:8080", "asset.virtual": "yes", "asset.public_facing": "false", "finding.severity": "high", "finding.raw_severity": "MODERATE", "finding.residual_risk": "low"}
    assert validate_values(values) == []
    assert values["asset.virtual"] is True and values["asset.public_facing"] is False
    assert values["finding.severity"] == "High" and values["finding.raw_severity"] == "Medium"


def test_critical_information_boolean():
    values = {"asset.critical_information": "YES"}
    assert validate_values(values) == []
    assert values["asset.critical_information"] is True


def test_semantics_reach_workbook_preview():
    output = BytesIO()
    definition = builtins()["assets"]
    workbook(definition, {}, [{"asset.name": "Example", "asset.ip": "bad-ip", "asset.virtual": "Yes"}]).save(output)
    preview = parse_workbook(output.getvalue(), definition)
    assert any("asset.ip" in error for error in preview["errors"])
    assert preview["rows"][0]["values"]["asset.virtual"] is True


def test_import_only_fields_do_not_block_or_appear_in_export():
    definition = builtins()["ppsm"]
    definition["metadata"] = [{"field": "system.owner", "label": "Import-only owner", "required": True, "direction": "import"}]
    definition["columns"].append({"field": "system.owner", "label": "Import-only column", "required": True, "direction": "import"})
    rows = [{"network.port": 443, "network.protocol": "TCP"}]
    assert missing_fields(definition, {}, rows) == []
    cells = [cell.value for row in workbook(definition, {}, rows).active for cell in row]
    assert "Import-only owner *" not in cells and "Import-only column" not in cells


def test_configurable_workbook_byte_limit(monkeypatch):
    monkeypatch.setenv("CATS_WORKBOOK_MAX_BYTES", "4")
    with pytest.raises(ValueError, match="configured 4-byte"):
        parse_workbook(b"12345", builtins()["ppsm"])


def test_external_workbook_detects_sheet_header_aliases_and_extra_columns():
    definition = builtins()["ppsm"]
    next(item for item in definition["columns"] if item["field"] == "network.port")["aliases"] = ["Destination Port"]
    book = Workbook()
    book.active.title = "Instructions"
    sheet = book.create_sheet("Organization inventory")
    sheet.append(["Network approval schedule"])
    sheet.append(["Protocol", "Unrelated notes", "Destination Port", "Data Service"])
    sheet.append(["TCP", "ignored", 443, "HTTPS"])
    output = BytesIO()
    book.save(output)
    result = parse_workbook(output.getvalue(), definition)
    assert result["worksheet"] == "Organization inventory"
    assert result["header_row"] == 2
    assert result["ignored"] == ["Unrelated notes"]
    assert result["rows"][0]["values"]["network.port"] == 443
    assert result["errors"] == []


def test_duplicate_mapped_columns_are_rejected():
    book = Workbook()
    sheet = book.active
    sheet.append(["Port", "network.port", "Protocol"])
    sheet.append([443, 443, "TCP"])
    output = BytesIO()
    book.save(output)
    with pytest.raises(ValueError, match="same field"):
        parse_workbook(output.getvalue(), builtins()["ppsm"])
