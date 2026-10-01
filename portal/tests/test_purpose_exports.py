"""Purpose-built downloads use server templates without changing the complete workbook."""
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AuditEvent, Group, PortalSetting, Service, ServiceArtifact, ServiceImage
from app.purpose_exports import CATALOG, default_template, import_heading_aliases, template_policy, validate_template
from test_portal import SessionLocal, add_user, csrf, new_client, setup_function


def _service():
    with SessionLocal() as db:
        service = Service(service_key="export-test", name="Export Test")
        db.add(service)
        db.flush()
        db.add(ServiceImage(service_id=service.id, image_reference="docker.io/library/nginx:1.27",
                            image_digest="sha256:abc", scan_status="complete"))
        db.commit()


def _sheet(response):
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content), data_only=False).active


def test_default_downloads_are_separate_from_complete_workbook():
    _service()
    client = new_client()
    for kind in CATALOG:
        sheet = _sheet(client.get(f"/services/export-test/exports/{kind}.xlsx"))
        assert sheet.title
        assert [cell.value for cell in sheet[1]] == [c["heading"] for c in default_template(kind) if c["enabled"]]
    assets = _sheet(client.get("/services/export-test/exports/asset_list.xlsx"))
    assert assets.max_row == 2
    assert "docker.io/library/nginx:1.27" in [cell.value for cell in assets[2]]
    assert client.get("/services/export-test/export.xlsx").status_code == 200


def test_overview_offers_each_purpose_export():
    _service()
    response = new_client().get("/services/export-test?overview=true", headers={"Accept": "application/vnd.cats.page+json"})
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["view"]["service"]["service_key"] == "export-test"
    assert data["can"]["service.export"][str(data["view"]["service"]["id"]) ] is True
    assert not data["view"]["version"]
    assert data["history_versions"] == ["Unknown"]


def test_asset_export_uses_full_inventory_and_omits_source_credentials():
    _service()
    with SessionLocal() as db:
        service = db.scalar(select(Service).where(Service.service_key == "export-test"))
        for index in range(25):
            db.add(ServiceImage(service_id=service.id,
                                image_reference=f"registry.internal/app:{index}"))
        db.add(ServiceArtifact(service_id=service.id, artifact_type="helm_chart",
                               artifact_name="chart", chart_name="demo", chart_version="1.2.3",
                               source_reference="https://user:secret@charts.internal/demo.tgz"))
        db.commit()
    sheet = _sheet(new_client().get("/services/export-test/exports/asset_list.xlsx"))
    assert sheet.max_row == 28  # Header, original image, 25 more images, and chart.
    assert "demo" in [cell.value for cell in sheet[sheet.max_row]]
    assert "secret" not in " ".join(str(cell.value) for row in sheet for cell in row)


def test_admin_can_reorder_rename_disable_and_reset_template():
    _service()
    client = new_client()
    assert client.get("/admin/configuration/export-templates/ppsm").status_code == 200
    columns = list(reversed(default_template("ppsm")))
    columns[0]["heading"] = "Version at Export"
    columns[0]["enabled"] = True
    response = client.post("/admin/configuration/export-templates/ppsm", data={
        "csrf_token": csrf(client), "field": [c["field"] for c in columns],
        "heading": [c["heading"] for c in columns],
        "enabled": [c["field"] for c in columns if c["enabled"]],
    }, follow_redirects=False)
    assert response.status_code == 303, response.text
    sheet = _sheet(client.get("/services/export-test/exports/ppsm.xlsx"))
    assert sheet["A1"].value == "Version at Export"
    with SessionLocal() as db:
        assert db.scalar(select(PortalSetting).where(PortalSetting.key == "purpose_export_template:ppsm"))
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == "export_template.updated"))
    assert client.post("/admin/configuration/export-templates/ppsm/reset",
                       data={"csrf_token": csrf(client)}, follow_redirects=False).status_code == 303
    assert _sheet(client.get("/services/export-test/exports/ppsm.xlsx"))["A1"].value == "Service"
    with SessionLocal() as db:
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == "export_template.reset"))


def test_template_validation_and_csrf():
    _service()
    client = new_client()
    columns = default_template("poam")
    columns[0]["heading"] = "=HYPERLINK(\"https://bad.invalid\")"
    try:
        validate_template("poam", columns)
        assert False, "formula heading must be rejected"
    except ValueError:
        pass
    response = client.post("/admin/configuration/export-templates/poam", data={
        "csrf_token": "invalid", "field": [c["field"] for c in columns],
        "heading": [c["heading"] for c in columns],
        "enabled": [c["field"] for c in columns if c["enabled"]],
    })
    assert response.status_code == 403
    assert client.get("/services/export-test/exports/unknown.xlsx").status_code == 404


def test_unknown_fields_and_non_admin_are_rejected():
    _service()
    add_user("assessor-export", "Assessor")
    assessor = new_client("assessor-export")
    assert assessor.get("/services/export-test/exports/poam.xlsx").status_code == 200
    assert assessor.get("/admin/configuration/export-templates/poam").status_code == 403
    assert assessor.post("/admin/configuration/export-templates/poam/reset",
                         data={"csrf_token": csrf(assessor)}).status_code == 403
    admin = new_client()
    columns = default_template("poam")
    columns[0]["field"] = "user.password_hash"
    response = admin.post("/admin/configuration/export-templates/poam", data={
        "csrf_token": csrf(admin), "field": [c["field"] for c in columns],
        "heading": [c["heading"] for c in columns],
        "enabled": [c["field"] for c in columns if c["enabled"]],
    })
    assert response.status_code == 422
    with SessionLocal() as db:
        assert db.scalar(select(PortalSetting).where(PortalSetting.key == "purpose_export_template:poam")) is None


def test_group_template_inheritance_default_and_service_export():
    _service()
    with SessionLocal() as db:
        parent = Group(name="Cybersecurity")
        child = Group(name="Program A")
        db.add_all([parent, child])
        db.flush()
        child.parent_id = parent.id
        service = db.scalar(select(Service).where(Service.service_key == "export-test"))
        service.groups.append(child)
        parent_id, child_id = parent.id, child.id
        db.commit()
    client = new_client()
    columns = default_template("ppsm")
    columns[0]["heading"] = "Parent Service"
    response = client.post(f"/admin/general-policy/export-templates/ppsm?group_id={parent_id}", data={
        "csrf_token": csrf(client), "field": [c["field"] for c in columns],
        "heading": [c["heading"] for c in columns],
        "enabled": [c["field"] for c in columns if c["enabled"]],
    }, follow_redirects=False)
    assert response.status_code == 303, response.text
    assert _sheet(client.get("/services/export-test/exports/ppsm.xlsx"))["A1"].value == "Parent Service"
    with SessionLocal() as db:
        assert template_policy(db, "ppsm", child_id)[1] == "Inherited from Cybersecurity"
        service = db.scalar(select(Service).where(Service.service_key == "export-test"))
        assert import_heading_aliases(db, "ppsm", service)["network.port"] == ["Port"]
    response = client.post(f"/admin/general-policy/export-templates/ppsm/reset?group_id={child_id}",
                           data={"csrf_token": csrf(client), "mode": "default"}, follow_redirects=False)
    assert response.status_code == 303
    assert _sheet(client.get("/services/export-test/exports/ppsm.xlsx"))["A1"].value == "Service"
    assert client.get(f"/admin/general-policy?group_id={child_id}").status_code == 200
    assert client.post(f"/admin/general-policy/group-parent?group_id={parent_id}",
                       data={"csrf_token": csrf(client), "parent_id": str(child_id)}).status_code == 422
